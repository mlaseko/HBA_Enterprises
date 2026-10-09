"""Regenerate the user guide's screenshots: app/static/help/*.webp and shots.json (their sizes).

    cd ffe && python tools/help_shots.py          # dev machine only: needs Chromium and Playwright for Node
    HELP_SHOTS_ONLY=items-table,items-list python tools/help_shots.py   # just these pictures; shots.json keeps the rest

Nothing here touches the real database. A fictional demo project ("Msasani Villa") is seeded into a SQLite file in a temp
directory, uvicorn is started on 127.0.0.1:8777 against it, tools/help_shots.js drives Chromium through the pages, the PDFs
are rendered with pypdfium2, and everything is written as WebP into app/static/help/. Re-run it whenever a page changes
enough for its picture in the guide to be wrong.

Env: HELP_SHOTS_CHROME = path of the chrome binary for Playwright (default: Playwright's own download),
     NODE_PATH = where `require('playwright')` is found when it is not installed next to this repo."""
import io
import openpyxl
import json
import os
import random
import subprocess
import sys
import tempfile
import time
import urllib.request

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
os.chdir(ROOT)
sys.path.insert(0, ROOT)
TMP = tempfile.mkdtemp(prefix="help-shots-")
PORT = int(os.environ.get("HELP_SHOTS_PORT", "8777"))
BASE = f"http://127.0.0.1:{PORT}"
ENV = dict(os.environ, DATABASE_URL=f"sqlite:///{TMP}/demo.db", LOCAL_UPLOAD_DIR=f"{TMP}/uploads", STORAGE_BACKEND="local",
           APP_PASSWORD="demo-pass", SECRET_KEY="help-shots-" + "x" * 24, ANTHROPIC_API_KEY="")
os.environ.update(ENV)
OUT = os.path.join(ROOT, "app", "static", "help")
PNG = os.path.join(TMP, "png")
os.makedirs(PNG, exist_ok=True)

from PIL import Image, ImageDraw, ImageFont, ImageColor  # noqa: E402
from reportlab.lib.pagesizes import A3, landscape  # noqa: E402
from reportlab.lib.units import mm  # noqa: E402
from reportlab.pdfgen import canvas  # noqa: E402
import pypdfium2 as pdfium  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from app.main import app  # noqa: E402
from app import drawings  # noqa: E402

random.seed(7)
FONT = "/usr/share/fonts/truetype/dejavu/DejaVuSerif.ttf"


def font(size, bold=False):
    for cand in ([FONT.replace(".ttf", "-Bold.ttf")] if bold else []) + [FONT]:
        try:
            return ImageFont.truetype(cand, size)
        except OSError:
            continue
    return ImageFont.load_default()


# ---- pictures: sample "photos" with a texture, mood images, a cover ---------------------------------------------------
def shade(color, d):
    r, g, b = ImageColor.getrgb(color)
    return tuple(max(0, min(255, v + d)) for v in (r, g, b))


def swatch(kind, color):
    W, H = 900, 675
    im = Image.new("RGB", (W, H), color)
    d = ImageDraw.Draw(im)
    if kind == "tile":
        for y in range(0, H, 168):
            for x in range(0, W, 225):
                d.rectangle([x, y, x + 225, y + 168], fill=shade(color, random.randint(-10, 10)), outline=shade(color, 40), width=5)
    elif kind == "wood":
        for i, y in enumerate(range(0, H, 84)):
            d.rectangle([0, y, W, y + 84], fill=shade(color, (-14 if i % 2 else 8)))
            for k in range(6):
                yy = y + random.randint(6, 78)
                d.line([(0, yy), (W, yy + random.randint(-6, 6))], fill=shade(color, -26), width=1)
            d.line([(0, y), (W, y)], fill=shade(color, -40), width=2)
    elif kind == "fabric":
        for x in range(-H, W, 9):
            d.line([(x, 0), (x + H, H)], fill=shade(color, -12), width=2)
        for x in range(0, W + H, 9):
            d.line([(x, 0), (x - H, H)], fill=shade(color, 10), width=1)
    elif kind == "metal":
        for y in range(H):
            t = abs((y / H) - 0.42)
            d.line([(0, y), (W, y)], fill=shade(color, int(60 - 160 * t)))
    elif kind == "lamp":
        im = Image.new("RGB", (W, H), shade(color, -120))
        d = ImageDraw.Draw(im)
        for r in range(260, 0, -8):
            d.ellipse([W / 2 - r, H / 2 - r * 0.85, W / 2 + r, H / 2 + r * 0.85], fill=shade(color, int(70 - r / 3)))
        d.rectangle([W / 2 - 6, H / 2 + 120, W / 2 + 6, H], fill=shade(color, -60))
    elif kind == "stone":
        for _ in range(60):
            x, y = random.randint(0, W), random.randint(0, H)
            r = random.randint(20, 120)
            d.ellipse([x - r, y - r / 2, x + r, y + r / 2], fill=shade(color, random.randint(-18, 14)))
    b = io.BytesIO()
    im.save(b, "JPEG", quality=86)
    return b.getvalue()


def gradient(c1, c2, W=1200, H=800, text=None):
    im = Image.new("RGB", (W, H))
    d = ImageDraw.Draw(im)
    a, b2 = ImageColor.getrgb(c1), ImageColor.getrgb(c2)
    for y in range(H):
        t = y / H
        d.line([(0, y), (W, y)], fill=tuple(int(a[i] * (1 - t) + b2[i] * t) for i in range(3)))
    for _ in range(5):
        x, y, r = random.randint(0, W), random.randint(0, H), random.randint(120, 320)
        d.ellipse([x - r, y - r, x + r, y + r], fill=shade(c2, 18))
    if text:
        f = font(96)
        tw = d.textlength(text, font=f)
        d.text(((W - tw) / 2, H / 2 - 70), text, font=f, fill="white")
        f2 = font(30)
        sub = "Furniture & fixture schedule"
        d.text(((W - d.textlength(sub, font=f2)) / 2, H / 2 + 50), sub, font=f2, fill="white")
    b = io.BytesIO()
    im.save(b, "JPEG", quality=84)
    return b.getvalue()


def receipt():
    im = Image.new("RGB", (700, 900), "#FBF8F1")
    d = ImageDraw.Draw(im)
    f, fb = font(26), font(30, True)
    d.text((60, 60), "FOSHAN TILE CO.", font=fb, fill="#222")
    for i, line in enumerate(["Receipt no. 2026-0812-31", "Deposit, Msasani Villa", "Amount: ¥ 8,000.00", "Paid via WeChat Pay", "Date 12-08-2026"]):
        d.text((60, 130 + i * 48), line, font=f, fill="#333")
    d.line([(60, 400), (640, 400)], fill="#999", width=2)
    d.text((60, 430), "Thank you", font=f, fill="#555")
    b = io.BytesIO()
    im.save(b, "JPEG", quality=80)
    return b.getvalue()


# ---- the architect's PDF: seven A3 pages with title blocks ---------------------------------------------------------------
PAGE_W, PAGE_H = landscape(A3)
MARGIN, TB_H = 15 * mm, 34 * mm                       # margin and the title-block strip at the bottom
PLAN_X0, PLAN_Y0 = MARGIN, MARGIN + TB_H + 8 * mm       # plan area, PDF coordinates (origin bottom-left)
PLAN_W, PLAN_H = PAGE_W - 2 * MARGIN, PAGE_H - MARGIN - PLAN_Y0

GROUND = [  # (code, label, x, y, w, h) as fractions of the plan area, y from the top; code None = drawn, no room entry
    ("GF-CAR", "CARPORT", 0.00, 0.00, 0.22, 0.42), (None, "STAIRS / STORE", 0.00, 0.42, 0.22, 0.58),
    ("GF-ENT", "ENTRANCE", 0.22, 0.00, 0.16, 0.20), ("GF-HAL", "HALL & CORRIDOR", 0.22, 0.20, 0.16, 0.80),
    ("GF-LIV", "LIVING ROOM", 0.38, 0.00, 0.34, 0.48), ("GF-DIN", "DINING ROOM", 0.72, 0.00, 0.28, 0.48),
    ("GF-OFF", "OFFICE", 0.38, 0.48, 0.17, 0.52), ("GF-GST", "GUEST BEDROOM", 0.55, 0.48, 0.17, 0.30),
    ("GF-GBA", "GUEST BATH", 0.55, 0.78, 0.17, 0.22), ("GF-KIT", "KITCHEN", 0.72, 0.48, 0.28, 0.32),
    ("GF-VER", "KITCHEN VERANDAH", 0.72, 0.80, 0.28, 0.20),
]
FIRST = [
    ("FF-MBR", "MASTER BEDROOM", 0.00, 0.00, 0.36, 0.50), ("FF-MEN", "MASTER ENSUITE", 0.00, 0.50, 0.20, 0.50),
    (None, "WARDROBE", 0.20, 0.50, 0.16, 0.50), ("FF-LAN", "LANDING & CORRIDOR", 0.36, 0.00, 0.14, 1.00),
    ("FF-BR2", "BEDROOM 2", 0.50, 0.00, 0.25, 0.50), ("FF-BR3", "BEDROOM 3", 0.75, 0.00, 0.25, 0.50),
    ("FF-FBA", "FAMILY BATHROOM", 0.50, 0.50, 0.25, 0.50), ("FF-BAL", "BALCONY", 0.75, 0.50, 0.25, 0.50),
]
ROOF = [("RF-TER", "ROOF TERRACE", 0.00, 0.00, 0.78, 1.00), (None, "PLANT", 0.78, 0.00, 0.22, 0.40), (None, "WATER TANKS", 0.78, 0.40, 0.22, 0.60)]
FURNITURE = {  # a few light shapes per room, fractions of the room box
    "GF-LIV": [(0.1, 0.55, 0.5, 0.18), (0.1, 0.2, 0.18, 0.3), (0.65, 0.25, 0.25, 0.5)], "GF-DIN": [(0.2, 0.3, 0.6, 0.4)],
    "FF-MBR": [(0.3, 0.1, 0.4, 0.5), (0.1, 0.15, 0.12, 0.2), (0.78, 0.15, 0.12, 0.2)], "GF-OFF": [(0.15, 0.3, 0.7, 0.2)],
    "FF-BR2": [(0.2, 0.15, 0.35, 0.6)], "FF-BR3": [(0.2, 0.15, 0.35, 0.6)], "GF-GST": [(0.2, 0.15, 0.45, 0.6)],
    "FF-MEN": [(0.1, 0.1, 0.8, 0.3)], "GF-KIT": [(0.05, 0.05, 0.9, 0.18)],
}


def page_frac(x, y, w, h):
    """A room box (fractions of the plan area) as fractions of the whole page image, for RoomPin."""
    px = (PLAN_X0 + x * PLAN_W) / PAGE_W
    py = (PAGE_H - (PLAN_Y0 + PLAN_H) + y * PLAN_H) / PAGE_H
    return round(px, 4), round(py, 4), round(w * PLAN_W / PAGE_W, 4), round(h * PLAN_H / PAGE_H, 4)


def title_block(c, title, sheet, rev="B"):
    c.setLineWidth(1.2)
    c.rect(MARGIN, MARGIN, PAGE_W - 2 * MARGIN, TB_H)
    cols = [MARGIN, MARGIN + 150 * mm, MARGIN + 290 * mm, MARGIN + 340 * mm, PAGE_W - MARGIN]
    for x in cols[1:-1]:
        c.line(x, MARGIN, x, MARGIN + TB_H)
    c.setFont("Helvetica-Bold", 11)
    c.drawString(cols[0] + 4 * mm, MARGIN + TB_H - 9 * mm, "PROJECT: PROPOSED RESIDENCE AT MSASANI PENINSULA")
    c.setFont("Helvetica", 9)
    c.drawString(cols[0] + 4 * mm, MARGIN + TB_H - 16 * mm, "CLIENT: THE MREMA FAMILY")
    c.drawString(cols[0] + 4 * mm, MARGIN + TB_H - 22 * mm, "ARCHITECT: STUDIO M ARCHITECTS   STRUCTURAL ENGINEER: PWANI CONSULT")
    c.drawString(cols[0] + 4 * mm, MARGIN + TB_H - 28 * mm, "DRAWN: AK   CHECKED: JM   STATUS: FOR CONSTRUCTION")
    c.setFont("Helvetica", 8)
    c.drawString(cols[1] + 4 * mm, MARGIN + TB_H - 7 * mm, "DRAWING TITLE")
    c.setFont("Helvetica-Bold", 13)
    c.drawString(cols[1] + 4 * mm, MARGIN + TB_H - 17 * mm, title)
    c.setFont("Helvetica", 8)
    c.drawString(cols[2] + 4 * mm, MARGIN + TB_H - 7 * mm, "DRAWING NO.")
    c.drawString(cols[3] + 4 * mm, MARGIN + TB_H - 7 * mm, "SCALE        DATE          REV")
    c.setFont("Helvetica-Bold", 16)
    c.drawString(cols[2] + 4 * mm, MARGIN + TB_H - 19 * mm, sheet)
    c.setFont("Helvetica", 10)
    c.drawString(cols[3] + 4 * mm, MARGIN + TB_H - 19 * mm, f"1:100        03-2026        {rev}")


def draw_plan(c, rooms, title, sheet, labels_size=13):
    c.setLineWidth(4)
    c.rect(PLAN_X0, PLAN_Y0, PLAN_W, PLAN_H)
    c.setLineWidth(2)
    for code, label, x, y, w, h in rooms:
        X, Y = PLAN_X0 + x * PLAN_W, PLAN_Y0 + PLAN_H - (y + h) * PLAN_H
        W, H = w * PLAN_W, h * PLAN_H
        c.rect(X, Y, W, H)
        c.setFillGray(0.86)
        for fx, fy, fw, fh in FURNITURE.get(code or "", []):
            c.rect(X + fx * W, Y + H - (fy + fh) * H, fw * W, fh * H, stroke=0, fill=1)
        c.setFillGray(0)
        c.setFont("Helvetica", labels_size if W > 60 * mm else 9)
        c.drawCentredString(X + W / 2, Y + H / 2 - 4, label)
    c.setFont("Helvetica", 8)
    c.drawString(PLAN_X0 + 3 * mm, PLAN_Y0 + PLAN_H + 2 * mm, title + "   SCALE 1:100")
    title_block(c, title, sheet)


def drawing_set_pdf() -> bytes:
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=(PAGE_W, PAGE_H))
    # 1. cover: the drawing list (must not be pre-ticked)
    c.setFont("Helvetica-Bold", 28)
    c.drawString(MARGIN, PAGE_H - 50 * mm, "PROPOSED RESIDENCE AT MSASANI PENINSULA")
    c.setFont("Helvetica-Bold", 14)
    c.drawString(MARGIN, PAGE_H - 70 * mm, "DRAWING LIST")
    c.setFont("Helvetica", 12)
    for i, line in enumerate(["A-100  SITE LAYOUT PLAN", "A-101  GROUND FLOOR PLAN (FURNITURE LAYOUT)", "A-102  FIRST FLOOR PLAN (FURNITURE LAYOUT)",
                              "A-103  ROOF PLAN", "A-201  ELEVATIONS", "A-301  SECTION A-A"]):
        c.drawString(MARGIN, PAGE_H - 82 * mm - i * 8 * mm, line)
    title_block(c, "COVER SHEET AND DRAWING LIST", "A-000", "A")
    c.showPage()
    # 2. site plan
    c.setLineWidth(2)
    c.rect(PLAN_X0, PLAN_Y0, PLAN_W, PLAN_H)
    c.setDash(6, 4)
    c.rect(PLAN_X0 + 40 * mm, PLAN_Y0 + 30 * mm, PLAN_W - 80 * mm, PLAN_H - 60 * mm)
    c.setDash()
    c.rect(PLAN_X0 + 110 * mm, PLAN_Y0 + 70 * mm, 180 * mm, 110 * mm)
    c.setFont("Helvetica", 13)
    c.drawCentredString(PLAN_X0 + 200 * mm, PLAN_Y0 + 125 * mm, "RESIDENCE")
    c.drawString(PLAN_X0 + 45 * mm, PLAN_Y0 + PLAN_H - 40 * mm, "PLOT 12  -  MSASANI PENINSULA")
    title_block(c, "SITE LAYOUT PLAN", "A-100")
    c.showPage()
    for rooms, title, sheet in ((GROUND, "GROUND FLOOR PLAN (FURNITURE LAYOUT)", "A-101"), (FIRST, "FIRST FLOOR PLAN (FURNITURE LAYOUT)", "A-102"),
                                (ROOF, "ROOF PLAN", "A-103")):
        draw_plan(c, rooms, title, sheet)
        c.showPage()
    # 6. elevations (not a plan)
    c.setLineWidth(2)
    for k in range(2):
        X = PLAN_X0 + k * (PLAN_W / 2 + 10 * mm)
        c.rect(X, PLAN_Y0 + 20 * mm, PLAN_W / 2 - 20 * mm, 120 * mm)
        for i in range(4):
            c.rect(X + 20 * mm + i * 40 * mm, PLAN_Y0 + 60 * mm, 22 * mm, 30 * mm)
        c.setFont("Helvetica", 11)
        c.drawString(X, PLAN_Y0 + 150 * mm, ["FRONT ELEVATION", "SIDE ELEVATION"][k] + "   SCALE 1:100")
    title_block(c, "ELEVATIONS", "A-201")
    c.showPage()
    # 7. section (refers to a plan: not a plan)
    c.rect(PLAN_X0, PLAN_Y0 + 20 * mm, PLAN_W, 130 * mm)
    c.line(PLAN_X0, PLAN_Y0 + 85 * mm, PLAN_X0 + PLAN_W, PLAN_Y0 + 85 * mm)
    c.setFont("Helvetica", 11)
    c.drawString(PLAN_X0, PLAN_Y0 + 160 * mm, "SECTION A-A   SCALE 1:50   FOR SECTION LINE SEE GROUND FLOOR PLAN")
    title_block(c, "SECTION A-A", "A-301")
    c.showPage()
    c.save()
    return buf.getvalue()


# ---- the demo project ------------------------------------------------------------------------------------------------------
ROOMS = [  # code, name, floor, kind
    ("GF-ENT", "Entrance", "Ground", "area"), ("GF-HAL", "Hall & corridors", "Ground", "area"), ("GF-LIV", "Living room", "Ground", "room"),
    ("GF-DIN", "Dining room", "Ground", "room"), ("GF-KIT", "Kitchen", "Ground", "room"), ("GF-GST", "Guest bedroom", "Ground", "room"),
    ("GF-GBA", "Guest bathroom", "Ground", "room"), ("GF-OFF", "Office", "Ground", "room"), ("GF-VER", "Kitchen verandah", "Ground", "area"),
    ("GF-CAR", "Carport", "Ground", "area"), ("FF-MBR", "Master bedroom", "First", "room"), ("FF-MEN", "Master ensuite", "First", "room"),
    ("FF-BR2", "Bedroom 2", "First", "room"), ("FF-BR3", "Bedroom 3", "First", "room"), ("FF-FBA", "Family bathroom", "First", "room"),
    ("FF-LAN", "Landing & corridor", "First", "area"), ("FF-BAL", "Balconies", "First", "area"), ("RF-TER", "Roof terrace", "Roof", "area"),
]
SUPPLIERS = [  # name, city, category, contact, phone, wechat, terms
    ("Foshan Tile Co.", "Foshan - Nanzhuang", "Tiles & flooring", "Ms. Lin", "+86 138 0000 1122", "lin_tiles", "30% deposit, 70% before loading"),
    ("Shunde Furniture Hub", "Foshan - Lecong", "Furniture & soft furnishing", "Mr. Chen", "+86 139 0000 3344", "chen_lecong", "40% deposit, 60% before loading"),
    ("Guangzhou Lighting Market", "Guangzhou - Panyu", "Lighting", "Ms. Zhao", "+86 136 0000 5566", "zhao_light", "Full payment on order"),
    ("Kaiping Sanitary Ware", "Kaiping - Shuikou", "Sanitary ware & taps", "Mr. Wu", "+86 137 0000 7788", "wu_shuikou", "30% deposit, 70% before loading"),
    ("Zhongshan Cabinet Works", "Zhongshan", "Cabinets, wardrobes, worktops", "Ms. Huang", "+86 135 0000 9900", "huang_zs", "50% deposit, 50% before loading"),
]
# room, category, name, spec, brand, size, finish, qty, unit, price, currency, supplier, status, lead, (kind, colour)
ITEMS = [
    ("GF-LIV", "Furniture", "3-seat sofa, linen", "Solid wood frame, feather-wrapped foam seat, removable washable covers", "", "2400 x 950 x 850 mm", "Oatmeal linen", 1, "pcs", 6800, "CNY", "Shunde Furniture Hub", "Ordered", "35 days", ("fabric", "#D9CDB8")),
    ("GF-LIV", "Lighting", "Arc floor lamp", "Dimmable, E27, weighted marble base", "", "H 2100 mm", "Brushed brass", 1, "pcs", 1450, "CNY", "Guangzhou Lighting Market", "Quoted", "2 weeks", ("lamp", "#D8B36A")),
    ("GF-LIV", "Curtains & Soft", "Sheer curtains, floor length", "Double width, wave heading, with track", "", "W 4.2 m x H 3.0 m", "Off-white", 12.6, "m²", 180, "CNY", "Shunde Furniture Hub", "To buy", "3 weeks", ("fabric", "#EFE9DD")),
    ("GF-LIV", "Flooring", "Engineered oak floor, 190 mm plank", "15/4 mm, brushed, matt lacquer, click system", "", "190 x 1900 mm", "Natural oak", 42, "m²", 320, "CNY", "Foshan Tile Co.", "Paid", "In stock", ("wood", "#C9A572")),
    ("GF-DIN", "Furniture", "Dining table, solid oak", "Solid oak top 40 mm, tapered legs, oil finish", "", "2200 x 1000 x 750 mm", "Natural oak", 1, "pcs", 5200, "CNY", "Shunde Furniture Hub", "Ordered", "35 days", ("wood", "#BE9A66")),
    ("GF-DIN", "Furniture", "Dining chair, upholstered", "Oak frame, bouclé seat, stackable", "", "W 480 mm", "Cream bouclé", 8, "pcs", 480, "CNY", "Shunde Furniture Hub", "Ordered", "35 days", ("fabric", "#E6DED0")),
    ("GF-DIN", "Lighting", "Linear pendant over table", "LED 3000K, dimmable, adjustable suspension", "", "L 1500 mm", "Matt black", 1, "pcs", 1900, "CNY", "Guangzhou Lighting Market", "Quoted", "2 weeks", ("lamp", "#C8C0B0")),
    ("GF-KIT", "Tiles", "Floor tile, large-format porcelain", "Rectified, R10, matt, 9 mm, 10% extra for cuts", "", "600 x 1200 mm", "Warm grey", 24, "m²", 165, "CNY", "Foshan Tile Co.", "Received", "In stock", ("tile", "#B9B0A4")),
    ("GF-KIT", "Cabinets", "Kitchen cabinets, lacquered", "Plywood carcass, soft-close hinges and drawers, handleless", "", "Per drawing A-101", "Sage green matt lacquer", 1, "set", 28000, "CNY", "Zhongshan Cabinet Works", "Paid", "45 days", ("wood", "#8FA08A")),
    ("GF-KIT", "Countertops", "Quartz worktop, 20 mm", "Mitred edge, undermount sink cut-out, polished", "", "6.4 m² per drawing", "Calacatta look", 6.4, "m²", 980, "CNY", "Zhongshan Cabinet Works", "Paid", "45 days", ("stone", "#E9E4DC")),
    ("GF-KIT", "Plumbing & Sanitary", "Kitchen mixer, pull-out", "Brass body, ceramic cartridge, 2 spray modes", "", "H 420 mm", "Brushed nickel", 1, "pcs", 620, "CNY", "Kaiping Sanitary Ware", "Shipped", "In stock", ("metal", "#B5B8BC")),
    ("GF-KIT", "Appliances", "Built-in oven, 60 cm", "72 L, pyrolytic, A+ energy", "", "595 x 595 mm", "Black glass", 1, "pcs", 3600, "CNY", "", "To buy", "", ("metal", "#3B3B3F")),
    ("GF-GBA", "Plumbing & Sanitary", "Wall-hung WC, rimless", "With concealed cistern and soft-close seat", "", "540 x 360 mm", "White", 1, "pcs", 1350, "CNY", "Kaiping Sanitary Ware", "Received", "In stock", ("metal", "#F2F2F0")),
    ("GF-GBA", "Tiles", "Wall tile, zellige look", "Glossy, hand-made look, 15% extra for cuts", "", "100 x 100 mm", "Sea green", 18, "m²", 210, "CNY", "Foshan Tile Co.", "Received", "In stock", ("tile", "#7FA79A")),
    ("GF-GBA", "Plumbing & Sanitary", "Basin mixer, brushed brass", "Single lever, PVD finish, with waste", "", "H 160 mm", "Brushed brass", 1, "pcs", 540, "CNY", "Kaiping Sanitary Ware", "Shipped", "In stock", ("metal", "#C9A24A")),
    ("GF-OFF", "Furniture", "Writing desk", "Oak veneer, two drawers, cable grommet", "", "1600 x 700 x 750 mm", "Oak", 1, "pcs", 2400, "CNY", "Shunde Furniture Hub", "Quoted", "35 days", ("wood", "#B48B5C")),
    ("GF-OFF", "Lighting", "Desk lamp", "LED, touch dimmer, USB port", "", "H 450 mm", "Matt black", 1, "pcs", 380, "CNY", "Guangzhou Lighting Market", "To buy", "In stock", ("lamp", "#A9A49B")),
    ("GF-ENT", "Lighting", "Entrance pendant, rattan", "E27, natural rattan shade", "", "Ø 500 mm", "Natural", 1, "pcs", 760, "CNY", "Guangzhou Lighting Market", "Ordered", "2 weeks", ("lamp", "#D4B27A")),
    ("GF-ENT", "Furniture", "Console table", "Solid oak, one drawer", "", "1200 x 350 x 800 mm", "Oak", 1, "pcs", 1900, "CNY", "Shunde Furniture Hub", "To buy", "35 days", ("wood", "#B8925E")),
    ("GF-HAL", "Tiles", "Corridor floor tile", "Same as kitchen floor, 10% extra for cuts", "", "600 x 1200 mm", "Warm grey", 16, "m²", 165, "CNY", "Foshan Tile Co.", "Received", "In stock", ("tile", "#B9B0A4")),
    ("GF-HAL", "Lighting", "Recessed downlight", "LED 7 W, 3000K, CRI 90, dimmable", "", "Ø 90 mm cut-out", "White", 14, "pcs", 65, "CNY", "Guangzhou Lighting Market", "Paid", "In stock", ("lamp", "#E8E4DA")),
    ("GF-CAR", "Lighting", "Outdoor wall light", "IP65, LED 3000K, up/down", "", "H 220 mm", "Anthracite", 2, "pcs", 290, "CNY", "Guangzhou Lighting Market", "To buy", "In stock", ("lamp", "#6B6E73")),
    ("GF-VER", "Furniture", "Outdoor dining set, teak", "Table + 6 chairs, grade A teak, weather cushions", "", "1800 x 900 mm", "Teak", 1, "set", 7400, "CNY", "Shunde Furniture Hub", "Quoted", "40 days", ("wood", "#A87E52")),
    ("FF-MBR", "Furniture", "King bed, upholstered headboard", "Solid wood slats, headboard H 1200 mm, bouclé", "", "1800 x 2000 mm", "Sand bouclé", 1, "pcs", 5600, "CNY", "Shunde Furniture Hub", "Ordered", "35 days", ("fabric", "#D7C7B0")),
    ("FF-MBR", "Furniture", "Bedside table", "One drawer, oak, soft-close", "", "500 x 400 x 550 mm", "Oak", 2, "pcs", 890, "CNY", "Shunde Furniture Hub", "Ordered", "35 days", ("wood", "#B48B5C")),
    ("FF-MBR", "Curtains & Soft", "Blackout curtains", "Triple weave, pinch pleat, with track", "", "W 3.6 m x H 2.8 m", "Warm taupe", 10, "m²", 220, "CNY", "Shunde Furniture Hub", "To buy", "3 weeks", ("fabric", "#B8A594")),
    ("FF-MEN", "Plumbing & Sanitary", "Freestanding bath", "Acrylic, 1700, with click-clack waste", "", "1700 x 800 x 580 mm", "Gloss white", 1, "pcs", 4800, "CNY", "Kaiping Sanitary Ware", "Paid", "3 weeks", ("metal", "#F4F4F2")),
    ("FF-MEN", "Plumbing & Sanitary", "Rain shower set", "Thermostatic, 300 mm head, hand shower", "", "Head Ø 300 mm", "Brushed brass", 1, "pcs", 1680, "CNY", "Kaiping Sanitary Ware", "Paid", "In stock", ("metal", "#C9A24A")),
    ("FF-MEN", "Tiles", "Floor tile, slip-resistant", "R11, matt, 10% extra for cuts", "", "600 x 600 mm", "Light stone", 9, "m²", 195, "CNY", "Foshan Tile Co.", "Shipped", "In stock", ("stone", "#D8D0C2")),
    ("FF-BR2", "Furniture", "Single bed with storage", "Two drawers, solid pine, lacquered", "", "1000 x 2000 mm", "White", 1, "pcs", 2900, "CNY", "Shunde Furniture Hub", "To buy", "35 days", ("wood", "#E9E3D8")),
    ("FF-BR3", "Furniture", "Wardrobe, sliding doors", "Melamine carcass, soft-close sliders, interior lighting", "", "2400 x 600 x 2400 mm", "Oak + mirror", 1, "pcs", 6200, "CNY", "Zhongshan Cabinet Works", "Quoted", "45 days", ("wood", "#B59364")),
    ("FF-FBA", "Plumbing & Sanitary", "Vanity unit with basin", "Wall-hung, 2 drawers, ceramic basin", "", "900 x 460 mm", "Oak + white", 1, "pcs", 3200, "CNY", "Kaiping Sanitary Ware", "Ordered", "3 weeks", ("stone", "#EDE8E0")),
    ("FF-LAN", "Lighting", "Landing pendant", "Three glass globes, dimmable", "", "L 900 mm", "Smoked glass / brass", 1, "pcs", 980, "CNY", "Guangzhou Lighting Market", "To buy", "2 weeks", ("lamp", "#9A8A74")),
    ("FF-BAL", "Furniture", "Balcony chairs, rope", "Aluminium frame, rope weave, outdoor", "", "W 620 mm", "Charcoal", 4, "pcs", 650, "CNY", "Shunde Furniture Hub", "To buy", "40 days", ("fabric", "#595C60")),
    ("RF-TER", "Furniture", "Sun loungers", "Teak, adjustable back, wheels", "", "2000 x 700 mm", "Teak", 2, "pcs", 1800, "CNY", "Shunde Furniture Hub", "To buy", "40 days", ("wood", "#A87E52")),
    ("", "Hardware", "Door handles, lever on rose", "Stainless 304, with privacy sets for bathrooms", "", "", "Brushed brass", 24, "set", 145, "CNY", "Zhongshan Cabinet Works", "Quoted", "In stock", ("metal", "#C9A24A")),
    ("", "Paint", "Interior emulsion, matt", "Low VOC, 2 coats, 18 L buckets", "", "18 L", "Warm white", 18, "pcs", 380, "CNY", "", "To buy", "", ("fabric", "#F1EDE4")),
    ("GF-GBA", "Lighting", "Mirror light", "LED 3000K, IP44", "", "L 600 mm", "Brushed brass", 1, "pcs", 42, "USD", "Guangzhou Lighting Market", "To buy", "In stock", ("lamp", "#D8B36A")),
]
PAYMENTS = [  # supplier, date, amount, kind, reference, note, receipt
    ("Foshan Tile Co.", "2026-08-12", 8000, "Deposit", "WeChat 8831", "Tiles + oak floor deposit", True),
    ("Foshan Tile Co.", "2026-09-02", 12736, "Balance", "Bank 20260902-7", "Balance before loading", False),
    ("Zhongshan Cabinet Works", "2026-08-20", 17000, "Deposit", "Bank 20260820-2", "Kitchen + worktop 50%", False),
    ("Shunde Furniture Hub", "2026-09-10", 9000, "Deposit", "WeChat 9120", "Sofa, table, chairs, bed 40%", False),
    ("Kaiping Sanitary Ware", "2026-09-14", 4500, "Deposit", "WeChat 9377", "Sanitary 30%", False),
]
CARTONS = [  # supplier, room, contents, codes, qty, L, W, H, kg, received, by_supplier
    ("Foshan Tile Co.", "GF-KIT", "Floor tiles 600x1200, 12 cartons on pallet", "GF-KIT-01", 12, 125, 65, 110, 480, True, False),
    ("Foshan Tile Co.", "GF-HAL", "Floor tiles 600x1200, 8 cartons", "GF-HAL-01", 8, 125, 65, 80, 320, True, False),
    ("Foshan Tile Co.", "GF-GBA", "Wall tiles 100x100, 6 cartons", "GF-GBA-02", 6, 60, 40, 50, 150, True, False),
    ("Foshan Tile Co.", "GF-LIV", "Engineered oak, 21 packs", "GF-LIV-04", 21, 195, 60, 60, 410, False, False),
    ("Kaiping Sanitary Ware", "GF-GBA", "Wall-hung WC + cistern frame", "GF-GBA-01", 1, 70, 45, 60, 38, False, True),
    ("Kaiping Sanitary Ware", "FF-MEN", "Freestanding bath in crate", "FF-MEN-01", 1, 180, 85, 70, 62, False, True),
    ("Kaiping Sanitary Ware", "FF-MEN", "Rain shower set, mixers", "FF-MEN-02 GF-KIT-04 GF-GBA-03", 3, 60, 40, 30, 14, False, True),
    ("Shunde Furniture Hub", "GF-DIN", "Dining table top + legs, crate", "GF-DIN-01", 1, 230, 110, 25, 95, False, False),
    ("Shunde Furniture Hub", "GF-DIN", "Dining chairs, 2 per carton", "GF-DIN-02", 4, 100, 55, 95, 22, False, False),
]


def seed() -> dict:
    ids = {}
    with TestClient(app) as c:
        c.post("/login", data={"password": ENV["APP_PASSWORD"], "next": "/"}, follow_redirects=False)
        c.post("/settings", data={"studio_name": "HBA Interiors", "studio_phone": "+255 700 000 000", "studio_email": "studio@example.com",
                                  "studio_website": "hba-interiors.example", "default_rate": "7.1"}, follow_redirects=False)
        r = c.post("/projects/new", data={"client_name": "The Mrema family", "name": "Msasani Villa", "address": "Plot 12, Msasani Peninsula, Dar es Salaam",
                                           "rate": "7.1", "budget_usd": "60000", "description": "A four-bedroom family house by the sea: warm wood, linen, brushed brass and sea-green tiles."},
                   follow_redirects=False)
        pid = int(r.headers["location"].split("/")[-1])
        ids["pid"] = pid
        for code, name, floor, kind in ROOMS:
            c.post(f"/p/{pid}/rooms", data={"code": code, "name": name, "floor": floor, "kind": kind}, follow_redirects=False)
        for name, city, cat, contact, phone, wechat, terms in SUPPLIERS:
            c.post("/suppliers/new", data={"name": name, "city": city, "category": cat, "contact": contact, "phone": phone, "wechat": wechat,
                                           "payment_terms": terms, "back": f"/p/{pid}/suppliers"}, follow_redirects=False)
        from app.db import SessionLocal
        from app.models import Room, Supplier, Item, Project, ProjectImage, DrawingSet, SupplierLink
        db = SessionLocal()
        rooms = {x.code: x.id for x in db.query(Room).filter(Room.project_id == pid)}
        sups = {x.name: x.id for x in db.query(Supplier)}
        db.close()
        ids.update(rooms=rooms, sups=sups)
        # the architect's PDF, picked once for the plans and uploaded again (unpicked) for the page-picker picture
        pdf = drawing_set_pdf()
        for k, want in ((1, False), (2, True), (3, True), (4, True), (5, True), (6, False), (7, False)):
            tb = drawings.read_title_block(drawings.page_text(pdf, k))
            assert tb["is_plan"] is want, (k, tb)
        c.post(f"/p/{pid}/images/plans", data={"floor": "", "sheet": ""}, files=[("files", ("Msasani Villa - Rev B.pdf", pdf, "application/pdf"))], follow_redirects=False)
        db = SessionLocal()
        set1 = db.query(DrawingSet).filter(DrawingSet.project_id == pid).order_by(DrawingSet.id).first().id
        db.close()
        c.post(f"/p/{pid}/images/sets/{set1}/pick", data={"page_3": "on", "floor_3": "Ground", "sheet_3": "A-101", "caption_3": "Furniture Layout",
                                                          "page_4": "on", "floor_4": "First", "sheet_4": "A-102", "caption_4": "Furniture Layout",
                                                          "page_5": "on", "floor_5": "Roof", "sheet_5": "A-103", "caption_5": ""}, follow_redirects=False)
        c.post(f"/p/{pid}/images/plans", data={"floor": "", "sheet": ""}, files=[("files", ("Msasani Villa - Rev C.pdf", pdf, "application/pdf"))], follow_redirects=False)
        db = SessionLocal()
        sets = [s.id for s in db.query(DrawingSet).filter(DrawingSet.project_id == pid).order_by(DrawingSet.id)]
        plans = {im.floor: im.id for im in db.query(ProjectImage).filter(ProjectImage.project_id == pid, ProjectImage.kind == "floorplan")}
        db.close()
        ids.update(set2=sets[1], plans=plans)
        # cover and mood board
        c.post(f"/p/{pid}/images", data={"kind": "cover", "caption": ""}, files=[("files", ("cover.jpg", gradient("#2B211D", "#B9593A", text="Msasani Villa"), "image/jpeg"))], follow_redirects=False)
        for i, (a, b2, cap) in enumerate((("#D9CDB8", "#B9593A", "Warm wood & linen"), ("#6F8B74", "#E7EFE8", "Sea-green tiles"), ("#C9A227", "#F8F0D6", "Brushed brass"))):
            c.post(f"/p/{pid}/images", data={"kind": "mood", "caption": cap}, files=[("files", (f"mood{i}.jpg", gradient(a, b2), "image/jpeg"))], follow_redirects=False)
        # items, with a sample photo each
        for room, cat, name, spec, brand, size, finish, qty, unit, price, cur, sup, status, lead, (kind, colour) in ITEMS:
            data = {"category": cat, "name": name, "spec": spec, "brand": brand, "size": size, "finish": finish, "lead_time": lead, "qty": str(qty), "unit": unit,
                    "unit_price": str(price), "price_currency": cur, "supplier_id": str(sups.get(sup, "")), "status": status, "notes": ""}
            if room:
                data["room_ids"] = [str(rooms[room])]
            c.post(f"/p/{pid}/items/new", data=data, files=[("photos", ("sample.jpg", swatch(kind, colour), "image/jpeg"))], follow_redirects=False)
        # the same product in several rooms: one line per room in the schedule, one line on the Item list view
        db = SessionLocal()
        beds = [x.id for x in db.query(Room).filter(Room.project_id == pid) if x.group == "bedroom"]
        baths = [x.id for x in db.query(Room).filter(Room.project_id == pid) if x.group == "bathroom"]
        db.close()
        for room_ids, cat, name, spec, size, finish, qty, price, sup, (kind, colour) in (
                (beds, "Lighting", "Bedside wall light", "7 W LED 2700 K, swing arm, switch on the base", "350 mm", "Brushed brass", 2, 420, "Guangzhou Lighting Market", ("lamp", "#B08D57")),
                (baths, "Plumbing & Sanitary", "Heated towel rail", "Electric, 60 W, with timer", "1200 x 500 mm", "Brushed brass", 1, 980, "Kaiping Sanitary Ware", ("metal", "#C8B98F"))):
            c.post(f"/p/{pid}/items/new", data={"room_ids": [str(x) for x in room_ids], "category": cat, "name": name, "spec": spec, "size": size, "finish": finish, "qty": str(qty),
                                               "unit": "pcs", "unit_price": str(price), "price_currency": "CNY", "supplier_id": str(sups.get(sup, "")), "status": "Quoted", "lead_time": "3 weeks", "notes": ""},
                   files=[("photos", ("sample.jpg", swatch(kind, colour), "image/jpeg"))], follow_redirects=False)
        # an edited export for the import preview picture: two prices, a status and one new row
        xl = openpyxl.load_workbook(io.BytesIO(c.get(f"/p/{pid}/export/items.xlsx").content))
        sl = xl["Shopping List"]
        at = {sl.cell(row=i, column=4).value: i for i in range(2, 120) if sl.cell(row=i, column=1).value}
        sl.cell(row=at["3-seat sofa, linen"], column=12, value=6400)
        sl.cell(row=at["Arc floor lamp"], column=15, value="Ordered")
        sl.cell(row=at["King bed, upholstered headboard"], column=12, value=5900)
        sl.append(["", "All bedrooms", "Lighting", "Reading light", None, "7 W LED, swing arm, switch on the base", "", "350 mm", "Brushed brass", 1, "pcs", 390,
                   "Guangzhou Lighting Market", "3 weeks", "To buy", "no", ""])
        ids["edited_xlsx"] = os.path.join(TMP, "edited-items.xlsx")
        xl.save(ids["edited_xlsx"])
        db = SessionLocal()
        items = db.query(Item).filter(Item.project_id == pid, Item.draft == False).order_by(Item.id).all()  # noqa: E712
        by_code = {i.code: i.id for i in items}
        db.close()
        ids["items"] = by_code
        ids["sofa"] = by_code["GF-LIV-01"]
        for sup, date, amount, kind, ref, note, rec in PAYMENTS:
            files = [("receipt", ("receipt.jpg", receipt(), "image/jpeg"))] if rec else []
            c.post(f"/p/{pid}/payments", data={"supplier_id": str(sups[sup]), "paid_on": date, "amount": str(amount), "kind": kind, "reference": ref, "note": note},
                   files=files, follow_redirects=False)
        # packing links for two suppliers; the supplier enters some boxes through their link
        for sup in ("Kaiping Sanitary Ware", "Foshan Tile Co."):
            c.post(f"/p/{pid}/suppliers/{sups[sup]}/link", data={"action": "create"}, follow_redirects=False)
        db = SessionLocal()
        links = {l.supplier_id: l.token for l in db.query(SupplierLink).filter(SupplierLink.project_id == pid)}
        db.close()
        ids["stoken"] = links[sups["Kaiping Sanitary Ware"]]
        for sup, room, contents, codes, qty, L, W, H, kg, received, by_sup in CARTONS:
            data = {"room_id": str(rooms[room]), "contents": contents, "item_codes": codes, "qty": str(qty), "length_cm": str(L), "width_cm": str(W), "height_cm": str(H), "weight_kg": str(kg)}
            if by_sup:
                c.post(f"/s/{links[sups[sup]]}/cartons", data=data, follow_redirects=False)
            else:
                c.post(f"/p/{pid}/cartons", data=dict(data, supplier_id=str(sups[sup]), notes=""), follow_redirects=False)
        from app.models import Carton
        db = SessionLocal()
        for ct, (sup, room, contents, codes, qty, L, W, H, kg, received, by_sup) in zip(db.query(Carton).filter(Carton.project_id == pid).order_by(Carton.id), CARTONS):
            ct.received = received
        db.commit()
        db.close()
        # room boxes on the plans and a few item dots
        for floor, layout in (("Ground", GROUND), ("First", FIRST), ("Roof", ROOF)):
            for code, label, x, y, w, h in layout:
                if code:
                    px, py, pw, ph = page_frac(x, y, w, h)
                    c.post(f"/p/{pid}/plan/pins", data={"image_id": plans[floor], "room_id": rooms[code], "x": px, "y": py, "w": pw, "h": ph}, headers={"Accept": "application/json"})
        box = {code: (x, y, w, h) for code, label, x, y, w, h in GROUND if code}
        dots = [("GF-KIT-01", 0.5, 0.75), ("GF-KIT-02", 0.5, 0.15), ("GF-KIT-03", 0.2, 0.15), ("GF-KIT-04", 0.8, 0.2),
                ("GF-LIV-01", 0.35, 0.64), ("GF-LIV-02", 0.12, 0.3), ("GF-LIV-04", 0.78, 0.5), ("GF-DIN-01", 0.5, 0.5), ("GF-DIN-02", 0.3, 0.5), ("GF-DIN-03", 0.5, 0.25),
                ("GF-ENT-01", 0.5, 0.5), ("GF-OFF-01", 0.5, 0.4), ("GF-GBA-01", 0.3, 0.5), ("GF-GBA-03", 0.7, 0.3)]
        for code, fx, fy in dots:
            bx, by, bw, bh = box[code[:6]]
            px, py, _, _ = page_frac(bx + bw * fx, by + bh * fy, 0, 0)
            c.post(f"/p/{pid}/plan/item-pins", data={"image_id": plans["Ground"], "item_id": by_code[code], "x": px, "y": py}, headers={"Accept": "application/json"})
        # two drafts from quick capture
        for kind, colour in (("fabric", "#C5B9A5"), ("lamp", "#D8B36A")):
            c.post(f"/p/{pid}/capture", data={"room_id": str(rooms["FF-BR2"]), "category": "", "n": "0"}, files=[("photo", ("shot.jpg", swatch(kind, colour), "image/jpeg"))],
                   headers={"Accept": "application/json"})
        # a second, smaller project for the projects list
        r = c.post("/projects/new", data={"client_name": "Ms. L. Kimaro", "name": "Oyster Bay Apartment", "address": "Toure Drive, Oyster Bay", "rate": "7.1", "budget_usd": "18000"},
                   follow_redirects=False)
        pid2 = int(r.headers["location"].split("/")[-1])
        for code, name, floor, kind in (("LIV", "Living room", "Apartment", "room"), ("BED", "Bedroom", "Apartment", "room"), ("KIT", "Kitchen", "Apartment", "room"), ("BAL", "Balcony", "Apartment", "area")):
            c.post(f"/p/{pid2}/rooms", data={"code": code, "name": name, "floor": floor, "kind": kind}, follow_redirects=False)
        db = SessionLocal()
        r2 = {x.code: x.id for x in db.query(Room).filter(Room.project_id == pid2)}
        db.close()
        for room, cat, name, qty, price, status, kind, colour in (("LIV", "Furniture", "2-seat sofa", 1, 4200, "Received", "fabric", "#A8B5A0"), ("LIV", "Lighting", "Floor lamp", 1, 980, "Received", "lamp", "#D8B36A"),
                                                                   ("BED", "Furniture", "Queen bed", 1, 3900, "Shipped", "wood", "#B48B5C"), ("KIT", "Tiles", "Splashback tile", 4, 240, "Received", "tile", "#7FA79A"),
                                                                   ("BAL", "Furniture", "Bistro set", 1, 1400, "To buy", "metal", "#6B6E73")):
            c.post(f"/p/{pid2}/items/new", data={"room_ids": [str(r2[room])], "category": cat, "name": name, "qty": str(qty), "unit": "pcs", "unit_price": str(price), "price_currency": "CNY", "status": status},
                   files=[("photos", ("sample.jpg", swatch(kind, colour), "image/jpeg"))], follow_redirects=False)
        c.post(f"/p/{pid2}/images", data={"kind": "cover", "caption": ""}, files=[("files", ("cover.jpg", gradient("#1F3A5F", "#6F8B74"), "image/jpeg"))], follow_redirects=False)
        db = SessionLocal()
        ids["token"] = db.get(Project, pid).client_token
        db.close()
    return ids


# ---- the PDFs, rendered page by page ------------------------------------------------------------------------------------
def pdf_png(data: bytes, name: str, pick, width=1100):
    """Render one page of a PDF to PNG/<name>.png. pick = page index, or a text the page must contain (first match after the cover)."""
    doc = pdfium.PdfDocument(data)
    try:
        idx = pick
        if isinstance(pick, str):
            idx = None
            for i in range(len(doc)):
                tp = doc[i].get_textpage()
                txt = tp.get_text_range() or ""
                tp.close()
                if pick.lower() in txt.lower() and i > 0:
                    idx = i
                    break
            assert idx is not None, (name, pick)
        if idx < 0:
            idx = len(doc) + idx
        page = doc[idx]
        w, h = page.get_size()
        img = page.render(scale=width / w).to_pil().convert("RGB")
        img.save(os.path.join(PNG, name + ".png"))
    finally:
        doc.close()


def render_pdfs(ids):
    pid, kit = ids["pid"], ids["rooms"]["GF-KIT"]
    with TestClient(app) as c:
        c.post("/login", data={"password": ENV["APP_PASSWORD"], "next": "/"}, follow_redirects=False)
        sched = c.get(f"/p/{pid}/export/schedule.pdf").content
        pdf_png(sched, "pdf-cover", 0)
        pdf_png(sched, "pdf-contents", 1)
        pdf_png(sched, "pdf-plan", "Ground floor")
        pdf_png(sched, "pdf-table", "Kitchen mixer, pull-out")
        pdf_png(sched, "pdf-summary", -1)
        floor = c.get(f"/p/{pid}/export/schedule.pdf?layout=floor").content
        pdf_png(floor, "pdf-floor", "Living room")
        pdf_png(c.get(f"/p/{pid}/export/room/{kit}.pdf").content, "pdf-checklist", 0)
        pdf_png(c.get(f"/p/{pid}/export/labels.pdf").content, "pdf-labels", 0, width=900)
        pdf_png(c.get(f"/p/{pid}/export/packing.pdf").content, "pdf-packing", 0)
        pdf_png(c.get(f"/p/{pid}/export/po/{ids['sups']['Foshan Tile Co.']}.pdf").content, "pdf-po", 0)


# ---- the browser, the pictures ----------------------------------------------------------------------------------------------
def start_server():
    proc = subprocess.Popen([sys.executable, "-m", "uvicorn", "app.main:app", "--host", "127.0.0.1", "--port", str(PORT), "--log-level", "warning"], env=ENV, cwd=ROOT)
    for _ in range(60):
        try:
            urllib.request.urlopen(BASE + "/login", timeout=1).read()
            return proc
        except Exception:
            time.sleep(0.5)
    proc.kill()
    raise SystemExit("the dev server did not start")


def to_webp():
    os.makedirs(OUT, exist_ok=True)
    sizes = {}
    if os.environ.get("HELP_SHOTS_ONLY"):  # a partial run keeps the sizes of the pictures it did not take
        try:
            with open(os.path.join(OUT, "shots.json")) as f:
                sizes = json.load(f)
        except OSError:
            pass
    for f in sorted(os.listdir(PNG)):
        if not f.endswith(".png"):
            continue
        name = f[:-4]
        im = Image.open(os.path.join(PNG, f)).convert("RGB")
        if im.width > 1100:
            im = im.resize((1100, round(im.height * 1100 / im.width)), Image.LANCZOS)
        im.save(os.path.join(OUT, name + ".webp"), "WEBP", quality=80, method=6)
        sizes[name] = [im.width, im.height]
    with open(os.path.join(OUT, "shots.json"), "w") as f:
        json.dump(sizes, f, indent=0, sort_keys=True)
    total = sum(os.path.getsize(os.path.join(OUT, x)) for x in os.listdir(OUT) if x.endswith(".webp"))
    print(f"{len(sizes)} pictures, {total / 1024:.0f} KB total -> {OUT}")


if __name__ == "__main__":
    ids = seed()
    print("seeded", {k: v for k, v in ids.items() if k not in ("items",)})
    with open(os.path.join(TMP, "ids.json"), "w") as f:
        json.dump(ids, f)
    render_pdfs(ids)
    srv = start_server()
    try:
        env = dict(os.environ)
        env.setdefault("HELP_SHOTS_CHROME", "")
        subprocess.run(["node", os.path.join(ROOT, "tools", "help_shots.js"), BASE, PNG, os.path.join(TMP, "ids.json"), ENV["APP_PASSWORD"]], check=True, env=env)
    finally:
        srv.terminate()
        srv.wait(timeout=10)
    to_webp()
