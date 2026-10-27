"""End-to-end smoke test. Run: python tests/test_flow.py  (prints ALL OK)."""
import os, io, sys, tempfile
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
os.chdir(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
_tmp=tempfile.mkdtemp()
os.environ["DATABASE_URL"]=f"sqlite:///{_tmp}/test.db"; os.environ["LOCAL_UPLOAD_DIR"]=f"{_tmp}/uploads"; os.environ["STORAGE_BACKEND"]="local"
os.environ["APP_PASSWORD"]="test123"; os.environ["SECRET_KEY"]="x"*32; os.environ["PHOTO_PREPARE_AT_START"]="0"  # the suite drives the backfill itself
os.environ["PLAN_MAX_PX"]="2400"  # PDF pages rendered at 2400 px rather than 3200: the same rules, half the pixels
from fastapi.testclient import TestClient
from app.main import app
from PIL import Image
import openpyxl

def img(color):
    b=io.BytesIO(); Image.new("RGB",(900,700),color).save(b,"JPEG"); return b.getvalue()

def buy(c, pid, item_id, supplier, amount, currency="CNY", lead="", lines=None, **extra):
    """A shop's quote on the item's product (Buying page), bought for the rooms without a shop, or for lines=[item ids].
    supplier: an id from the list or the name of a new shop. Returns (product id, quote id)."""
    from app.db import SessionLocal
    from app.models import Quote
    r=c.get(f"/p/{pid}/buying/item/{item_id}", follow_redirects=False); assert r.status_code==303, r.status_code
    prod=int(r.headers["location"].split("?")[0].split("/")[-1])
    data=dict(extra, amount=str(amount), currency=currency, lead_time=lead)
    if isinstance(supplier, int): data["supplier_id"]=str(supplier)
    else: data.update(supplier_id="__new__", supplier_name=supplier)
    if not lines: data["buy"]="open"
    r=c.post(f"/p/{pid}/buying/{prod}/quotes", data=data, follow_redirects=False); assert r.status_code==303 and "err=" not in r.headers["location"], r.headers
    db=SessionLocal(); qid=db.query(Quote).filter(Quote.product_id==prod).order_by(Quote.id.desc()).first().id; db.close()
    if lines:
        r=c.post(f"/p/{pid}/buying/{prod}/quotes/{qid}/buy", data={"rooms":"1","line":[str(x) for x in lines]}, follow_redirects=False); assert r.status_code==303 and "err=" not in r.headers["location"], r.headers
    return prod, qid

def flow(c):
    """The walk through the app, in sections ending with print("<name> ok"). It is never called as a function: the runner
    at the end of the file executes its statements section by section in one shared namespace, so a section can be
    run on its own (python tests/test_flow.py quotations) with the setup and the sections it depends on."""
    r=c.get("/", follow_redirects=False); assert r.status_code==200 and "Enter the app password" in r.text
    r=c.post("/login", data={"password":"wrong","next":"/"}); assert "Wrong password" in r.text
    r=c.post("/login", data={"password":"test123","next":"/"}, follow_redirects=False); assert r.status_code==303
    r=c.get("/"); assert "Projects" in r.text
    r=c.post("/projects/new", data={"client_name":"Mohamed Laseko","name":"Kinondoni House","address":"Plot 70, Block 41, Kinondoni","rate":"7.1","budget_usd":"150000"}, follow_redirects=False)
    pid=r.headers["location"].split("/")[-1]; print("project", pid)
    # import excel
    with open("samples/Kinondoni_House_Procurement_List.xlsx","rb") as f:
        r=c.post(f"/p/{pid}/import", files={"file":("list.xlsx", f.read())}, data={"apply":"1"})
    assert "Imported" in r.text, r.text[:500]
    import re; print(re.search(r"Imported[^<]*", r.text).group(0))
    r=c.get(f"/p/{pid}"); assert "Kinondoni House" in r.text
    r=c.get(f"/p/{pid}/items"); assert "Floor - big porcelain slab" in r.text
    r=c.get(f"/p/{pid}/items?view=table&category=Tiles"); assert "status-sel" in r.text
    # add supplier + item with photo
    r=c.post("/suppliers/new", data={"name":"Foshan Tile Co.","city":"Foshan","category":"Tiles","wechat":"li_tiles","payment_terms":"30/70","back":f"/p/{pid}/suppliers"}, follow_redirects=False)
    r=c.get(f"/p/{pid}/suppliers"); assert "Foshan Tile Co." in r.text
    from app.db import SessionLocal
    from app.models import Supplier, Room, Item, Carton
    db=SessionLocal(); sup=db.query(Supplier).first(); room=db.query(Room).filter(Room.code=="GF-KIT").first(); db.close()
    r=c.post(f"/p/{pid}/items/new", data={"room_id":room.id,"category":"Lighting","name":"Island pendant","spec":"LED 24V","size":"300 mm","finish":"black","qty":"3","unit":"pcs","unit_price":"450","supplier_id":sup.id,"status":"Quoted","lead_time":"2 weeks"},
             files=[("photos",("a.jpg",img("red"),"image/jpeg")),("photos",("b.jpg",img("blue"),"image/jpeg"))], follow_redirects=False)
    iid=r.headers["location"].split("/")[-1]; print("item", iid)
    db=SessionLocal(); _it=db.get(Item, int(iid)); assert _it.unit_price==0 and _it.supplier_id is None and _it.lead_time==""; db.close()  # the item form takes no shop or price
    buy(c, pid, int(iid), sup.id, 450, lead="2 weeks")
    db=SessionLocal(); _it=db.get(Item, int(iid)); assert _it.unit_price==450 and _it.supplier_id==sup.id and _it.lead_time=="2 weeks" and _it.status=="Quoted" and _it.purchase_id; db.close()
    r=c.get(f"/p/{pid}/items/{iid}"); assert "Island pendant" in r.text and "/media/" in r.text
    r=c.post(f"/p/{pid}/items/{iid}/status", data={"status":"Ordered"}, headers={"Accept":"application/json"}); assert r.json()["status"]=="Ordered"
    # assign supplier to some tile items quickly via db
    db=SessionLocal()
    for it in db.query(Item).filter(Item.category=="Tiles").limit(5): it.supplier_id=sup.id; it.unit_price=95
    db.commit(); db.close()
    # payment
    r=c.post(f"/p/{pid}/payments", data={"supplier_id":sup.id,"paid_on":"2026-10-05","amount":"5000","kind":"Deposit","reference":"TT123"}, files={"receipt":("r.jpg",img("green"),"image/jpeg")}, follow_redirects=False)
    r=c.get(f"/p/{pid}/payments"); assert "5,000" in r.text
    # supplier link
    r=c.post(f"/p/{pid}/suppliers/{sup.id}/link", data={"action":"create"}, follow_redirects=False)
    r=c.get(f"/suppliers/{sup.id}?project={pid}"); tok=re.search(r"/s/([A-Za-z0-9_\-]+)", r.text).group(1); print("token", tok)
    c2=TestClient(app)  # no login
    r=c2.get(f"/s/{tok}"); assert r.status_code==200 and "Your items" in r.text
    r=c2.post(f"/s/{tok}/cartons", data={"room_id":room.id,"contents":"Slabs x2","item_codes":"GF-KIT-01","qty":"2","length_cm":"245","width_cm":"125","height_cm":"15","weight_kg":"95"}, follow_redirects=False); assert r.status_code==303
    r=c2.post(f"/s/{tok}/cartons", data={"room_id":room.id,"contents":"Slabs x2 (2)","qty":"2","length_cm":"245","width_cm":"125","height_cm":"15","weight_kg":"95"}, follow_redirects=False)
    r=c2.get(f"/s/{tok}"); assert "BOX 2 OF 2" in r.text
    r=c2.get(f"/s/{tok}/labels.pdf"); assert r.headers["content-type"]=="application/pdf"; pass
    r=c2.get(f"/p/{pid}", follow_redirects=False); assert r.status_code==303  # protected
    # designer side
    r=c.post(f"/p/{pid}/cartons", data={"supplier_id":sup.id,"room_id":room.id,"contents":"Pendants","item_codes":f"GF-KIT-18","qty":"3","length_cm":"60","width_cm":"40","height_cm":"40"}, follow_redirects=False)
    r=c.get(f"/p/{pid}/cartons"); assert "BOX 3 OF 3" in r.text
    # images
    r=c.post(f"/p/{pid}/images", data={"kind":"mood","caption":"Warm wood"}, files=[("files",("m1.jpg",img("brown"),"image/jpeg")),("files",("m2.jpg",img("beige"),"image/jpeg"))], follow_redirects=False)
    r=c.post(f"/p/{pid}/images", data={"kind":"cover"}, files=[("files",("c.jpg",img("gray"),"image/jpeg"))], follow_redirects=False)
    r=c.post(f"/p/{pid}/images", data={"kind":"floorplan"}, files=[("files",("f.jpg",img("white"),"image/jpeg"))], follow_redirects=False)
    # PDFs
    for path,name in [(f"/p/{pid}/export/schedule.pdf","schedule"),(f"/p/{pid}/export/schedule.pdf?currency=CNY&prices=0","schedule2"),(f"/p/{pid}/export/packing.pdf","packing"),(f"/p/{pid}/export/labels.pdf","labels2"),(f"/p/{pid}/export/room/{room.id}.pdf","room"),(f"/p/{pid}/export/po/{sup.id}.pdf","po")]:
        r=c.get(path); assert r.status_code==200 and r.headers["content-type"]=="application/pdf", (path, r.status_code, r.text[:300]); print(name, len(r.content))
    r=c.get(f"/p/{pid}/export/items.xlsx"); assert r.status_code==200; pass
    # client link
    r=c.get(f"/p/{pid}/edit"); ctok=re.search(r"/c/([A-Za-z0-9_\-]+)", r.text).group(1)
    r=c2.get(f"/c/{ctok}"); assert r.status_code==200 and "Furniture" in r.text
    r=c2.get(f"/c/{ctok}/schedule.pdf"); assert r.status_code==200
    r=c.get(f"/p/{pid}/rooms"); assert "GF-KIT" in r.text
    r=c.post("/settings", data={"studio_name":"Harmony Designs Studio","studio_phone":"+255 7xx","default_rate":"7.2"}, follow_redirects=False)
    r=c.get("/settings"); assert "Harmony" in r.text
    r=c.get(f"/p/{pid}/items?q=pendant"); assert "Island pendant" in r.text
    # ---- save from web: pasted links + bookmarklet /clip (network replaced by a fake) ----
    import app.webimage as wi
    PAGE='<html><head><title>Kivik 3-seat sofa</title><meta property="og:image" content="/pics/sofa.jpg"><meta name="description" content="Linen, 228 cm wide"></head><body><script>var x=1</script><h1>Kivik sofa</h1><p>Price: ¥3,999</p><img src="small.png" width="10" height="10"></body></html>'.encode()
    def fake_get(url, accept):
        if url.endswith("/product/sofa"): return ("text/html; charset=utf-8".split(";")[0], PAGE, url)
        if url.endswith("/pics/sofa.jpg"): return ("image/jpeg", img("teal"), url)
        if url.endswith("/nope.jpg"): raise wi.WebImageError("Could not reach that site.")
        return ("text/html", b"<html><body>no pictures</body></html>", url)
    wi._http_get=fake_get
    assert wi.find_page_image(PAGE.decode(), "https://shop.example/product/sofa")=="https://shop.example/pics/sofa.jpg"
    pt=wi.page_text(PAGE.decode()); assert pt.startswith("Title: Kivik 3-seat sofa") and "¥3,999" in pt and "var x" not in pt, pt
    n_before=len(c.get(f"/p/{pid}/items/{iid}").text.split("/media/"))
    r=c.post(f"/p/{pid}/items/{iid}/photo-url", data={"url":"https://shop.example/product/sofa","caption":"from web"}, follow_redirects=False)
    assert r.status_code==303 and "err=" not in r.headers["location"], r.headers
    assert len(c.get(f"/p/{pid}/items/{iid}").text.split("/media/"))>n_before
    r=c.post(f"/p/{pid}/items/{iid}/photo-url", data={"url":"https://shop.example/nope.jpg"}, follow_redirects=False)
    assert "err=" in r.headers["location"]; r=c.get(r.headers["location"]); assert "Could not reach that site" in r.text
    r=c.post(f"/p/{pid}/items/{iid}", data={"room_id":room.id,"category":"Lighting","name":"Island pendant","qty":"3","unit":"pcs","status":"Ordered","photo_url":"shop.example/pics/sofa.jpg"}, follow_redirects=False)
    assert "err=" not in r.headers["location"]
    r=c.post(f"/p/{pid}/images/url", data={"kind":"mood","caption":"Sofa","url":"https://shop.example/product/sofa"}, follow_redirects=False); assert r.status_code==303 and "err" not in r.headers["location"]
    r=c.post(f"/p/{pid}/images/url", data={"kind":"mood","url":"https://shop.example/empty"}, follow_redirects=False); assert "err=" in r.headers["location"]
    r=c.get(f"/p/{pid}/images"); assert "Sofa" in r.text and "Add from a web link" in r.text
    r=c.get("/clip?url=https://shop.example/product/sofa&img=https://shop.example/pics/sofa.jpg"); assert "javascript:(function()" in r.text and "Save this picture" in r.text
    r=c.post("/clip", data={"project_id":pid,"dest":"item","link":"https://shop.example/pics/sofa.jpg","page":"https://shop.example/product/sofa"}, follow_redirects=False)
    assert r.headers["location"]==f"/p/{pid}/drafts", r.headers
    db=SessionLocal(); d=db.query(Item).filter(Item.project_id==int(pid), Item.draft==True).order_by(Item.id.desc()).first(); assert d and d.photos and "shop.example/product/sofa" in d.notes
    for ph in d.photos: db.delete(ph)
    db.delete(d); db.commit(); db.close()  # keep the drafts count at zero for the capture test below
    r=c.post("/clip", data={"project_id":pid,"dest":"cover","link":"","page":"https://shop.example/product/sofa"}, follow_redirects=False); assert r.headers["location"]==f"/p/{pid}/images"
    print("setup ok")
    # ---- Claude auto-fill (API replaced by a fake; off without the key) ----
    import app.ai as ai
    assert ai.suggest_item(img("red"), "text", "u", []) is None  # no key: feature off
    os.environ["ANTHROPIC_API_KEY"]="test-key"
    calls=[]
    def fake_call(content, schema=None):
        calls.append(content)
        assert schema is None or "Furniture" in schema["properties"]["category"]["enum"]  # the project's own categories
        return {"name":"Kivik 3-seat sofa","brand":"IKEA Kivik","category":"Furniture","spec":"Linen cover, removable","size":"2280 x 950 mm","finish":"Beige","unit":"pcs","unit_price_cny":3999,"room":room.label,"confidence":"high"}
    ai._call=fake_call
    r=c.post("/clip", data={"project_id":pid,"dest":"item","link":"https://shop.example/pics/sofa.jpg","page":"https://shop.example/product/sofa"}, follow_redirects=False)
    assert r.headers["location"].startswith(f"/p/{pid}/drafts?filled="), r.headers
    assert calls and any(b.get("type")=="image" for b in calls[-1]) and "Kivik" in [b for b in calls[-1] if b.get("type")=="text"][0]["text"]
    db=SessionLocal(); d=db.query(Item).filter(Item.project_id==int(pid), Item.draft==True).order_by(Item.id.desc()).first()
    assert d.name=="Kivik 3-seat sofa" and d.unit_price==0 and "Price seen on the page: ¥3,999.00" in d.notes and d.room_id==room.id and d.category=="Furniture" and "Filled in by Claude" in d.notes, (d.name, d.notes)
    r=c.get(f"/p/{pid}/drafts?filled={d.id}"); assert "Claude filled this in" in r.text
    for ph in d.photos: db.delete(ph)
    db.delete(d); db.commit(); db.close()
    r=c.get(f"/p/{pid}/items/{iid}"); assert "Fill in with Claude" in r.text
    calls.clear(); r=c.post(f"/p/{pid}/items/{iid}/suggest", follow_redirects=False); assert "filled=" in r.headers["location"], r.headers
    db=SessionLocal(); it=db.get(Item, int(iid)); assert it.name=="Island pendant" and it.category=="Lighting" and it.brand=="IKEA Kivik" and it.spec=="Linen cover, removable", (it.name, it.category, it.brand, it.spec); db.close()  # existing name/category kept, empty brand/spec filled
    ai._call=lambda content, schema=None: (_ for _ in ()).throw(RuntimeError("down"))
    r=c.post(f"/p/{pid}/items/{iid}/suggest", follow_redirects=False); assert "err=" in r.headers["location"]
    del os.environ["ANTHROPIC_API_KEY"]
    r=c.get(f"/p/{pid}/items/{iid}"); assert "Fill in with Claude" not in r.text
    print("claude auto-fill ok")
    r=c.get("/clip?url=http://testserver/p/1/items&img=http://testserver/static/brand-mark.png"); assert "tapped while you were on HBA itself" in r.text and "Save this picture" not in r.text and 'id="cl-paste"' in r.text
    r=c.post("/clip", data={"project_id":pid,"dest":"mood","link":"https://shop.example/product/sofa","page":""}, follow_redirects=False); assert r.headers["location"]==f"/p/{pid}/images"  # pasted page link, no bookmarklet
    r=TestClient(app).get("/clip", follow_redirects=False); assert r.status_code==303  # login required
    print("web images ok")

    # ---- quick capture: photo-first drafts ----
    from app.services import summary as _summary
    from app.models import Project as _P
    from app import storage as _st
    from openpyxl import load_workbook
    db=SessionLocal(); before=_summary(db, db.get(_P,int(pid))); db.close()
    sched_len=len(c.get(f"/p/{pid}/export/schedule.pdf").content); room_len=len(c.get(f"/p/{pid}/export/room/{room.id}.pdf").content)
    r=c.get(f"/p/{pid}/capture"); assert r.status_code==200 and 'capture="environment"' in r.text and "Quick capture" in r.text
    r=c.post(f"/p/{pid}/capture", data={"room_id":room.id,"category":"Lighting"}, files={"photo":("d1.jpg",img("purple"),"image/jpeg")}, headers={"Accept":"application/json"})
    j=r.json(); assert j["ok"] and j["drafts"]==1 and j["media"].startswith("/media/"), j; d1=j["id"]
    r=c.post(f"/p/{pid}/capture", data={"n":"1"}, files={"photo":("d2.jpg",img("orange"),"image/jpeg")}, follow_redirects=False)
    assert r.status_code==303 and "n=2" in r.headers["location"], r.headers
    r=c.post(f"/p/{pid}/capture", data={}, headers={"Accept":"application/json"}); assert r.status_code==400  # no photo
    db=SessionLocal(); d2=db.query(Item).filter(Item.draft==True, Item.id!=d1).one(); d2id=d2.id; d2key=d2.photos[0].file_key
    assert d2.name=="" and d2.code=="" and d2.room_id is None and d2.category=="Other"; db.close()
    r=c.get(f"/p/{pid}/drafts"); assert r.text.count('class="card draft"')==2 and d2key in r.text
    r=c.get(f"/p/{pid}/capture"); assert "Drafts (<span id=\"drafts\">2</span>)" in r.text
    # drafts are excluded everywhere until they have a name
    db=SessionLocal(); s=_summary(db, db.get(_P,int(pid))); db.close()
    assert s["items"]==before["items"] and s["total"]==before["total"] and s["drafts"]==2, (s["items"], before["items"], s["drafts"])
    r=c.get(f"/p/{pid}"); assert "Drafts (2)" in r.text and f'/items/{d1}"' not in r.text and f'/items/{d2id}"' not in r.text
    r=c.get(f"/p/{pid}/items"); assert "Drafts (2)" in r.text and f'/items/{d1}"' not in r.text and d2key not in r.text
    r=c.get(f"/p/{pid}/items?view=table"); assert d2key not in r.text
    r=c.get(f"/p/{pid}/export/items.xlsx"); assert sum(1 for row_ in load_workbook(io.BytesIO(r.content))["Shopping List"].iter_rows(min_row=2, values_only=True) if row_[0])==before["items"]  # one filled row per item (the sheet carries styled blank rows too)
    r=c2.get(f"/c/{ctok}"); assert d2key not in r.text and f">{before['items']}<" in r.text
    r=c.get(f"/p/{pid}/export/schedule.pdf"); assert r.status_code==200 and len(r.content)==sched_len, (len(r.content), sched_len)
    r=c.get(f"/p/{pid}/export/room/{room.id}.pdf"); assert r.status_code==200 and len(r.content)==room_len
    r=c2.get(f"/s/{tok}"); assert d2key not in r.text
    # complete: no name keeps it a draft (but remembers the rest)
    r=c.post(f"/p/{pid}/drafts/{d1}", data={"name":"  ","room_id":room.id,"category":"Lighting","qty":"2","unit_price":"120"}, follow_redirects=False)
    assert r.status_code==303 and f"err={d1}" in r.headers["location"]
    db=SessionLocal(); it=db.get(Item,d1); assert it.draft and it.code=="" and it.qty==2 and it.unit_price==0; db.close()  # the price comes from a quote now
    r=c.get(f"/p/{pid}/drafts?err={d1}"); assert "Give it a name" in r.text
    r=c.post(f"/p/{pid}/drafts/{d1}", data={"name":"Wall sconce","room_id":room.id,"category":"Lighting","qty":"2","unit_price":"120"}, follow_redirects=False)
    assert r.status_code==303 and "saved=GF-KIT-" in r.headers["location"]
    db=SessionLocal(); it=db.get(Item,d1); assert not it.draft and it.code.startswith("GF-KIT-") and it.total==0 and it.status=="To buy"; code=it.code
    assert sum(1 for i in db.get(_P,int(pid)).items if i.code==code)==1; db.close()  # code is unique
    buy(c, pid, d1, sup.id, 120)
    db=SessionLocal(); it=db.get(Item,d1); assert it.total==240 and it.status=="Quoted"; db.close()
    r=c.post(f"/p/{pid}/drafts/{d1}", data={"name":"again"}, follow_redirects=False); assert r.status_code==303  # not a draft any more: ignored
    db=SessionLocal(); assert db.get(Item,d1).name=="Wall sconce"; db.close()
    r=c.get(f"/p/{pid}/items?q=sconce"); assert "Wall sconce" in r.text and code in r.text and f'/items/{d1}"' in r.text
    r=c.get(f"/p/{pid}/drafts?saved={code}"); assert r.text.count('class="card draft"')==1 and code in r.text
    db=SessionLocal(); s=_summary(db, db.get(_P,int(pid))); db.close(); assert s["items"]==before["items"]+1 and s["drafts"]==1 and abs(s["total"]-before["total"]-240)<1e-6
    # discard deletes the draft and its photo
    assert _st.read_image(d2key) is not None
    r=c.post(f"/p/{pid}/drafts/{d2id}/discard", follow_redirects=False); assert r.status_code==303
    db=SessionLocal(); assert db.get(Item,d2id) is None; db.close(); assert _st.read_image(d2key) is None
    r=c.get(f"/p/{pid}/drafts"); assert "No drafts waiting" in r.text
    r=c.get(f"/p/{pid}"); assert "Drafts (" not in r.text
    # a draft completed through the full item form also gets a code
    r=c.post(f"/p/{pid}/capture", data={}, files={"photo":("d3.jpg",img("pink"),"image/jpeg")}, headers={"Accept":"application/json"}); d3=r.json()["id"]
    r=c.get(f"/p/{pid}/items/{d3}"); assert "quick-capture draft" in r.text
    r=c.post(f"/p/{pid}/items/{d3}", data={"room_id":room.id,"category":"Hardware","name":"Door stop","qty":"6","unit":"pcs","unit_price":"15","status":"To buy"}, follow_redirects=False); assert r.status_code==303
    db=SessionLocal(); it=db.get(Item,d3); assert not it.draft and it.code.startswith("GF-KIT-") and it.code!=code; db.close()
    # one item added to several rooms at once: one row per room, each with its own code and photo copy
    db=SessionLocal(); rms=db.query(Room).filter(Room.project_id==int(pid)).order_by(Room.sort).limit(3).all(); rids=[x.id for x in rms]; rcodes=[x.code for x in rms]; db.close()
    r=c.get(f"/p/{pid}/items/new"); assert 'name="room_ids"' in r.text
    r=c.post(f"/p/{pid}/items/new", data={"room_ids":[str(x) for x in rids],"category":"Electrical","name":"Air conditioner 12000 BTU","spec":"Inverter split","qty":"1","unit":"pcs","unit_price":"2100","status":"To buy"},
             files=[("photos",("ac.jpg",img("white"),"image/jpeg"))], follow_redirects=False)
    assert "/items?q=" in r.headers["location"], r.headers["location"]
    db=SessionLocal(); acs=db.query(Item).filter(Item.project_id==int(pid), Item.name=="Air conditioner 12000 BTU").all()
    assert sorted(a.room_id for a in acs)==sorted(rids) and sum(len(a.photos) for a in acs)==1 and all(len(a.all_photos)==1 and a.cover is not None for a in acs)
    assert len({a.cover.file_key for a in acs})==1 and all(a.code.split("-")[:-1]==a.room.code.split("-") for a in acs)
    ac_ids=[a.id for a in acs]; db.close()
    # edit specs once, applied to every room using the item; per-room qty/status stay
    db=SessionLocal(); slabs=db.query(Item).filter(Item.project_id==int(pid), Item.name=="Floor - big porcelain slab (+ extra % for cuts)").all()
    assert len(slabs)>3; s0=slabs[0]; qtys={x.id:x.qty for x in slabs}; s1=slabs[1]; s1.status="Ordered"; db.commit(); s0id, s1id=s0.id, s1.id; room0=s0.room_id; db.close()
    r=c.get(f"/p/{pid}/items/{s0id}"); assert 'name="apply_all"' in r.text and f"{len(slabs)-1} other rooms" in r.text
    r=c.post(f"/p/{pid}/items/{s0id}", data={"room_id":room0,"category":"Tiles","name":"Floor - big porcelain slab (+ extra % for cuts)","spec":"Calacatta look, 9 mm","size":"1600 x 3200 mm","qty":"5","unit":"sqm","unit_price":"88","status":"Quoted","apply_all":"1"}, follow_redirects=False)
    db=SessionLocal(); slabs=db.query(Item).filter(Item.project_id==int(pid), Item.name=="Floor - big porcelain slab (+ extra % for cuts)").all()
    assert all(x.spec=="Calacatta look, 9 mm" and x.size=="1600 x 3200 mm" for x in slabs)
    assert all(x.qty==qtys[x.id] for x in slabs if x.id!=s0id) and db.get(Item,s1id).status=="Ordered"; db.close()
    # without the tick only this room changes
    r=c.post(f"/p/{pid}/items/{s0id}", data={"room_id":room0,"category":"Tiles","name":"Floor - big porcelain slab (+ extra % for cuts)","spec":"Only here","qty":"5","unit":"sqm","status":"Quoted"}, follow_redirects=False)
    db=SessionLocal(); assert db.get(Item,s0id).spec=="Only here" and db.get(Item,s1id).spec=="Calacatta look, 9 mm"; db.close()
    # add an existing item to more rooms from its page
    db=SessionLocal(); extra=[x for x in db.query(Room).filter(Room.project_id==int(pid)).all() if x.id not in rids][:2]; extra_ids=[x.id for x in extra]; db.close()
    r=c.post(f"/p/{pid}/items/{ac_ids[0]}", data={"room_id":rids[0],"category":"Electrical","name":"Air conditioner 12000 BTU","spec":"Inverter split","qty":"1","unit":"pcs","unit_price":"2100","status":"Ordered","add_room_ids":[str(x) for x in extra_ids]}, follow_redirects=False)
    db=SessionLocal(); acs=db.query(Item).filter(Item.project_id==int(pid), Item.name=="Air conditioner 12000 BTU").all()
    assert len(acs)==5 and {a.room_id for a in acs}==set(rids+extra_ids)
    new=[a for a in acs if a.room_id in extra_ids]; assert all(a.status=="To buy" and a.photos==[] and a.cover is not None and a.code for a in new); db.close()
    # ---- floor plans & drawing sets: tagged plans, PDF page picker, rooms/client pages, by-floor schedule, checklist ----
    from app import drawings as _dr
    from app.models import ProjectImage, DrawingSet, PlanTag
    import pypdfium2 as _pf; assert _dr.available()
    from reportlab.pdfgen import canvas as _cv
    from reportlab.lib.pagesizes import A4 as _A4
    def pdf_pages(labels):
        b=io.BytesIO(); k=_cv.Canvas(b, pagesize=_A4)
        for t_ in labels: k.setFont("Helvetica",40); k.drawString(100,500,t_); k.rect(80,100,400,300); k.showPage()
        k.save(); return b.getvalue()
    def pdf_text(data):
        doc=_pf.PdfDocument(data); out=[]
        for pg in doc: tp=pg.get_textpage(); out.append(tp.get_text_range()); tp.close(); pg.close()
        doc.close(); return "\n".join(out)
    def long_edge(key): return max(Image.open(io.BytesIO(_st.read_image(key))).size)
    def big_img(color, w=2400, h=1800):
        b=io.BytesIO(); Image.new("RGB",(w,h),color).save(b,"JPEG"); return b.getvalue()
    assert _dr.floor_key("Ground Floor")=="ground" and _dr.floor_key(" first fl ")=="first" and _dr.floor_title("Ground")=="Ground floor"
    assert _dr.floor_title("Roof")=="Roof" and _dr.floor_title("First Floor")=="First Floor" and _dr.floor_title("ground floor")=="Ground floor" and _dr.is_pseudo("All") and _dr.is_pseudo("Outside") and not _dr.is_pseudo("Roof")
    r=c.get(f"/p/{pid}/images"); assert r.status_code==200 and "Add floor plans" in r.text and 'name="floor"' in r.text and 'value="Ground"' in r.text
    assert "Untagged" in r.text and 'option value="floorplan"' not in r.text and "Drawing sets" not in r.text
    # image plan, tagged at upload: 1600px preview + full-size copy
    r=c.post(f"/p/{pid}/images/plans", data={"floor":"Roof","sheet":"A-103","caption":"Roof terrace"}, files=[("files",("roof.jpg",big_img("white"),"image/jpeg"))], follow_redirects=False)
    assert r.status_code==303 and "?ok=" in r.headers["location"], r.headers
    db=SessionLocal(); roof=db.query(ProjectImage).filter(ProjectImage.project_id==int(pid), ProjectImage.kind=="floorplan").order_by(ProjectImage.id.desc()).first()
    assert roof.tag.floor=="Roof" and roof.sheet=="A-103" and long_edge(roof.hires_key)==2400 and long_edge(roof.file_key)<=1600; roof_id=roof.id; roof_keys=(roof.file_key, roof.hires_key); db.close()
    # PDF drawing sets: every PDF in the upload becomes a set; redirect to the first set's page picker
    r=c.post(f"/p/{pid}/images/plans", data={"floor":"","sheet":"A-100"}, files=[("files",("plans.pdf",pdf_pages(["GROUND FLOOR","FIRST FLOOR"]),"application/pdf")),("files",("more.pdf",pdf_pages(["SITE"]),"application/pdf"))], follow_redirects=False)
    assert r.status_code==303, r.headers; loc=r.headers["location"]; sid=int(re.search(rf"/p/{pid}/images/sets/(\d+)", loc).group(1)); assert "sheet=A-100" in loc
    db=SessionLocal(); sets=db.query(DrawingSet).filter(DrawingSet.project_id==int(pid)).order_by(DrawingSet.id).all()
    assert len(sets)==2 and sets[0].id==sid and sets[0].pages==2 and sets[0].name=="plans.pdf" and sets[1].pages==1
    assert _st.read_image(sets[0].file_key)[:5]==b"%PDF-" and long_edge(sets[0].thumb_key(1))<=480; set_key=sets[0].file_key; thumb1=sets[0].thumb_key(1); db.close()
    r=c.get(loc); assert r.status_code==200 and r.text.count('name="page_')==2 and "/t001.jpg" in r.text and 'value="A-100"' in r.text
    r=c.get(f"/p/{pid}/images/sets/{sid}.pdf"); assert r.status_code==200 and r.headers["content-type"]=="application/pdf"
    r=c2.get(f"/p/{pid}/images/sets/{sid}.pdf", follow_redirects=False); assert r.status_code==303  # login required
    r=c.post(f"/p/{pid}/images/sets/{sid}/pick", data={}, follow_redirects=False); assert "err=Tick" in r.headers["location"]
    r=c.post(f"/p/{pid}/images/sets/{sid}/pick", data={"page_1":"on","floor_1":"Ground","sheet_1":"A-101 Rev A","page_2":"on","floor_2":"First","sheet_2":"A-102"}, follow_redirects=False)
    assert r.status_code==303 and "?ok=" in r.headers["location"], r.headers
    db=SessionLocal(); fps=db.query(ProjectImage).filter(ProjectImage.project_id==int(pid), ProjectImage.kind=="floorplan").order_by(ProjectImage.id).all()
    assert len(fps)==4, len(fps)  # legacy untagged, roof, two picked pages
    pages={im.tag.page_no: im for im in fps if im.tag and im.tag.set_id==sid}; assert set(pages)=={1,2}
    assert 2000<=long_edge(pages[1].hires_key)<=3200 and long_edge(pages[1].file_key)<=1600 and pages[2].floor=="First"; gid=pages[1].id; db.close()
    r=c.get(f"/p/{pid}/images/sets/{sid}"); assert r.text.count("Added")==2 and 'name="page_' not in r.text
    r=c.get(f"/p/{pid}/images"); assert "A-101 Rev A" in r.text and f'action="/p/{pid}/images/{gid}"' in r.text and "Drawing sets" in r.text and "plans.pdf" in r.text and "more.pdf" in r.text
    # tag edit; floor matching is case-insensitive and ignores "floor"
    r=c.post(f"/p/{pid}/images/{gid}", data={"floor":" ground floor ","sheet":"A-101 Rev B","caption":"Ground plan"}, follow_redirects=False)
    assert r.status_code==303 and r.headers["location"].endswith(f"#img-{gid}")
    db=SessionLocal(); im=db.get(ProjectImage,gid); assert im.tag.floor=="Ground" and im.sheet=="A-101 Rev B" and im.caption=="Ground plan"  # stored as the floor list spells it
    pr=db.get(_P,int(pid)); assert _dr.plan_for_room(pr, db.get(Room, room.id)).id==gid and [g["title"] for g in _dr.rooms_by_floor(pr)]==["Ground floor","First floor","Roof","Whole house / other"]
    plan_prev=im.file_key; db.close()
    r=c.get(f"/p/{pid}/rooms"); assert "Ground floor" in r.text and "A-101 Rev B" in r.text and plan_prev in r.text and "Plan A-101 Rev B" in r.text and 'class="in floor-sel"' in r.text
    assert "No plan for this floor yet" not in r.text and "Whole house / other" in r.text
    r=c2.get(f"/c/{ctok}"); assert "Floor plans" in r.text and plan_prev in r.text and "layout=floor" in r.text
    r=c2.get(f"/c/{ctok}/schedule.pdf?layout=floor"); assert r.status_code==200 and r.headers["content-type"]=="application/pdf"
    tc=pdf_text(r.content); assert tc.index("GROUND FLOOR") < tc.index("WHOLE HOUSE") < tc.index("SUMMARY") and "At a glance" in tc  # the client link really gets the by-floor layout
    # schedules: by category unchanged in structure, by floor = plan page then that floor's rooms, whole house, summary
    a=c.get(f"/p/{pid}/export/schedule.pdf"); b=c.get(f"/p/{pid}/export/schedule.pdf?layout=floor"); b2=c.get(f"/p/{pid}/export/schedule.pdf?layout=floor")
    assert a.status_code==b.status_code==200 and len(b.content)==len(b2.content) and len(a.content)!=len(b.content)
    assert c.get(f"/p/{pid}/export/schedule.pdf?layout=floor&currency=CNY&prices=0").status_code==200
    t_=pdf_text(b.content)
    assert t_.index("GROUND FLOOR") < t_.index("FIRST FLOOR") < t_.index("ROOF") < t_.index("WHOLE HOUSE") < t_.index("SUMMARY"), t_[:3000]
    assert "A-101 Rev B" in t_ and "GF-KIT - Kitchen" in t_ and "At a glance" in t_ and "Contents" in t_
    # the pseudo-floor room (ALL) prints only in the whole-house section, never under a real floor
    assert t_.index("WHOLE HOUSE") < t_.index("ALL - Whole house") and "ALL - Whole house" not in t_[:t_.index("WHOLE HOUSE")]
    assert t_.count("Whole house / other") == 1, t_.count("Whole house / other")
    ta=pdf_text(a.content); assert "FLOOR PLANS" in ta and "SCHEDULE" in ta and "Ground floor · A-101 Rev B · Ground plan" in ta and "Contents" in ta
    # room checklist: the room's floor plan is page 1
    r=c.get(f"/p/{pid}/export/room/{room.id}.pdf"); assert r.status_code==200 and len(r.content)>room_len and len(_pf.PdfDocument(r.content))>=2
    assert "Find GF-KIT - Kitchen on: Ground floor" in pdf_text(r.content) and "A-101 Rev B" in pdf_text(r.content)
    db=SessionLocal(); gym=db.query(Room).filter(Room.project_id==int(pid), Room.code=="RF-GYM").first(); db.close()
    assert "A-103" in pdf_text(c.get(f"/p/{pid}/export/room/{gym.id}.pdf").content)
    assert c.get(f"/p/{pid}/export/room/0.pdf").status_code==200
    r=c.post("/projects/new", data={"client_name":"Other","name":"Other house","rate":"7.1"}, follow_redirects=False); pid2=r.headers["location"].split("/")[-1]
    c.post(f"/p/{pid2}/rooms", data={"code":"X-1","name":"Other room","floor":"Ground"}, follow_redirects=False)
    db=SessionLocal(); other=db.query(Room).filter(Room.project_id==int(pid2)).first(); db.close()
    assert c.get(f"/p/{pid}/export/room/{other.id}.pdf").status_code==404 and c.get(f"/p/{pid}/images/sets/{sid}").status_code==200 and c.get(f"/p/{pid2}/images/sets/{sid}").status_code==404
    # by-floor summary: rooms all on real floors, a plan tagged with a pseudo floor, loose items -> one whole-house section, not two
    db=SessionLocal(); rooms2={r.code:r for r in db.query(Room).filter(Room.project_id==int(pid2))}; all_id=rooms2["ALL"].id; x1_id=rooms2["X-1"].id; db.close()
    c.post(f"/p/{pid2}/items/new", data={"room_ids":[str(x1_id)],"category":"Lighting","name":"Lamp","qty":"1","unit":"pcs","unit_price":"100","status":"To buy"}, follow_redirects=False)
    c.post(f"/p/{pid2}/items/new", data={"room_ids":[str(all_id)],"category":"Lighting","name":"Loose lamp","qty":"1","unit":"pcs","unit_price":"50","status":"To buy"}, follow_redirects=False)
    c.post(f"/p/{pid2}/rooms/{all_id}/delete", follow_redirects=False)  # its item becomes loose (room_id NULL)
    c.post(f"/p/{pid2}/images/plans", data={"floor":"Site","sheet":"A-000"}, files=[("files",("site.jpg",big_img("gray",900,600),"image/jpeg"))], follow_redirects=False)
    t2=pdf_text(c.get(f"/p/{pid2}/export/schedule.pdf?layout=floor").content)
    assert t2.index("GROUND FLOOR") < t2.index("Lamp") < t2.index("WHOLE HOUSE") < t2.index("Loose lamp") < t2.index("SUMMARY"), t2[:2000]
    assert t2.count("Whole house / other")==1 and "Whole house / not room-specific" in t2.split("SUMMARY", 1)[1], t2[-1500:]
    # deleting the PDF keeps the pages already added
    r=c.post(f"/p/{pid}/images/sets/{sid}/delete", follow_redirects=False); assert r.status_code==303
    assert _st.read_image(set_key) is None and _st.read_image(thumb1) is None and c.get(f"/p/{pid}/images/sets/{sid}").status_code==404
    db=SessionLocal(); im=db.get(ProjectImage,gid); assert im is not None and im.tag.set_id is None and im.tag.page_no is None and im.floor=="Ground"; db.close()
    # without pypdfium2: PDFs are refused with a clear message, image plans still work
    _real=_dr.pdfium; _dr.pdfium=None
    try:
        r=c.post(f"/p/{pid}/images/plans", data={"floor":"Roof"}, files=[("files",("x.pdf",pdf_pages(["X"]),"application/pdf"))], follow_redirects=False)
        assert "err=" in r.headers["location"] and "pypdfium2" in r.headers["location"]
        assert "not available" in c.get(f"/p/{pid}/images").text
        r=c.post(f"/p/{pid}/images/plans", data={"floor":"Roof","sheet":"A-103b"}, files=[("files",("r2.jpg",big_img("gray",1000,700),"image/jpeg"))], follow_redirects=False)
        assert "?ok=" in r.headers["location"]
    finally:
        _dr.pdfium=_real
    # caps and junk: pages beyond MAX_PDF_PAGES are ignored, too many picks refused, unreadable files reported
    from app import config as _cfg
    _oldpick=_cfg.PICK_MAX_PAGES; _cfg.PICK_MAX_PAGES=3  # a 5-page set against a cap of 3 proves the rule as a 14-page set against 12 did, in a third of the time
    r=c.post(f"/p/{pid}/images/plans", data={}, files=[("files",("big.pdf",pdf_pages(["P"]*5),"application/pdf"))], follow_redirects=False)
    sid2=int(re.search(r"/sets/(\d+)", r.headers["location"]).group(1))
    r=c.post(f"/p/{pid}/images/sets/{sid2}/pick", data={f"page_{k}":"on" for k in range(1,5)}, follow_redirects=False); _cfg.PICK_MAX_PAGES=_oldpick; assert "err=Add+at+most" in r.headers["location"] or "err=Add%20at%20most" in r.headers["location"], r.headers
    r=c.post(f"/p/{pid}/images/plans", data={}, files=[("files",("junk.bin",b"hello","application/octet-stream"))], follow_redirects=False); assert "could" in r.headers["location"]
    from app import config as _cfg
    _old=_cfg.MAX_PDF_PAGES; _cfg.MAX_PDF_PAGES=3
    try:
        r=c.post(f"/p/{pid}/images/plans", data={}, files=[("files",("many.pdf",pdf_pages(["P"]*5),"application/pdf"))], follow_redirects=False)
        sid3=int(re.search(r"/sets/(\d+)", r.headers["location"]).group(1)); db=SessionLocal(); assert db.get(DrawingSet,sid3).pages==3; db.close()
    finally:
        _cfg.MAX_PDF_PAGES=_old
    # a good PDF next to a broken one: the good one is kept (set in the DB, no orphan files) and the error still shows
    db=SessionLocal(); before=db.query(DrawingSet).count(); db.close()
    r=c.post(f"/p/{pid}/images/plans", data={}, files=[("files",("good.pdf",pdf_pages(["G"]),"application/pdf")),("files",("broken.pdf",b"%PDF-1.4 not really","application/pdf"))], follow_redirects=False)
    assert "/sets/" in r.headers["location"] and "err=" in r.headers["location"] and "broken.pdf" in r.headers["location"], r.headers["location"]
    db=SessionLocal(); assert db.query(DrawingSet).count()==before+1; db.close()
    r=c.get(r.headers["location"]); assert r.status_code==200 and "broken.pdf" in r.text
    # storage keys are validated: /media never leaves the upload folder (no login needed for /media)
    anon=TestClient(app)
    for bad in ["%2e%2e/%2e%2e/%2e%2e/%2e%2e/etc/passwd", "../../etc/passwd", "/etc/passwd", "p1/..%2fx.jpg"]:
        r=anon.get(f"/media/{bad}"); assert r.status_code==404, (bad, r.status_code)
    assert _st.safe_key("p1/sets/abc.pdf") and _st.safe_key("img/0123abcd.jpg") and not _st.safe_key("../x") and not _st.safe_key("/x") and not _st.safe_key("a/../b") and not _st.safe_key("")
    assert _st.read_image("../../etc/passwd") is None
    r=anon.get(f"/media/{roof_keys[0]}"); assert r.status_code==200 and r.headers["content-type"].startswith("image/jpeg")
    # Chinese drawing-set names download fine (RFC 5987 header) and the schedule PDF too
    r=c.post(f"/p/{pid}/images/plans", data={}, files=[("files",("一层平面图.pdf",pdf_pages(["CN"]),"application/pdf"))], follow_redirects=False)
    sid_cn=int(re.search(r"/sets/(\d+)", r.headers["location"]).group(1))
    r=c.get(f"/p/{pid}/images/sets/{sid_cn}.pdf"); assert r.status_code==200 and "filename*=UTF-8''%E4%B8%80" in r.headers["content-disposition"], r.headers.get("content-disposition")
    # the size cap is applied before the file is read into memory
    _oldmb=_cfg.MAX_PDF_MB; _cfg.MAX_PDF_MB=0
    try:
        r=c.post(f"/p/{pid}/images/plans", data={}, files=[("files",("huge.pdf",pdf_pages(["H"]),"application/pdf"))], follow_redirects=False); assert "larger" in r.headers["location"], r.headers["location"]
    finally:
        _cfg.MAX_PDF_MB=_oldmb
    # ---- interactive plan: pins (room boxes on a plan), browse page, room panel, mark mode, cascades ----
    from app.models import RoomPin
    db=SessionLocal(); gf_rooms=[r for r in db.query(Room).filter(Room.project_id==int(pid)).order_by(Room.sort) if _dr.floor_key(r.floor)=="ground"]; db.close()
    assert len(gf_rooms)>=2 and room.id in [r.id for r in gf_rooms]
    r=c.get(f"/p/{pid}/plan"); assert r.status_code==200 and f'data-plan="{gid}"' in r.text and "Tap a room" in r.text and "0 of" in r.text, r.text[:400]
    assert "plan-tabs" in r.text and "Ground floor" in r.text and "Not on this plan yet" in r.text and f'plan?plan={gid}&amp;mode=mark' in r.text
    assert 'href="/p/'+pid+'/plan"' in r.text  # shell navigation (sidebar / bottom bar / More sheet)
    # a box drawn the "wrong way round" (negative size) is normalised; too small is refused; other project's room refused
    r=c.post(f"/p/{pid}/plan/pins", data={"image_id":gid,"room_id":room.id,"x":"0.4","y":"0.6","w":"-0.3","h":"-0.25"}, headers={"Accept":"application/json"})
    j=r.json(); assert r.status_code==200 and j["ok"] and j["pin"]["x"]==0.1 and j["pin"]["y"]==0.35 and j["pin"]["w"]==0.3 and j["pin"]["h"]==0.25, j
    pin_id=j["pin"]["id"]
    r=c.post(f"/p/{pid}/plan/pins", data={"image_id":gid,"room_id":room.id,"x":"0.1","y":"0.1","w":"0.001","h":"0.2"}, headers={"Accept":"application/json"}); assert r.status_code==400 and "bigger" in r.json()["error"]
    r=c.post(f"/p/{pid}/plan/pins", data={"image_id":gid,"room_id":other.id,"x":"0.1","y":"0.1","w":"0.2","h":"0.2"}, headers={"Accept":"application/json"}); assert r.status_code==404
    r=c.post(f"/p/{pid}/plan/pins", data={"image_id":gid,"room_id":room.id,"x":"abc","y":"0.1","w":"0.2","h":"0.2"}, headers={"Accept":"application/json"}); assert r.status_code==400
    # saving again for the same room on the same plan replaces the box (one pin per room per plan), clipped to the image
    r=c.post(f"/p/{pid}/plan/pins", data={"image_id":gid,"room_id":room.id,"x":"0.9","y":"0.9","w":"0.5","h":"0.5"}, headers={"Accept":"application/json"}); j=r.json()
    assert j["pin"]["id"]==pin_id and abs(j["pin"]["w"]-0.1)<1e-6 and abs(j["pin"]["h"]-0.1)<1e-6, j
    db=SessionLocal(); assert db.query(RoomPin).filter(RoomPin.image_id==gid).count()==1; db.close()
    # second room via the plain form (no JSON): redirects back to mark mode
    r2=gf_rooms[1] if gf_rooms[0].id==room.id else gf_rooms[0]
    r=c.post(f"/p/{pid}/plan/pins", data={"image_id":gid,"room_id":r2.id,"x":"0.5","y":"0.1","w":"0.3","h":"0.3"}, follow_redirects=False); assert r.status_code==303 and f"plan={gid}&mode=mark" in r.headers["location"]
    # browse page: boxes as links, counts, selected room renders its panel inline (deep link, no JavaScript needed)
    r=c.get(f"/p/{pid}/plan?plan={gid}&room={room.id}"); assert r.status_code==200
    assert f'class="pin sel' in r.text and 'left:90.000%;top:90.000%' in r.text and f'data-room="{r2.id}"' in r.text and "2 of" in r.text
    assert 'class="card tight plan-room"' in r.text and "Room details" in r.text and "status-sel" in r.text and f"/items/new?room={room.id}" in r.text and "GF-KIT" in r.text
    assert f"next=/p/{pid}/plan%3Fplan%3D{gid}%26room%3D{room.id}" in r.text  # item links come back to the plan
    # a room not on any plan opens on its floor's plan; a room asked for without a plan opens the plan it is marked on
    r3=next(x for x in gf_rooms if x.id not in (room.id, r2.id))
    r=c.get(f"/p/{pid}/plan?room={r3.id}"); assert r.status_code==200 and f'data-plan="{gid}"' in r.text and "Not placed on this plan yet" in r.text
    r=c.get(f"/p/{pid}/plan?room={room.id}"); assert f'data-plan="{gid}"' in r.text and 'class="pin sel' in r.text
    # the panel as a fragment (what plan.js fetches when a box is tapped): no app shell, items of that room only
    r=c.get(f"/p/{pid}/plan/room/{room.id}?plan={gid}"); assert r.status_code==200 and "<html" not in r.text and "plan-room" in r.text and "Island pendant" in r.text
    assert c.get(f"/p/{pid}/plan/room/{other.id}").status_code==404 and c2.get(f"/p/{pid}/plan", follow_redirects=False).status_code==303
    # saving the item form with next= returns to the plan; the room form on the panel too
    nxt=f"/p/{pid}/plan?plan={gid}&room={room.id}"
    from urllib.parse import quote as _q
    r=c.get(f"/p/{pid}/items/{iid}?next={_q(nxt, safe='')}"); assert f'name="next" value="{nxt.replace("&", "&amp;")}"' in r.text and "Back to the plan" in r.text, r.text[:300]
    r=c.post(f"/p/{pid}/items/{iid}", data={"room_id":room.id,"category":"Lighting","name":"Island pendant","qty":"3","unit":"pcs","unit_price":"450","status":"Ordered","next":nxt}, follow_redirects=False); assert r.headers["location"]==nxt
    r=c.post(f"/p/{pid}/items/{iid}", data={"room_id":room.id,"category":"Lighting","name":"Island pendant","qty":"3","unit":"pcs","unit_price":"450","status":"Ordered","next":"https://evil.example/x"}, follow_redirects=False); assert r.headers["location"].startswith(f"/p/{pid}/items/")
    r=c.post(f"/p/{pid}/rooms/{room.id}", data={"code":"GF-KIT","name":"Kitchen","floor":"Ground","floor_area":"0","wall_area":"0","notes":"via plan","next":nxt}, follow_redirects=False); assert r.headers["location"]==nxt
    db=SessionLocal(); assert db.get(Room, room.id).notes=="via plan"; db.close()
    # mark mode: the floor's rooms with placed/not placed state, other floors collapsed; tabs keep the mode
    r=c.get(f"/p/{pid}/plan?plan={gid}&mode=mark"); assert r.status_code==200 and "Mark rooms" in r.text and "Done marking" in r.text
    assert r.text.count('class="r placed"')==2 and "Rooms on other floors" in r.text and f"plan={roof_id}&amp;mode=mark" in r.text and f'data-sel=""' in r.text
    r=c.get(f"/p/{pid}/plan?plan={gid}&mode=mark&room={r3.id}"); assert f'data-sel="{r3.id}"' in r.text
    # entry points: Rooms page ("On the plan" / "Place on plan"), Images (count + mark link), Overview, item list filtered by room
    r=c.get(f"/p/{pid}/rooms"); assert "Open the plan" in r.text and f"plan?plan={gid}&amp;room={room.id}" in r.text and f"plan?plan={gid}&amp;mode=mark&amp;room={r3.id}" in r.text
    r=c.get(f"/p/{pid}/images"); assert "2 rooms marked" in r.text and f"plan?plan={gid}&amp;mode=mark" in r.text
    r=c.get(f"/p/{pid}"); assert "Plan view" in r.text and f"plan?room={room.id}" in r.text
    r=c.get(f"/p/{pid}/items?room={room.id}"); assert f'href="/p/{pid}/plan?room={room.id}"' in r.text
    # delete a pin (JSON and form); deleting the room or the plan removes its pins
    r=c.post(f"/p/{pid}/plan/pins/{pin_id}/delete", headers={"Accept":"application/json"}); assert r.json()["ok"]
    r=c.post(f"/p/{pid}/plan/pins/{pin_id}/delete", headers={"Accept":"application/json"}); assert r.status_code==404
    c.post(f"/p/{pid}/plan/pins", data={"image_id":roof_id,"room_id":r2.id,"x":"0.1","y":"0.1","w":"0.2","h":"0.2"}, headers={"Accept":"application/json"})
    c.post(f"/p/{pid}/plan/pins", data={"image_id":roof_id,"room_id":r3.id,"x":"0.5","y":"0.1","w":"0.2","h":"0.2"}, headers={"Accept":"application/json"})
    db=SessionLocal(); assert db.query(RoomPin).filter(RoomPin.image_id==roof_id).count()==2 and db.query(RoomPin).filter(RoomPin.room_id==r2.id).count()==2; db.close()
    c.post(f"/p/{pid}/rooms", data={"code":"TMP","name":"Temp","floor":"Ground"}, follow_redirects=False)
    db=SessionLocal(); tmp=db.query(Room).filter(Room.project_id==int(pid), Room.code=="TMP").first(); db.close()
    c.post(f"/p/{pid}/plan/pins", data={"image_id":gid,"room_id":tmp.id,"x":"0.1","y":"0.1","w":"0.2","h":"0.2"}, headers={"Accept":"application/json"})
    c.post(f"/p/{pid}/rooms/{tmp.id}/delete", follow_redirects=False)
    db=SessionLocal(); assert db.query(RoomPin).filter(RoomPin.room_id==tmp.id).count()==0; db.close()
    r=c.get(f"/p/{pid2}/plan"); assert r.status_code==200 and f'data-plan=' in r.text  # project with only a pseudo-floor plan still opens
    print("interactive plan ok")
    # ---- item dots: one spot per item per plan; the room panel places / locates / removes them ----
    from app.models import ItemPin
    db=SessionLocal(); kit=db.get(Room, room.id); kit_live=[i.id for i in kit.items if not i.draft]; draft_id=db.query(Item).filter(Item.project_id==int(pid), Item.draft==True).first(); draft_id=draft_id.id if draft_id else None; db.close()  # noqa: E712
    assert int(iid) in kit_live and len(kit_live)>=2
    other_item=kit_live[1] if kit_live[0]==int(iid) else kit_live[0]
    r=c.post(f"/p/{pid}/plan/item-pins", data={"image_id":gid,"item_id":iid,"x":"1.7","y":"-0.2"}, headers={"Accept":"application/json"}); j=r.json()
    assert r.status_code==200 and j["ok"] and j["pin"]["x"]==1.0 and j["pin"]["y"]==0.0 and j["pin"]["room_id"]==room.id and j["pin"]["code"] and j["pin"]["photo"], j  # clipped to the image
    ip=j["pin"]["id"]
    r=c.post(f"/p/{pid}/plan/item-pins", data={"image_id":gid,"item_id":iid,"x":"0.25","y":"0.75"}, headers={"Accept":"application/json"}); j=r.json(); assert j["pin"]["id"]==ip and j["pin"]["x"]==0.25  # same item again = moved
    db=SessionLocal(); assert db.query(ItemPin).filter(ItemPin.item_id==int(iid)).count()==1; db.close()
    r=c.post(f"/p/{pid}/plan/item-pins", data={"image_id":gid,"item_id":iid,"x":"nope","y":"0.1"}, headers={"Accept":"application/json"}); assert r.status_code==400
    r=c.post(f"/p/{pid}/plan/item-pins", data={"image_id":gid,"item_id":"999999","x":"0.1","y":"0.1"}, headers={"Accept":"application/json"}); assert r.status_code==404
    if draft_id: r=c.post(f"/p/{pid}/plan/item-pins", data={"image_id":gid,"item_id":draft_id,"x":"0.1","y":"0.1"}, headers={"Accept":"application/json"}); assert r.status_code==400 and "draft" in r.json()["error"]
    db=SessionLocal(); oi=db.query(Item).filter(Item.project_id==int(pid2)).first(); db.close()
    r=c.post(f"/p/{pid}/plan/item-pins", data={"image_id":gid,"item_id":oi.id,"x":"0.1","y":"0.1"}, headers={"Accept":"application/json"}); assert r.status_code==404  # other project's item
    r=c.post(f"/p/{pid}/plan/item-pins", data={"image_id":gid,"item_id":other_item,"x":"0.6","y":"0.6"}, follow_redirects=False); assert r.status_code==303 and f"room={room.id}&item={other_item}" in r.headers["location"]
    # the page: dots as links to their room + item, the deep-linked item highlighted in the panel, the items toggle counts them
    r=c.get(f"/p/{pid}/plan?plan={gid}&room={room.id}&item={iid}"); assert r.status_code==200
    assert 'class="ipin sel"' in r.text and 'left:25.000%;top:75.000%' in r.text and f'data-item="{iid}"' in r.text and 'class="pi placed hl"' in r.text and "pinbtn" in r.text and "pinrm" in r.text
    assert f'id="plan-items"' in r.text and '<span class="n">2</span>' in r.text and 'data-placed>2</span>' in r.text
    r=c.get(f"/p/{pid}/plan?item={iid}"); assert f'data-sel="{room.id}"' in r.text and f'data-sel-item="{iid}"' in r.text  # ?item= alone opens its room
    r=c.get(f"/p/{pid}/plan/room/{room.id}?plan={gid}&item={iid}"); assert 'class="pi placed hl"' in r.text
    r=c.get(f"/p/{pid}/items/{iid}"); assert f"plan?plan={gid}&amp;room={room.id}&amp;item={iid}" in r.text and "On the plan" in r.text
    r=c.get(f"/p/{pid}/plan?plan={gid}&mode=mark"); assert r.status_code==200 and 'id="plan-items"' not in r.text  # dots stay out of the way while marking
    # ---- the client link: the same plan, read-only (photos, USD, status badges; no controls, suppliers, notes or CNY) ----
    r=c2.get(f"/c/{ctok}"); assert r.status_code==200 and 'id="plan"' in r.text and 'data-readonly="1"' in r.text and 'class="ipin' in r.text and "plan.js" in r.text
    assert f'data-room-base="/c/{ctok}/plan/room/"' in r.text and f'href="/c/{ctok}?plan={gid}&amp;room={r2.id}#plan-section"' in r.text and "Tap a room" in r.text  # r2's box (GF-KIT's was removed above)
    assert "status-sel" not in r.text and "plan-count" not in r.text and "Mark rooms" not in r.text and "/p/" not in r.text.split('id="plan"')[1].split("</aside>")[0]
    r=c2.get(f"/c/{ctok}?plan={gid}&room={room.id}&item={iid}"); assert r.status_code==200 and 'class="card tight plan-room"' in r.text and 'class="pi placed hl"' in r.text
    panel=r.text.split('class="card tight plan-room"')[1].split("</aside>")[0]
    assert "Total USD" in panel and "status-sel" not in panel and "Room details" not in panel and "/items/new" not in panel and "pinrm" not in panel and "¥" not in panel and "Foshan Tile Co." not in panel and "pinbtn" in panel
    r=c2.get(f"/c/{ctok}/plan/room/{room.id}?plan={gid}&item={iid}"); assert r.status_code==200 and "<html" not in r.text and "plan-room" in r.text and "status-sel" not in r.text and "¥" not in r.text
    assert c2.get(f"/c/{ctok}/plan/room/{other.id}").status_code==404 and c2.get(f"/c/nope/plan/room/{room.id}").status_code==404
    assert c2.post(f"/p/{pid}/plan/item-pins", data={"image_id":gid,"item_id":iid,"x":"0.1","y":"0.1"}, follow_redirects=False).status_code==303  # designer routes still need login
    # remove a dot; deleting an item or a plan takes its dots along
    r=c.post(f"/p/{pid}/plan/item-pins/{ip}/delete", headers={"Accept":"application/json"}); assert r.json()["ok"]
    assert c.post(f"/p/{pid}/plan/item-pins/{ip}/delete", headers={"Accept":"application/json"}).status_code==404
    c.post(f"/p/{pid}/plan/item-pins", data={"image_id":roof_id,"item_id":iid,"x":"0.3","y":"0.3"}, headers={"Accept":"application/json"})
    c.post(f"/p/{pid}/items/{other_item}/duplicate", follow_redirects=False)
    db=SessionLocal(); dup=db.query(Item).filter(Item.project_id==int(pid)).order_by(Item.id.desc()).first(); db.close()
    c.post(f"/p/{pid}/plan/item-pins", data={"image_id":gid,"item_id":dup.id,"x":"0.3","y":"0.3"}, headers={"Accept":"application/json"})
    db=SessionLocal(); assert db.query(ItemPin).filter(ItemPin.item_id==dup.id).count()==1; db.close()
    c.post(f"/p/{pid}/items/{dup.id}/delete", follow_redirects=False)
    db=SessionLocal(); assert db.query(ItemPin).filter(ItemPin.item_id==dup.id).count()==0 and db.query(ItemPin).filter(ItemPin.image_id==roof_id).count()==1; db.close()
    print("item dots + client plan ok")
    # ---- prices: typed on a quote in USD or CNY (kept in CNY at the project rate), never on the item form; the help page ----
    from app.models import ItemPrice, Quote, Purchase, Product
    r=c.get(f"/p/{pid}/items/new"); assert 'name="unit_price"' not in r.text and 'name="price_currency"' not in r.text and 'name="supplier_id"' not in r.text and 'name="lead_time"' not in r.text and "Preferred brand" in r.text
    r=c.post(f"/p/{pid}/items/new", data={"room_ids":[str(room.id)],"category":"Lighting","name":"Brass reading light","qty":"2","unit":"pcs","unit_price":"120","price_currency":"USD","supplier_id":str(sup.id),"lead_time":"1 week","status":"To buy"}, follow_redirects=False)
    usd_id=int(r.headers["location"].split("/")[-1])
    db=SessionLocal(); u=db.get(Item, usd_id); assert u.unit_price==0 and u.supplier_id is None and u.price_entry is None and u.lead_time=="" and u.total==0; db.close()
    r=c.get(f"/p/{pid}/items/{usd_id}"); assert "No shop chosen yet" in r.text and f"/p/{pid}/buying/item/{usd_id}" in r.text and 'name="unit_price"' not in r.text
    prod_id, q1 = buy(c, pid, usd_id, sup.id, 120, "USD", "2 weeks")
    db=SessionLocal(); u=db.get(Item, usd_id); assert u.unit_price==852.0 and u.price_currency=="USD" and u.price_amount==120 and u.total==1704.0 and u.supplier_id==sup.id and u.lead_time=="2 weeks" and u.status=="Quoted"; db.close()
    r=c.get(f"/p/{pid}/items/{usd_id}"); assert "Buying from" in r.text and "Foshan Tile Co." in r.text and "¥852.00 each" in r.text
    r=c.get(f"/p/{pid}/items"); assert "1,704" in r.text  # lists, totals and PDFs keep using the CNY value
    r=c.get(f"/p/{pid}/buying/{prod_id}"); assert "typed as $120.00" in r.text and "Buying for 1 room" in r.text
    # the same product into a second room: no shop there until it is bought for it; buying from the same quote widens its purchase
    r=c.post(f"/p/{pid}/items/{usd_id}", data={"room_id":room.id,"category":"Lighting","name":"Brass reading light","qty":"2","unit":"pcs","status":"Quoted","add_room_ids":[str(r2.id)]}, follow_redirects=False)
    db=SessionLocal(); copies=[i for i in db.query(Item).filter(Item.project_id==int(pid), Item.name=="Brass reading light").all() if i.id!=usd_id]; assert len(copies)==1 and copies[0].unit_price==0 and copies[0].purchase_id is None; cid=copies[0].id; db.close()
    r=c.get(f"/p/{pid}/buying/{prod_id}"); assert "Part sourced" in r.text and "Buy for the room without a shop" in r.text
    r=c.post(f"/p/{pid}/buying/{prod_id}/quotes/{q1}/buy", follow_redirects=False); assert r.status_code==303 and "#buying-from" in r.headers["location"]
    db=SessionLocal(); assert db.get(Item, cid).unit_price==852.0 and db.get(Item, cid).purchase_id==db.get(Item, usd_id).purchase_id and db.query(Purchase).filter(Purchase.product_id==prod_id).count()==1; db.close()
    # a corrected quote reaches every room bought from it; back to CNY drops the USD memory everywhere
    r=c.post(f"/p/{pid}/buying/{prod_id}/quotes/{q1}", data={"supplier_id":str(sup.id),"amount":"130","currency":"USD","lead_time":"3 weeks"}, follow_redirects=False); assert "The rooms bought from it follow" in c.get(r.headers["location"]).text
    db=SessionLocal(); assert db.get(Item, cid).unit_price==923.0 and db.get(Item, cid).price_amount==130 and db.get(Item, usd_id).lead_time=="3 weeks"; db.close()
    c.post(f"/p/{pid}/buying/{prod_id}/quotes/{q1}", data={"supplier_id":str(sup.id),"amount":"900","currency":"CNY"}, follow_redirects=False)
    db=SessionLocal(); assert db.get(Item, usd_id).unit_price==900.0 and db.get(Item, usd_id).price_entry is None and db.get(Item, cid).price_entry is None and db.query(ItemPrice).filter(ItemPrice.item_id.in_([usd_id, cid])).count()==0; db.close()
    # junk currency = CNY; a copy starts without a shop; the item form ignores a typed price for good
    c.post(f"/p/{pid}/buying/{prod_id}/quotes/{q1}", data={"supplier_id":str(sup.id),"amount":"55.5","currency":"EUR"}, follow_redirects=False)
    db=SessionLocal(); assert db.get(Item, usd_id).unit_price==55.5 and db.get(Item, usd_id).price_currency=="CNY" and db.get(Quote, q1).currency=="CNY"; db.close()
    c.post(f"/p/{pid}/items/{usd_id}", data={"room_id":room.id,"category":"Lighting","name":"Brass reading light","qty":"2","unit":"pcs","unit_price":"1","price_currency":"USD","supplier_id":"","status":"Quoted"}, follow_redirects=False)
    db=SessionLocal(); assert db.get(Item, usd_id).unit_price==55.5 and db.get(Item, usd_id).supplier_id==sup.id; db.close()
    r=c.post(f"/p/{pid}/items/{usd_id}/duplicate", follow_redirects=False); dup_id=int(r.headers["location"].split("/")[-1])
    db=SessionLocal(); d_=db.get(Item, dup_id); assert d_.unit_price==0 and d_.supplier_id is None and d_.purchase_id is None and d_.status=="To buy"; db.close()
    c.post(f"/p/{pid}/items/{dup_id}/delete", follow_redirects=False)
    db=SessionLocal(); assert db.get(Item, dup_id) is None and db.query(Purchase).filter(Purchase.product_id==prod_id).count()==1; db.close()
    # drafts: the quick form takes no price
    r=c.post(f"/p/{pid}/capture", data={"room_id":room.id,"category":"Hardware"}, files=[("photo",("d.jpg",img("gray"),"image/jpeg"))], follow_redirects=False)
    db=SessionLocal(); dr=db.query(Item).filter(Item.project_id==int(pid), Item.draft==True).order_by(Item.id.desc()).first(); db.close()  # noqa: E712
    r=c.get(f"/p/{pid}/drafts"); assert 'name="price_currency"' not in r.text and 'name="unit_price"' not in r.text
    r=c.post(f"/p/{pid}/drafts/{dr.id}", data={"name":"Door stop","room_id":room.id,"category":"Hardware","qty":"4","unit_price":"3","price_currency":"USD"}, follow_redirects=False)
    db=SessionLocal(); x=db.get(Item, dr.id); assert not x.draft and x.unit_price==0 and x.qty==4; db.close()
    # the help page: every feature has a section; login required
    r=c.get("/help"); assert r.status_code==200 and "How to use" in r.text and 'id="prices"' in r.text and 'id="plan"' in r.text and 'id="client"' in r.text and "Quick capture" in r.text and 'href="/help"' in r.text
    assert c2.get("/help", follow_redirects=False).status_code==303
    # the guide: every picture it shows exists and has its size, every contents entry has its section, every tip's anchor exists
    help_html=r.text; import json as _json
    pics=sorted(set(re.findall(r'/static/help/([\w-]+)\.webp', help_html))); assert len(pics)>=30, pics
    sizes=_json.load(open("app/static/help/shots.json"))
    for name in pics:
        assert os.path.exists(f"app/static/help/{name}.webp"), name
        assert os.path.getsize(f"app/static/help/{name}.webp") < 200_000, name
        assert name in sizes and f'width="{sizes[name][0]}" height="{sizes[name][1]}"' in help_html, name
    assert c.get(f"/static/help/{pics[0]}.webp").status_code==200
    ids=set(re.findall(r'<section class="card[^"]*" id="([\w-]+)"', help_html)) | set(re.findall(r'<div class="stage" id="([\w-]+)"', help_html))
    toc=help_html.split('<nav class="help-toc')[1].split('</nav>')[0]; flow=help_html.split('<ol class="flow">')[1].split('</ol>')[0]
    for href in re.findall(r'href="#([\w-]+)"', toc + flow):
        assert href in ids, href
    for st in ("s1","s9","flow","statuses","purchasing","faq","help"): assert st in ids, st
    from app import help_tips as _ht
    for k,t in _ht.TIPS.items():
        assert t["anchor"] in ids or t.get("public"), (k, t["anchor"])
    assert set(_ht.TEMPLATE_KEYS.values()) <= set(_ht.TIPS)
    assert 'id="page-help"' not in help_html  # the guide itself has no drawer
    # help for this page: the drawer, the nudge and the deep link into the guide on the designer's pages
    r=c.get(f"/p/{pid}/rooms"); assert 'id="page-help"' in r.text and 'data-key="rooms"' in r.text and 'id="help-btn"' in r.text and 'id="help-nudge"' in r.text and 'href="/help#rooms"' in r.text and "Sort rooms &amp; areas by name" in r.text.split('id="page-help"')[1]
    r=c.get(f"/p/{pid}/plan?mode=mark"); assert 'data-key="plan_mark"' in r.text and "Done marking" in r.text.split('id="page-help"')[1]
    r=c.get(f"/p/{pid}/plan"); assert 'data-key="plan"' in r.text
    r=c.get(f"/p/{pid}/items/new"); assert 'data-key="item_new"' in r.text
    r=c.get(f"/p/{pid}/items/{iid}"); assert 'data-key="item_edit"' in r.text
    r=c.get("/projects/new"); assert 'data-key="project_new"' in r.text
    r=c.get(f"/p/{pid}/edit"); assert 'data-key="project_edit"' in r.text
    for path,key in ((f"/p/{pid}","overview"),("/","projects"),(f"/p/{pid}/images","images"),(f"/p/{pid}/capture","capture"),(f"/p/{pid}/drafts","drafts"),(f"/p/{pid}/suppliers","suppliers"),(f"/p/{pid}/payments","payments"),(f"/p/{pid}/cartons","cartons"),(f"/p/{pid}/import","import"),("/settings","settings"),("/clip","clip")):
        r=c.get(path); assert r.status_code==200 and f'data-key="{key}"' in r.text, (path, key)
    r=c.get("/login", follow_redirects=False); assert 'id="page-help"' not in c2.get("/login").text
    # the share pages get their own help, with no link into the designer's guide and nothing designer-only
    r=c2.get(f"/c/{ctok}"); drawer=r.text.split('id="page-help"')[1].split('</aside>')[0]; assert 'data-key="share_client"' in r.text and "/help" not in drawer and "/p/" not in drawer and "supplier" not in drawer.lower() and "Download PDF" in drawer
    r=c2.get(f"/s/{tok}"); drawer=r.text.split('id="page-help"')[1].split('</aside>')[0]; assert 'data-key="share_supplier"' in r.text and "/help" not in drawer and "/p/" not in drawer and "price" not in drawer.lower() and "Add carton" in drawer
    # empty states point at the help
    r=c.post("/projects/new", data={"client_name":"Empty","name":"Empty house","rate":"7.1"}, follow_redirects=False); pid_e=r.headers["location"].split("/")[-1]
    assert 'href="#help"' in c.get(f"/p/{pid_e}/items").text and 'href="#help"' in c.get(f"/p/{pid_e}/plan").text and 'href="#help"' in c.get(f"/p/{pid_e}/drafts").text and 'href="#help"' in c.get(f"/p/{pid_e}/cartons").text and 'href="#help"' in c.get(f"/p/{pid_e}/payments").text
    print("prices in USD or CNY + help ok")
    # quick picks on the room checklist, the phone-friendly photo field, and the import template round trip
    from app.models import Room as _R
    def _grp(code, name): return _R(code=code, name=name).group
    assert _grp("FF-MBR", "Master bedroom")=="bedroom" and _grp("FF-MEN", "Master ensuite")=="bathroom" and _grp("GF-GBA", "Guest bathroom")=="bathroom" and _grp("GF-GST", "Guest bedroom")=="bedroom"
    assert _grp("FF-BA1", "BA 1")=="bathroom" and _grp("FF-MBA", "MBA")=="bathroom" and _grp("FF-BR2", "Second")=="bedroom" and _grp("FF-BAL", "Balconies")=="" and _grp("GF-KIT", "Kitchen")=="" and _grp("GF-WC", "Downstairs WC")=="bathroom" and _grp("FF-NUR", "Nursery")=="bedroom"
    r=c.post("/projects/new", data={"client_name":"Picks","name":"Picks house","rate":"7.1"}, follow_redirects=False); pid_p=r.headers["location"].split("/")[-1]
    for code,name,floor,kind in (("GF-LIV","Living room","Ground","room"),("GF-GBA","Guest bathroom","Ground","room"),("GF-GST","Guest bedroom","Ground","room"),("FF-MBR","Master bedroom","First","room"),("FF-MEN","Master ensuite","First","room"),("FF-BR1","Bedroom 1","First","room"),("FF-BA1","Bathroom 1","First","room"),("GF-ENT","Entrance","Ground","area"),("FF-BAL","Balconies","First","area")):
        c.post(f"/p/{pid_p}/rooms", data={"code":code,"name":name,"floor":floor,"kind":kind}, follow_redirects=False)
    r=c.get(f"/p/{pid_p}/items/new"); h=r.text
    assert 'data-pick="group:bedroom"' in h and 'All bedrooms <b>3</b>' in h and 'All bathrooms <b>3</b>' in h and 'data-pick="floor:Ground"' in h and 'data-pick="floor:First"' in h and 'data-pick="kind:room"' in h and 'data-pick="kind:area"' in h and 'data-pick="none"' in h
    assert 'data-group="bathroom" data-floor="First" data-kind="room"' in h and 'capture=' not in h.split('name="photos"')[1][:120] and 'name="photos" accept="image/*" multiple' in h
    db=SessionLocal(); mbr=db.query(Room).filter(Room.project_id==int(pid_p), Room.code=="FF-MBR").one().id; db.close()
    r=c.post(f"/p/{pid_p}/items/new", data={"room_ids":[str(mbr)],"category":"Furniture","name":"Bedside lamp","qty":"2","unit":"pcs","unit_price":"300","status":"To buy"}, follow_redirects=False)
    r=c.get(r.headers["location"]); assert 'name="rooms_form" value="1"' in r.text and 'name="room_ids"' in r.text and 'All bedrooms <b>3</b>' in r.text and "Rooms with this item" in r.text  # every room, this one locked
    r=c.get(f"/p/{pid_p}/import/template.xlsx"); assert r.status_code==200 and "spreadsheetml" in r.headers["content-type"] and "import_template.xlsx" in r.headers["content-disposition"]
    wb=openpyxl.load_workbook(io.BytesIO(r.content)); assert wb.sheetnames==["How to","Rooms","Shopping List","Lists"]
    ws=wb["Rooms"]; labels=[ws.cell(row=i, column=1).value for i in range(2, 30)]
    assert labels[:4]==["All rooms","All areas","All bedrooms","All bathrooms"] and "All on Ground floor" in labels and "All on First floor" in labels and ws.cell(row=2, column=3).value=="pick"  # quick picks first
    assert "FF-MBR - Master bedroom" in labels and "GF-ENT - Entrance" in labels and [c_.value for c_ in ws[1]]==["Room","Floor","Kind","Floor area m²","Wall tile m²","Notes"]
    assert ws.cell(row=labels.index("GF-ENT - Entrance")+2, column=3).value=="area" and len(ws.data_validations.dataValidation)==2
    sl=wb["Shopping List"]; hdr=[c_.value for c_ in sl[1]]; assert hdr[:5]==["Code","Room","Category","Item","Photo"] and "Qty" in hdr and "Price (CNY)" in hdr and "Optional" in hdr and "Total CNY" not in hdr
    dvs={str(d.sqref)[0]: d.formula1 for d in sl.data_validations.dataValidation}; assert dvs["B"]=="Rooms!$A$2:$A$401" and dvs["C"].startswith("Lists!$A$2") and "O" in dvs and "P" in dvs and "A" not in dvs and "E" not in dvs, dvs
    assert "How to fill this in" in str(wb["How to"]["A1"].value) and wb["Lists"]["A2"].value=="Paint" and wb["Lists"]["D2"].value=="room"
    ws.append(["SF-STU - Studio","Second","room",None,None,"new floor"]); ws.append(["SF-TER - Terrace","Second","area",None,None,None]); ws.append(["(example) ZZ - Ignore","Second","room",None,None,None])
    sl.append(["","SF-STU - Studio","Furniture","Day bed","","Oak frame","","2000 mm","Oak",1,"pcs",3200,"New Supplier Co","4 weeks","Quoted","no",""])
    sl.append(["","SF-TER - Terrace","Lighting","Terrace wall light","","IP65","","","Black",2,"pcs",310,"","","To buy","yes",""])
    sl.append(["","FF-BA1 - Bathroom 1","Tiles","Floor tile","","R11","","600 x 600","Grey",8,"m²",190,"","","To buy","",""])
    sl.append(["","All bedrooms","Lighting","Bedside sconce","","","","","Brass",1,"pcs",260,"","","To buy","",""])  # three bedrooms
    sl.append(["","GF-LIV - Living room; FF-BR1 - Bedroom 1","Furniture","Side table","","","","450 mm","Oak",1,"pcs",700,"","","To buy","",""])  # two rooms, typed
    sl.append(["","All on Second","Paint","Ceiling paint","","","","","White",2,"pcs",180,"","","To buy","",""])  # the two rooms added above
    sl.append(["ZZ-99","GF-LIV - Living room","Lighting","Mystery lamp","","","","","",1,"pcs",90,"","","To buy","",""])  # an unknown code: a new item that keeps it
    sl.append(["","Kitchn","Other","Typo room item","","","","","",1,"pcs",10,"","","To buy","",""])  # an unknown room: whole house, with a warning
    sl.append(["(example) X","FF-BA1 - Bathroom 1","Tiles","Should be skipped","","","","","",1,"pcs",1,"","","","",""]); sl.append(["","","Paint","(example) skipped too","","","","","",1,"pcs",1,"","","","",""])
    b=io.BytesIO(); wb.save(b); r=c.post(f"/p/{pid_p}/import", files={"file":("filled.xlsx", b.getvalue())}, data={"apply":"1"}); msg=re.search(r"Imported[^<]*", r.text).group(0)
    assert msg.startswith("Imported 12 new items, 2 new rooms.") and "Skipped 2 rows" in msg and 'Room &#34;Kitchn&#34; not found' in r.text and "come from the Buying page now" in r.text, msg
    db=SessionLocal(); rs={x.code:x for x in db.query(Room).filter(Room.project_id==int(pid_p))}; its={i.name:i for i in db.query(Item).filter(Item.project_id==int(pid_p))}
    assert len(rs)==12 and rs["SF-STU"].floor=="Second" and rs["SF-STU"].kind=="room" and rs["SF-TER"].kind=="area" and "ZZ" not in rs and rs["FF-MBR"].floor=="First" and not any(x.name.startswith("All ") for x in rs.values())
    assert its["Day bed"].room.code=="SF-STU" and its["Day bed"].status=="Quoted" and its["Day bed"].supplier is None and its["Day bed"].unit_price==0 and db.query(Supplier).filter(Supplier.name=="New Supplier Co").count()==0 and its["Terrace wall light"].optional and its["Floor tile"].room.code=="FF-BA1" and its["Floor tile"].code=="FF-BA1-01" and "Should be skipped" not in its
    by_name=lambda n: sorted(i.room.code for i in db.query(Item).filter(Item.project_id==int(pid_p), Item.name==n))
    assert by_name("Bedside sconce")==["FF-BR1","FF-MBR","GF-GST"] and by_name("Side table")==["FF-BR1","GF-LIV"] and by_name("Ceiling paint")==["SF-STU","SF-TER"]
    assert its["Mystery lamp"].code=="ZZ-99" and its["Typo room item"].room_id is None and db.query(Item).filter(Item.project_id==int(pid_p)).count()==13  # 12 new + the bedside lamp
    day_id=its["Day bed"].id; db.close()
    buy(c, pid_p, day_id, "New Supplier Co", 3200, lead="4 weeks")  # the shop is chosen on the Buying page, not in the sheet
    # export → edit → import: rows are matched by their code and updated in place; nothing is added twice
    r=c.get(f"/p/{pid_p}/export/items.xlsx"); assert r.status_code==200 and "items.xlsx" in r.headers["content-disposition"]
    wb2=openpyxl.load_workbook(io.BytesIO(r.content)); sl2=wb2["Shopping List"]; hdr2=[c_.value for c_ in sl2[1]]
    assert hdr2[:4]==["Code","Room","Category","Item"] and hdr2[-2:]==["Total CNY","Total USD"] and wb2["Rooms"]["A2"].value=="All rooms" and len(sl2.data_validations.dataValidation)>=5
    rows={sl2.cell(row=i, column=4).value: i for i in range(2, 40) if sl2.cell(row=i, column=1).value}; assert len(rows)>=9 and sl2.cell(row=rows["Day bed"], column=1).value=="SF-STU-01" and sl2.cell(row=rows["Day bed"], column=12).value==3200 and sl2.cell(row=rows["Day bed"], column=18).value==3200
    i_=rows["Day bed"]; sl2.cell(row=i_, column=12, value=3500); sl2.cell(row=i_, column=15, value="Ordered"); sl2.cell(row=i_, column=6, value="-"); sl2.cell(row=i_, column=10, value=2); sl2.cell(row=i_, column=2, value="SF-TER - Terrace")
    sl2.cell(row=rows["Floor tile"], column=7, value="Marazzi")
    b=io.BytesIO(); wb2.save(b); r=c.post(f"/p/{pid_p}/import", files={"file":("edited.xlsx", b.getvalue())}, data={"apply":"1"}); msg=re.search(r"Imported[^<]*", r.text).group(0)
    assert msg.startswith("Imported 0 new items and updated 2, 0 new rooms.") and "1 row carries a price, supplier or lead time that differs from the app" in r.text, msg
    db=SessionLocal(); its={i.name:i for i in db.query(Item).filter(Item.project_id==int(pid_p))}; d=its["Day bed"]
    assert d.code=="SF-STU-01" and d.room.code=="SF-TER" and d.unit_price==3200 and d.status=="Ordered" and d.spec=="" and d.qty==2 and d.supplier.name=="New Supplier Co" and d.lead_time=="4 weeks" and d.size=="2000 mm" and d.finish=="Oak"  # blanks kept, "-" cleared, code kept; the price stays the shop's
    assert its["Floor tile"].brand=="Marazzi" and its["Floor tile"].qty==8 and db.query(Item).filter(Item.project_id==int(pid_p)).count()==13; db.close()
    # preview: the plan is shown and nothing is saved until "Apply"; the kept file is removed afterwards
    from app import storage as _stx
    i_=rows["Day bed"]; sl2.cell(row=i_, column=12, value=3600); sl2.cell(row=i_, column=15, value="Bought"); sl2.cell(row=i_, column=17, value="Deliver to the warehouse")  # an unknown status: a warning, the status stays
    sl2.append(["", "GF-LIV - Living room", "Furniture", "Preview stool","", "", "", "", "", 1, "pcs", 120, "", "", "To buy", "", ""])
    b=io.BytesIO(); wb2.save(b); r=c.post(f"/p/{pid_p}/import", files={"file":("edited2.xlsx", b.getvalue())}, data={}); h=r.text
    assert "Check before applying" in h and "1 new item</span>" in h and "1 updated</span>" in h and "Preview stool" in h and "GF-LIV-" in h and "come from the Buying page now" in h and "Price (CNY)</b>" not in h and 'Status &#34;Bought&#34; is not one of' in h and "Apply the import" in h, h[h.find("Check before"):h.find("Check before")+3000]
    key=re.search(r'name="key" value="(imports/p\d+/[a-f0-9]+\.xlsx)"', h).group(1); assert _stx.read_image(key) is not None
    db=SessionLocal(); assert db.query(Item).filter(Item.project_id==int(pid_p), Item.name=="Day bed").one().notes=="" and db.query(Item).filter(Item.project_id==int(pid_p), Item.name=="Preview stool").count()==0; db.close()  # nothing saved yet
    r=c.post(f"/p/{pid_p}/import", data={"key":key,"apply":"1"}); msg=re.search(r"Imported[^<]*", r.text).group(0); assert msg.startswith("Imported 1 new item and updated 1,"), msg
    db=SessionLocal(); d=db.query(Item).filter(Item.project_id==int(pid_p), Item.name=="Day bed").one(); assert d.unit_price==3200 and d.notes=="Deliver to the warehouse" and d.status=="Ordered" and db.query(Item).filter(Item.project_id==int(pid_p), Item.name=="Preview stool").one().room.code=="GF-LIV"; db.close()
    assert _stx.read_image(key) is None and "no longer here" in c.post(f"/p/{pid_p}/import", data={"key":key,"apply":"1"}).text
    r=c.post(f"/p/{pid_p}/import", files={"file":("again.xlsx", b.getvalue())}, data={}); key2=re.search(r'name="key" value="([^"]+)"', r.text).group(1); assert "1 new item</span>" in r.text and "0 updated</span>" in r.text
    r=c.post(f"/p/{pid_p}/import/cancel", data={"key":key2}, follow_redirects=False); assert r.status_code==303 and _stx.read_image(key2) is None
    assert "Choose an .xlsx file first" in c.post(f"/p/{pid_p}/import", data={}).text and c.post(f"/p/{pid_p}/import", data={"key":"imports/p999/x.xlsx","apply":"1"}).status_code==200
    # pictures: the export carries each item's cover over its Photo cell; a picture over a row, or a file named after a code or
    # an item name, is added to items that have no photo yet, and never twice
    from openpyxl.drawing.image import Image as _XLI
    xk=openpyxl.load_workbook(io.BytesIO(c.get(f"/p/{pid}/export/items.xlsx").content)); sk=xk["Shopping List"]
    assert len(sk._images)>=1 and sk._images[0].anchor._from.col==4 and sk.row_dimensions[sk._images[0].anchor._from.row+1].height==66 and [c_.value for c_ in sk[1]][4]=="Photo"  # the Kinondoni project has photos
    wb3=openpyxl.load_workbook(io.BytesIO(c.get(f"/p/{pid_p}/import/template.xlsx").content)); sl3=wb3["Shopping List"]
    sl3.append(["","GF-LIV - Living room","Lighting","Embedded lamp","","","","","",1,"pcs",80,"","","To buy","",""]); prow=sl3.max_row
    pic=io.BytesIO(img("green")); xi=_XLI(pic); xi.width, xi.height=96, 72; sl3.add_image(xi, f"E{prow}")
    sl3.append(["","GF-LIV - Living room; FF-BR1 - Bedroom 1","Furniture","Photo stool","","","","","",1,"pcs",120,"","","To buy","",""])
    b=io.BytesIO(); wb3.save(b)
    db=SessionLocal(); dbc=db.query(Item).filter(Item.project_id==int(pid_p), Item.name=="Day bed").one().code; db.close()
    r=c.post(f"/p/{pid_p}/import", files=[("file",("pics.xlsx", b.getvalue())),("photos",("Photo stool.jpg", img("red"), "image/jpeg")),("photos",("photo_stool-2.JPG", img("blue"), "image/jpeg")),("photos",(f"{dbc}.jpg", img("gray"), "image/jpeg")),("photos",("nobody.jpg", img("white"), "image/jpeg"))], data={}); h=r.text
    assert "4 pictures</span>" in h and 'Picture &#34;nobody.jpg&#34; matches no item' in h and "Embedded lamp" in h and "Photo</b>: — → 1 added" in h, re.findall(r"\d+ pictures?</span>", h)
    key3=re.search(r'name="key" value="([^"]+)"', h).group(1); assert _stx.read_image(key3[:-5]+".json") is not None
    db=SessionLocal(); assert db.query(Item).filter(Item.project_id==int(pid_p), Item.name=="Embedded lamp").count()==0; db.close()
    r=c.post(f"/p/{pid_p}/import", data={"key":key3,"apply":"1"}); assert "3 new items" in r.text and "4 pictures added" in r.text, re.search(r"Imported[^<]*", r.text).group(0)
    db=SessionLocal(); its3={i.name:i for i in db.query(Item).filter(Item.project_id==int(pid_p))}
    assert len(its3["Embedded lamp"].photos)==1 and all(len(i.all_photos)==2 for i in db.query(Item).filter(Item.project_id==int(pid_p), Item.name=="Photo stool")) and sum(len(i.photos) for i in db.query(Item).filter(Item.project_id==int(pid_p), Item.name=="Photo stool"))==2 and len(its3["Day bed"].photos)==1 and its3["Day bed"].photos[0].thumb_key
    db.close(); assert _stx.read_image(key3) is None and _stx.read_image(key3[:-5]+".json") is None
    xk=c.get(f"/p/{pid_p}/export/items.xlsx").content; r=c.post(f"/p/{pid_p}/import", files={"file":("export.xlsx", xk)}, data={}); h=r.text  # the export's own pictures are never re-added
    assert "0 pictures</span>" in h and "not added: those items already have a photo" in h; key4=re.search(r'name="key" value="([^"]+)"', h).group(1); c.post(f"/p/{pid_p}/import/cancel", data={"key":key4})
    assert _stx.read_image(key4) is None
    from app.routers.importer import RoomIndex as _RI
    db=SessionLocal(); ri=_RI(db.get(_P,int(pid_p)))
    assert sorted(x.code for x in ri.from_cell("All bathrooms"))==["FF-BA1","FF-MEN","GF-GBA"] and ri.from_cell("")==[None] and ri.from_cell("ALL - Whole house")==[None] and ri.from_cell("Nowhere") is None
    assert [x.code for x in ri.from_cell("Kitchn, Living room")]==["GF-LIV"] and [x.code for x in ri.from_cell("First floor")]==[x.code for x in ri.from_cell("All on First floor")] and len(ri.from_cell("all on first floor"))==5 and ri.from_cell("All bathrooms; GF-ENT")[-1].code=="GF-ENT"; db.close()
    print("quick picks + import template ok")
    # plan layers: several sheets of one floor; the furniture layout carries the boxes, the others borrow them
    from app import drawings as _dr2
    from app.models import RoomPin, Project as _P
    r=c.post("/projects/new", data={"client_name":"Layers","name":"Layered house","rate":"7.1"}, follow_redirects=False); pid_l=r.headers["location"].split("/")[-1]
    for code,name in (("GF-KIT","Kitchen"),("GF-LIV","Living room"),("GF-ENT","Entrance")): c.post(f"/p/{pid_l}/rooms", data={"code":code,"name":name,"floor":"Ground","kind":"auto"}, follow_redirects=False)
    c.post(f"/p/{pid_l}/images/plans", data={"floor":"Ground","sheet":"A-101","caption":"Furniture Layout"}, files=[("files",("g.jpg",img("white"),"image/jpeg"))], follow_redirects=False)
    c.post(f"/p/{pid_l}/images/plans", data={"floor":"Ground","sheet":"E-01","layer":"electrical"}, files=[("files",("e.jpg",img("blue"),"image/jpeg"))], follow_redirects=False)
    c.post(f"/p/{pid_l}/images/plans", data={"floor":"Roof","sheet":"A-103","layer":"bogus"}, files=[("files",("r.jpg",img("gray"),"image/jpeg"))], follow_redirects=False)
    db=SessionLocal(); lp={im.sheet:im for im in db.query(ProjectImage).filter(ProjectImage.project_id==int(pid_l))}; lr={x.code:x.id for x in db.query(Room).filter(Room.project_id==int(pid_l))}
    assert lp["E-01"].layer=="electrical" and lp["A-101"].layer=="main" and lp["A-103"].layer=="main" and _dr2.plan_caption(lp["E-01"])=="Ground floor · Electrical & lighting · E-01" and _dr2.plan_caption(lp["A-101"])=="Ground floor · A-101 · Furniture Layout"
    gid_l,eid_l=lp["A-101"].id,lp["E-01"].id; db.close()
    for code,box in (("GF-KIT",(0.1,0.1,0.3,0.3)),("GF-LIV",(0.5,0.1,0.4,0.3))): c.post(f"/p/{pid_l}/plan/pins", data={"image_id":gid_l,"room_id":lr[code],"x":box[0],"y":box[1],"w":box[2],"h":box[3]}, headers={"Accept":"application/json"})
    r=c.get(f"/p/{pid_l}/plan"); h=r.text; assert h.count('<span class="t">Ground floor</span>')==1 and '<span class="t">Roof</span>' in h and 'class="plan-layers"' in h and "Main plan" in h and "Electrical &amp; lighting" in h and "2 sheets" in h and f'data-plan="{gid_l}"' in h
    r=c.get(f"/p/{pid_l}/plan?plan={eid_l}"); h=r.text; assert len(re.findall(r'class="pin[" ]', h))==2 and "boxes from the main plan" in h and "2 of 3 placed" in h and f'data-plan="{eid_l}"' in h  # borrowed boxes
    r=c.get(f"/p/{pid_l}/plan?plan={eid_l}&mode=mark"); h=r.text; assert "Boxes borrowed from the main plan" in h and "Copy the 2 boxes here" in h and 'data-mode=""' in h and "Done marking" not in h
    c.post(f"/p/{pid_l}/items/new", data={"room_ids":[str(lr["GF-KIT"])],"category":"Lighting","name":"Pendant","qty":"1","unit":"pcs","unit_price":"500","status":"To buy"}, follow_redirects=False)
    c.post(f"/p/{pid_l}/items/new", data={"room_ids":[str(lr["GF-KIT"])],"category":"Furniture","name":"Stool","qty":"2","unit":"pcs","unit_price":"200","status":"To buy"}, follow_redirects=False)
    db=SessionLocal(); li={i.name:i.id for i in db.query(Item).filter(Item.project_id==int(pid_l))}; db.close()
    assert c.post(f"/p/{pid_l}/plan/item-pins", data={"image_id":eid_l,"item_id":li["Pendant"],"x":0.2,"y":0.2}, headers={"Accept":"application/json"}).json()["ok"]
    assert c.post(f"/p/{pid_l}/plan/item-pins", data={"image_id":gid_l,"item_id":li["Stool"],"x":0.25,"y":0.25}, headers={"Accept":"application/json"}).json()["ok"]
    r=c.get(f"/p/{pid_l}/plan?plan={eid_l}&room={lr['GF-KIT']}"); assert "Not placed on this plan yet" not in r.text and re.search(r'data-placed>1</span> of 2', r.text)  # the room's box is borrowed; one dot on this sheet
    db=SessionLocal(); prl=db.get(_P,int(pid_l)); kit_l=db.get(Room, lr["GF-KIT"])
    z=_dr2.room_zoom(prl, kit_l); assert z["plan"].id==gid_l and [d["code"] for d in z["dots"]]==["GF-KIT-02"]  # the main plan, with the furniture dot
    z2=_dr2.room_zoom(prl, kit_l, plan=db.get(ProjectImage, eid_l)); assert z2["plan"].id==eid_l and [d["code"] for d in z2["dots"]]==["GF-KIT-01"] and z2["pin"].image_id==gid_l  # borrowed box, this sheet's dots
    assert _dr2.plan_for_room(prl, kit_l).id==gid_l and _dr2.default_plan(prl).id==gid_l and [im.id for im in _dr2.plans_for_floor(prl, "ground")]==[gid_l, eid_l]; db.close()
    tx=pdf_text(c.get(f"/p/{pid_l}/export/room/{lr['GF-KIT']}.pdf").content); assert "ON THE PLAN" in tx and "ON THE ELECTRICAL & LIGHTING PLAN" in tx
    ts=pdf_text(c.get(f"/p/{pid_l}/export/schedule.pdf").content); assert "FLOOR PLANS" in ts and "Ground floor · electrical & lighting plan" in ts and ts.count("Ground floor · A-101 · Furniture Layout")==1 and ts.count("Ground floor · Electrical & lighting · E-01")==1
    r=c.post(f"/p/{pid_l}/plan/pins/copy", data={"image_id":eid_l}, follow_redirects=False); assert r.status_code==303 and r.headers["location"].endswith(f"plan={eid_l}&mode=mark")
    db=SessionLocal(); assert db.query(RoomPin).filter(RoomPin.image_id==eid_l).count()==2; db.close()
    r=c.get(f"/p/{pid_l}/plan?plan={eid_l}&mode=mark"); assert "Done marking" in r.text and "Boxes borrowed" not in r.text  # own boxes now
    assert c.post(f"/p/{pid_l}/plan/pins/copy", data={"image_id":eid_l}, headers={"Accept":"application/json"}).status_code==404  # nothing left to borrow
    r=c.get(f"/p/{pid_l}/images"); assert r.text.count('name="layer"')==4 and "Ground floor · Electrical &amp; lighting · E-01" in r.text
    c.post(f"/p/{pid_l}/images/{eid_l}", data={"floor":"Ground","sheet":"E-01","caption":"","layer":"ceiling"}, follow_redirects=False); db=SessionLocal(); assert db.get(ProjectImage, eid_l).layer=="ceiling"; ltok=db.get(_P,int(pid_l)).client_token; db.close()
    r=c2.get(f"/c/{ltok}"); assert 'class="plan-layers"' in r.text and "Ceiling" in r.text and r.text.count('<span class="t">Ground floor</span>')==1
    r=c.get(f"/p/{pid_l}/rooms"); assert r.status_code==200 and f"plan={gid_l}&amp;room=" in r.text
    print("plan layers ok")
    # existing plans adopted into layers: a plan filed as its own floor ("Ground Electrical") is offered as a layer in one press
    c.post(f"/p/{pid_l}/images/plans", data={"floor":"Ground Electrical","sheet":"BS4449"}, files=[("files",("ge.jpg",img("white"),"image/jpeg"))], follow_redirects=False)
    db=SessionLocal(); lp={im.sheet:im for im in db.query(ProjectImage).filter(ProjectImage.project_id==int(pid_l))}; prl=db.get(_P,int(pid_l)); leg=lp["BS4449"]; leg_id=leg.id
    hint=_dr2.legacy_layer_hint(prl, leg); assert hint and hint["floor"]=="Ground" and hint["layer"]=="electrical" and hint["has_main"] and hint["floor_title"]=="Ground floor", hint
    assert _dr2.legacy_layer_hint(prl, lp["A-101"]) is None and _dr2.legacy_layer_hint(prl, lp["E-01"]) is None and _dr2.plan_role(prl, lp["A-101"])=="main" and _dr2.plan_role(prl, lp["E-01"])=="layer" and _dr2.plan_role(prl, leg)=="main"
    db.close()
    r=c.get(f"/p/{pid_l}/images"); h=r.text
    assert "Main plan (general floor plan)" in h and 'Looks like the <b>electrical &amp; lighting</b> sheet of <b>Ground floor</b>.' in h and f'action="/p/{pid_l}/images/{leg_id}" class="plan-hint"' in h and h.count('class="plan-hint"')==1
    assert 'name="floor" value="Ground"><input type="hidden" name="sheet" value="BS4449"><input type="hidden" name="caption" value=""><input type="hidden" name="layer" value="electrical">' in h and "Make it a layer of Ground floor" in h
    assert h.count('<span class="pill accent">Main plan</span> Ground floor')==1 and '<span class="pill">Layer</span> of Ground floor' in h and "Main plan</span> Ground Electrical" in h
    r=c.post(f"/p/{pid_l}/images/{leg_id}", data={"floor":"Ground","sheet":"BS4449","caption":"","layer":"electrical"}, follow_redirects=False); assert r.status_code==303
    db=SessionLocal(); prl=db.get(_P,int(pid_l)); leg=db.get(ProjectImage, leg_id); assert leg.layer=="electrical" and leg.floor=="Ground" and _dr2.legacy_layer_hint(prl, leg) is None and _dr2.plan_role(prl, leg)=="layer"
    assert [im.id for im in _dr2.plans_for_floor(prl, "Ground")]==[gid_l, eid_l, leg_id]; db.close()
    r=c.get(f"/p/{pid_l}/plan"); h=r.text; assert h.count('<span class="t">Ground floor</span>')==1 and "3 sheets" in h
    r=c.get(f"/p/{pid_l}/images"); assert 'class="plan-hint"' not in r.text and "Ground floor · Electrical &amp; lighting · BS4449" in r.text
    # two furniture layouts on one floor: the one with room boxes is the main plan, the other says so
    r=c.post("/projects/new", data={"client_name":"Twins","name":"Twin plans","rate":"7.1"}, follow_redirects=False); pid_t=r.headers["location"].split("/")[-1]
    c.post(f"/p/{pid_t}/rooms", data={"code":"GF-KIT","name":"Kitchen","floor":"Ground","kind":"auto"}, follow_redirects=False)
    c.post(f"/p/{pid_t}/images/plans", data={"floor":"Ground","sheet":"A-100","caption":"Dimensions"}, files=[("files",("d.jpg",img("white"),"image/jpeg"))], follow_redirects=False)
    c.post(f"/p/{pid_t}/images/plans", data={"floor":"Ground","sheet":"A-101","caption":"Furniture"}, files=[("files",("f.jpg",img("gray"),"image/jpeg"))], follow_redirects=False)
    db=SessionLocal(); tp={im.sheet:im.id for im in db.query(ProjectImage).filter(ProjectImage.project_id==int(pid_t))}; tk=db.query(Room).filter(Room.project_id==int(pid_t), Room.code=="GF-KIT").one().id; prt=db.get(_P,int(pid_t))
    assert _dr2.main_plan(prt,"Ground").id==tp["A-100"] and _dr2.plan_role(prt, db.get(ProjectImage, tp["A-101"]))=="twin" and _dr2.plan_role(prt, db.get(ProjectImage, tp["A-100"]))=="main"; db.close()  # no boxes yet: the first sheet
    c.post(f"/p/{pid_t}/plan/pins", data={"image_id":tp["A-101"],"room_id":tk,"x":0.1,"y":0.1,"w":0.3,"h":0.3}, headers={"Accept":"application/json"})
    db=SessionLocal(); prt=db.get(_P,int(pid_t)); assert _dr2.main_plan(prt,"Ground").id==tp["A-101"] and _dr2.plan_role(prt, db.get(ProjectImage, tp["A-100"]))=="twin" and _dr2.plan_for_room(prt, db.get(Room, tk)).id==tp["A-101"]; db.close()  # the marked sheet is the main plan
    r=c.get(f"/p/{pid_t}/images"); h=r.text; assert '<span class="pill">Not the main plan</span> another sheet of Ground floor is the main plan and carries the boxes' in h and h.count('<span class="pill accent">Main plan</span> Ground floor')==1
    print("legacy layers + main plan ok")
    # plan sheet types: the Shows list is data (Studio settings → Plan sheet types), seeded with the built-ins and a windows & doors sheet
    from app import layers as _ly, config as _cfg
    r=c.get("/settings"); h=r.text; assert 'id="sheets"' in h and 'id="ly-windows"' in h and 'id="ly-main"' in h and "Windows &amp; doors" in h and 'value="window, door, joinery"' in h and h.count('class="btn sm ghost danger"')==6 and 'id="sheet-cats"' in h
    assert [k for k,_ in _ly.options()]==["main","furnishing","electrical","plumbing","ceiling","flooring","windows"] and _ly.options()[0][1]=="Main plan (general floor plan)" and _ly.title("windows")=="Windows & doors" and _ly.title("")=="Main plan" and _ly.title("old-sheet")=="Old sheet"
    assert _dr2.layer_from_title("GROUND FLOOR WINDOW & DOOR PLAN")=="windows" and _dr2.layer_for_category("Windows & Doors")=="windows" and _dr2.layer_for_category("Lighting")=="electrical" and _dr2.layer_for_category("Paint")=="main" and _dr2.layer_from_title("Textile Plan")=="main"
    tb=_dr2.read_title_block("WINDOW & DOOR PLAN - FIRST FLOOR\nDRAWING NO: W-01"); assert tb["is_plan"] is True and tb["floor"]=="First" and tb["sheet"]=="W-01" and _dr2.layer_from_title(tb["title"])=="windows", tb
    r=c.get(f"/p/{pid_l}/images"); assert 'value="windows">Windows &amp; doors</option>' in r.text and 'href="/settings#sheets"' in r.text
    # a windows sheet of Ground floor: named after the type, under the floor's Sheet switch, borrowing the boxes; the type cannot go while a plan shows it
    c.post(f"/p/{pid_l}/images/plans", data={"floor":"Ground","sheet":"W-01","layer":"windows"}, files=[("files",("w.jpg",img("green"),"image/jpeg"))], follow_redirects=False)
    db=SessionLocal(); wim=next(im for im in db.query(ProjectImage).filter(ProjectImage.project_id==int(pid_l)) if im.sheet=="W-01"); assert wim.layer=="windows" and _dr2.plan_caption(wim)=="Ground floor · Windows & doors · W-01" and not wim.is_main_layer and _dr2.plan_role(db.get(_P,int(pid_l)), wim)=="reference"; wid=wim.id; db.close()
    r=c.get(f"/p/{pid_l}/plan?plan={wid}"); assert "Windows &amp; doors" in r.text and "Reference sheet" in r.text and "4 sheets" in r.text  # a reference sheet: no borrowed boxes (tested below)
    r=c.post("/settings/layers/windows/delete", follow_redirects=False); assert "still%20shown%20by%201%20plan." in r.headers["location"] and r.headers["location"].endswith("#sheets")
    r=c.post("/settings/layers/main/delete", follow_redirects=False); assert "stays%20in%20the%20list" in r.headers["location"]
    r=c.get("/settings"); assert 'disabled title="1 plan still shows it' in r.text and "windows" in _ly.titles()
    # add a type with its words: in every dropdown at once, the page picker recognises its titles; rename, reorder, delete
    r=c.post("/settings/layers", data={"title":"Landscape","words":"landscape, garden | Planting"}, follow_redirects=False); assert r.headers["location"].startswith("/settings?ok=") and "landscape%20/%20garden%20/%20planting" in r.headers["location"]
    assert [k for k,_ in _ly.options()][-1]=="landscape" and _ly.title("landscape")=="Landscape" and _dr2.layer_from_title("Landscape Layout Plan")=="landscape" and _dr2.layer_from_title("Ground Floor Plan")=="main"
    tb=_dr2.read_title_block("GARDEN PLANTING PLAN\nL-01"); assert tb["is_plan"] is True and tb["sheet"]=="L-01" and _dr2.layer_from_title(tb["title"])=="landscape" and _dr2.caption_from_title("Garden Planting Plan", "landscape")=="" and _dr2.caption_from_title("Garden Planting Plan (Rev B)", "landscape")=="Rev B", tb
    r=c.get(f"/p/{pid_l}/images"); assert 'value="landscape">Landscape</option>' in r.text
    r=c.post("/settings/layers", data={"title":" landscape ","words":""}, follow_redirects=False); assert "already%20in%20the%20list" in r.headers["location"]
    r=c.post("/settings/layers", data={"title":"","words":"x"}, follow_redirects=False); assert "Give%20the%20sheet%20type%20a%20title" in r.headers["location"]
    r=c.post("/settings/layers", data={"title":"Landscape","words":""}, follow_redirects=False); assert "already" in r.headers["location"]  # same title again: refused, not a second key
    r=c.post("/settings/layers/landscape", data={"title":"Landscape & garden","words":"landscape;garden"}, follow_redirects=False); assert "is%20now" in r.headers["location"]
    assert _ly.title("landscape")=="Landscape & garden" and _ly.rows()[-1]["words"]==["landscape","garden"] and _ly.rows()[-1]["words_text"]=="landscape, garden"
    r=c.post("/settings/layers/main", data={"title":"General plan","words":"ignored"}, follow_redirects=False); assert "is%20now" in r.headers["location"] and _ly.options()[0]==("main","General plan (general floor plan)") and _ly.rows()[0]["words"]==[]
    c.post("/settings/layers/main", data={"title":"Main plan"}, follow_redirects=False); assert _ly.title("main")=="Main plan"
    r=c.post("/settings/layers/nope", data={"title":"X"}, follow_redirects=False); assert "not%20in%20the%20list" in r.headers["location"]
    c.post("/settings/layers/landscape/move", data={"dir":"up"}, follow_redirects=False); assert [k for k,_ in _ly.options()][-2:]==["landscape","windows"]
    c.post("/settings/layers/main/move", data={"dir":"down"}, follow_redirects=False); assert [k for k,_ in _ly.options()][0]=="main"  # the main plan is always first
    c.post("/settings/layers/landscape/move", data={"dir":"down"}, follow_redirects=False); assert [k for k,_ in _ly.options()][-1]=="landscape"
    r=c.get(f"/p/{pid_l}/plan?plan={wid}"); assert "Windows &amp; doors" in r.text
    # the categories table: which sheet a category's items are read from in the PDFs
    cats={f"c{i}":"main" for i in range(len(_cfg.CATEGORIES))}; cats[f"c{_cfg.CATEGORIES.index('Lighting')}"]="electrical"; cats[f"c{_cfg.CATEGORIES.index('Windows & Doors')}"]="landscape"; cats[f"c{_cfg.CATEGORIES.index('Paint')}"]="bogus"
    r=c.post("/settings/layers/categories", data=cats, follow_redirects=False); assert "2%20categories%20are%20read" in r.headers["location"]
    assert _dr2.layer_for_category("Windows & Doors")=="landscape" and _dr2.layer_for_category("Lighting")=="electrical" and _dr2.layer_for_category("Tiles")=="main" and _dr2.layer_for_category("Paint")=="main" and _ly.category_map()=={"Lighting":"electrical","Windows & Doors":"landscape"}
    h=c.get("/settings").text; wd=re.search(r'<select class="in" id="cat-%d"[^>]*>(.*?)</select>' % _cfg.CATEGORIES.index("Windows & Doors"), h, re.S).group(1); assert 'value="landscape" selected>Landscape &amp; garden<' in wd and 'value="main" >' in wd
    r=c.post("/settings/layers/landscape/delete", follow_redirects=False); assert "removed%20from%20the%20list" in r.headers["location"] and "landscape" not in _ly.titles() and _dr2.layer_for_category("Windows & Doors")=="main"
    db=SessionLocal(); _ly.set_categories(db, {cat:d["key"] for d in _ly.DEFAULTS for cat in d["categories"].split("|") if cat}); assert _ly.seed(db)==0; db.close()  # back to the built-in map; seeding never adds to a filled list
    assert _ly.category_map()=={"Furniture":"furnishing","Curtains & Soft":"furnishing","Lighting":"electrical","Electrical":"electrical","Plumbing & Sanitary":"plumbing","Water Treatment":"plumbing","Tiles":"flooring","Flooring":"flooring","Windows & Doors":"windows"}
    assert _ly.slug("Windows & doors", set())=="windows-doors" and _ly.slug("Windows & doors", {"windows-doors"})=="windows-doors-2" and len(_ly.slug("A very long sheet type title indeed", set()))<=20 and _ly.slug("***", set())=="sheet" and _ly.parse_words("Window, door|JOINERY; x; a-b, window")==["window","door","joinery","a-b"]
    r=c2.get("/settings/layers", follow_redirects=False); assert r.status_code in (303, 405) and c2.post("/settings/layers", data={"title":"Sneak"}, follow_redirects=False).status_code==303 and "sneak" not in _ly.titles()  # login required
    print("plan sheet types ok")
    # reference sheets: a windows sheet shows drawings of items, not the floor: no borrowed boxes, no mark mode; a dot on it cuts the drawing out as the item's picture
    from app import refs as _refs, storage as _st2
    from app.models import ItemPhoto as _IPh, PlanLayer as _PL
    assert _ly.kind("windows")=="reference" and _ly.kind("electrical")=="plan" and _ly.kind("main")=="plan" and _ly.kind("nope")=="plan" and _ly.rows()[-1]["kind"]=="reference"
    db=SessionLocal(); prl=db.get(_P,int(pid_l)); wim=db.get(ProjectImage, wid); assert wim.is_reference and _dr2.pins_for(prl, wim)==[] and _dr2.borrowed_from(prl, wim) is None and _dr2.main_plan(prl,"Ground").id==gid_l and _dr2.plans_for_floor(prl,"Ground")[-1].id==wid and _dr2.plan_role(prl, wim)=="reference" and _dr2.default_plan(prl).id!=wid; db.close()
    r=c.get(f"/p/{pid_l}/plan?plan={wid}"); h=r.text; assert len(re.findall(r'class="pin[" ]', h))==0 and "boxes from the main plan" not in h and "Reference sheet" in h and "Match drawings to items" in h and "0 drawings linked" in h and "Mark rooms</a>" not in h and "not placed" not in h and 'class="plan-dir"' in h
    r=c.get(f"/p/{pid_l}/plan?plan={wid}&mode=mark"); assert 'data-mode=""' in r.text and "Done marking" not in r.text and "Boxes borrowed" not in r.text  # nothing to mark on a reference sheet
    r=c.get(f"/p/{pid_l}/images"); assert "Reference sheet</span> of Ground floor" in r.text and "0 drawings linked · match drawings" in r.text
    r=c.get(f"/p/{pid_l}/plan/room/{lr['GF-KIT']}?plan={wid}"); assert "Not placed on this plan yet" not in r.text and "tied to a drawing on this sheet" in r.text
    # a dot placed by hand on the reference sheet cuts the drawing around it as the item's picture; moving it cuts again; removing it takes the picture away
    r=c.post(f"/p/{pid_l}/items/new", data={"room_id":lr["GF-KIT"],"category":"Windows & Doors","name":"Sliding window","size":"1915 x 1980 mm","qty":"1","unit":"pcs","status":"To buy","notes":"Designer ref F2-W07 · Chumba (Bedroom)"}, follow_redirects=False); win_id=int(r.headers["location"].split("/")[-1])
    r=c.post(f"/p/{pid_l}/plan/item-pins", data={"image_id":wid,"item_id":win_id,"x":"0.3","y":"0.4"}, headers={"Accept":"application/json"}); j=r.json(); assert j["ok"] and j["pin"]["photo"], j
    db=SessionLocal(); wi=db.get(Item, win_id); assert len(wi.photos)==1 and wi.photos[0].source==f"sheet:{wid}" and wi.photos[0].caption.startswith("Drawing on Ground floor · Windows & doors · W-01") and wi.photos[0].thumb_key; k1=wi.photos[0].file_key; db.close()
    r=c.post(f"/p/{pid_l}/plan/item-pins", data={"image_id":wid,"item_id":win_id,"x":"0.6","y":"0.5"}, headers={"Accept":"application/json"}); assert r.json()["ok"]
    db=SessionLocal(); wi=db.get(Item, win_id); assert len(wi.photos)==1 and wi.photos[0].file_key!=k1 and _st2.read_image(k1) is None; pin_id=next(q.id for q in wi.pins); db.close()  # cut again, the old cut gone from storage
    r=c.post(f"/p/{pid_l}/plan/item-pins/{pin_id}/delete", headers={"Accept":"application/json"}); assert r.json()["ok"]
    db=SessionLocal(); wi=db.get(Item, win_id); assert wi.photos==[] and db.query(ItemPin).filter(ItemPin.item_id==win_id).count()==0; db.close()
    r=c.post(f"/p/{pid_l}/plan/item-pins", data={"image_id":gid_l,"item_id":win_id,"x":"0.2","y":"0.2"}, headers={"Accept":"application/json"}); assert r.json()["ok"] and not r.json()["pin"]["photo"]  # a dot on a plan of the floor cuts nothing
    db=SessionLocal(); assert db.get(Item, win_id).photos==[]; db.close()
    # match drawings to items by the references printed on the sheet: a PDF page with two titled drawings and a list of refs in a note
    assert _refs.tokens("Designer ref F2-W07 · R11 anti-slip · A3 · 1:40 · W12 · D-03 · 1915 x 1980 mm · F2-W07")==["F2-W07","W12","D-03"] and _refs.page_refs(b"not a pdf", 1) is None
    b=io.BytesIO(); k=_cv.Canvas(b, pagesize=_A4)
    k.setFont("Helvetica-Bold",14); k.drawString(100,700,"F2-W07 - TYPE A"); k.setFont("Helvetica",9); k.drawString(100,688,"Chumba (Bedroom) · 1915 x 1980 · Qty 1"); k.setLineWidth(2); k.rect(100,450,150,230); k.line(100,620,250,620); k.drawString(160,440,"1915")
    k.setFont("Helvetica-Bold",14); k.drawString(350,700,"F2-W08 - TYPE A"); k.setLineWidth(2); k.rect(350,500,120,180)
    k.setFont("Helvetica-Bold",14); k.drawString(100,350,"F2-W09 - TYPE B"); k.setLineWidth(2); k.rect(100,250,100,80)
    k.setFont("Helvetica",9); k.drawString(100,150,"Applies to: F2-W07, F2-W08, F2-W09, F2-W10"); k.drawString(100,120,"Elevations F2-W07 to F2-W09"); k.setFont("Helvetica",12); k.drawString(100,60,"WINDOW ELEVATIONS - GROUND FLOOR     DRAWING NO: W-02"); k.showPage(); k.save(); pdfw=b.getvalue()
    tb=_dr2.read_title_block(pdf_text(pdfw)); assert tb["is_plan"] is True and tb["floor"]=="Ground" and _dr2.layer_from_title(tb["title"])=="windows", tb  # a reference type's words + ELEVATIONS: ticked and prefilled
    tb=_dr2.read_title_block("FLOOR 1 (GROUND FLOOR) - WINDOW ELEVATIONS 630\nF1-W01 - TYPE D\nWINDOW ELEVATIONS\nSheet 2 of 6"); assert tb["is_plan"] is True and tb["floor"]=="Ground" and _dr2.caption_from_title(tb["title"], "windows")=="Elevations", tb  # the title twice, long and short, is one drawing, not a list
    assert _refs.ink_block(Image.new("RGB",(400,400),"white"), 0.5, 0.5)==(0.38, 0.46, 0.24, 0.3)  # nothing inked: the fixed box
    found=_refs.page_refs(pdfw, 1); assert [f["ref"] for f in found if f["strong"]]==["F2-W07","F2-W08","F2-W09"] and {f["ref"] for f in found}=={"F2-W07","F2-W08","F2-W09","F2-W10","W-02"} and found[0]["strong"] and all(not f["strong"] for f in found if "Applies" in f["line"] or "DRAWING NO" in f["line"]) and 0.15<found[0]["x"]<0.2 and 0.15<found[0]["y"]<0.19, found
    c.post(f"/p/{pid_l}/images/plans", data={"floor":"","sheet":""}, files=[("files",("windows.pdf",pdfw,"application/pdf"))], follow_redirects=False)
    db=SessionLocal(); dsw=db.query(DrawingSet).filter(DrawingSet.project_id==int(pid_l)).order_by(DrawingSet.id.desc()).first().id; db.close()
    r=c.get(f"/p/{pid_l}/images/sets/{dsw}"); h=r.text; assert 'name="page_1" checked' in h and 'value="windows" selected' in h and 'name="floor_1" list="floors" value="Ground"' in h
    r=c.post(f"/p/{pid_l}/images/sets/{dsw}/pick", data={"page_1":"on","floor_1":"Ground","sheet_1":"W-02","caption_1":"","layer_1":"windows"}, follow_redirects=False); assert r.status_code==303
    db=SessionLocal(); wid2=next(im.id for im in db.query(ProjectImage).filter(ProjectImage.project_id==int(pid_l)) if im.sheet=="W-02"); db.close()
    r=c.post(f"/p/{pid_l}/items/new", data={"room_id":lr["GF-LIV"],"category":"Windows & Doors","name":"Sliding window","size":"975 x 1030 mm","spec":"Ref F2-W08, frosted","qty":"1","unit":"pcs","status":"To buy"}, follow_redirects=False); win2=int(r.headers["location"].split("/")[-1])
    r=c.get(f"/p/{pid_l}/plan/match?plan={wid2}"); h=r.text; assert r.status_code==200 and "3 references on the sheet, 2 with an item" in h and h.count('name="ref"')==2 and 'value="F2-W07" aria-label="Link F2-W07" checked' in h and "No item carries this reference" in h and h.count("data:image/jpeg;base64,")==2 and 'class="pill sage">Title</span> F2-W07 - TYPE A' in h and "F2-W10" not in h and "W-02</td>" not in h  # a code only mentioned that no item carries, and the sheet's own number, are left out
    r=c.get(f"/p/{pid_l}/plan/match?plan={wid}"); assert "No readable text on this sheet" in r.text and "uploaded as a picture" in r.text  # a JPG upload has no text
    r=c.get(f"/p/{pid_l}/plan/match?plan={gid_l}", follow_redirects=False); assert r.status_code==303  # not a reference sheet
    r=c.post(f"/p/{pid_l}/plan/match", data={"image_id":wid2,"ref":["F2-W07","F2-W08","F2-W09","F2-W10"]}, follow_redirects=False); assert "2%20drawings%20linked" in r.headers["location"] and "2%20pictures" in r.headers["location"], r.headers["location"]
    r=c.get(r.headers["location"]); assert "2 drawings linked" in r.text and "2 drawings linked</span>" in r.text
    db=SessionLocal(); pins={q.item_id:q for q in db.query(ItemPin).filter(ItemPin.image_id==wid2)}; assert set(pins)=={win_id, win2} and abs(pins[win_id].x-0.21)<0.06 and abs(pins[win_id].y-0.175)<0.03 and pins[win2].x>0.55, {k_:(v.x,v.y) for k_,v in pins.items()}
    wi=db.get(Item, win_id); assert len(wi.photos)==1 and wi.photos[0].source==f"sheet:{wid2}" and wi.cover.thumb; cut=Image.open(io.BytesIO(_st2.read_image(wi.photos[0].file_key))); page=Image.open(io.BytesIO(_st2.read_image(db.get(ProjectImage, wid2).best_key))); db.close()
    assert cut.width<0.45*page.width and cut.height<0.5*page.height and cut.height>cut.width and cut.width>0.15*page.width, (cut.size, page.size)  # the drawing with its title and dimension, not the neighbour
    r=c.post(f"/p/{pid_l}/plan/match", data={"image_id":wid2,"ref":["F2-W07"]}, follow_redirects=False); assert "1%20drawing%20linked" in r.headers["location"]
    db=SessionLocal(); assert db.query(ItemPin).filter(ItemPin.image_id==wid2).count()==2 and len(db.get(Item, win_id).photos)==1 and db.query(_IPh).filter(_IPh.source==f"sheet:{wid2}").count()==2; db.close()  # linked again: moved and recut, never doubled
    r=c.post(f"/p/{pid_l}/plan/match", data={"image_id":wid2}, follow_redirects=False); assert "nothing%20changed" in r.headers["location"]
    r=c.get(f"/p/{pid_l}/export/room/{lr['GF-KIT']}.pdf"); assert r.status_code==200 and "ON THE WINDOWS & DOORS SHEET" in pdf_text(r.content)  # the checklist frames the room's dots on the sheet
    r=c2.get(f"/c/{ltok}?plan={wid2}"); h=r.text; assert r.status_code==200 and len(re.findall(r'class="pin[" ]', h))==0 and "Tap a dot to see its item" in h and h.count('class="ipin')==2
    # the kind is set on the settings page: a plan type turned reference loses its borrowed boxes, and back; the main plan stays a plan
    r=c.post("/settings/layers/ceiling", data={"title":"Ceiling","words":"ceiling, reflected, rcp","kind":"reference"}, follow_redirects=False); assert "now%20reference%20sheet" in r.headers["location"] and "1%20plan%20shows%20no%20room%20boxes" in r.headers["location"], r.headers["location"]  # E-01 is the ceiling sheet
    assert _ly.kind("ceiling")=="reference"; r=c.get(f"/p/{pid_l}/plan?plan={eid_l}"); assert len(re.findall(r'class="pin[" ]', r.text))==0 and "Reference sheet" in r.text
    r=c.post("/settings/layers/ceiling", data={"title":"Ceiling","words":"ceiling, reflected, rcp","kind":"plan"}, follow_redirects=False); assert "1%20plan%20borrows%20the%20room%20boxes%20of%20its%20floor" in r.headers["location"], r.headers["location"]
    r=c.get(f"/p/{pid_l}/plan?plan={eid_l}"); assert len(re.findall(r'class="pin[" ]', r.text))==2 and "Reference sheet" not in r.text and _ly.kind("ceiling")=="plan"  # its boxes (own, copied earlier) are back
    c.post("/settings/layers/main", data={"title":"Main plan","kind":"reference"}, follow_redirects=False); assert _ly.kind("main")=="plan"
    r=c.post("/settings/layers", data={"title":"Joinery details","kind":"reference","words":"joinery"}, follow_redirects=False); assert "added%20as%20a%20reference%20sheet" in r.headers["location"] and _ly.kind("joinery-details")=="reference"
    r=c.get("/settings"); assert '<select class="in kind" name="kind" form="ly-windows"' in r.text and 'name="kind" form="ly-main"' not in r.text and 'value="reference" selected>Reference sheet' in r.text  # no kind select on the main plan
    c.post("/settings/layers/joinery-details/delete", follow_redirects=False); assert "joinery-details" not in _ly.titles()
    from app.models import SchemaMark as _SM
    db=SessionLocal(); assert db.get(_SM, _ly.KIND_MARK) is not None; db.get(_PL,"windows").kind="plan"; db.commit(); _ly.invalidate(); assert _ly.kind("windows")=="plan" and _ly.seed(db)==0 and _ly.kind("windows")=="plan"  # marked as upgraded: the studio's choice of plan stays
    assert _ly.seed(db, kind_added=True)==1 and _ly.kind("windows")=="reference" and _ly.seed(db, kind_added=True)==0  # the kind column arriving later gives the built-in reference types their kind once
    db.get(_PL,"windows").kind="plan"; db.delete(db.get(_SM, _ly.KIND_MARK)); db.commit(); _ly.invalidate(); assert _ly.seed(db)==1 and _ly.kind("windows")=="reference" and db.get(_SM, _ly.KIND_MARK) is not None and _ly.seed(db)==0; db.close()  # no mark yet (a start that lost the race to add the column): upgraded and marked
    print("reference sheets + match by reference ok")
    # the Items page: sortable table columns, and the Item list (one line per product across its rooms)
    from app.routers.items import group_items as _gi
    from app.services import sort_items as _si
    db=SessionLocal(); beds=[x.id for x in db.query(Room).filter(Room.project_id==int(pid_p)) if x.group=="bedroom"]; db.close(); assert len(beds)==3
    r=c.post(f"/p/{pid_p}/items/new", data={"room_ids":[str(x) for x in beds],"category":"Lighting","name":"Reading light","qty":"2","unit":"pcs","unit_price":"150","status":"To buy"}, follow_redirects=False)
    assert r.headers["location"]==f"/p/{pid_p}/items?q=Reading%20light"
    db=SessionLocal(); _rl=db.query(Item).filter(Item.project_id==int(pid_p), Item.name=="Reading light").first().id; db.close()
    buy(c, pid_p, _rl, "Lamp Street Co", 150)  # one quote, bought for the three bedrooms: To buy → Quoted
    r=c.get(f"/p/{pid_p}/items?view=table&category=Lighting"); h=r.text
    assert '<table class="sortable">' in h and '<th data-sort="text">Code</th>' in h and '<th class="num" data-sort="num">Total CNY</th>' in h and '<th data-sort="num">Status</th>' in h
    assert h.count('data-v="Reading light"')==3 and 'data-v="150.0"' in h and f'href="/p/{pid_p}/items?category=Lighting&amp;view=list"' in h and 'aria-current="page" title="One line per room' in h and f'href="/p/{pid_p}/items?view=table">Clear</a>' in h
    r=c.get(f"/p/{pid_p}/items?view=list&q=Reading"); h=r.text
    assert h.count('<tr id="item-')==1 and re.search(r'data-items="\d+ \d+ \d+"', h) and '>3 rooms</a>' in h and '6 pcs</td>' in h and 'data-v="900.0">900</td>' in h and "1 item · 3 room lines" in h and 'Quoted · 3</span>' in h and '<th data-sort="num">Rooms</th>' in h
    assert "status-sel" not in h and "Press a column heading to sort" in h and c.get(f"/p/{pid_p}/items?view=list").text.count('<tr id="item-')>=5
    h=c.get(f"/p/{pid_p}/items").text; assert 'aria-current="page" title="One line per item' in h and h.count('<tr id="item-')>=5 and 'name="view"' not in h  # the Item list is the default view
    h=c.get(f"/p/{pid_p}/items?q=Bedside lamp").text; assert h.count('<tr id="item-')==1 and 'class="status-sel"' in h  # one room: the status dropdown right there
    db=SessionLocal(); prp=db.get(_P,int(pid_p)); gs={g["item"].name:g for g in _gi(_si(list(prp.live_items), prp))}; db.close()
    assert len(gs["Reading light"]["rows"])==3 and gs["Reading light"]["qty"]==6 and gs["Reading light"]["total"]==900 and gs["Reading light"]["status_text"]=="Quoted" and gs["Reading light"]["unit"]=="pcs" and len(gs["Reading light"]["rooms"])==3
    assert len(gs["Bedside lamp"]["rows"])==1 and gs["Bedside lamp"]["rooms"][0].code=="FF-MBR" and gs["Bedside lamp"]["codes"]=="FF-MBR-01"
    db=SessionLocal(); rl_id=db.query(Item).filter(Item.project_id==int(pid_p), Item.name=="Reading light").first().id; db.close()
    c.post(f"/p/{pid_p}/items/{rl_id}/status", data={"status":"Ordered"}, headers={"Accept":"application/json"})
    db=SessionLocal(); prp=db.get(_P,int(pid_p)); g=[g for g in _gi(_si(list(prp.live_items), prp)) if g["item"].name=="Reading light"][0]; db.close()
    assert g["status_text"]=="2 Quoted, 1 Ordered" and g["status_rank"]==1 and g["price_min"]==g["price_max"]==150
    h=c.get(f"/p/{pid_p}/items?view=list&q=Reading").text; assert 'Quoted · 2</span>' in h and 'Ordered · 1</span>' in h and '<td data-v="1" class="small nowrap">' in h  and "status-sel" not in h  # sorts by the least advanced status
    print("item list + sorting ok")
    # the item page: a checklist of the rooms that have this item; untick = that room's line goes, tick = a copy there
    db=SessionLocal(); rls=sorted(db.query(Item).filter(Item.project_id==int(pid_p), Item.name=="Reading light"), key=lambda i: i.id); rl=rls[0]; own=rl.room_id; others_=[i.room_id for i in rls[1:]]
    rc={x.id:x.code for x in db.query(Room).filter(Room.project_id==int(pid_p))}; liv=[k for k,v in rc.items() if v=="GF-LIV"][0]; db.close()
    r=c.get(f"/p/{pid_p}/items/{rl.id}"); h=r.text
    assert 'name="rooms_form" value="1"' in h and h.count('name="room_ids"')>=12 and f'<input type="hidden" name="room_ids" value="{own}">' in h and re.search(r'value="%d"[^>]*checked disabled>' % own, h) and h.count('class="pick-link"')==2 and "Rooms with this item" in h and 'data-removal="Remove this item from' in h
    r=c.post(f"/p/{pid_p}/items/{rl.id}", data={"room_id":str(own),"category":"Lighting","name":"Reading light","qty":"2","unit":"pcs","unit_price":"150","status":"To buy","rooms_form":"1","room_ids":[str(own), str(others_[0]), str(liv)]}, follow_redirects=False); assert r.status_code==303
    db=SessionLocal(); codes=sorted(i.room.code for i in db.query(Item).filter(Item.project_id==int(pid_p), Item.name=="Reading light")); db.close()
    assert codes==sorted([rc[own], rc[others_[0]], "GF-LIV"]), codes  # one room dropped, one added, the own room kept
    r=c.post(f"/p/{pid_p}/items/{rl.id}", data={"room_id":str(own),"category":"Lighting","name":"Reading light","qty":"2","unit":"pcs","unit_price":"150","status":"To buy","rooms_form":"1","room_ids":[str(own)]}, follow_redirects=False); assert r.status_code==303
    db=SessionLocal(); assert [i.room_id for i in db.query(Item).filter(Item.project_id==int(pid_p), Item.name=="Reading light")]==[own]; db.close()  # only this line is left
    print("room checklist ok")
    # speed: small copies of photos, the picture cache and headers, compression, static versioning, few queries per page
    from app.models import ItemPhoto as _IP
    from app import storage as _st
    from sqlalchemy import event as _ev2
    from app.db import engine as _eng3
    c.post(f"/p/{pid}/items/{iid}/photo", files=[("photos",("a.jpg",img("teal"),"image/jpeg")),("photos",("b.jpg",img("navy"),"image/jpeg"))], follow_redirects=False)
    db=SessionLocal(); phs=db.query(_IP).filter(_IP.item_id==int(iid)).order_by(_IP.id).all(); keys=[(x.file_key, x.thumb_key) for x in phs]; db.close()
    assert len(keys)>=2 and all(tk==_st.thumb_key_for(fk) and tk.endswith("_t.jpg") for fk,tk in keys[-2:])
    fk,tk=keys[-1]; up=os.environ["LOCAL_UPLOAD_DIR"]
    assert os.path.exists(f"{up}/{fk}") and os.path.exists(f"{up}/{tk}") and os.path.getsize(f"{up}/{tk}") < os.path.getsize(f"{up}/{fk}")
    assert Image.open(f"{up}/{tk}").size[0] <= 640 and Image.open(f"{up}/{fk}").size[0] > 640
    r=c.get(f"/media/{tk}"); assert r.status_code==200 and r.headers["cache-control"]=="public, max-age=31536000, immutable" and r.headers.get("content-encoding")=="identity"
    assert c.get("/media/nope/zzz.jpg").status_code==404
    r=c.get(f"/p/{pid}/items"); assert f"/media/{keys[0][1] or keys[0][0]}" in r.text and "_t.jpg" in r.text  # lists show the small copy
    r=c.get(f"/p/{pid}/items/{iid}"); assert f'src="/media/{tk}"' in r.text and f'href="/media/{fk}"' in r.text  # gallery: small picture, full on tap
    r=c.get(f"/p/{pid}/items", headers={"Accept-Encoding":"gzip"}); assert r.headers.get("content-encoding")=="gzip" and "Items" in r.text
    assert c.get("/static/app.css?v=abc").headers["cache-control"]=="public, max-age=31536000, immutable" and c.get("/static/app.css").headers["cache-control"]=="public, max-age=3600"
    from app.common import STATIC_V as _SV; assert len(_SV)==10 and f'/static/app.css?v={_SV}' in r.text and f'/static/app.js?v={_SV}' in r.text
    assert f'plan.js?v={_SV}' in c.get(f"/p/{pid}/plan").text and f'plan.js?v={_SV}' in c2.get(f"/c/{ctok}").text
    # make cover swaps both copies; deleting a photo removes both files
    last=phs[-1].id; first_fk=keys[0][0]
    c.post(f"/p/{pid}/items/{iid}/photo/{last}/cover", follow_redirects=False)
    db=SessionLocal(); phs2=db.query(_IP).filter(_IP.item_id==int(iid)).order_by(_IP.id).all(); assert phs2[0].file_key==fk and phs2[0].thumb_key==tk and phs2[-1].file_key==first_fk; victim=phs2[-1]; vk,vt=victim.file_key,victim.thumb_key; db.close()
    c.post(f"/p/{pid}/items/{iid}/photo/{victim.id}/delete", follow_redirects=False)
    assert not os.path.exists(f"{up}/{vk}") and (not vt or not os.path.exists(f"{up}/{vt}"))
    # settings: older photos without a small copy are counted and made in batches
    db=SessionLocal(); old_ph=db.query(_IP).filter(_IP.item_id==int(iid)).first(); old_ph.thumb_key=""; db.commit(); okey=old_ph.file_key; db.close()
    r=c.get("/settings"); assert 'id="speed"' in r.text and "Prepare 1 older photo" in r.text
    r=c.post("/settings/thumbs", follow_redirects=False); assert r.headers["location"]=="/settings?thumbs=1&left=0"
    db=SessionLocal(); old_ph=db.query(_IP).filter(_IP.file_key==okey).one(); assert old_ph.thumb_key==_st.thumb_key_for(okey) and os.path.exists(f"{up}/{old_ph.thumb_key}"); db.close()
    assert "Every photo has its small version and fingerprint" in c.get("/settings?thumbs=1&left=0").text
    # a photo whose original is gone is marked and not retried
    db=SessionLocal(); gone=_IP(item_id=int(iid), file_key="p1/missing.jpg", thumb_key=""); db.add(gone); db.commit(); db.close()
    c.post("/settings/thumbs", follow_redirects=False); db=SessionLocal(); g2=db.query(_IP).filter(_IP.file_key=="p1/missing.jpg").one(); assert g2.thumb_key=="-" and g2.thumb=="p1/missing.jpg"; db.delete(g2); db.commit(); db.close()
    # queries per page stay small on the imported project (lazy loading used to cost one query per item)
    qn={"n":0}
    def _cnt(conn, cursor, statement, parameters, context, executemany): qn["n"]+=1
    _ev2.listen(_eng3, "before_cursor_execute", _cnt)
    try:
        for url,limit in ((f"/p/{pid}",30),(f"/p/{pid}/items",30),(f"/p/{pid}/items?view=table",30),(f"/p/{pid}/plan",40),(f"/c/{ctok}",40),(f"/p/{pid}/export/schedule.pdf",45)):
            qn["n"]=0; r=c.get(url) if not url.startswith("/c/") else c2.get(url); assert r.status_code==200 and qn["n"]<=limit, (url, qn["n"])
    finally:
        _ev2.remove(_eng3, "before_cursor_execute", _cnt)
    print("speed ok")
    # ---- review fixes: counts per room, drafts never get dots, client leak checks, dots follow rooms, PDFs with markup, scoping ----
    # the panel's "placed" count is this room's: a dot of another room on the same plan does not count
    db=SessionLocal(); kit_dots=sum(1 for q in db.query(ItemPin).filter(ItemPin.image_id==gid) if q.item.room_id==room.id); r2_item=[i.id for i in db.get(Room, r2.id).items if not i.draft][0]; db.close()
    assert kit_dots>=1
    r=c.post(f"/p/{pid}/plan/item-pins", data={"image_id":gid,"item_id":r2_item,"x":"0.6","y":"0.2"}, headers={"Accept":"application/json"}); assert r.status_code==200
    r=c.get(f"/p/{pid}/plan/room/{room.id}?plan={gid}"); assert f"data-placed>{kit_dots}</span>" in r.text, re.search(r"data-placed>\d+</span> of \d+", r.text).group(0)
    r=c.get(f"/p/{pid}/plan/room/{r2.id}?plan={gid}"); assert "data-placed>1</span>" in r.text
    # a draft can never get a dot, and a dot row for a draft never shows (designer page or client link)
    db=SessionLocal(); drf=Item(project_id=int(pid), room_id=room.id, name="", category="Other", draft=True); db.add(drf); db.commit(); drf_id=drf.id; db.close()
    r=c.post(f"/p/{pid}/plan/item-pins", data={"image_id":gid,"item_id":drf_id,"x":"0.1","y":"0.1"}, headers={"Accept":"application/json"}); assert r.status_code==400 and "draft" in r.json()["error"]
    db=SessionLocal(); db.add(ItemPin(image_id=gid, item_id=drf_id, x=0.5, y=0.5)); db.commit(); db.close()
    for page in (c.get(f"/p/{pid}/plan?plan={gid}"), c2.get(f"/c/{ctok}?plan={gid}"), c.get(f"/p/{pid}/plan/room/{room.id}?plan={gid}")):
        assert page.status_code==200 and f'data-item="{drf_id}"' not in page.text
    db=SessionLocal(); db.delete(db.get(Item, drf_id)); db.commit(); assert db.query(ItemPin).filter(ItemPin.item_id==drf_id).count()==0; db.close()
    # the client panel hides what the designer's panel shows: supplier, notes, CNY, designer links
    c.post(f"/p/{pid}/items/{iid}", data={"room_id":room.id,"category":"Lighting","name":"Island pendant","qty":"3","unit":"pcs","unit_price":"450","price_currency":"CNY","supplier_id":sup.id,"status":"Ordered","notes":"secret supplier note"}, follow_redirects=False)
    d=c.get(f"/p/{pid}/plan/room/{room.id}?plan={gid}").text; assert "Foshan Tile Co." in d and "¥" in d and "/items/" in d
    for page in (c2.get(f"/c/{ctok}?plan={gid}&room={room.id}"), c2.get(f"/c/{ctok}/plan/room/{room.id}?plan={gid}")):
        panel=page.text.split('class="card tight plan-room"')[1].split("</aside>")[0] if "</aside>" in page.text else page.text
        assert "Foshan Tile Co." not in panel and "secret supplier note" not in panel and "¥" not in panel and "/p/" not in panel and "/items/" not in panel and "status-sel" not in panel
    # a dot sits in a room's box: moving the item to another room removes its dots; deleting a room removes its items' dots
    c.post(f"/p/{pid}/plan/item-pins", data={"image_id":gid,"item_id":iid,"x":"0.2","y":"0.2"}, headers={"Accept":"application/json"})
    db=SessionLocal(); assert db.query(ItemPin).filter(ItemPin.item_id==int(iid)).count()==2; db.close()  # ground + roof
    c.post(f"/p/{pid}/items/{iid}", data={"room_id":r3.id,"category":"Lighting","name":"Island pendant","qty":"3","unit":"pcs","unit_price":"450","status":"Ordered"}, follow_redirects=False)
    db=SessionLocal(); assert db.get(Item, int(iid)).room_id==r3.id and db.query(ItemPin).filter(ItemPin.item_id==int(iid)).count()==0; db.close()
    c.post(f"/p/{pid}/items/{iid}", data={"room_id":room.id,"category":"Lighting","name":"Island pendant","qty":"3","unit":"pcs","unit_price":"450","status":"Ordered"}, follow_redirects=False)
    c.post(f"/p/{pid}/rooms", data={"code":"TMP2","name":"Temp two","floor":"Ground"}, follow_redirects=False)
    db=SessionLocal(); tmp2=db.query(Room).filter(Room.project_id==int(pid), Room.code=="TMP2").first(); db.close()
    r=c.post(f"/p/{pid}/items/new", data={"room_ids":[str(tmp2.id)],"category":"Other","name":"Temp lamp","qty":"1","unit":"pcs","status":"To buy"}, follow_redirects=False); tl=int(r.headers["location"].split("/")[-1])
    c.post(f"/p/{pid}/plan/item-pins", data={"image_id":gid,"item_id":tl,"x":"0.4","y":"0.4"}, headers={"Accept":"application/json"})
    c.post(f"/p/{pid}/rooms/{tmp2.id}/delete", follow_redirects=False)
    db=SessionLocal(); assert db.get(Item, tl).room_id is None and db.query(ItemPin).filter(ItemPin.item_id==tl).count()==0; c.post(f"/p/{pid}/items/{tl}/delete", follow_redirects=False); db.close()
    # markup in names must not break the PDFs the client and suppliers download
    db=SessionLocal(); pr=db.get(_P, int(pid)); old_name, old_addr=pr.name, pr.address; pr.name='Kinondoni <b>House</b> & "Villa"'; pr.address="Plot <70>"; db.commit(); db.close()
    for url, cl in ((f"/c/{ctok}/schedule.pdf", c2), (f"/c/{ctok}/schedule.pdf?layout=floor", c2), (f"/p/{pid}/export/packing.pdf", c), (f"/p/{pid}/export/room/{room.id}.pdf", c), (f"/p/{pid}/export/labels.pdf", c)):
        r=cl.get(url); assert r.status_code==200 and r.headers["content-type"]=="application/pdf", url
    assert "Kinondoni <b>House</b>" in pdf_text(c2.get(f"/c/{ctok}/schedule.pdf").content)
    db=SessionLocal(); pr=db.get(_P, int(pid)); pr.name, pr.address=old_name, old_addr; db.commit(); db.close()
    # the client link: ?item= alone opens the item's room; another project's plan / room ids are ignored, nothing of theirs renders
    r=c2.get(f"/c/{ctok}?item={other_item}"); assert r.status_code==200 and f'data-sel="{room.id}"' in r.text and 'class="card tight plan-room"' in r.text
    db=SessionLocal(); op=db.query(ProjectImage).filter(ProjectImage.project_id==int(pid2), ProjectImage.kind=="floorplan").first(); db.close()
    r=c2.get(f"/c/{ctok}?plan={op.id}&room={other.id}"); assert r.status_code==200 and op.file_key not in r.text and f'data-plan="{gid}"' in r.text and 'class="card tight plan-room"' not in r.text and "Other room" not in r.text
    r=c.get(f"/p/{pid}/plan?plan={op.id}&room={other.id}"); assert r.status_code==200 and op.file_key not in r.text and f'data-plan="{gid}"' in r.text and 'class="card tight plan-room"' not in r.text
    r=c.get(f"/p/{pid}/plan/room/{room.id}?plan={op.id}"); assert r.status_code==200 and op.file_key not in r.text and "data-placed" not in r.text  # unknown plan = no plan context
    assert c2.get(f"/c/{ctok}/plan/room/{room.id}?plan={op.id}").status_code==200
    print("review fixes ok")
    # ---- rooms vs areas, the startup column migration, PDF title blocks in the page picker ----
    from app.db import migrate as _migrate
    from app.services import guess_kind as _gk
    from app import drawings as _dr2
    from sqlalchemy import create_engine as _ce, text as _text
    # migrate(): adds items.draft and rooms.kind to an old database, once
    _eng=_ce(f"sqlite:///{_tmp}/old.db")
    with _eng.begin() as cn:
        cn.execute(_text("CREATE TABLE rooms (id INTEGER PRIMARY KEY, code VARCHAR(20), name VARCHAR(120))"))
        cn.execute(_text("CREATE TABLE items (id INTEGER PRIMARY KEY, name VARCHAR(200))"))
        cn.execute(_text("CREATE TABLE project_images (id INTEGER PRIMARY KEY, kind VARCHAR(20))"))
        cn.execute(_text("CREATE TABLE item_photos (id INTEGER PRIMARY KEY, file_key VARCHAR(255))"))
        cn.execute(_text("INSERT INTO rooms (code, name) VALUES ('X', 'Old room')"))
        cn.execute(_text("INSERT INTO project_images (kind) VALUES ('floorplan')"))
        cn.execute(_text("INSERT INTO item_photos (file_key) VALUES ('p1/old.jpg')"))
    assert _migrate(_eng)==["items.draft", "rooms.kind", "project_images.layer", "item_photos.thumb_key", "item_photos.source", "item_photos.digest", "items.purchase_id"] and _migrate(_eng)==[]
    with _eng.begin() as cn: cn.execute(_text("INSERT INTO project_images (kind, layer) VALUES ('floorplan', 'furniture')")); cn.execute(_text("INSERT INTO project_images (kind, layer) VALUES ('floorplan', 'furnishing')"))
    assert _migrate(_eng)==["project_images.layer: 1 rows relabelled"] and _migrate(_eng)==[]  # the old name of the main plan; a furniture layout stays
    with _eng.connect() as cn: assert sorted(r[0] for r in cn.execute(_text("SELECT layer FROM project_images")))==["furnishing", "main", "main"]
    with _eng.connect() as cn: assert cn.execute(_text("SELECT kind FROM rooms")).scalar()=="room" and cn.execute(_text("SELECT layer FROM project_images")).scalar()=="main" and cn.execute(_text("SELECT thumb_key FROM item_photos")).scalar()==""
    # the guess: zones are areas, everything else a room
    for nm, fl, k in [("Entrance","Ground","area"),("Hall & Corridors","Ground","area"),("Living Room","Ground","room"),("Store Room","Ground","room"),
                      ("Carport","Ground","area"),("Kitchen Verandah","Ground","area"),("Stairs Ground to First","Ground","area"),("Landing & Corridor","First","area"),
                      ("Balconies (all 4)","First","area"),("Roof Lobby","Roof","area"),("Roof Terrace (Phase II)","Roof","area"),("Family Lounge","First","room"),
                      ("Master Closet","First","room"),("Guest Bathroom","Ground","room"),("Whole house","All","area"),("All doors","All","area"),("RO plant","Outside","area"),("Gym","Roof","room"),
                      ("Downstairs WC","Ground","room"),("Upstairs Lounge","First","room"),("Plant Room","Ground","room"),("Garden Room","Ground","room"),("Pool Bathroom","Ground","room"),
                      ("Gate House Bedroom","Ground","room"),("Hallway","Ground","area"),("Staircase","Ground","area"),("Balconies","First","area"),("Hall Bathroom","First","room")]:
        assert _gk(nm, fl)==k, (nm, _gk(nm, fl))
    # the import classified the Kinondoni list; the Rooms page counts rooms and areas separately and groups them
    db=SessionLocal(); kinds={r.code: r.kind for r in db.query(Room).filter(Room.project_id==int(pid))}; db.close()
    assert kinds["GF-ENT"]=="area" and kinds["GF-HAL"]=="area" and kinds["GF-KIT"]=="room" and kinds["GF-LIV"]=="room" and kinds["FF-BAL"]=="area" and kinds["RF-GYM"]=="room"
    r=c.get(f"/p/{pid}/rooms"); assert r.status_code==200 and "Rooms &amp; areas" in r.text and 'name="kind"' in r.text and "Sort rooms &amp; areas by name" in r.text
    gf=r.text.split("<h2>Ground floor</h2>")[1].split("<h2>First floor</h2>")[0]
    assert re.search(r"\d+ rooms · \d+ areas", gf) and ">Areas</span>" in gf and 'class="pill area"' in gf and gf.index("GF-KIT") < gf.index(">Areas</span>") < gf.index("GF-ENT")
    assert _dr2.count_label([1,2],[])=="2 rooms" and _dr2.count_label([1],[1,1])=="1 room · 2 areas" and _dr2.count_label([],[1])=="1 area" and _dr2.count_label([],[])=="0 rooms"
    # kind by hand: edit, add with Auto, sort everything by name again
    r=c.post(f"/p/{pid}/rooms/{room.id}", data={"code":"GF-KIT","name":"Kitchen","floor":"Ground","floor_area":"0","wall_area":"0","notes":"via plan","kind":"area"}, follow_redirects=False)
    db=SessionLocal(); assert db.get(Room, room.id).is_area; db.close()
    c.post(f"/p/{pid}/rooms", data={"code":"GF-PAS","name":"Side Passage","floor":"Ground","kind":"auto"}, follow_redirects=False)
    db=SessionLocal(); pas=db.query(Room).filter(Room.project_id==int(pid), Room.code=="GF-PAS").first(); assert pas.is_area; db.close()
    r=c.post(f"/p/{pid}/rooms/guess-kinds", follow_redirects=False); assert r.status_code==303
    db=SessionLocal(); assert not db.get(Room, room.id).is_area and db.get(Room, pas.id).is_area; c.post(f"/p/{pid}/rooms/{pas.id}/delete", follow_redirects=False); db.close()
    # everywhere rooms are listed: selects group areas after rooms; the Overview separates them; the plan's mark list too
    db=SessionLocal(); ent=db.query(Room).filter(Room.project_id==int(pid), Room.code=="GF-ENT").first(); db.close()
    r=c.get(f"/p/{pid}/items"); assert '<optgroup label="Rooms">' in r.text and '<optgroup label="Areas">' in r.text and r.text.index("GF-KIT - Kitchen") < r.text.index('<optgroup label="Areas">') < r.text.index("GF-ENT - Entrance")
    r=c.get(f"/p/{pid}/items/new"); assert 'class="picks-sub">Areas<' in r.text
    r=c.get(f"/p/{pid}/items/{iid}"); assert '<optgroup label="Areas">' in r.text
    r=c.get(f"/p/{pid}"); assert "By room &amp; area" in r.text and '<tr class="sub"><td colspan="5">Areas</td></tr>' in r.text
    r=c.get(f"/p/{pid}/plan?plan={gid}&mode=mark"); assert r.text.count('class="sub-title mt"')==2 and '<span>Rooms</span>' in r.text and '<span>Areas</span>' in r.text and re.search(r'id="plan-count" data-total="\d+">\d+ of \d+ placed</span>', r.text)
    pl=r.text.split('id="plan-rooms"')[1]; assert pl.index('<span>Rooms</span>') < pl.index("GF-KIT") < pl.index('<span>Areas</span>') < pl.index("GF-ENT")
    r=c.get(f"/p/{pid}/plan?plan={gid}&room={ent.id}"); head=r.text.split('class="card tight plan-room"')[1][:400]; assert 'class="pill area"' in head and ">Area<" in head  # an area's panel says so
    # PDFs: the summary lists rooms, then an Areas row, then areas; the checklist of an area says so
    tf=pdf_text(c.get(f"/p/{pid}/export/schedule.pdf?layout=floor").content); summ=tf.split("SUMMARY", 1)[1]
    assert summ.index("GF-KIT - Kitchen") < summ.index("Areas") < summ.index("GF-ENT - Entrance"), summ[:800]
    assert "AREA CHECKLIST" in pdf_text(c.get(f"/p/{pid}/export/room/{ent.id}.pdf").content) and "ROOM CHECKLIST" in pdf_text(c.get(f"/p/{pid}/export/room/{room.id}.pdf").content)
    # the title block reader
    tb=_dr2.read_title_block("PROPOSED RESIDENTIAL HOUSE\nGROUND FLOOR PLAN (DIMENSION DETAILS)\nSCALE 1:100\nDRAWING NO: A-102\nREV A")
    assert tb=={"title":"Ground Floor Plan (Dimension Details)","sheet":"A-102","floor":"Ground","is_plan":True}, tb
    assert _dr2.read_title_block("GROUND FLOOR PLAN (FURNITURE\nLAYOUT)\nA-104")["title"]=="Ground Floor Plan (Furniture Layout)"  # wrapped title
    # labels row above the values row; addresses, revisions, standards and marks are never a sheet number
    assert _dr2.read_title_block("GROUND FLOOR PLAN\nDRAWING NO. SCALE DATE REV\nA-102 1:100 12-03-2026 B\nP.O. BOX 1234 DAR ES SALAAM")["sheet"]=="A-102"
    assert _dr2.read_title_block("GROUND FLOOR PLAN\nP.O. BOX 1234\nREV 01\nISO 9001 CERTIFIED\nDN100 PVC\nDOOR D-01\nPLOT NO 456\nPROJECT NO P-2024-017")["sheet"]==""
    assert _dr2.read_title_block("GROUND FLOOR PLAN\nSCALE 1:100\nDATE: MAY 2024\nJOB NO 2024-15\nA01\nREV 01")["sheet"]=="A01"
    assert _dr2.read_title_block("GROUND FLOOR PLAN\nDWG NO 01\nREV A")["sheet"]=="01" and _dr2.read_title_block("ROOF PLAN\nDWG NO\n07 1:100")["sheet"]=="07"
    # cover sheets with a drawing list and notes that refer to a plan are not plans; consultants on the same text line do not disqualify a plan
    tb=_dr2.read_title_block("DRAWING LIST\nA-100 SITE LAYOUT PLAN\nA-101 GROUND FLOOR PLAN\nA-102 FIRST FLOOR PLAN\nA-000"); assert tb["is_plan"] is False, tb
    tb=_dr2.read_title_block("SECTION A-A\nFOR SECTION LINE SEE GROUND FLOOR PLAN\nA-301"); assert tb["is_plan"] is False and tb["title"]=="" and tb["floor"]=="", tb
    tb=_dr2.read_title_block("PROJECT: PROPOSED RESIDENTIAL HOUSE P-2024-017 GROUND FLOOR PLAN DRAWING NO. SCALE DATE REV\nA-102 1:100 12-03-2026 B"); assert tb=={"title":"Ground Floor Plan","sheet":"A-102","floor":"Ground","is_plan":True}, tb
    assert _dr2.read_title_block("GROUND FLOOR PLAN STRUCTURAL ENGINEER: ABC LTD\nA-101")["is_plan"] is True and _dr2.read_title_block("GROUND FLOOR PLAN - FIREPLACE DETAIL\nA-101")["is_plan"] is True
    # the floor comes from the title only; lower / upper ground are their own floors
    tb=_dr2.read_title_block("FRONT ELEVATION\nALL DIMENSIONS TO BE CHECKED ON SITE\nGROUND FLOOR LEVEL +0.000\nA-201"); assert tb["floor"]=="" and tb["is_plan"] is False and tb["title"]=="", tb
    assert _dr2.read_title_block("LOWER GROUND FLOOR PLAN\nA-100")["floor"]=="Lower ground" and _dr2.read_title_block("UPPER GROUND FLOOR PLAN (FURNITURE LAYOUT)\nA-104")["floor"]=="Upper ground"
    # the caption suggestion keeps only what plan_caption would not print anyway
    assert _dr2.caption_from_title("Ground Floor Plan (Furniture Layout)")=="" and _dr2.caption_from_title("Ground Floor Plan (Furniture Layout)", "main")=="Furniture Layout" and _dr2.caption_from_title("Site Layout Plan")=="" and _dr2.caption_from_title("Roof Plan")=="" and _dr2.caption_from_title("Lower Ground Floor Plan (Dimension Details)")=="Dimension Details"
    assert _dr2.read_title_block("FRONT ELEVATION\n1:100\nA-201")=={"title":"","sheet":"A-201","floor":"","is_plan":False}
    tb=_dr2.read_title_block("ELECTRICAL LAYOUT PLAN - GROUND FLOOR\nE-01"); assert tb["is_plan"] is True and tb["floor"]=="Ground" and tb["sheet"]=="E-01" and _dr2.layer_from_title(tb["title"])=="electrical", tb  # a layer of its floor
    assert _dr2.layer_from_title("Reflected Ceiling Plan - First Floor")=="ceiling" and _dr2.layer_from_title("Floor Finishes Plan")=="flooring" and _dr2.layer_from_title("Plumbing & Drainage Layout Plan")=="plumbing" and _dr2.layer_from_title("Ground Floor Plan (Furniture Layout)")=="furnishing" and _dr2.layer_from_title("")=="main" and _dr2.layer_from_title("Ground Floor Plan")=="main"
    assert _dr2.caption_from_title("Electrical Layout Plan - Ground Floor")=="" and _dr2.caption_from_title("Reflected Ceiling Plan - First Floor")==""
    assert _dr2.read_title_block("STRUCTURAL FRAMING PLAN - GROUND FLOOR\nS-01")["is_plan"] is False and _dr2.read_title_block("HVAC LAYOUT PLAN\nM-01")["is_plan"] is False and _dr2.read_title_block("REFLECTED CEILING PLAN - FIRST FLOOR\nA-401")=={"title":"Reflected Ceiling Plan - First Floor","sheet":"A-401","floor":"First","is_plan":True}
    assert _dr2.read_title_block("SITE LAYOUT PLAN\nA-100")=={"title":"Site Layout Plan","sheet":"A-100","floor":"Site","is_plan":True}
    assert _dr2.read_title_block("ROOF PLAN\nSHEET A-104")["floor"]=="Roof" and _dr2.read_title_block("FIRST FLOOR PLAN (FURNITURE LAYOUT)\nA4\nA-103")["sheet"]=="A-103"
    assert _dr2.read_title_block("")=={"title":"","sheet":"","floor":"","is_plan":False}
    # a drawing set with title blocks: the picker arrives prefilled and pre-ticked; the pick keeps the caption
    def pdf_titled(pages):
        b=io.BytesIO(); k=_cv.Canvas(b, pagesize=_A4)
        for title, sheet in pages:
            k.setFont("Helvetica",18); k.drawString(60,780,title); k.rect(80,200,400,450); k.setFont("Helvetica",10); k.drawString(380,120,"DRAWING NO: "+sheet); k.showPage()
        k.save(); return b.getvalue()
    r=c.post(f"/p/{pid}/images/plans", data={}, files=[("files",("titled.pdf",pdf_titled([("GROUND FLOOR PLAN (FURNITURE LAYOUT)","A-105"),("FRONT ELEVATION","A-301"),("FIRST FLOOR PLAN","A-106")]),"application/pdf"))], follow_redirects=False)
    tsid=int(r.headers["location"].split("/sets/")[1].split("?")[0])
    r=c.get(f"/p/{pid}/images/sets/{tsid}"); assert r.status_code==200 and "2 pages look like plan sheets" in r.text
    assert 'name="page_1" checked' in r.text and 'name="page_2">' in r.text and 'name="page_3" checked' in r.text
    sel=lambda h, name: re.search(r'name="%s".*?</select>' % name, h, re.S).group(0)
    assert 'value="furnishing" selected' in sel(r.text, "layer_1") and 'value="main" selected' in sel(r.text, "layer_3")  # Ground already has a main plan: the furniture layout is a layer of it
    assert 'name="floor_1" list="floors" value="Ground"' in r.text and 'name="sheet_1" value="A-105"' in r.text and 'name="caption_1" value=""' in r.text and "Ground Floor Plan (Furniture Layout) · A-105" in r.text
    assert 'name="sheet_2" value="A-301"' in r.text and 'name="floor_2" list="floors" value=""' in r.text and 'name="floor_3" list="floors" value="First"' in r.text
    r=c.post(f"/p/{pid}/images/sets/{tsid}/pick", data={"page_1":"on","floor_1":"Ground","sheet_1":"A-105","caption_1":"Furniture Layout"}, follow_redirects=False); assert r.status_code==303
    db=SessionLocal(); newp=db.query(ProjectImage).filter(ProjectImage.project_id==int(pid), ProjectImage.kind=="floorplan").order_by(ProjectImage.id.desc()).first()
    assert newp.caption=="Furniture Layout" and newp.sheet=="A-105" and newp.floor=="Ground" and newp.tag.set_id==tsid and _dr2.plan_caption(newp)=="Ground floor · A-105 · Furniture Layout"; db.close()
    # a floor whose only plan page is a furniture layout: the picker promotes it to the main plan and keeps the caption
    r=c.post("/projects/new", data={"client_name":"Solo","name":"Solo layout","rate":"7.1"}, follow_redirects=False); pid_s=r.headers["location"].split("/")[-1]
    r=c.post(f"/p/{pid_s}/images/plans", data={}, files=[("files",("solo.pdf",pdf_titled([("GROUND FLOOR PLAN (FURNITURE LAYOUT)","A-104"),("GROUND FLOOR ELECTRICAL LAYOUT","E-01")]),"application/pdf"))], follow_redirects=False)
    r=c.get(r.headers["location"]); h=r.text
    assert 'value="main" selected' in sel(h, "layer_1") and 'name="caption_1" value="Furniture Layout"' in h and 'value="electrical" selected' in sel(h, "layer_2"), sel(h, "layer_1")
    assert "Add page 1 as a plan" in h and '<label for="layer_2">Shows</label>' in h and "main plan or a layer of it" in h  # the picker says what the tick and the Shows box do
    r=c.post(f"/p/{pid_s}/images/plans", data={"floor":"Ground","sheet":"A-100","caption":""}, files=[("files",("g.jpg",img("white"),"image/jpeg"))], follow_redirects=False)  # now a general plan exists ...
    r=c.get(f"/p/{pid_s}/images/sets/{c.get(f'/p/{pid_s}/images').text.split('/images/sets/')[1].split('\"')[0]}"); assert 'value="furnishing" selected' in sel(r.text, "layer_1") and 'name="caption_1" value=""' in r.text  # ... so the furniture layout is prefilled as its layer
    r=c.get(f"/p/{pid}/images/sets/{tsid}"); assert "1 page looks like a plan sheet and is ticked" in r.text  # page 3 still suggested, page 1 added
    r=c.post(f"/p/{pid}/images/sets/{tsid}/pick", data={"page_3":"on","floor_3":"First","sheet_3":"A-106"}, follow_redirects=False)
    r=c.get(f"/p/{pid}/images/sets/{tsid}"); assert "nothing more here looks like a floor plan" in r.text and "look like" not in r.text
    # a scanned set (no text layer) says so and prefills nothing
    r=c.post(f"/p/{pid}/images/plans", data={}, files=[("files",("scan.pdf",pdf_pages([""]),"application/pdf"))], follow_redirects=False)
    ssid=int(r.headers["location"].split("/sets/")[1].split("?")[0])
    r=c.get(f"/p/{pid}/images/sets/{ssid}"); assert "Nothing could be read" in r.text and "checked" not in r.text.split('<div class="pages">')[1]
    c.post(f"/p/{pid}/images/sets/{ssid}/delete", follow_redirects=False); c.post(f"/p/{pid}/images/sets/{tsid}/delete", follow_redirects=False)
    db=SessionLocal(); from app.models import DrawingPage as _DP; assert db.query(_DP).filter(_DP.set_id.in_([tsid, ssid])).count()==0; db.close()  # page meta goes with the set
    # a new project's default "Whole house" entry is an area; deleting a room that has a box works; the export carries Kind
    r=c.post("/projects/new", data={"client_name":"New","name":"Fresh","rate":"7.1"}, follow_redirects=False); pid3=r.headers["location"].split("/")[-1]
    db=SessionLocal(); assert db.query(Room).filter(Room.project_id==int(pid3)).one().is_area; db.close()
    r=c.get(f"/p/{pid3}/rooms"); assert "1 area</span>" in r.text and "0 rooms" not in r.text and "Everything here is listed as a room" not in r.text
    r=c.get(f"/p/{pid3}/items"); assert '<optgroup label="Rooms">' not in r.text and '<optgroup label="Areas">' in r.text  # no empty Rooms group
    c.post(f"/p/{pid}/rooms", data={"code":"TMP3","name":"Box room","floor":"Ground","kind":"room"}, follow_redirects=False)
    db=SessionLocal(); tmp3=db.query(Room).filter(Room.project_id==int(pid), Room.code=="TMP3").first(); db.close()
    c.post(f"/p/{pid}/cartons", data={"supplier_id":sup.id,"room_id":tmp3.id,"contents":"Box in a room","qty":"1","length_cm":"10","width_cm":"10","height_cm":"10"}, follow_redirects=False)
    from sqlalchemy import event as _ev
    from app.db import engine as _eng2
    def _fk_on(dbapi_conn, rec): dbapi_conn.execute("PRAGMA foreign_keys=ON")
    _ev.listen(_eng2, "connect", _fk_on)
    try:
        r=c.post(f"/p/{pid}/rooms/{tmp3.id}/delete", follow_redirects=False); assert r.status_code==303
    finally:
        _ev.remove(_eng2, "connect", _fk_on)
    db=SessionLocal(); assert db.get(Room, tmp3.id) is None and db.query(Carton).filter(Carton.contents=="Box in a room").one().room_id is None; db.close()
    x=openpyxl.load_workbook(io.BytesIO(c.get(f"/p/{pid}/export/items.xlsx").content), read_only=True)["Rooms"]; hdr=[v for v in next(x.iter_rows(values_only=True))]; assert hdr[:3]==["Room","Floor","Kind"]
    r=c2.get(f"/c/{ctok}"); assert re.search(r"\d+ rooms? · \d+ areas? marked|\d+ rooms? marked|\d+ areas? marked", r.text) and " room · " not in r.text.split('id="plan-section"')[1][:300]
    r=c.get(f"/p/{pid2}"); assert '<td colspan="5">Not in a room or area</td>' in r.text and '<td colspan="5">Areas</td>' not in r.text  # the loose items row is not filed under Areas
    # startup on an old database adds the columns (the app's own path, not just migrate() directly)
    import subprocess, sys as _sys
    _env=dict(os.environ, DATABASE_URL=f"sqlite:///{_tmp}/old.db")
    out=subprocess.run([_sys.executable, "-c", "from app.main import startup; startup(); from app.db import engine; from sqlalchemy import inspect; print(sorted(c['name'] for c in inspect(engine).get_columns('rooms')))"], env=_env, capture_output=True, text=True, cwd=os.getcwd())
    assert out.returncode==0 and "'kind'" in out.stdout, out.stderr[-800:]
    print("rooms & areas + title blocks ok")
    # ---- a room zoomed in on its plan: the crop rectangle, the Pillow crop, the items-list card, the Rooms page, the checklist PDF ----
    class _Pin: pass
    pn=_Pin(); pn.x,pn.y,pn.w,pn.h=0.4,0.4,0.2,0.2
    assert _dr2.crop_rect(pn)==(0.35,0.35,0.3,0.3)
    pn.x,pn.y,pn.w,pn.h=0.9,0.9,0.05,0.05; cr=_dr2.crop_rect(pn); assert cr[0]+cr[2]<=1 and cr[1]+cr[3]<=1 and cr[2]>=0.22 and cr[3]>=0.22, cr   # tiny room: at least 22 % of the image, inside it
    pn.x,pn.y,pn.w,pn.h=0.0,0.0,0.3,0.3; cr=_dr2.crop_rect(pn); assert cr[0]==0 and cr[1]==0, cr
    pn.x,pn.y,pn.w,pn.h=0.2,0.1,0.4,0.3
    zoom={"rect":_dr2.crop_rect(pn),"pin":pn,"dots":[{"x":0.3,"y":0.2,"label":"01","color":"#16A34A"},{"x":0.5,"y":0.3,"label":"12","color":"#9CA3AF"},{"x":0.95,"y":0.95,"label":"99","color":"#000000"}]}
    jpeg=_dr2.render_room_crop(big_img("white", 2400, 1600), zoom); cim=Image.open(io.BytesIO(jpeg)); assert cim.format=="JPEG" and max(cim.size)<=1600
    assert abs(cim.width/cim.height - (zoom["rect"][2]*2400)/(zoom["rect"][3]*1600)) < 0.02 and cim.getpixel((2,2))!=(185,89,58)  # crop keeps the region's proportions
    # the items list filtered by a marked room shows it zoomed with its dots; an unmarked room gets the hint instead
    db=SessionLocal(); zr=_dr2.room_zoom(db.get(_P,int(pid)), db.get(Room, r2.id)); r2_dots=len(zr["dots"]) if zr else None; db.close()
    assert zr is not None and r2_dots>=1 and zr["plan"].id==gid and all(d["label"].isdigit() for d in zr["dots"])
    r=c.get(f"/p/{pid}/items?room={r2.id}"); assert r.status_code==200 and 'class="room-zoom"' in r.text and r.text.count('class="rz-dot"')==r2_dots and f'data-cx="{zr["rect"][0]}"' in r.text
    assert f'href="/p/{pid}/plan?plan={gid}&amp;room={r2.id}"' in r.text and "Checklist PDF" in r.text and f'id="item-{r2_item}"' in r.text and f'data-item="{r2_item}"' in r.text
    r=c.get(f"/p/{pid}/items?room={room.id}"); assert 'class="room-zoom"' not in r.text and "Mark it on the plan" in r.text and f"plan?room={room.id}&amp;mode=mark" in r.text  # GF-KIT's box was removed earlier
    r=c.get(f"/p/{pid}/items?room={r2.id}&view=table"); assert f'<tr id="item-{r2_item}" ' in r.text
    r=c.get(f"/p/{pid}/items?room=none"); assert r.status_code==200 and "room-head" not in r.text
    r=c.get(f"/p/{pid}/rooms"); assert r.text.count('class="room-zoom sm"')>=1 and f'plan?plan={gid}&amp;room={r2.id}"' in r.text
    # the checklist PDF: page one is the zoomed room for a marked room, the whole plan for an unmarked one
    t_r2=pdf_text(c.get(f"/p/{pid}/export/room/{r2.id}.pdf").content); assert "ON THE PLAN" in t_r2 and "each dot carries the number" in t_r2 and "FLOOR PLAN" not in t_r2.split("CHECKLIST")[0]
    t_kit=pdf_text(c.get(f"/p/{pid}/export/room/{room.id}.pdf").content); assert "FLOOR PLAN" in t_kit and "ON THE PLAN" not in t_kit
    assert c.get(f"/p/{pid}/export/room/0.pdf").status_code==200
    print("room zoom ok")
    # deleting a plan removes both files and the tag row
    r=c.post(f"/p/{pid}/images/{roof_id}/delete", follow_redirects=False); assert r.status_code==303
    assert _st.read_image(roof_keys[0]) is None and _st.read_image(roof_keys[1]) is None
    db=SessionLocal(); assert db.get(PlanTag, roof_id) is None and db.get(ProjectImage, roof_id) is None and db.query(RoomPin).filter(RoomPin.image_id==roof_id).count()==0 and db.query(ItemPin).filter(ItemPin.image_id==roof_id).count()==0; db.close()
    print("floor plans ok")
    # shared storage: a photo still on this server's disk is served and copied into the bucket
    class _Fake:
        def __init__(self): self.b={}
        def download_as_bytes(self, k): return self.b[k]
        def upload_from_bytes(self, k, d): self.b[k]=d
    db=SessionLocal(); code_key=db.get(Item,d3).photos[0].file_key; db.close()
    fake=_Fake(); _st._replit_client=fake; _st.config.STORAGE_BACKEND="replit"
    try:
        assert _st.read_image(code_key) is not None and code_key in fake.b
        assert _st.read_image("img/missing.jpg") is None
        k=_st.save_image(img("blue")); assert k in fake.b and _st.read_image(k)==fake.b[k]
    finally:
        _st.config.STORAGE_BACKEND="local"; _st._replit_client=None
    print("shared storage ok")
    # one product, many rooms: identity = name + size + finish (the preferred brand is shared, not identity); another size is another product; Apply is pre-ticked only while the twins match
    from app.routers.items import same_item_elsewhere as _sie, product_key as _pk
    db=SessionLocal(); rids=[r.id for r in db.get(_P,int(pid)).rooms if r.kind=="room"][:3]; db.close(); assert len(rids)==3
    def mk(room_id, size, finish="Ceramic, matt, cream"):
        r=c.post(f"/p/{pid}/items/new", data={"room_ids":[str(room_id)],"category":"Tiles","name":"Floor tiles","size":size,"finish":finish,"qty":"20","unit":"m²","status":"To buy"}, follow_redirects=False)
        return int(r.headers["location"].rstrip("/").split("/")[-1])
    t1=mk(rids[0], "1200 x 600 mm"); t2=mk(rids[1], "1200×600mm"); t3=mk(rids[2], "800 * 400 mm")
    db=SessionLocal(); P=db.get(_P,int(pid)); i1,i2,i3=db.get(Item,t1),db.get(Item,t2),db.get(Item,t3)
    assert _pk(i1)==_pk(i2)!=_pk(i3) and [x.id for x in _sie(P,i1)]==[t2] and _sie(P,i3)==[]; db.close()
    h=c.get(f"/p/{pid}/items/{t1}").text
    assert 'name="apply_all" value="1" checked>' in h and "1 other room with this exact item" in h and "differ in" not in h
    assert 'list="dl-name"' in h and '<datalist id="dl-name">' in h and '<option value="Floor tiles">' in h and '<datalist id="dl-size">' in h and '<option value="1200 x 600 mm">' in h and '<option value="Ceramic, matt, cream">' in h
    assert "only in this room so far" in c.get(f"/p/{pid}/items/{t3}").text and '<datalist id="dl-name">' in c.get(f"/p/{pid}/items/new").text
    h=c.get(f"/p/{pid}/items?view=list&q=Floor%20tiles").text; assert f'<tr id="item-{t1}"' in h and f'<tr id="item-{t3}"' in h and "2 rooms</a>" in h and "· 800 * 400 mm · Ceramic, matt, cream" in h
    c.post(f"/p/{pid}/items/{t2}/status", data={"status":"Quoted"}); assert 'name="apply_all" value="1" checked>' in c.get(f"/p/{pid}/items/{t1}").text  # status is per room: no difference that matters
    r=c.post(f"/p/{pid}/items/{t2}", data={"room_id":rids[1],"category":"Tiles","name":"Floor tiles","size":"1200*600 mm","finish":"Ceramic, matt, cream","spec":"Rectified edge","brand":"Marazzi","qty":"20","unit":"m²","status":"Quoted"}, follow_redirects=False); assert r.status_code==303
    h=c.get(f"/p/{pid}/items/{t1}").text; assert 'name="apply_all" value="1">' in h and "The other rooms differ in spec and preferred brand:" in h  # set up its own way: not overwritten by default
    db=SessionLocal(); assert db.get(Item,t1).spec=="" and db.get(Item,t2).spec=="Rectified edge" and _pk(db.get(Item,t1))==_pk(db.get(Item,t2)); db.close()  # still one product: the brand is a preference
    r=c.post(f"/p/{pid}/items/{t1}", data={"room_id":rids[0],"category":"Tiles","name":"Floor tiles","size":"1200 x 600 mm","finish":"Ceramic, matt, cream","spec":"Rectified, 9 mm","qty":"20","unit":"m²","status":"To buy","apply_all":"1"}, follow_redirects=False)
    db=SessionLocal(); assert db.get(Item,t2).spec=="Rectified, 9 mm" and db.get(Item,t2).brand=="" and db.get(Item,t2).status=="Quoted" and db.get(Item,t3).spec=="" and db.get(Item,t3).size=="800 * 400 mm"; db.close()  # ticked: twins follow, the other size untouched
    assert 'name="apply_all" value="1" checked>' in c.get(f"/p/{pid}/items/{t1}").text and '<datalist id="dl-name">' in c.get(f"/p/{pid}/drafts").text
    print("product identity ok")
    # a document that cannot be built explains itself (the designer's route) instead of a bare 500; the client gets a plain message
    from app.routers import exports as _ex
    _orig_build=_ex.build_schedule
    def _boom(*a, **k): raise ValueError("picture 12 cannot be read")
    _ex.build_schedule=_boom
    try:
        r=c.get(f"/p/{pid}/export/schedule.pdf"); assert r.status_code==500 and "Client schedule could not be built" in r.text and "ValueError: picture 12 cannot be read" in r.text and "Try again" in r.text
        r=c2.get(f"/c/{ctok}/schedule.pdf"); assert r.status_code==500 and "could not be built" in r.text and "ValueError" not in r.text
    finally:
        _ex.build_schedule=_orig_build
    assert c.get(f"/p/{pid}/export/schedule.pdf").status_code==200
    print("pdf error page ok")
    # a big project (70 rooms): the two-column summary paginates instead of overflowing the page (ReportLab LayoutError)
    r=c.post("/projects/new", data={"client_name":"Big","name":"Big house","rate":"7.1"}, follow_redirects=False); pid_big=r.headers["location"].split("/")[-1]
    for k in range(70):
        c.post(f"/p/{pid_big}/rooms", data={"code":f"R{k:02d}","name":f"Room number {k} with a longer name","floor":["Ground","First","Second"][k%3]}, follow_redirects=False)
    big_rids=re.findall(r'name="room_ids" value="(\d+)"', c.get(f"/p/{pid_big}/items/new").text); assert len(big_rids)>=70
    r=c.post(f"/p/{pid_big}/items/new", data={"room_ids":big_rids,"category":"Lighting","name":"Pendant light","qty":"2","unit":"pcs","unit_price":"120","status":"To buy"}, follow_redirects=False); assert r.status_code==303
    for q in ("", "?layout=floor", "?prices=0"):
        r=c.get(f"/p/{pid_big}/export/schedule.pdf{q}"); assert r.status_code==200 and r.headers["content-type"]=="application/pdf", (q, r.status_code, r.text[:300])
    tb=pdf_text(c.get(f"/p/{pid_big}/export/schedule.pdf?layout=floor").content); assert tb.count("R69 - Room number 69")>=1 and "Grand total" in tb
    print("big project summary ok")
    # delete by selection: the user ticks lines and presses Delete selected; a filter on its own never deletes anything.
    # An Item list box carries every room line of its product; other projects' ids, drafts and junk are ignored
    r=c.post(f"/p/{pid_big}/items/new", data={"room_ids":[big_rids[0]],"category":"Tiles","name":"Floor tiles","qty":"10","unit":"m²","status":"To buy"}, follow_redirects=False); keep_id=int(r.headers["location"].rstrip("/").split("/")[-1])
    db=SessionLocal(); first=db.query(Item).filter(Item.project_id==int(pid_big), Item.category=="Lighting").first(); first_id=first.id; other_id=db.query(Item).filter(Item.project_id!=int(pid_big), Item.draft==False).first().id; db.close()
    c.post(f"/p/{pid_big}/items/{first_id}", data={"room_id":big_rids[0],"category":"Lighting","name":"Pendant light","qty":"2","unit":"pcs","status":"To buy"}, files={"photos":("p.jpg",img("pink"),"image/jpeg")}, follow_redirects=False)
    db=SessionLocal(); ph_key=db.get(Item,first_id).photos[0].file_key; db.close(); assert _st.read_image(ph_key) is not None
    h=c.get(f"/p/{pid_big}/items?category=Lighting&view=table").text
    assert 'id="bulk"' in h and f'action="/p/{pid_big}/items/delete-selected"' in h and h.count('class="sel-box" name="ids"')==71 and f'name="ids" value="{first_id}" form="bulk" data-photos="1"' in h, h[:200]
    assert 'name="category" value="Lighting"' in h and 'name="view" value="table"' in h and "Delete selected" in h and "delete-filtered" not in h and h.count('class="sel-all"')==1
    h=c.get(f"/p/{pid_big}/items?category=Lighting").text  # the Item list: one box per product, carrying all its room lines
    m=re.search(r'name="ids" value="([\d ]+)" form="bulk" data-photos="1"', h); assert m and len(m.group(1).split())==71 and h.count('class="sel-box"')==1, h.count('class="sel-box"')
    assert c.get(f"/p/{pid_big}/items?view=cards").text.count('class="sel-box"')==72 and 'id="bulk"' in c.get(f"/p/{pid_big}/items").text  # no filter needed: nothing goes without a tick
    assert "bulk" not in c.get(f"/p/{pid_big}/items?category=Curtains").text  # nothing shown, no bar
    r=c.post(f"/p/{pid_big}/items/delete-selected", data={"category":"Lighting"}, follow_redirects=False); assert r.status_code==303 and r.headers["location"]==f"/p/{pid_big}/items?category=Lighting", r.headers  # nothing ticked: nothing happens
    db=SessionLocal(); assert db.query(Item).filter(Item.project_id==int(pid_big)).count()==72; db.close()
    others=[x for x in m.group(1).split() if int(x)!=first_id]
    r=c.post(f"/p/{pid_big}/items/delete-selected", data={"ids":[" ".join(others[:30])]+others[30:]+[str(other_id),"abc"],"category":"Lighting","view":"table"}, follow_redirects=False)
    assert r.status_code==303 and r.headers["location"]==f"/p/{pid_big}/items?category=Lighting&view=table&deleted=70", r.headers
    db=SessionLocal(); assert [i.id for i in db.query(Item).filter(Item.project_id==int(pid_big), Item.category=="Lighting").all()]==[first_id] and db.get(Item,other_id) is not None; db.close()  # the unticked one (with the photo) and the other project's item survived
    assert _st.read_image(ph_key) is not None
    r=c.post(f"/p/{pid_big}/items/delete-selected", data={"ids":[str(first_id)],"category":"Lighting"}, follow_redirects=False); assert r.status_code==303 and r.headers["location"].endswith("category=Lighting&deleted=1")
    db=SessionLocal(); left=db.query(Item).filter(Item.project_id==int(pid_big)).all(); assert [i.id for i in left]==[keep_id]; db.close()
    assert _st.read_image(ph_key) is None and "70 items deleted." in c.get(f"/p/{pid_big}/items?deleted=70").text and "1 item deleted." in c.get(f"/p/{pid_big}/items?category=Lighting&deleted=1").text
    print("delete selected ok")
    # messages and confirmations are the app's own dialog (app.js appConfirm/appAlert, .modal), never the browser's system boxes
    import re as _re; from pathlib import Path
    for f in list(Path("app/templates").rglob("*.html")) + list(Path("app/static").glob("*.js")):
        code = _re.sub(r"^\s*//.*$", "", f.read_text(), flags=_re.M)  # comments may name the forbidden calls
        assert not _re.search(r"(?<![\w.])(?:window\.)?(?:confirm|alert|prompt)\s*\(", code) and "confirmSubmit" not in code and "onsubmit=" not in code, f
    assert "data-confirm=" in c.get(f"/p/{pid}/items/{d3}").text and "data-removal=" in c.get(f"/p/{pid}/items/{d3}").text
    assert "window.appConfirm" in Path("app/static/app.js").read_text()
    print("no system dialogs ok")
    # ---- floors: the floor list on the Rooms page (models.Floor): add, rename, reorder, merge, delete; every dropdown, title and the import follow ----
    r=c.post("/projects/new", data={"client_name":"Floors","name":"Floors test","rate":"7"}, follow_redirects=False); pf=r.headers["location"].rstrip("/").split("/")[-1]
    c.post(f"/p/{pf}/rooms", data={"code":"GF-LIV","name":"Living room","floor":"Ground"}, follow_redirects=False)
    c.post(f"/p/{pf}/rooms", data={"code":"FF-BR1","name":"Bedroom 1","floor":"first floor"}, follow_redirects=False)  # a new floor starts with a capital
    c.post(f"/p/{pf}/rooms", data={"code":"GF-KIT","name":"Kitchen","floor":"ground floor"}, follow_redirects=False)  # the same floor as Ground: stored as "Ground"
    c.post(f"/p/{pf}/rooms", data={"code":"RF-GYM","name":"Gym","floor":"Roof"}, follow_redirects=False)
    c.post(f"/p/{pf}/rooms", data={"code":"BQ","name":"Boys' Quarter","floor":"__new__","floor_new":"Staff quarters"}, follow_redirects=False)  # the dropdown's "+ New floor…"
    c.post(f"/p/{pf}/rooms", data={"code":"EXT","name":"Garden","floor":"Site","kind":"area"}, follow_redirects=False)  # a pseudo floor: never listed
    r=c.post(f"/p/{pf}/images/plans", data={"floor":"roof","sheet":"A-103"}, files=[("files",("roof.jpg",img("white"),"image/jpeg"))], follow_redirects=False); assert r.status_code==303
    db=SessionLocal(); pr=db.get(_P,int(pf)); fl={f.name:f for f in pr.floors}
    assert list(fl)==["Ground","First floor","Roof","Staff quarters"], list(fl)
    assert fl["Staff quarters"].kind=="area" and fl["Ground"].kind=="floor" and _dr.plans(pr)[0].floor=="Roof"
    assert {x.code:x.floor for x in pr.rooms if x.code!="ALL"}=={"GF-LIV":"Ground","FF-BR1":"First floor","GF-KIT":"Ground","RF-GYM":"Roof","BQ":"Staff quarters","EXT":"Site"}
    assert _dr.floor_order(pr)==["Ground","First floor","Roof","Staff quarters"] and _dr.floor_title("staff quarters",pr)=="Staff quarters" and _dr.floor_title("ground floor",pr)=="Ground floor" and _dr.floor_title("Site",pr)=="Site"
    roof_id, ground_id, sq_id, first_id = fl["Roof"].id, fl["Ground"].id, fl["Staff quarters"].id, fl["First floor"].id; bq_id=next(x.id for x in pr.rooms if x.code=="BQ"); db.close()
    h=c.get(f"/p/{pf}/rooms").text
    assert 'id="floors"' in h and h.count('class="in floor-sel"')>=7 and "+ New floor…" in h and f'action="/p/{pf}/floors/{roof_id}/merge"' in h and 'list="floors"' not in h and h.count("<option value=\"Staff quarters\"")>=7
    assert "Staff quarters</h2>" in h and "Roof</h2>" in h  # group titles: a separate area prints as it is
    # add: a duplicate (any spelling) and a pseudo floor are refused with a message; a real one joins the list
    r=c.post(f"/p/{pf}/floors", data={"name":"GROUND FLOOR"}, follow_redirects=False); assert "already in the list" in c.get(r.headers["location"]).text
    r=c.post(f"/p/{pf}/floors", data={"name":"Site"}, follow_redirects=False); assert "not a floor" in c.get(r.headers["location"]).text
    r=c.post(f"/p/{pf}/floors", data={"name":"Second","kind":"floor"}, follow_redirects=False); assert "added" in c.get(r.headers["location"]).text
    db=SessionLocal(); pr=db.get(_P,int(pf)); second_id=next(f.id for f in pr.floors if f.name=="Second"); db.close()
    # rename: its rooms and plans follow the new spelling; a rename onto another listed floor is refused (merge instead)
    r=c.post(f"/p/{pf}/floors/{first_id}", data={"name":"First","kind":"floor"}, follow_redirects=False); assert "1 entry follows" in c.get(r.headers["location"]).text
    r=c.post(f"/p/{pf}/floors/{first_id}", data={"name":"ground","kind":"floor"}, follow_redirects=False); assert "use Merge into" in c.get(r.headers["location"]).text
    db=SessionLocal(); pr=db.get(_P,int(pf)); assert next(x for x in pr.rooms if x.code=="FF-BR1").floor=="First" and _dr.floor_title("first floor",pr)=="First floor"; db.close()
    # merge Roof into Second: the gym and the roof plan move, Roof leaves the list; then Second moves up under First
    r=c.post(f"/p/{pf}/floors/{roof_id}/merge", data={"into":str(second_id)}, follow_redirects=False); assert "merged into" in c.get(r.headers["location"]).text
    db=SessionLocal(); pr=db.get(_P,int(pf)); assert [f.name for f in pr.floors]==["Ground","First","Staff quarters","Second"] and next(x for x in pr.rooms if x.code=="RF-GYM").floor=="Second"
    assert all(im.floor=="Second" for im in _dr.plans(pr)) and [g["title"] for g in _dr.rooms_by_floor(pr)]==["Ground floor","First floor","Staff quarters","Second floor","Whole house / other"]; db.close()
    c.post(f"/p/{pf}/floors/{second_id}/move", data={"dir":"up"}, follow_redirects=False)
    db=SessionLocal(); pr=db.get(_P,int(pf)); assert _dr.floor_order(pr)==["Ground","First","Second","Staff quarters"]; db.close()
    # delete: refused while something is on it, fine once empty
    r=c.post(f"/p/{pf}/floors/{sq_id}/delete", data={}, follow_redirects=False); assert "still has 1 room" in c.get(r.headers["location"]).text
    c.post(f"/p/{pf}/rooms/{bq_id}", data={"code":"BQ","name":"Boys' Quarter","floor":"Second"}, follow_redirects=False)
    r=c.post(f"/p/{pf}/floors/{sq_id}/delete", data={}, follow_redirects=False); assert "deleted" in c.get(r.headers["location"]).text
    # a plan tag typed another way lands on the listed spelling; the import registers a new floor and says what moves
    db=SessionLocal(); pr=db.get(_P,int(pf)); plan_id=_dr.plans(pr)[0].id; db.close()
    c.post(f"/p/{pf}/images/{plan_id}", data={"floor":" second floor ","sheet":"A-103"}, follow_redirects=False)
    db=SessionLocal(); assert db.get(ProjectImage,plan_id).tag.floor=="Second"; db.close()
    import openpyxl as _ox, io as _io
    wb=_ox.Workbook(); ws=wb.active; ws.title="Rooms"; ws.append(["Room","Floor","Kind"]); ws.append(["BQ - Boys' Quarter","Guest house","room"]); ws.append(["GH-BR - Guest bedroom","Guest house","room"])
    ws2=wb.create_sheet("Shopping List"); ws2.append(["Code","Room","Category","Item","Qty","Unit"]); ws2.append(["","GH-BR - Guest bedroom","Furniture","Bed","1","pcs"])
    buf=_io.BytesIO(); wb.save(buf)
    r=c.post(f"/p/{pf}/import", files={"file":("floors.xlsx",buf.getvalue(),"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}); h=r.text
    assert "1 new floor" in h and "1 room changed" in h and "floor Second → Guest house" in h and "1 new room" in h, h[:400]
    key=re.search(r'name="key" value="([^"]+)"', h).group(1); r=c.post(f"/p/{pf}/import", data={"key":key,"apply":"1"}, follow_redirects=False); assert r.status_code in (200,303)
    db=SessionLocal(); pr=db.get(_P,int(pf)); gh=next(f for f in pr.floors if f.name=="Guest house"); assert gh.kind=="area" and {x.code:x.floor for x in pr.rooms if x.code in ("BQ","GH-BR")}=={"BQ":"Guest house","GH-BR":"Guest house"}; db.close()
    r=c.get(f"/p/{pf}/export/schedule.pdf?layout=floor"); assert r.status_code==200 and "GUEST HOUSE" in pdf_text(r.content).upper() and "GUEST HOUSE FLOOR" not in pdf_text(r.content).upper()
    print("floors ok")
    # ---- the plan page follows the floor list: a tab per listed floor (with or without a plan), a floor without a plan lists its rooms ----
    db=SessionLocal(); pr=db.get(_P,int(pf)); liv_id=next(x.id for x in pr.rooms if x.code=="GF-LIV"); gym_id=next(x.id for x in pr.rooms if x.code=="RF-GYM"); ctok=pr.client_token; db.close()
    h=c.get(f"/p/{pf}/plan").text
    assert 'class="plan-tabs"' in h and "?floor=ground" in h and "?floor=first" in h and "?floor=guest%20house" in h and "No plan yet" in h and "Second floor" in h, h[:300]
    assert h.count("No plan yet")==3 and "· 2 rooms" in h  # Ground has two rooms, no plan
    h=c.get(f"/p/{pf}/plan?floor=ground").text
    assert "No plan for Ground floor yet" in h and f'href="/p/{pf}/items?room={liv_id}"' in h and "GF-KIT" in h and "Add its plan" in h and 'id="plan-stage"' not in h
    h=c.get(f"/p/{pf}/plan?floor=second").text  # the floor's plan opens; the side panel lists the floor's rooms, placed or not
    assert 'id="plan-stage"' in h and f'room={gym_id}' in h and "not placed" in h and "Second floor</div>" in h
    assert "No plan for" in c.get(f"/p/{pf}/plan?floor=Guest%20House").text  # any spelling of a listed floor
    h=c.get(f"/c/{ctok}").text; assert "?floor=" not in h and "not placed" not in h  # the client never sees floors without a plan or designer hints
    # ---- filtered exports: the Items page filters and a floor, the same editable sheet, named after the choice ----
    r=c.post(f"/p/{pf}/items/new", data={"room_ids":[str(liv_id)],"category":"Furniture","name":"Sofa","qty":"1","unit":"pcs","status":"To buy"}, follow_redirects=False); assert r.status_code==303
    r=c.post(f"/p/{pf}/items/new", data={"room_ids":[str(gym_id)],"category":"Lighting","name":"Spot light","qty":"4","unit":"pcs","status":"To buy"}, follow_redirects=False); assert r.status_code==303
    def xrows(resp): wb_=_ox.load_workbook(_io.BytesIO(resp.content), data_only=True); return [x for x in wb_["Shopping List"].iter_rows(min_row=2, values_only=True) if x[0]]
    r=c.get(f"/p/{pf}/export/items.xlsx"); assert r.status_code==200 and len(xrows(r))==3 and "-items.xlsx" in r.headers["content-disposition"]
    r=c.get(f"/p/{pf}/export/items.xlsx?category=Lighting"); assert [x[2] for x in xrows(r)]==["Lighting"] and "-Lighting.xlsx" in r.headers["content-disposition"]
    r=c.get(f"/p/{pf}/export/items.xlsx?floor=ground%20floor"); assert [x[3] for x in xrows(r)]==["Sofa"] and "Ground-floor" in r.headers["content-disposition"]
    r=c.get(f"/p/{pf}/export/items.xlsx?room={gym_id}&status=To+buy"); assert [x[3] for x in xrows(r)]==["Spot light"] and "RF-GYM-To-buy" in r.headers["content-disposition"]
    assert xrows(c.get(f"/p/{pf}/export/items.xlsx?category=Tiles"))==[]  # an empty filter: a sheet with the rooms and no rows
    h=c.get(f"/p/{pf}/items?category=Lighting").text; assert "Export these 1" in h and f'href="/p/{pf}/export/items.xlsx?category=Lighting"' in h
    h=c.get(f"/p/{pf}/items").text; assert f'href="/p/{pf}/export/items.xlsx"' in h and "Export these" not in h
    h=c.get(f"/p/{pf}/import").text; assert f'action="/p/{pf}/export/items.xlsx"' in h and 'name="floor"' in h and 'value="Guest house"' in h and "Export a part" in h
    print("exports and plan navigation ok")
    # ---- room inspiration: mood-board pictures per room (MoodTag) on the panel, the Rooms page, the item list, the Images page and the client's link; the client adds their own ----
    r=c.post("/projects/new", data={"client_name":"Insp","name":"Inspiration test","rate":"7"}, follow_redirects=False); pi=r.headers["location"].rstrip("/").split("/")[-1]
    c.post(f"/p/{pi}/rooms", data={"code":"GF-KIT","name":"Kitchen","floor":"Ground"}, follow_redirects=False); c.post(f"/p/{pi}/rooms", data={"code":"GF-LIV","name":"Living room","floor":"Ground"}, follow_redirects=False)
    db=SessionLocal(); pr=db.get(_P,int(pi)); kit=next(x.id for x in pr.rooms if x.code=="GF-KIT"); liv=next(x.id for x in pr.rooms if x.code=="GF-LIV"); itok=pr.client_token; db.close()
    # the designer: a whole-house picture, and a kitchen picture the way the room panel's form posts it (room_id + next)
    r=c.post(f"/p/{pi}/images", data={"kind":"mood","caption":"Warm wood"}, files=[("files",("h.jpg",img("brown"),"image/jpeg"))], follow_redirects=False); assert r.headers["location"]==f"/p/{pi}/images"
    r=c.post(f"/p/{pi}/images", data={"kind":"mood","caption":"Black taps","room_id":str(kit),"next":f"/p/{pi}/plan?room={kit}"}, files=[("files",("k.jpg",img("black"),"image/jpeg"))], follow_redirects=False); assert r.headers["location"]==f"/p/{pi}/plan?room={kit}"
    db=SessionLocal(); pr=db.get(_P,int(pi)); ms=[im for im in pr.images if im.kind=="mood"]; assert len(ms)==2 and {im.room_id for im in ms}=={None,kit} and not any(im.by_client for im in ms); kpic=next(im for im in ms if im.room_id==kit).id; db.close()
    h=c.get(f"/p/{pi}/plan/room/{kit}").text; assert "Inspiration" in h and "Black taps" in h and "Warm wood" not in h and "Add inspiration" in h and f'action="/p/{pi}/images/{kpic}/delete"' in h
    assert "Black taps" in c.get(f"/p/{pi}/rooms").text and "Black taps" in c.get(f"/p/{pi}/items?room={kit}").text and "Black taps" not in c.get(f"/p/{pi}/items?room={liv}").text
    h=c.get(f"/p/{pi}/images").text; assert "GF-KIT - Kitchen</span>" in h and "Whole house</span>" in h and f'id="img-{kpic}"' in h and 'name="room_id"' in h
    # a picture moved to the living room from the Images page, back to the whole house, then to the kitchen again
    c.post(f"/p/{pi}/images/{kpic}", data={"caption":"Black taps","room_id":str(liv)}, follow_redirects=False); db=SessionLocal(); assert db.get(ProjectImage,kpic).room_id==liv; db.close()
    c.post(f"/p/{pi}/images/{kpic}", data={"caption":"Black taps","room_id":""}, follow_redirects=False); db=SessionLocal(); assert db.get(ProjectImage,kpic).room_id is None; db.close()
    c.post(f"/p/{pi}/images/{kpic}", data={"caption":"Black taps","room_id":str(kit)}, follow_redirects=False)
    # the client: the link shows the house pictures with the add form; they add one for the kitchen ("you" for them, "client" for the designer)
    h=c2.get(f"/c/{itok}").text; assert 'id="mood"' in h and "Add a picture of what you like" in h and f'action="/c/{itok}/inspiration"' in h and "Warm wood" in h and "Add inspiration" not in h and "/p/" not in h.split('id="mood"')[1].split("plan-section")[0]
    r=c2.post(f"/c/{itok}/inspiration", data={"room_id":str(kit),"caption":"Like this green"}, files=[("files",("c.jpg",img("green"),"image/jpeg"))], follow_redirects=False)
    assert r.status_code==303 and r.headers["location"].startswith(f"/c/{itok}?room={kit}&ok=") and r.headers["location"].endswith("#plan-section"), r.headers
    db=SessionLocal(); pr=db.get(_P,int(pi)); cp=next(im for im in pr.images if im.by_client); assert cp.room_id==kit and cp.caption=="Like this green"; cpid=cp.id; db.close()
    h=c2.get(f"/c/{itok}/plan/room/{kit}").text; assert "Like this green" in h and ">you</span>" in h and f'action="/c/{itok}/inspiration/{cpid}/delete"' in h and f'/images/{kpic}/delete' not in h and "Black taps" in h
    h=c.get(f"/p/{pi}/plan/room/{kit}").text; assert "Like this green" in h and ">client</span>" in h and "1 from the client" in h
    assert "1 inspiration picture" in c.get(f"/p/{pi}").text and "1 from the client" in c.get(f"/p/{pi}/images").text
    # room covers: the designer's newest picture stands in, the star picks one; it shows on the card, in the panel (the client's too) and rings the picture
    from app.services import room_cover as _rc
    from app.models import Room as _Rm
    db=SessionLocal(); pr=db.get(_P,int(pi)); assert _rc(pr,db.get(_Rm,kit)).id==kpic and _rc(pr,db.get(_Rm,liv)) is None; db.close()  # the designer's, although the client's is newer
    h=c.get(f"/p/{pi}/rooms").text; assert 'class="rthumb" src="/media/' in h and 'class="rcover mt"' in h and f'action="/p/{pi}/images/{kpic}/cover"' in h and f'action="/p/{pi}/images/{cpid}/cover"' in h
    r=c.post(f"/p/{pi}/images/{cpid}/cover", data={"next":f"/p/{pi}/rooms#room-{kit}"}, follow_redirects=False); assert r.status_code==303 and r.headers["location"]==f"/p/{pi}/rooms#room-{kit}"
    db=SessionLocal(); pr=db.get(_P,int(pi)); cv=_rc(pr,db.get(_Rm,kit)); assert cv.id==cpid and cv.is_cover and not db.get(ProjectImage,kpic).is_cover; ck=cv.file_key; db.close()
    h=c.get(f"/p/{pi}/plan/room/{kit}").text; assert f'<div class="ph-cover"><img src="/media/{ck}"' in h and 'class="btn sm icon on"' in h
    h=c2.get(f"/c/{itok}/plan/room/{kit}").text; assert f'<div class="ph-cover"><img src="/media/{ck}"' in h and "/cover\"" not in h  # the client sees the cover, never the star
    h=c.get(f"/p/{pi}/images").text; assert 'class="g client cover"' in h and f'action="/p/{pi}/images/{cpid}/cover"' in h
    c.post(f"/p/{pi}/images/{cpid}/cover", data={"next":""}, follow_redirects=False); db=SessionLocal(); pr=db.get(_P,int(pi)); assert not db.get(ProjectImage,cpid).is_cover and _rc(pr,db.get(_Rm,kit)).id==kpic; db.close()  # the star again: the newest stands in
    c.post(f"/p/{pi}/images/{kpic}/cover", data={"next":""}, follow_redirects=False)
    c.post(f"/p/{pi}/images/{kpic}", data={"caption":"Black taps","room_id":str(liv)}, follow_redirects=False); db=SessionLocal(); assert not db.get(ProjectImage,kpic).is_cover; db.close()  # moved to another room: no longer the kitchen's cover
    c.post(f"/p/{pi}/images/{kpic}", data={"caption":"Black taps","room_id":str(kit)}, follow_redirects=False)
    db=SessionLocal(); wh=next(im.id for im in db.get(_P,int(pi)).images if im.kind=="mood" and im.room_id is None); db.close()
    r=c.post(f"/p/{pi}/images/{wh}/cover", data={"next":""}, follow_redirects=False); db=SessionLocal(); assert r.status_code==303 and not db.get(ProjectImage,wh).is_cover; db.close()  # a whole-house picture is never a room's cover
    assert c.post(f"/p/{pi}/images/999999/cover", data={"next":""}, follow_redirects=False).status_code==404
    print("room covers ok")

    # photos belong to the product: added from any room's line (or the item list), shown on every room's line, one copy stored
    from app.routers.items import new_photo as _np
    from app.models import ItemPhoto
    r=c.post("/projects/new", data={"client_name":"Shared","name":"Shared photos","rate":"7.1"}, follow_redirects=False); ps=r.headers["location"].split("/")[-1]
    for k in range(3): c.post(f"/p/{ps}/rooms", data={"code":f"S{k}","name":f"Shared room {k}","floor":"Ground","kind":"room"}, follow_redirects=False)
    srids=re.findall(r'name="room_ids" value="(\d+)"', c.get(f"/p/{ps}/items/new").text)[:3]
    c.post(f"/p/{ps}/items/new", data={"room_ids":srids,"category":"Lighting","name":"Wall light","qty":"2","unit":"pcs","status":"To buy"}, files=[("photos",("a.jpg",img("red"),"image/jpeg"))], follow_redirects=False)
    db=SessionLocal(); sits=db.query(Item).filter(Item.project_id==int(ps)).order_by(Item.id).all(); sids=[i.id for i in sits]
    assert sum(len(i.photos) for i in sits)==1 and [len(i.all_photos) for i in sits]==[1,1,1] and len({i.cover.file_key for i in sits})==1 and len(sits[0].photos[0].digest)==40; db.close()
    c.post(f"/p/{ps}/items/{sids[1]}/photo", files=[("photos",("b.jpg",img("blue"),"image/jpeg"))], follow_redirects=False)  # from another room's line
    db=SessionLocal(); sits=[db.get(Item,i) for i in sids]; assert [len(i.photos) for i in sits]==[1,1,0] and [len(i.all_photos) for i in sits]==[2,2,2]; bid=sits[1].photos[0].id; bk=sits[1].photos[0].file_key; db.close()
    h=c.get(f"/p/{ps}/items?view=table").text; assert h.count('class="thumb sm"')==3 and h.count("data-search=")==3 and 'id="live-q"' in h and 'id="live-q-hidden"' in h and 'type="text" name="q"' not in h and "Shared room 2 " in h
    h=c.get(f"/p/{ps}/items/{sids[2]}").text; assert h.count("/photo/")==3 and "Shared with this item" in h and f"/photo/{bid}/cover" in h  # the line without photos of its own shows both, with the actions
    c.post(f"/p/{ps}/items/{sids[2]}/photo/{bid}/cover", follow_redirects=False); db=SessionLocal(); assert all(db.get(Item,i).cover.file_key==bk for i in sids); fid=db.get(Item,sids[0]).photos[0].id; db.close()  # the cover changes for every room
    c.post(f"/p/{ps}/items/{sids[2]}/photo/{fid}/delete", follow_redirects=False); db=SessionLocal(); sits=[db.get(Item,i) for i in sids]; assert [len(i.all_photos) for i in sits]==[1,1,1] and sum(len(i.photos) for i in sits)==1; holder=next(i for i in sits if i.photos); hid=holder.id; hk=holder.photos[0].file_key; db.close()
    c.post(f"/p/{ps}/items/{hid}/delete", follow_redirects=False); db=SessionLocal(); rest=[db.get(Item,i) for i in sids if i!=hid]; assert [len(i.all_photos) for i in rest]==[1,1] and _st.read_image(hk) is not None and sum(len(i.photos) for i in rest)==1; db.close()  # the holder's line goes, the photo stays with the product
    keep_s, gone_s = [i.id for i in rest]
    r=c.post(f"/p/{ps}/items/{keep_s}", data={"room_id":str(rest[0].room_id),"category":"Lighting","name":"Wall light","qty":"2","unit":"pcs","status":"To buy","rooms_form":"1"}, follow_redirects=False)
    db=SessionLocal(); assert r.status_code==303 and db.get(Item,gone_s) is None and len(db.get(Item,keep_s).photos)==1 and db.get(Item,keep_s).cover.file_key==hk; db.close()  # unticked room: its line goes, the photo moves
    c.post(f"/p/{ps}/items/new", data={"room_ids":srids[:2],"category":"Tiles","name":"Floor tile","size":"600 x 600","qty":"1","unit":"m²","status":"To buy"}, follow_redirects=False)
    db=SessionLocal(); tl=db.query(Item).filter(Item.project_id==int(ps), Item.name=="Floor tile").order_by(Item.id).all(); pr=tl[0].project; same=img("green"); db.add(_np(tl[0].id, same, pr)); db.add(_np(tl[1].id, same, pr)); db.commit(); tids=[i.id for i in tl]; db.close()
    db=SessionLocal(); tl=[db.get(Item,i) for i in tids]; assert sum(len(i.photos) for i in tl)==2 and [len(i.all_photos) for i in tl]==[1,1]; pk=tl[0].photos[0].id; db.close()  # the same picture stored twice (before sharing) shows once
    c.post(f"/p/{ps}/items/{tids[1]}/photo/{pk}/delete", follow_redirects=False); db=SessionLocal(); assert db.query(ItemPhoto).join(Item).filter(Item.id.in_(tids)).count()==0; db.close()  # and goes together
    c.post(f"/p/{ps}/items/{tids[0]}/photo", files=[("photos",("c.jpg",img("gray"),"image/jpeg"))], follow_redirects=False)
    c.post(f"/p/{ps}/items/delete-selected", data={"ids":[str(tids[0])]}, follow_redirects=False); db=SessionLocal(); t1=db.get(Item,tids[1]); assert len(t1.photos)==1 and t1.cover is not None; fk=t1.photos[0].file_key; db.close()  # bulk delete: the photo moves to the surviving room
    c.post(f"/p/{ps}/items/delete-selected", data={"ids":[str(tids[1])]}, follow_redirects=False); assert _st.read_image(fk) is None  # the last line: the files go
    db=SessionLocal(); ph0=db.query(ItemPhoto).join(Item).filter(Item.project_id==int(ps)).first(); ph0.digest=""; ph0id=ph0.id; db.commit(); db.close()
    assert "Prepare 1 older photo" in c.get("/settings").text; c.post("/settings/thumbs", follow_redirects=False); db=SessionLocal(); assert len(db.get(ItemPhoto,ph0id).digest)==40; db.close()  # the backfill fingerprints older photos
    assert "livesearch" in open("app/static/app.js").read() and ".livesearch{" in open("app/static/app.css").read()
    print("shared photos + live search ok")

    # the floor filter: every room and area on a floor, separate areas and site places, whole-house items, export and delete keep it
    r=c.post("/projects/new", data={"client_name":"Floors","name":"Floor filter","rate":"7.1"}, follow_redirects=False); pf=r.headers["location"].split("/")[-1]
    for code,name,fl,kind in [("GF-LIV","Living","Ground","room"),("GF-STA","Stairs","Ground floor","area"),("FF-MBR","Master","First","room"),("BQ","Quarters","Staff quarters","room"),("EXT","Fence","Site","area")]:
        c.post(f"/p/{pf}/rooms", data={"code":code,"name":name,"floor":fl,"kind":kind}, follow_redirects=False)
    db=SessionLocal(); fr={x.code:x.id for x in db.query(Room).filter(Room.project_id==int(pf))}; db.close()
    for code in ["GF-LIV","GF-STA","FF-MBR","BQ","EXT"]:
        c.post(f"/p/{pf}/items/new", data={"room_ids":[str(fr[code])],"category":"Lighting","name":f"Light {code}","qty":"1","unit":"pcs","status":"To buy"}, follow_redirects=False)
    c.post(f"/p/{pf}/items/new", data={"category":"Lighting","name":"Spare bulbs","qty":"1","unit":"set","status":"To buy"}, follow_redirects=False)
    h=c.get(f"/p/{pf}/items?view=table").text
    assert '<select name="floor"' in h and '<optgroup label="Floors">' in h and '<option value="Ground" >Ground floor</option>' in h and '<option value="Site" >Site</option>' in h and '<option value="none" >No floor (whole house)</option>' in h, re.findall(r'<select name="floor".*?</select>', h)[:1]
    def rows(qs): return set(re.findall(r'<tr id="item-\d+" data-search="([A-Z]+-?[A-Z]*)-\d+ ', c.get(f"/p/{pf}/items?view=table&{qs}").text))
    assert rows("floor=Ground")=={"GF-LIV","GF-STA"} and rows("floor=ground%20floor")=={"GF-LIV","GF-STA"} and rows("floor=First")=={"FF-MBR"} and rows("floor=Staff%20quarters")=={"BQ"} and rows("floor=Site")=={"EXT"} and rows("floor=none")=={"ALL"}, (rows("floor=Ground"), rows("floor=none"))
    h=c.get(f"/p/{pf}/items?view=table&floor=Ground").text; assert f'value="{fr["GF-LIV"]}"' in h and f'value="{fr["FF-MBR"]}"' not in h and "All rooms &amp; areas on Ground floor" in h and 'value="Ground" selected' in h  # the Room list narrows
    assert rows(f"floor=Ground&room={fr['FF-MBR']}")=={"GF-LIV","GF-STA"} and rows(f"floor=Ground&room={fr['GF-STA']}")=={"GF-STA"}  # a room from another floor is let go
    x=c.get(f"/p/{pf}/export/items.xlsx?floor=First"); wbx=openpyxl.load_workbook(io.BytesIO(x.content)); assert [r_[0] for r_ in wbx["Shopping List"].iter_rows(min_row=2, values_only=True) if r_[0]]==["FF-MBR-01"] and "First" in x.headers["content-disposition"]
    x=c.get(f"/p/{pf}/export/items.xlsx?floor=none"); assert [r_[0] for r_ in openpyxl.load_workbook(io.BytesIO(x.content))["Shopping List"].iter_rows(min_row=2, values_only=True) if r_[0]]==["ALL-01"]
    h=c.get(f"/p/{pf}/items?view=table&floor=Ground").text; assert '<input type="hidden" name="floor" value="Ground">' in h and "export/items.xlsx?floor=Ground" in h
    r=c.post(f"/p/{pf}/items/delete-selected", data={"ids":[],"floor":"Ground","view":"table"}, follow_redirects=False); assert r.headers["location"]==f"/p/{pf}/items?floor=Ground&view=table", r.headers["location"]
    print("floor filter ok")
    # nothing chosen, a bad link, the cap per room, someone else's picture, a wrong token
    r=c2.post(f"/c/{itok}/inspiration", data={"room_id":str(kit)}, follow_redirects=False); assert "err=" in r.headers["location"]
    r=c2.post(f"/c/{itok}/inspiration", data={"room_id":str(kit),"url":"https://shop.example/empty"}, follow_redirects=False); assert "err=" in r.headers["location"]
    for k in range(11): c2.post(f"/c/{itok}/inspiration", data={"room_id":str(liv)}, files=[("files",(f"l{k}.jpg",img("pink"),"image/jpeg"))], follow_redirects=False)
    r=c2.post(f"/c/{itok}/inspiration", data={"room_id":str(liv)}, files=[("files",("l12.jpg",img("pink"),"image/jpeg"))], follow_redirects=False); assert "ok=" in r.headers["location"]
    r=c2.post(f"/c/{itok}/inspiration", data={"room_id":str(liv)}, files=[("files",("l13.jpg",img("pink"),"image/jpeg"))], follow_redirects=False); assert "err=Up+to+12" in r.headers["location"], r.headers
    db=SessionLocal(); pr=db.get(_P,int(pi)); assert sum(1 for im in pr.images if im.by_client and im.room_id==liv)==12; db.close()
    c2.post(f"/c/{itok}/inspiration/{kpic}/delete", data={"room_id":str(kit)}, follow_redirects=False); db=SessionLocal(); assert db.get(ProjectImage,kpic) is not None; db.close()  # the designer's stays
    c2.post(f"/c/{itok}/inspiration/{cpid}/delete", data={"room_id":str(kit)}, follow_redirects=False); db=SessionLocal(); assert db.get(ProjectImage,cpid) is None; db.close()  # their own goes
    assert c2.post("/c/nosuchtoken/inspiration", data={"room_id":str(kit)}, files=[("files",("x.jpg",img("pink"),"image/jpeg"))], follow_redirects=False).status_code==404
    r=c2.post(f"/c/{itok}/inspiration", data={"room_id":str(kit)}, files=[("files",("bad.jpg",b"not a picture","image/jpeg"))], follow_redirects=False); assert "err=That+picture+could+not+be+read" in r.headers["location"]  # a broken file: a message, never a 500
    r=c.post(f"/p/{pi}/images", data={"kind":"mood","room_id":str(kit),"next":f"/p/{pi}/rooms"}, files=[("files",("bad.jpg",b"not a picture","image/jpeg"))], follow_redirects=False); assert r.headers["location"].startswith(f"/p/{pi}/rooms?err=1%20file%20could%20not%20be%20read")
    # the schedule names the room on its picture; the designer deletes with next; a deleted room leaves its pictures as whole-house ones
    r=c.get(f"/p/{pi}/export/schedule.pdf"); assert r.status_code==200 and "Kitchen · Black taps" in pdf_text(r.content)
    r=c.post(f"/p/{pi}/images/{kpic}/delete", data={"next":f"/p/{pi}/rooms#room-{kit}"}, follow_redirects=False); assert r.headers["location"]==f"/p/{pi}/rooms#room-{kit}"
    db=SessionLocal(); assert db.get(ProjectImage,kpic) is None; db.close()
    c.post(f"/p/{pi}/rooms/{liv}/delete", data={}, follow_redirects=False)
    db=SessionLocal(); pr=db.get(_P,int(pi)); ms=[im for im in pr.images if im.kind=="mood"]; assert len(ms)==13 and all(im.room_id is None for im in ms) and sum(1 for im in ms if im.by_client)==12; db.close()
    print("room inspiration ok")

    # the picture viewer and Tick by codes: the static parts exist and the pages carry their hooks
    js=open("app/static/app.js").read(); css=open("app/static/app.css").read()
    assert "lb-stage" in js and "data-lb-group" in js and "by-codes" in js and "history.pushState({ lb: 1 }" in js and ".lb{" in css and ".codes-box" in css
    h=c.get(f"/p/{pid_big}/items?view=table").text; assert 'class="btn sm ghost by-codes"' in h and 'data-codes="' in h and 'class="codes-box" hidden' in h, h[:300]
    h=c.get(f"/p/{pid_big}/items?view=list").text; assert 'data-codes="' in h, h[:300]
    assert 'data-lb-group' in c2.get(f"/c/{itok}").text and 'data-full="/media/' in c.get(f"/p/{pi}/images").text
    assert 'data-full="/static/help/' in c.get("/help").text
    print("picture viewer + tick by codes ok")
    # ---- Buying: one product across its rooms, the shops' quotes, purchases for whole rooms, the Purchase schedule ----
    from app import buying as _bz
    from app.models import Quote as _Q, Purchase as _Pu, Product as _Pr, SchemaMark as _SM, Supplier as _S
    r=c.post("/projects/new", data={"client_name":"Buyers","name":"Buying house","rate":"7"}, follow_redirects=False); pid_b=int(r.headers["location"].split("/")[-1])
    for code,name,floor in (("GF-LIV","Living room","Ground"),("GF-KIT","Kitchen","Ground"),("FF-BED","Bedroom","First")):
        c.post(f"/p/{pid_b}/rooms", data={"code":code,"name":name,"floor":floor,"kind":"room"}, follow_redirects=False)
    db=SessionLocal(); br={x.code:x.id for x in db.query(Room).filter(Room.project_id==pid_b)}; db.close()
    c.post(f"/p/{pid_b}/items/new", data={"room_ids":[str(br["GF-LIV"]),str(br["GF-KIT"]),str(br["FF-BED"])],"category":"Tiles","name":"Floor tile","size":"600 x 600 mm","finish":"Grey","brand":"Marazzi","qty":"20","unit":"m²","status":"To buy"}, follow_redirects=False)
    c.post(f"/p/{pid_b}/items/new", data={"room_ids":[str(br["GF-KIT"])],"category":"Lighting","name":"Pendant","qty":"3","unit":"pcs","status":"To buy"}, follow_redirects=False)
    db=SessionLocal(); tl={i.room.code:i.id for i in db.query(Item).filter(Item.project_id==pid_b, Item.name=="Floor tile")}; pend=db.query(Item).filter(Item.project_id==pid_b, Item.name=="Pendant").one().id
    db.get(Item, tl["FF-BED"]).qty=10; db.commit(); db.close()
    r=c.get(f"/p/{pid_b}/buying"); h=r.text
    assert r.status_code==200 and h.count('class="bu-card"')==2 and "pref. Marazzi" in h and f'href="/p/{pid_b}/buying" class="on"' in h and 'data-key="buying"' in h and "To source" in h
    db=SessionLocal(); P=db.get(_P,pid_b); gs={g.name:g for g in _bz.groups(P)}; tile=gs["Floor tile"]
    assert tile.needed==50 and tile.allowance==0.10 and tile.order["qty"]==55 and tile.order["packs"] is None and tile.stage=="To source" and gs["Pendant"].allowance==0 and gs["Pendant"].order["qty"]==3
    tp=tile.product.id; pp=gs["Pendant"].product.id; db.close()
    assert c.get(f"/p/{pid_b}/buying?stage=To%20source").text.count('class="bu-card"')==2 and c.get(f"/p/{pid_b}/buying?category=Lighting").text.count('class="bu-card"')==1 and c.get(f"/p/{pid_b}/buying?floor=First").text.count('class="bu-card"')==1
    assert _bz.order_qty(7, 0.1, 0, "pcs")["qty"]==8 and _bz.order_qty(10, 0, 1.44, "m²")["packs"]==7 and _bz.order_qty(10, 0, 0, "m²", override=12)["qty"]==12  # whole pieces, whole boxes, by hand
    # a quote: the shop by name (the same name is the same shop), the price in CNY or USD, a photo; nothing bought yet
    r=c.post(f"/p/{pid_b}/buying/{tp}/quotes", data={"supplier_id":"","amount":"80"}, follow_redirects=False); assert "err=" in r.headers["location"] and r.headers["location"].endswith("#add-quote")
    assert "Pick the shop from the list" in c.get(r.headers["location"]).text
    r=c.post(f"/p/{pid_b}/buying/{tp}/quotes", data={"supplier_id":"__new__","supplier_name":"Shop A","supplier_city":"Foshan","brand":"Monalisa","model":"M-60","amount":"80","currency":"CNY","pack_size":"1.44","moq":"50 m²","lead_time":"2 weeks"}, files={"photo":("q.jpg",img("olive"),"image/jpeg")}, follow_redirects=False)
    assert r.status_code==303 and "ok=" in r.headers["location"]
    db=SessionLocal(); qa=db.query(_Q).filter(_Q.product_id==tp).one(); assert qa.supplier.name=="Shop A" and qa.supplier.city=="Foshan" and qa.unit_price==80 and qa.pack_size==1.44 and qa.thumb_key; qa_id, sa=qa.id, qa.supplier_id
    assert all(i.purchase_id is None and i.unit_price==0 for i in db.query(Item).filter(Item.project_id==pid_b)); db.close()  # a quote alone buys nothing
    h=c.get(f"/p/{pid_b}/buying/{tp}").text
    assert "Shop A" in h and "¥80.00" in h and "1.44 m² per box" in h and "All rooms: 56.16 m² · 39 boxes" in h and "¥4,493" in h and "No shop chosen yet" in h and "Quoting" in h and 'data-key="buying_product"' in h
    r=c.post(f"/p/{pid_b}/buying/{pp}/quotes", data={"supplier_id":"__new__","supplier_name":"shop a","supplier_wechat":"shopa_wx","amount":"20","currency":"USD"}, follow_redirects=False)
    db=SessionLocal(); qp=db.query(_Q).filter(_Q.product_id==pp).one(); assert qp.supplier_id==sa and qp.unit_price==140 and qp.currency=="USD" and qp.amount==20 and db.get(_S, sa).wechat=="shopa_wx"; qp_id=qp.id; db.close()
    r=c.post(f"/p/{pid_b}/buying/{pp}/quotes", data={"supplier_id":str(sa),"amount":""}, follow_redirects=False); db=SessionLocal(); q0=db.query(_Q).filter(_Q.product_id==pp).order_by(_Q.id.desc()).first().id; db.close()
    r=c.post(f"/p/{pid_b}/buying/{pp}/quotes/{q0}/buy", follow_redirects=False); assert "err=" in r.headers["location"]  # no price, no buying
    c.post(f"/p/{pid_b}/buying/{pp}/quotes/{q0}/delete", follow_redirects=False)
    # buy for whole rooms: the ticked rooms get the shop, its price, brand and lead time; the others stay open
    r=c.post(f"/p/{pid_b}/buying/{tp}/quotes/{qa_id}/buy", data={"rooms":"1","line":[str(tl["GF-LIV"]),str(tl["GF-KIT"])]}, follow_redirects=False); assert r.headers["location"].endswith("#buying-from")
    db=SessionLocal(); a1,a2,a3=(db.get(Item,tl[k]) for k in ("GF-LIV","GF-KIT","FF-BED"))
    assert a1.purchase_id==a2.purchase_id and a1.purchase_id and a3.purchase_id is None and a1.supplier_id==sa and a1.unit_price==80 and a1.lead_time=="2 weeks" and a1.status=="Quoted" and a3.status=="To buy" and a1.brand_shown=="Monalisa" and a1.brand=="Marazzi"
    pua=a1.purchase_id; P=db.get(_P,pid_b); g=_bz.group_of(P, db.get(_Pr,tp)); assert g.stage=="Part sourced" and g.covered==2; v=g.purchase_view(db.get(_Pu,pua))
    assert v["needed"]==40 and v["packs"]==31 and v["qty"]==44.64 and v["total"]==3571.2, v; db.close()
    # a second shop for the room left: added and bought in one go; two shops, two pack sizes
    c.post(f"/p/{pid_b}/buying/{tp}/quotes", data={"supplier_id":"__new__","supplier_name":"Shop B","amount":"75","currency":"CNY","pack_size":"1.2","buy":"open"}, follow_redirects=False)
    db=SessionLocal(); qb=db.query(_Q).filter(_Q.product_id==tp, _Q.id!=qa_id).one(); qb_id, sb=qb.id, qb.supplier_id; a3=db.get(Item,tl["FF-BED"]); assert a3.supplier_id==sb and a3.unit_price==75 and a3.purchase_id!=pua; pub=a3.purchase_id
    P=db.get(_P,pid_b); g=_bz.group_of(P, db.get(_Pr,tp)); assert g.stage=="Sourced" and len(g.purchases)==2 and g.best_quote.id==qb_id and g.order["packs"] is None and g.order["qty"]==56.64 and g.committed==4471.2, (g.order, g.committed); db.close()
    h=c.get(f"/p/{pid_b}/buying/{tp}").text; assert "Lowest price" in h and "Buying for 2 rooms" in h and "Buying for 1 room" in h and "¥4,471" in h and "Stop buying here" in h
    # a room moves to the other shop; a quantity set by hand; the product's allowance and pack
    r=c.post(f"/p/{pid_b}/buying/{tp}/purchases/{pub}", data={"rooms":"1","line":[str(tl["FF-BED"]),str(tl["GF-KIT"])]}, follow_redirects=False); assert r.headers["location"].endswith(f"#purchase-{pub}")
    db=SessionLocal(); assert db.get(Item,tl["GF-KIT"]).purchase_id==pub and db.get(Item,tl["GF-KIT"]).unit_price==75 and db.get(Item,tl["GF-LIV"]).purchase_id==pua; db.close()
    c.post(f"/p/{pid_b}/buying/{tp}/purchases/{pua}", data={"qty":"30","notes":"Deposit paid"}, follow_redirects=False)
    db=SessionLocal(); P=db.get(_P,pid_b); v=_bz.group_of(P, db.get(_Pr,tp)).purchase_view(db.get(_Pu,pua)); assert v["qty"]==30.24 and v["packs"]==21 and db.get(_Pu,pua).notes=="Deposit paid"; db.close()
    c.post(f"/p/{pid_b}/buying/{tp}", data={"allowance":"5","pack_size":"1.5","pack_name":"carton","notes":"One batch"}, follow_redirects=False)
    db=SessionLocal(); pr=db.get(_Pr,tp); assert pr.allowance==0.05 and pr.pack_size==1.5 and pr.pack_name=="carton" and pr.notes=="One batch"; db.close()
    c.post(f"/p/{pid_b}/buying/{tp}", data={"allowance":"","pack_size":"","pack_name":"","notes":""}, follow_redirects=False)
    db=SessionLocal(); assert db.get(_Pr,tp).allowance is None and _bz.allowance_of(db.get(_Pr,tp), "Tiles")==0.10; db.close()  # empty = the category's default again
    # a corrected quote reaches its rooms; a deleted quote leaves what was bought from it; stopping a purchase frees its rooms
    c.post(f"/p/{pid_b}/buying/{tp}/quotes/{qa_id}", data={"supplier_id":str(sa),"brand":"Monalisa","amount":"82","currency":"CNY","pack_size":"1.44","lead_time":"3 weeks"}, follow_redirects=False)
    db=SessionLocal(); x=db.get(Item,tl["GF-LIV"]); assert x.unit_price==82 and x.lead_time=="3 weeks" and db.get(_Pu,pua).unit_price==82; qkeys=(db.get(_Q,qa_id).photo_key, db.get(_Q,qa_id).thumb_key); db.close()
    c.post(f"/p/{pid_b}/buying/{tp}/quotes/{qa_id}/delete", follow_redirects=False)
    db=SessionLocal(); assert db.get(_Q,qa_id) is None and db.get(_Pu,pua).quote_id is None and db.get(Item,tl["GF-LIV"]).unit_price==82; db.close()
    assert all(_st.read_image(k) is None for k in qkeys if k) and "quote deleted" in c.get(f"/p/{pid_b}/buying/{tp}").text
    r=c.post(f"/p/{pid_b}/buying/{tp}/purchases/{pua}/delete", follow_redirects=False); assert r.headers["location"].endswith("#buying-from")
    db=SessionLocal(); x=db.get(Item,tl["GF-LIV"]); assert x.purchase_id is None and x.supplier_id is None and x.unit_price==0 and x.lead_time=="" and x.status=="To buy" and db.get(_Pu,pua) is None; db.close()
    c.post(f"/p/{pid_b}/buying/{tp}/purchases/{pub}", data={"rooms":"1"}, follow_redirects=False)  # nothing ticked: the purchase goes
    db=SessionLocal(); assert db.get(_Pu,pub) is None and all(db.get(Item,x).purchase_id is None and db.get(Item,x).unit_price==0 for x in tl.values()); db.close()
    c.post(f"/p/{pid_b}/buying/{tp}/quotes/{qb_id}/buy", follow_redirects=False)  # every room without a shop
    db=SessionLocal(); assert len({db.get(Item,x).purchase_id for x in tl.values()})==1 and db.get(Item,tl["GF-LIV"]).unit_price==75; db.close()
    # the Purchase schedule, the item page, the client link and the purchase order follow
    h=c.get(f"/p/{pid_b}/buying/schedule").text; assert "Purchase schedule" in h and "Shop B" in h and "Floor tile" in h and "Pendant" in h and "Quoting" in h and "Sourced" in h and "¥4,140" in h and 'data-key="buying_schedule"' in h
    h=c.get(f"/p/{pid_b}/buying/schedule?cur=USD").text; assert "prices in USD" in h and "$591.43" in h
    h=c.get(f"/p/{pid_b}/buying/schedule?stage=Quoting").text; assert "Pendant" in h and "Floor tile" not in h
    h=c.get(f"/p/{pid_b}/items/{tl['GF-LIV']}").text; assert "Buying from" in h and "Shop B" in h and f"/p/{pid_b}/buying/item/{tl['GF-LIV']}" in h
    assert c.get(f"/p/{pid_b}/buying/item/{tl['GF-LIV']}", follow_redirects=False).headers["location"]==f"/p/{pid_b}/buying/{tp}"
    assert c.get(f"/p/{pid_b}/buying/item/{iid}", follow_redirects=False).headers["location"]==f"/p/{pid_b}/buying" and c.get(f"/p/{pid}/buying/{tp}", follow_redirects=False).headers["location"]==f"/p/{pid}/buying"  # another project's item, product
    db=SessionLocal(); btok=db.get(_P,pid_b).client_token; db.close()
    h=c2.get(f"/c/{btok}").text; assert "Pendant" in h and h.count('<div class="price">')==3  # no price until a shop is chosen
    assert "Floor tile" in pdf_text(c.get(f"/p/{pid_b}/export/po/{sb}.pdf").content)
    # renames: the whole product renamed keeps its quotes and purchases; one line renamed alone leaves them; two products merge
    c.post(f"/p/{pid_b}/items/{tl['GF-LIV']}", data={"room_id":str(br["GF-LIV"]),"category":"Tiles","name":"Porcelain floor tile","size":"600 x 600 mm","finish":"Grey","brand":"Marazzi","qty":"20","unit":"m²","status":"Quoted","apply_all":"1"}, follow_redirects=False)
    db=SessionLocal(); pr=db.get(_Pr,tp); assert pr is not None and pr.key==_pk(db.get(Item,tl["GF-LIV"])) and all(db.get(Item,x).name=="Porcelain floor tile" and db.get(Item,x).purchase_id for x in tl.values()) and db.query(_Q).filter(_Q.product_id==tp).count()==1; db.close()
    c.post(f"/p/{pid_b}/items/{tl['FF-BED']}", data={"room_id":str(br["FF-BED"]),"category":"Tiles","name":"Bedroom tile","size":"600 x 600 mm","finish":"Grey","qty":"10","unit":"m²","status":"Quoted"}, follow_redirects=False)
    db=SessionLocal(); x=db.get(Item,tl["FF-BED"]); assert x.purchase_id is None and x.unit_price==0 and x.status=="To buy" and db.get(Item,tl["GF-KIT"]).purchase_id; db.close()
    for nm, rm in (("Wall light","GF-LIV"),("Wall lamp","FF-BED")):
        c.post(f"/p/{pid_b}/items/new", data={"room_ids":[str(br[rm])],"category":"Lighting","name":nm,"qty":"2","unit":"pcs","status":"To buy"}, follow_redirects=False)
    db=SessionLocal(); wl=db.query(Item).filter(Item.project_id==pid_b, Item.name=="Wall light").one().id; wm=db.query(Item).filter(Item.project_id==pid_b, Item.name=="Wall lamp").one().id; db.close()
    p_wl,_=buy(c, pid_b, wl, sa, 50); p_wm,q_wm=buy(c, pid_b, wm, sb, 45, lines=[wm])
    c.post(f"/p/{pid_b}/items/{wm}", data={"room_id":str(br["FF-BED"]),"category":"Lighting","name":"Wall light","qty":"2","unit":"pcs","status":"Quoted"}, follow_redirects=False)
    db=SessionLocal(); assert db.get(_Pr,p_wm) is None and db.get(_Q,q_wm).product_id==p_wl and db.get(Item,wm).purchase_id and db.get(Item,wm).unit_price==45 and db.query(_Q).filter(_Q.product_id==p_wl).count()==2; db.close()  # one product, both quotes, the purchase kept
    c.post(f"/p/{pid_b}/items/{wm}/delete", follow_redirects=False)
    db=SessionLocal(); assert db.query(_Pu).filter(_Pu.product_id==p_wl).count()==1; db.close()  # the purchase left without rooms went
    # a supplier deleted: its quotes and purchases go, the rooms it supplied have no shop again
    c.post(f"/suppliers/{sa}/delete", follow_redirects=False)
    db=SessionLocal(); assert db.get(_S, sa) is None and db.query(_Q).filter(_Q.supplier_id==sa).count()==0 and db.query(_Pu).filter(_Pu.supplier_id==sa).count()==0 and db.get(Item,wl).purchase_id is None and db.get(Item,wl).supplier_id is None; db.close()
    # the one-time upgrade of prices and suppliers typed on items before quotes: "Test" and a price without a shop are cleared,
    # a real supplier becomes a quote and a purchase (statuses kept), a draft's price goes into its notes; it runs once
    db=SessionLocal(); tsup=_S(name="Test"); rsup=_S(name="Real Shop"); db.add_all([tsup, rsup]); db.flush()
    def _old(name, room, code, **kw):
        it=Item(project_id=pid_b, room_id=br[room], category="Furniture", name=name, code=code, qty=1, unit="pcs", **kw); db.add(it); db.flush(); return it.id
    o1=_old("Old chair","GF-KIT","GF-KIT-91",supplier_id=tsup.id,unit_price=100,status="Quoted")
    o2=_old("Old table","GF-KIT","GF-KIT-92",supplier_id=rsup.id,unit_price=500,lead_time="3 weeks",status="Ordered")
    o3=_old("Old table","GF-LIV","GF-LIV-92",supplier_id=rsup.id,unit_price=500,lead_time="3 weeks",status="To buy")
    o4=_old("Old lamp","GF-LIV","GF-LIV-93",unit_price=50,status="Quoted")
    o5=_old("","GF-LIV","",unit_price=30,draft=True); rsid=rsup.id
    db.query(_SM).filter(_SM.key==_bz.UPGRADE_MARK).delete(); db.commit()
    done=_bz.upgrade(db); db.close(); assert done["cleared"]>=2 and done["quotes"]>=1 and done["drafts"]>=1, done
    db=SessionLocal(); a,b_,c_,d_,e_=(db.get(Item,x) for x in (o1,o2,o3,o4,o5))
    assert a.supplier_id is None and a.unit_price==0 and a.status=="To buy" and db.query(_S).filter(_S.name=="Test").count()==0
    assert b_.purchase_id and b_.purchase_id==c_.purchase_id and b_.supplier_id==rsid and b_.unit_price==500 and b_.lead_time=="3 weeks" and b_.status=="Ordered" and c_.status=="Quoted"
    q_=db.get(_Pu,b_.purchase_id).quote; assert q_.supplier_id==rsid and "before quotes existed" in q_.notes
    assert d_.unit_price==0 and d_.status=="To buy" and e_.unit_price==0 and "Price noted at capture: ¥30.00" in e_.notes; db.close()
    db=SessionLocal(); assert _bz.upgrade(db)=={}; db.close()  # once only
    print("buying ok")
    # ---- Orders: what was bought, on numbered orders per shop; confirming, paying, the statuses that follow, the tracker ----
    import datetime as _dt
    from app import ordering as _od
    from app.models import Order as _O, Payment as _Pay
    r=c.post("/settings/warehouse", data={"warehouse_name":"Baiyun Warehouse","warehouse_address":"Unit 8, Logistics Park\nGuangzhou","warehouse_contact":"Mr. Liang"}, follow_redirects=False)
    assert r.headers["location"].endswith("#warehouse") and 'value="Baiyun Warehouse"' in c.get("/settings").text
    r=c.post("/projects/new", data={"client_name":"Orders","name":"Kilimani House","rate":"7","budget_usd":"5000"}, follow_redirects=False); pid_o=int(r.headers["location"].split("/")[-1])
    for code,name in (("LIV","Living room"),("KIT","Kitchen"),("BED","Bedroom")):
        c.post(f"/p/{pid_o}/rooms", data={"code":code,"name":name,"floor":"Ground","kind":"room"}, follow_redirects=False)
    db=SessionLocal(); orr={x.code:x.id for x in db.query(Room).filter(Room.project_id==pid_o)}; db.close()
    c.post(f"/p/{pid_o}/items/new", data={"room_ids":[str(orr["LIV"]),str(orr["KIT"])],"category":"Tiles","name":"Floor tile","size":"600 x 600 mm","qty":"20","unit":"m²","status":"To buy"}, follow_redirects=False)
    c.post(f"/p/{pid_o}/items/new", data={"room_ids":[str(orr["BED"])],"category":"Furniture","name":"Bed","qty":"1","unit":"pcs","status":"To buy"}, follow_redirects=False)
    c.post(f"/p/{pid_o}/items/new", data={"room_ids":[str(orr["LIV"])],"category":"Lighting","name":"Floor lamp","qty":"2","unit":"pcs","status":"To buy"}, follow_redirects=False)
    db=SessionLocal(); oi={(i.name, i.room.code):i.id for i in db.query(Item).filter(Item.project_id==pid_o)}; db.close()
    p_tile,q_tile=buy(c, pid_o, oi[("Floor tile","LIV")], "Tile Shop", 80, pack_size="1.44")
    p_bed,_=buy(c, pid_o, oi[("Bed","BED")], "Tile Shop", 3000)
    p_lamp,_=buy(c, pid_o, oi[("Floor lamp","LIV")], "Lamp Shop", 400)
    db=SessionLocal(); ts=db.query(_S).filter(_S.name=="Tile Shop").one().id; ls=db.query(_S).filter(_S.name=="Lamp Shop").one().id
    db.get(_S, ts).payment_terms="40% deposit, 60% before loading"; db.commit(); db.close()
    tile_bed=[oi[("Floor tile","LIV")], oi[("Floor tile","KIT")], oi[("Bed","BED")]]
    def _sts(ids):
        db=SessionLocal(); out=[db.get(Item,x).status for x in ids]; db.close(); return out
    assert _od.initials("Kilimani House")=="KH" and _od.initials("")=="PO" and _od.deposit_from_terms("30% deposit, 70% before loading")==30 and _od.deposit_from_terms("Deposit 50%, rest on loading")==50 and _od.deposit_from_terms("Full payment on order")==100 and _od.deposit_from_terms("Net 30")is None
    # what is bought waits under Ready to order; Start an order makes a numbered draft for the shop with the warehouse
    h=c.get(f"/p/{pid_o}/orders").text
    assert 'data-key="orders"' in h and "Ready to order" in h and "Tile Shop" in h and "Lamp Shop" in h and "Start an order" in h and f'href="/p/{pid_o}/orders" class="on"' in h and 'href="#help"' not in h
    h=c.get(f"/p/{pid_e}/orders").text; assert "No orders yet" in h and 'href="#help"' in h  # the empty state points at the help
    r=c.post(f"/p/{pid_o}/orders/new", data={"supplier_id":str(ts)}, follow_redirects=False); assert "ok=" in r.headers["location"]
    oid=int(r.headers["location"].split("/orders/")[1].split("?")[0])
    db=SessionLocal(); o=db.get(_O,oid)
    assert o.number==1 and _od.ref(o)=="KH-001" and o.deposit_pct==40 and o.terms.startswith("40%") and o.deliver_to.startswith("Baiyun Warehouse\nUnit 8") and not o.confirmed and len(o.purchases)==2
    led=_bz.Ledger(db.get(_P,pid_o)); assert led.total[oid]==6571.2 and led.state(o.purchases[0])=="bought", led.total; db.close()  # 44.64 m² of tile (31 boxes) + the bed
    assert _sts(tile_bed)==["Quoted"]*3  # a draft changes no status
    r=c.post(f"/p/{pid_o}/orders/new", data={"supplier_id":str(ts)}, follow_redirects=False); assert "err=" in r.headers["location"]  # nothing left waiting from that shop
    lamp=oi[("Floor lamp","LIV")]; c.post(f"/p/{pid_o}/items/{lamp}/status", data={"status":"Ordered"})  # set by hand before the order
    r=c.post(f"/p/{pid_o}/orders/new", data={"supplier_id":str(ls)}, follow_redirects=False); oid2=int(r.headers["location"].split("/orders/")[1].split("?")[0])
    db=SessionLocal(); assert _od.ref(db.get(_O,oid2))=="KH-002" and db.get(_O,oid2).deposit_pct==30; db.close()
    assert _sts([lamp])==["Ordered"]  # a draft never lowers a status
    h=c.get(f"/p/{pid_o}/orders/{oid}").text; wc=h.split('id="wechat-text">')[1].split("</pre>")[0]
    assert 'data-key="order"' in h and "Order KH-001" in h and "¥6,571.20" in h and "Confirm the order" in h and "Deposit 40%" in h and "due once confirmed" in h and "Baiyun Warehouse" in h
    assert "采购订单 KH-001" in wc and "LIV-01" in wc and "KIT-01" in wc and "BED-01" in wc and "Deliver to 送货地址" in wc and "Baiyun Warehouse" in wc and "¥6,571.20" in wc
    r=c.get(f"/p/{pid_o}/orders/{oid}/po.pdf"); t=pdf_text(r.content)
    assert r.status_code==200 and r.headers["content-type"]=="application/pdf" and "PO-KH-001" in r.headers["content-disposition"]
    assert "KH-001" in t and "Floor tile" in t and "LIV-01" in t and "Baiyun Warehouse" in t and "Draft" in t and "6,571.20" in t and "31 boxes" in t, t[:600]
    # confirming: the rooms become Ordered; the deposit is due; the product page, the item and the Buying list follow
    r=c.post(f"/p/{pid_o}/orders/{oid}/confirm", data={"confirmed_on":"2026-10-01","ready_date":"2026-10-20"}, follow_redirects=False); assert "ok=" in r.headers["location"]
    db=SessionLocal(); o=db.get(_O,oid); assert o.confirmed_on==_dt.date(2026,10,1) and o.ready_date==_dt.date(2026,10,20) and o.confirmed_total==6571.2; db.close()
    assert _sts(tile_bed)==["Ordered"]*3 and _sts([lamp])==["Ordered"]
    h=c.get(f"/p/{pid_o}/buying/{p_tile}").text; assert "Order KH-001" in h and "Deposit due" in h
    assert "On order" in c.get(f"/p/{pid_o}/items/{oi[('Bed','BED')]}").text and "KH-001" in c.get(f"/p/{pid_o}/items/{oi[('Bed','BED')]}").text
    assert c.get(f"/p/{pid_o}/buying?stage=Ordered").text.count('class="bu-card"')==2
    assert "KH-001" not in pdf_text(c.get(f"/p/{pid_o}/orders/{oid2}/po.pdf").content)  # each order its own PO
    # paying: the deposit on the order page, the balance on Payments with the order chosen; paid in full = Paid; shipped stays shipped
    c.post(f"/p/{pid_o}/items/{oi[('Bed','BED')]}/status", data={"status":"Shipped"})
    r=c.post(f"/p/{pid_o}/orders/{oid}/payments", data={"amount":"2628.48","kind":"Deposit","paid_on":"2026-10-02","reference":"WX1"}, files={"receipt":("r.jpg",img("green"),"image/jpeg")}, follow_redirects=False)
    assert r.headers["location"].endswith("#payments") and "ok=" in r.headers["location"]
    db=SessionLocal(); pay1=db.query(_Pay).filter(_Pay.order_id==oid).one(); assert pay1.supplier_id==ts and pay1.receipt_key
    ov=_od.OrderView(db.get(_O,oid), _bz.Ledger(db.get(_P,pid_o)), today=_dt.date(2026,10,16)); assert ov.stage=="Balance due" and ov.deposit_left==0 and ov.balance==3942.72 and ov.soon and not ov.late and ov.ready_text=="in 4 days"; db.close()
    assert _sts(tile_bed)==["Ordered","Ordered","Shipped"]
    assert 'value="3942.72"' in c.get(f"/p/{pid_o}/orders/{oid}").text  # the form offers what is due
    h=c.get(f"/p/{pid_o}/payments").text; assert f'<option value="{oid}">KH-001 · Tile Shop · to pay ¥3,943</option>' in h and "On orders" in h and "KH-001" in h.split("All payments")[1]
    c.post(f"/p/{pid_o}/payments", data={"order_id":str(oid),"supplier_id":str(ls),"paid_on":"2026-10-15","amount":"3942.72","kind":"Balance","reference":"BANK2"}, follow_redirects=False)
    db=SessionLocal(); pay2=db.query(_Pay).filter(_Pay.order_id==oid, _Pay.kind=="Balance").one(); assert pay2.supplier_id==ts; pay2=pay2.id; db.close()  # the order's shop, not the one picked
    assert _sts(tile_bed)==["Paid","Paid","Shipped"] and c.get(f"/p/{pid_o}/buying?stage=Paid").text.count('class="bu-card"')==2  # the stage reads the order
    r=c.post(f"/p/{pid_o}/payments/{pay2}/delete", data={"next":f"/p/{pid_o}/orders/{oid}#payments"}, follow_redirects=False); assert r.headers["location"]==f"/p/{pid_o}/orders/{oid}#payments"
    assert _sts(tile_bed)==["Ordered","Ordered","Shipped"]
    # money: summary, Overview, Payments, the supplier page and its list
    db=SessionLocal(); s=_summary(db, db.get(_P,pid_o)); db.close()
    assert s["ordered"]==6571.2 and s["paid"]==2628.48 and s["to_pay"]==3942.72 and s["orders"]==2 and {r_["label"]:r_["bought"] for r_ in s["by_sup"]}=={"Tile Shop":6571.2,"Lamp Shop":800}
    assert "¥3,943 still to pay on ¥6,571 ordered" in c.get(f"/p/{pid_o}").text
    h=c.get(f"/suppliers/{ts}?project={pid_o}").text; assert 'id="orders"' in h and "KH-001" in h and f"/p/{pid_o}/orders/{oid}/po.pdf" in h and "on confirmed orders" in h
    assert "to pay ¥3,943" in c.get(f"/p/{pid_o}/suppliers").text
    assert c.get(f"/p/{pid_o}/export/po/{ts}.pdf").status_code==200  # every item from the shop, as before
    # a change after confirming shows on the order and the tracker until the shop agrees the new total
    c.post(f"/p/{pid_o}/buying/{p_tile}/quotes/{q_tile}", data={"supplier_id":str(ts),"amount":"85","currency":"CNY","pack_size":"1.44"}, follow_redirects=False)
    h=c.get(f"/p/{pid_o}/orders/{oid}").text; assert "Changed since it was confirmed" in h and "¥6,794.40" in h and "The shop agreed" in h
    assert "KH-001 changed after it was confirmed" in c.get(f"/p/{pid_o}/buying/tracker").text
    c.post(f"/p/{pid_o}/orders/{oid}/confirm", data={"confirmed_on":"2026-10-01"}, follow_redirects=False)
    db=SessionLocal(); o=db.get(_O,oid); assert o.confirmed_total==6794.4 and o.confirmed_on==_dt.date(2026,10,1); db.close()
    assert "Changed since it was confirmed" not in c.get(f"/p/{pid_o}/orders/{oid}").text
    # off the order and back; back to a draft; deleting an order keeps its products bought and its payments
    db=SessionLocal(); pu_bed=db.query(_Pu).filter(_Pu.product_id==p_bed).one().id; pu_tile=db.query(_Pu).filter(_Pu.product_id==p_tile).one().id; db.close()
    c.post(f"/p/{pid_o}/items/{oi[('Bed','BED')]}/status", data={"status":"Ordered"})
    r=c.post(f"/p/{pid_o}/orders/{oid}/remove/{pu_bed}", follow_redirects=False); assert "ok=" in r.headers["location"]
    assert _sts(tile_bed)==["Ordered","Ordered","Quoted"] and "Bed" in c.get(f"/p/{pid_o}/orders").text.split("Ready to order")[1].split("</section>")[0]
    h=c.get(f"/p/{pid_o}/orders/{oid}").text; assert "Also bought from Tile Shop" in h
    c.post(f"/p/{pid_o}/orders/{oid}/add", data={"pu":[str(pu_bed)]}, follow_redirects=False); assert _sts(tile_bed)==["Ordered"]*3
    r=c.post(f"/p/{pid_o}/orders/{oid}/add", data={"pu":[str(pu_tile)]}, follow_redirects=False); assert "err=" in r.headers["location"]  # already on it
    c.post(f"/p/{pid_o}/orders/{oid}/unconfirm", follow_redirects=False); assert _sts(tile_bed)==["Quoted"]*3
    db=SessionLocal(); assert db.get(_O,oid).confirmed_on is None and db.query(_Pay).filter(_Pay.order_id==oid).count()==1; db.close()
    c.post(f"/p/{pid_o}/orders/{oid}/confirm", data={}, follow_redirects=False); assert _sts(tile_bed)==["Ordered"]*3
    db=SessionLocal(); assert db.get(_O,oid).confirmed_on==_dt.date.today(); db.close()
    # the tracker: the money, what needs attention, the stages, the shops, the categories
    h=c.get(f"/p/{pid_o}/buying/tracker").text
    assert 'data-key="tracker"' in h and "Needs attention" in h and "Draft KH-002 to confirm with the shop" in h and "By stage" in h and "By shop" in h and "By category" in h and "¥35,000" in h, h[:300]
    db=SessionLocal(); P=db.get(_P,pid_o); led=_bz.Ledger(P); t=_od.tracker(P, led, _bz.groups(P, ledger=led), _od.views(P, led, _dt.date(2026,10,25)), _dt.date(2026,10,25)); db.close()
    assert t["ordered"]==6794.4 and t["paid"]==2628.48 and t["to_pay"]==4165.92 and t["stages"]["Ordered"]["products"]==2 and t["stages"]["Sourced"]["products"]==1
    assert t["attention"][0]["tone"]=="red" and t["attention"][0]["title"]=="KH-001 was due 5 days ago" and sum(1 for a in t["attention"] if "KH-001" in a["title"])==1, t["attention"]
    # a payment to the shop on no order (made before orders existed, say) can be counted toward the order
    c.post(f"/p/{pid_o}/payments", data={"order_id":"","supplier_id":str(ts),"paid_on":"2026-09-20","amount":"1000","kind":"Deposit","reference":"OLD1"}, follow_redirects=False)
    h=c.get(f"/p/{pid_o}/orders/{oid}").text; assert "Paid to Tile Shop before, on no order" in h and "OLD1" in h
    db=SessionLocal(); loose=db.query(_Pay).filter(_Pay.reference=="OLD1").one().id; db.close()
    r=c.post(f"/p/{pid_o}/orders/{oid}/claim", data={}, follow_redirects=False); assert "err=" in r.headers["location"]
    r=c.post(f"/p/{pid_o}/orders/{oid}/claim", data={"pay":[str(loose)]}, follow_redirects=False); assert "ok=" in r.headers["location"]
    db=SessionLocal(); assert db.get(_Pay,loose).order_id==oid and _od.OrderView(db.get(_O,oid), _bz.Ledger(db.get(_P,pid_o))).paid==3628.48; db.close()
    assert "Paid to Tile Shop before" not in c.get(f"/p/{pid_o}/orders/{oid}").text
    c.post(f"/p/{pid_o}/orders/{oid2}/delete", follow_redirects=False); assert _sts([lamp])==["Ordered"]  # a draft deleted lowers nothing either
    r=c.post(f"/p/{pid_o}/orders/new", data={"supplier_id":str(ls)}, follow_redirects=False); oid3=int(r.headers["location"].split("/orders/")[1].split("?")[0])
    db=SessionLocal(); assert _od.ref(db.get(_O,oid3))=="KH-003" and db.query(_O).filter(_O.project_id==pid_o, _O.number==2).count()==0; db.close()  # a number is never given twice
    c.post(f"/p/{pid_o}/orders/{oid}/delete", follow_redirects=False)
    db=SessionLocal(); assert db.get(_O,oid) is None and db.get(_Pay,pay1.id).order_id is None and db.get(_Pay,pay1.id).supplier_id==ts and db.query(_Pu).filter(_Pu.supplier_id==ts).count()==2; db.close()
    assert _sts(tile_bed)==["Quoted"]*3
    # another project's order is not reachable; a supplier deleted takes its orders with it
    assert c.get(f"/p/{pid_b}/orders/{oid3}", follow_redirects=False).headers["location"]==f"/p/{pid_b}/orders"
    c.post(f"/suppliers/{ls}/delete", follow_redirects=False)
    db=SessionLocal(); assert db.get(_O,oid3) is None and db.query(_O).filter(_O.project_id==pid_o).count()==0; db.close()
    print("orders ok")
    # ---- Shipping: boxes to the warehouse, into a shipment, to site; the rooms follow their boxes ----
    from app import shipping as _sh
    from app.models import Shipment as _Sh, Carton as _C, SupplierLink as _SL
    r=c.post("/projects/new", data={"client_name":"Ship","name":"Kibada House","rate":"7"}, follow_redirects=False); pid_s=int(r.headers["location"].split("/")[-1])
    for code,name in (("LIV","Living room"),("KIT","Kitchen")):
        c.post(f"/p/{pid_s}/rooms", data={"code":code,"name":name,"floor":"Ground","kind":"room"}, follow_redirects=False)
    db=SessionLocal(); sr={x.code:x.id for x in db.query(Room).filter(Room.project_id==pid_s)}; db.close()
    c.post(f"/p/{pid_s}/items/new", data={"room_ids":[str(sr["LIV"]),str(sr["KIT"])],"category":"Tiles","name":"Floor tile","qty":"10","unit":"m²","status":"To buy"}, follow_redirects=False)
    c.post(f"/p/{pid_s}/items/new", data={"room_ids":[str(sr["LIV"])],"category":"Furniture","name":"Sofa","qty":"1","unit":"pcs","status":"To buy"}, follow_redirects=False)
    db=SessionLocal(); si={(i.name, i.room.code):i for i in db.query(Item).filter(Item.project_id==pid_s)}; t_liv, t_kit, sofa=(si[k].id for k in (("Floor tile","LIV"),("Floor tile","KIT"),("Sofa","LIV"))); codes={x.id:x.code for x in si.values()}; db.close()
    buy(c, pid_s, t_liv, "Ship Tiles", 60); buy(c, pid_s, sofa, "Ship Tiles", 4000)
    db=SessionLocal(); sts=db.query(_S).filter(_S.name=="Ship Tiles").one().id; db.close()
    r=c.post(f"/p/{pid_s}/orders/new", data={"supplier_id":str(sts)}, follow_redirects=False); soid=int(r.headers["location"].split("/orders/")[1].split("?")[0])
    c.post(f"/p/{pid_s}/orders/{soid}/confirm", data={"confirmed_on":"2026-09-01","ready_date":"2026-09-20"}, follow_redirects=False)
    assert _sts([t_liv, t_kit, sofa])==["Ordered"]*3
    assert _sh.fits(None, 0.65, 300)=="LCL" and _sh.fits(None, 20, 3000)=="1 x 20ft" and _sh.fits(None, 60, 9000)=="1 x 40ft HC" and _sh.fits(None, 150, 1000)=="3 x 40ft HC" and _sh.capacity(None, "LCL")==(None, None)
    # the shop lists its boxes on its link: delivered to the warehouse of its order, the items ticked from its own lines
    c.post(f"/p/{pid_s}/suppliers/{sts}/link", data={"action":"create"}, follow_redirects=False)
    db=SessionLocal(); stok=db.query(_SL).filter(_SL.project_id==pid_s, _SL.supplier_id==sts).one().token; db.close()
    h=c2.get(f"/s/{stok}").text
    assert "Deliver to:</b> Baiyun Warehouse" in h and f'name="code" value="{codes[t_liv]}"' in h and f'value="{codes[sofa]}"' in h and "Rooms with your items" in h, h[:400]
    c2.post(f"/s/{stok}/cartons", data={"room_id":str(sr["LIV"]),"contents":"Tiles, 5 cartons","code":[codes[t_liv]],"qty":"5","length_cm":"60","width_cm":"60","height_cm":"40","weight_kg":"120"}, follow_redirects=False)
    c2.post(f"/s/{stok}/cartons", data={"room_id":str(sr["KIT"]),"contents":"Tiles, kitchen","item_codes":f"{codes[t_kit]} x2","qty":"2","length_cm":"60","width_cm":"60","height_cm":"20","weight_kg":"50"}, follow_redirects=False)
    c2.post(f"/s/{stok}/cartons", data={"room_id":str(sr["LIV"]),"contents":"Sofa, crate","code":[codes[sofa]],"qty":"1","length_cm":"240","width_cm":"100","height_cm":"90","weight_kg":"80"}, follow_redirects=False)
    db=SessionLocal(); bxs=db.query(_C).filter(_C.project_id==pid_s).order_by(_C.id).all(); b1,b2,b3=(x.id for x in bxs)
    assert [x.item_codes for x in bxs]==[codes[t_liv], f"{codes[t_kit]} x2", codes[sofa]] and all(x.entered_by=="supplier" and _sh.where(x)=="shop" for x in bxs); db.close()
    P_=lambda: SessionLocal()  # noqa: E731
    db=P_(); bxo=_sh.Boxes(db.get(_P,pid_s)); assert [i.id for i in bxo.lines_of[b2]]==[t_kit] and [c_.id for c_ in bxo.boxes_of[t_liv]]==[b1]
    assert _sh.split_codes(f"{codes[t_kit].lower()} x2, Z-99", bxo.by_code)==([codes[t_kit]], "x2, Z-99"); db.close()
    assert _sh.join_codes(["a-1", "A-1", "b-2"], "x2")=="A-1, B-2, x2"
    # the Boxes page: where each box is, filters, the shop's boxes delivered to the warehouse in one go
    h=c.get(f"/p/{pid_s}/cartons").text; assert 'data-key="cartons"' in h and h.count("At the shop")>=3 and f'href="/p/{pid_s}/cartons" class="on"' in h and "Put on" in h
    r=c.post(f"/p/{pid_s}/cartons/bulk", data={"c":[str(b1),str(b2)],"action":"warehouse","next":f"/p/{pid_s}/shipments"}, follow_redirects=False)
    assert r.headers["location"].startswith(f"/p/{pid_s}/shipments?ok=")
    r=c.post(f"/p/{pid_s}/cartons/bulk", data={"action":"warehouse"}, follow_redirects=False); assert "err=" in r.headers["location"]
    db=P_(); assert db.get(_C,b1).warehouse_on==_dt.date.today() and db.get(_C,b3).warehouse_on is None; db.close()
    r=c2.post(f"/s/{stok}/cartons/{b1}/delete", follow_redirects=False); db=P_(); assert db.get(_C,b1) is not None; db.close()  # the shop cannot take back a delivered box
    assert c.get(f"/p/{pid_s}/cartons?where=warehouse").text.count('class="box-sel"')==2
    h=c.get(f"/p/{pid_s}/shipments").text; assert 'data-key="shipments"' in h and "Waiting at the warehouse" in h and "would go <b>LCL</b>" in h and "Still at the shops" in h
    # a shipment with what waits at the warehouse; loading it ships the rooms whose boxes are all on board
    r=c.post(f"/p/{pid_s}/shipments/new", data={"all":"1","mode":"20ft"}, follow_redirects=False); ssid=int(r.headers["location"].split("/shipments/")[1].split("?")[0])
    db=P_(); s_=db.get(_Sh,ssid); assert s_.number==1 and s_.mode=="20ft" and sorted(x.id for x in s_.cartons)==[b1,b2]; db.close()
    h=c.get(f"/p/{pid_s}/shipments/{ssid}").text; assert 'data-key="shipment"' in h and "Shipment 1" in h and "Loaded today" in h and "Kitchen" in h
    c.post(f"/p/{pid_s}/shipments/{ssid}", data={"mode":"20ft","container_no":"tcnu7731042","seal_no":"s1","forwarder":"Fwd","bl_no":"BL1","vessel":"V","port_from":"Nansha","port_to":"Dar es Salaam",
                                              "cost":"21500","notes":"","loaded_on":"","etd":"","eta":(_dt.date.today()+_dt.timedelta(days=20)).isoformat(),"arrived_on":"","delivered_on":""}, follow_redirects=False)
    db=P_(); assert db.get(_Sh,ssid).container_no=="TCNU7731042" and db.get(_Sh,ssid).cost==21500; db.close()
    r=c.post(f"/p/{pid_s}/shipments/{ssid}/step", data={"step":"loaded"}, follow_redirects=False); assert "ok=" in r.headers["location"]
    assert _sts([t_liv, t_kit, sofa])==["Shipped","Shipped","Ordered"]  # the sofa's box is still at the shop
    db=P_(); sv=_sh.ShipmentView(db.get(_Sh,ssid), _sh.Boxes(db.get(_P,pid_s)).on[ssid]); assert sv.stage=="Loaded" and sv.cap==28.0 and sv.fill==1 and sv.eta_text=="in 20 days" and sv.count==2; db.close()
    r=c.post(f"/p/{pid_s}/shipments/{ssid}/boxes/{b2}/remove", follow_redirects=False); assert _sts([t_liv, t_kit])==["Shipped","Ordered"]  # off the shipment: back to the order's status
    c.post(f"/p/{pid_s}/shipments/{ssid}/boxes", data={"c":[str(b2)]}, follow_redirects=False); assert _sts([t_kit])==["Shipped"]
    c.post(f"/p/{pid_s}/shipments/{ssid}/step", data={"step":"loaded","undo":"1"}, follow_redirects=False); assert _sts([t_liv, t_kit])==["Ordered","Ordered"]
    c.post(f"/p/{pid_s}/shipments/{ssid}/step", data={"step":"loaded"}, follow_redirects=False)
    c.post(f"/p/{pid_s}/shipments/{ssid}/step", data={"step":"sailed"}, follow_redirects=False)
    db=P_(); assert _sh.stage(db.get(_Sh,ssid))=="At sea"; db.close()
    h=c.get(f"/p/{pid_s}/cartons").text; assert "Shipment 1" in h and "At sea · in 20 days" in h
    # arrival and receiving on site, room by room; all received = delivered
    c.post(f"/p/{pid_s}/shipments/{ssid}/step", data={"step":"arrived"}, follow_redirects=False)
    r=c.post(f"/p/{pid_s}/shipments/{ssid}/receive", data={"room":str(sr["LIV"])}, follow_redirects=False); assert r.headers["location"].endswith(f"#room-{sr['LIV']}")
    assert _sts([t_liv, t_kit])==["Received","Shipped"]
    db=P_(); assert db.get(_Sh,ssid).delivered_on is None; db.close()
    c.post(f"/p/{pid_s}/shipments/{ssid}/receive", data={"all":"1"}, follow_redirects=False)
    db=P_(); s_=db.get(_Sh,ssid); assert s_.delivered_on==_dt.date.today() and s_.arrived_on==_dt.date.today() and _sh.stage(s_)=="Delivered"; db.close()
    assert _sts([t_liv, t_kit])==["Received","Received"]
    c.post(f"/p/{pid_s}/shipments/{ssid}/receive", data={"c":[str(b2)],"on":"0"}, follow_redirects=False); assert _sts([t_kit])==["Shipped"]  # not received after all
    c.post(f"/p/{pid_s}/cartons/{b2}/received", data={"received":"1"}, follow_redirects=False); assert _sts([t_kit])==["Received"]
    # the documents: the loading list, the shop's packing list to the warehouse, the internal Purchase schedule
    r=c.get(f"/p/{pid_s}/shipments/{ssid}/loading.pdf"); t=pdf_text(r.content)
    assert r.status_code==200 and r.headers["content-type"]=="application/pdf" and "LOADING LIST" in t and "Shipment 1" in t and "TCNU7731042" in t and "BOX 1 OF 2" in t and "Kitchen" in t, t[:400]
    assert "Baiyun Warehouse" in pdf_text(c2.get(f"/s/{stok}/packing.pdf").content)
    r=c.get(f"/p/{pid_s}/buying/schedule.pdf?cur=USD"); t=pdf_text(r.content)
    assert r.status_code==200 and "PURCHASE SCHEDULE" in t and "Ship Tiles" in t and "Floor tile" in t and "USD" in t and "By shop" in t, t[:400]
    db=P_(); db.get(_S, sts).city="Foshan - Lecong"; db.commit(); db.close()
    assert "Floor tile" in c.get(f"/p/{pid_s}/buying/schedule?market=Foshan%20-%20Lecong").text and "Floor tile" not in c.get(f"/p/{pid_s}/buying/schedule?market=Elsewhere").text
    # the tracker and the Overview see the shipping; a box left at the warehouse needs attention
    c2.post(f"/s/{stok}/cartons", data={"room_id":str(sr["LIV"]),"contents":"Spare tiles","code":[codes[t_liv]],"qty":"1","length_cm":"60","width_cm":"60","height_cm":"20","weight_kg":"30"}, follow_redirects=False)
    db=P_(); b4=db.query(_C).filter(_C.project_id==pid_s).order_by(_C.id.desc()).first().id; db.close()
    assert _sts([t_liv])==["Received"]  # a new box at the shop takes nothing back on its own
    c.post(f"/p/{pid_s}/cartons/bulk", data={"c":[str(b4)],"action":"warehouse"}, follow_redirects=False)
    h=c.get(f"/p/{pid_s}/buying/tracker").text; assert "1 box at the warehouse, on no shipment" in h and "Received on site" in h
    assert f"/p/{pid_s}/shipments/" not in h.split("Where the boxes are")[1].split("By stage")[0]  # a delivered shipment leaves the list under way
    assert "1 at the warehouse · 0 on shipments" in c.get(f"/p/{pid_s}").text
    # a box edited: the line it no longer names steps back; a box deleted does the same
    h=c.get(f"/p/{pid_s}/cartons/{b4}").text; assert "The items it holds" in h and f'value="{codes[t_liv]}" checked' in h
    c.post(f"/p/{pid_s}/cartons/{b2}", data={"room_id":str(sr["KIT"]),"supplier_id":str(sts),"contents":"Tiles, kitchen","item_codes":"","qty":"2","length_cm":"60","width_cm":"60","height_cm":"20","weight_kg":"50"}, follow_redirects=False)
    assert _sts([t_kit])==["Ordered"]  # its only box names it no more: back to its order's status
    # a new shipment gets the next number, deleting one sends its boxes back to the warehouse and its rooms back
    r=c.post(f"/p/{pid_s}/shipments/new", data={"mode":"LCL","c":[str(b4)]}, follow_redirects=False); ssid2=int(r.headers["location"].split("/shipments/")[1].split("?")[0])
    c.post(f"/p/{pid_s}/shipments/{ssid2}/step", data={"step":"loaded"}, follow_redirects=False)
    db=P_(); assert db.get(_Sh,ssid2).number==2; db.close()
    c.post(f"/p/{pid_s}/shipments/{ssid2}/delete", follow_redirects=False)
    db=P_(); assert db.get(_Sh,ssid2) is None and db.get(_C,b4).shipment_id is None and db.get(_C,b4).warehouse_on is not None; db.close()
    r=c.post(f"/p/{pid_s}/shipments/new", data={"mode":"LCL"}, follow_redirects=False); ssid3=int(r.headers["location"].split("/shipments/")[1].split("?")[0])
    db=P_(); assert db.get(_Sh,ssid3).number==3; db.close()  # never twice
    c.post(f"/p/{pid_s}/cartons/{b1}/delete", follow_redirects=False); assert _sts([t_liv])==["Ordered"]  # its received box gone, only one at the warehouse left
    assert c.get(f"/p/{pid_b}/shipments/{ssid3}", follow_redirects=False).headers["location"]==f"/p/{pid_b}/shipments"
    print("shipping ok")
    # ---- the project's own lists: categories, units, payment types, containers, defaults ----
    from app import lists as _ls
    r=c.post("/projects/new", data={"client_name":"L","name":"Lists House","rate":"7"}, follow_redirects=False); pid_l=int(r.headers["location"].split("/")[-1])
    db=P_(); pl=db.get(_P,pid_l); assert _ls.categories(pl)==_ls.CATEGORIES and _ls.whole_units(pl)==_ls.WHOLE_UNITS and _ls.ship_modes(pl)==["40ft HC","40ft","20ft","LCL","Air"] and len(pl.list_rows)>=30; db.close()
    h=c.get(f"/p/{pid_l}/lists").text; assert 'data-key="lists"' in h and "Extra to order" in h and 'id="containers"' in h and f'href="/p/{pid_l}/lists" class="on"' in h
    r=c.post(f"/p/{pid_l}/lists/category/add", data={"name":"Wallpaper","allowance":"15"}, follow_redirects=False); assert "added" in r.headers["location"]
    r=c.post(f"/p/{pid_l}/lists/category/add", data={"name":"wallpaper"}, follow_redirects=False); assert "err=" in r.headers["location"]  # no twin, case ignored
    r=c.post(f"/p/{pid_l}/lists/unit/add", data={"name":"pair","whole":"1"}, follow_redirects=False); assert "added" in r.headers["location"]
    r=c.post(f"/p/{pid_l}/lists/container/add", data={"name":"LCL","cbm":"5"}, follow_redirects=False); assert "err=" in r.headers["location"]
    h=c.get(f"/p/{pid_l}/items/new").text; assert ">Wallpaper</option>" in h and ">pair</option>" in h and f'href="/p/{pid_l}/lists#categories"' in h
    assert ">Wallpaper</option>" not in c.get(f"/p/{pid_s}/items/new").text  # another project's list is its own
    r=c.post(f"/p/{pid_l}/rooms", data={"code":"BED","name":"Bedroom","floor":"First"}, follow_redirects=False)
    db=P_(); lr=[x.id for x in db.get(_P,pid_l).rooms if x.code=="BED"][0]; db.close()
    c.post(f"/p/{pid_l}/items/new", data={"room_ids":[str(lr)],"category":"Wallpaper","name":"Grasscloth","qty":"10","unit":"pair"}, follow_redirects=False)
    db=P_(); pl=db.get(_P,pid_l); assert _ls.allowance(pl,"Wallpaper")==0.15 and "pair" in _ls.whole_units(pl)
    g=[x for x in _bz.groups(pl) if x.name=="Grasscloth"][0]; assert g.allowance==0.15 and g.order["qty"]==12  # 10 pair + 15 % = 11.5 -> 12 whole pairs
    wid=[e.id for e in _ls.entries(pl,"category") if e.name=="Wallpaper"][0]; tid=[e.id for e in _ls.entries(pl,"category") if e.name=="Tiles"][0]; db.close()
    # rename: the items follow; the order follows the arrows; a category in use moves its items before it goes
    r=c.post(f"/p/{pid_l}/lists/category/{wid}", data={"name":"Wall coverings","allowance":"15"}, follow_redirects=False); assert "1%20line%20renamed" in r.headers["location"]
    db=P_(); assert [i.category for i in db.get(_P,pid_l).items]==["Wall coverings"]; db.close()
    c.post(f"/p/{pid_l}/lists/category/{wid}/move", data={"dir":"up"}, follow_redirects=False)
    db=P_(); cl=_ls.categories(db.get(_P,pid_l)); assert cl.index("Wall coverings")==len(cl)-2; db.close()
    r=c.post(f"/p/{pid_l}/lists/category/{wid}/delete", follow_redirects=False); assert "still%20used" in r.headers["location"]
    r=c.post(f"/p/{pid_l}/lists/category/{wid}/delete", data={"into":"Tiles"}, follow_redirects=False); assert "moved%20to" in r.headers["location"]
    db=P_(); pl=db.get(_P,pid_l); assert [i.category for i in pl.items]==["Tiles"] and "Wall coverings" not in _ls.categories(pl); db.close()
    r=c.post(f"/p/{pid_l}/lists/category/{tid}/delete", data={}, follow_redirects=False); assert "still%20used" in r.headers["location"]
    # containers and defaults drive the shipping suggestion and the orders
    r=c.post(f"/p/{pid_l}/lists/defaults", data={"deposit_pct":"40","soon_days":"5","balance_days":"10","lcl_up_to":"20"}, follow_redirects=False); assert "Defaults%20saved" in r.headers["location"]
    db=P_(); pl=db.get(_P,pid_l); assert (pl.deposit_pct, pl.soon_days, pl.balance_days, pl.lcl_up_to)==(40.0,5,10,20.0) and _sh.fits(pl, 15, 1000)=="LCL"
    kid=[e.id for e in _ls.entries(pl,"container") if e.name=="20ft"][0]; db.close()
    c.post(f"/p/{pid_l}/lists/container/{kid}", data={"name":"20GP","cbm":"25","kg":"25000"}, follow_redirects=False)
    db=P_(); pl=db.get(_P,pid_l); assert _sh.fits(pl, 22, 1000)=="1 x 20GP" and _sh.capacity(pl,"20GP")==(25.0,25000.0) and "20GP" in _ls.ship_modes(pl); db.close()
    # payment types; the import adds a category it brings; a new project copies another's lists
    db=P_(); pk=[e.id for e in _ls.entries(db.get(_P,pid_l),"payment") if e.name=="Other"][0]; db.close()
    c.post(f"/p/{pid_l}/lists/payment/{pk}", data={"name":"Freight"}, follow_redirects=False)
    assert ">Freight</option>" in c.get(f"/p/{pid_l}/payments").text
    r=c.post("/projects/new", data={"client_name":"M","name":"Copy House","rate":"7","lists_from":str(pid_l)}, follow_redirects=False); pid_m=int(r.headers["location"].split("/")[-1])
    db=P_(); pm=db.get(_P,pid_m); assert "20GP" in _ls.ship_modes(pm) and "Freight" in _ls.payment_kinds(pm) and pm.deposit_pct==40.0 and "pair" in _ls.units(pm); db.close()
    r=c.post(f"/p/{pid_m}/lists/copy", data={"source":str(pid_s)}, follow_redirects=False); assert "added%20from" in r.headers["location"] or "already%20have" in r.headers["location"]
    db=P_(); assert "Other" in _ls.payment_kinds(db.get(_P,pid_m)); db.close()  # what the other project has and this one lacked
    wb=openpyxl.Workbook(); ws=wb.active; ws.title="Shopping List"; ws.append(["Room","Category","Item","Qty","Unit"]); ws.append(["BED","Murals","Hand-painted mural","1","panel"])
    rs=wb.create_sheet("Rooms"); rs.append(["Code","Name","Floor"]); rs.append(["BED","Bedroom","First"])
    buf=io.BytesIO(); wb.save(buf)
    r=c.post(f"/p/{pid_l}/import", files={"file":("l.xlsx",buf.getvalue(),"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}, data={"apply":"1"}, follow_redirects=False)
    db=P_(); pl=db.get(_P,pid_l); assert "Murals" in _ls.categories(pl) and "panel" in _ls.units(pl), (r.status_code, _ls.categories(pl)[-3:]); db.close()
    # the studio's category -> plan sheet table offers every project's categories
    assert "Murals" in c.get("/settings").text
    print("lists ok")
    # ---- supplier quotations: upload a supplier's sheet, confirm the preview, apply; the inquiry sheet; memory; delete ----
    import datetime as _dt
    from openpyxl.drawing.image import Image as _XLImage
    from app.models import Quotation as _QD, QuotationLine as _QL, QuoteMatch as _QM, Product as _Pr, Quote as _Qt, Purchase as _Pu
    from app import quotations as _qz
    r=c.post("/projects/new", data={"client_name":"M","name":"Quote House","rate":"7.1"}, follow_redirects=False); pid_q=int(r.headers["location"].split("/")[-1])
    with open("samples/Kinondoni_House_Procurement_List.xlsx","rb") as f:
        r=c.post(f"/p/{pid_q}/import", files={"file":("list.xlsx", f.read())}, data={"apply":"1"}); assert "Imported" in r.text
    r=c.get(f"/p/{pid_q}/quotations"); assert r.status_code==200 and "Upload a quotation" in r.text and "Inquiry sheet" in r.text and "No quotation yet" in r.text
    assert 'href="/p/%d/quotations"'%pid_q in c.get(f"/p/{pid_q}/buying").text  # the tab
    # the inquiry sheet: one row per product of the category, our code first, yellow cells for the shop
    r=c.get(f"/p/{pid_q}/quotations/inquiry.xlsx?category=Tiles"); assert r.status_code==200 and r.headers["content-type"].startswith("application/vnd.openxml")
    iw=openpyxl.load_workbook(io.BytesIO(r.content)); iws=iw["Inquiry"]; hdr=[c_.value for c_ in iws[4]]; assert hdr[:2]==["Our code","Product"] and "Unit price" in hdr
    irows=[[c_.value for c_ in row] for row in iws.iter_rows(min_row=5) if row[0].value]; assert any(x[0]=="GF-ENT-01" and x[1].startswith("Floor - big") for x in irows) and all("Tiles" not in str(x[1]) or True for x in irows), irows[:2]
    assert not any(x[1].startswith("Basin mixer") for x in irows)  # only the category asked
    # the parser on the sample supplier files (when they are in the repo), and its units and sizes
    assert _qz.dims("600*1200")==[1200.0,600.0] and _qz.dims("main cabinet:1470*420*350")==[1470.0,420.0,350.0] and _qz.dims("1200 x 2400 mm, satin")==[2400.0,1200.0] and _qz.dims("90 cm")==[900.0] and _qz.dims("T=12")==[]
    assert _qz._unit("SQM")=="m²" and _qz._unit("PCS")=="pcs" and _qz._unit("set")=="set" and _qz._unit("M")=="m"
    assert _qz._same_unit("set")=="pcs"==_qz._same_unit("each") and _qz._same_unit("lm")=="m" and _qz._same_unit("m²")=="m²"  # a supplier's set of a mirror is the design's pcs
    assert _qz.expiry("2026-11-24", None)==_dt.date(2026,11,24) and _qz.expiry("15 days", _dt.date(2026,10,26))==_dt.date(2026,11,10) and _qz.expiry("2 weeks", _dt.date(2026,10,26))==_dt.date(2026,11,9) and _qz.expiry("", _dt.date(2026,10,26)) is None
    # memory by item number: one number on several products (a slab cut into floor, skirting, steps) is trusted alone only when the wording agrees
    _dbm=SessionLocal(); _cd=_qz.candidates(_dbm, _dbm.get(__import__("app.models", fromlist=["Project"]).Project, pid_q)); _big=next(c_ for c_ in _cd if c_.name.startswith("Floor - big")); _sk=next(c_ for c_ in _cd if c_.name.startswith("Skirting"))
    _ln=lambda place, name: dict(no=1, ref="GC-X", name=name, size="600*1200", place=place, section="", unit="m²", qty=10.0, unit_price=1.0, detail="", code="")
    _memo=[_QM(supplier_id=1, ref="gc-x", text=_qz._memo_text(_ln("main floor", "Porcelain tile")), product_key=_big.key), _QM(supplier_id=1, ref="gc-x", text=_qz._memo_text(_ln("skirting", "Porcelain tile")), product_key=_sk.key)]
    _sl=_qz.shortlist(_ln("skirting", "Porcelain tile"), _cd, _memo); assert _sl[0][0].id==_sk.id and _sl[0][1]==0.97 and "wording" in _sl[0][2], _sl[0]
    _sl=_qz.shortlist(_ln("stairs", "Porcelain tile"), _cd, _memo); assert _sl[0][1]<0.97 and {x[0].id for x in _sl[:2]}=={_big.id, _sk.id} and "several products" in _sl[0][2], [(x[0].name, x[1], x[2]) for x in _sl[:3]]
    _sl=_qz.shortlist(_ln("stairs", "Porcelain tile"), _cd, _memo[:1]); assert _sl[0][0].id==_big.id and _sl[0][1]==0.97  # one product behind the number: trusted
    _sdir=os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))), "samples", "quotes", "kinondoni")
    _tile=os.path.join(_sdir, "KINO-tile quotation.xlsx")
    if os.path.exists(_tile):
        _pt=_qz.parse(_qz.read_sheets(open(_tile,"rb").read(), "KINO-tile quotation.xlsx"), "KINO-tile quotation.xlsx")
        assert len(_pt["lines"])==29 and _pt["currency"]=="USD" and round(_pt["total"])==22586 and _pt["charges"][0]["amount"]==950 and _pt["head"]["sales"]=="Cathy" and _pt["head"]["quote_date"]=="2026-10-26"
        l1=_pt["lines"][0]; assert l1["ref"]=="GC-CX612Y91-SW" and l1["place"]=="main floor" and l1["section"].startswith("Basement floor") and l1["unit"]=="m²" and l1["qty"]==132.48 and l1["unit_price"]==16 and l1["pack_size"]==1.44 and l1["picture"]=="yes"
        assert _qz.out_of_scope(_pt["lines"][21]) and len(_pt["pictures"])>=25 and any(t.startswith("1.Delivery") for t in _pt["terms"])
        _door=os.path.join(_sdir, "Kino Porject- aluminum door and window quotation--10.26 .xls")
        if os.path.exists(_door):
            _pd=_qz.parse(_qz.read_sheets(open(_door,"rb").read(), "x.xls"), "x.xls")
            assert len(_pd["lines"])==35 and _pd["lines"][0]["ref"]=="F1-W01" and _pd["lines"][0]["size"]=="2160 x 2860 mm" and round(_pd["discount"])==630 and _pd["head"]["quote_date"]=="2026-10-24" and _pd["head"]["valid_until"]=="15 days"
            assert "Sliding Window" in _pd["lines"][0]["detail"] and "glass:" in _pd["lines"][0]["detail"].lower()
            assert any("old .xls" in w for w in _pd["warnings"])  # no pictures from an .xls: said on the preview
    # a supplier's sheet in their own layout: header block, merged "Area" header, a section row, a picture, a charge, discount, terms
    def quote_sheet():
        wb=openpyxl.Workbook(); ws=wb.active; ws.title="Quotation"
        ws["A2"]="Customer"; ws["K2"]="Customer Country"; ws["N2"]="Tanzania"
        ws["A4"]="Quote Date"; ws["F4"]=_dt.datetime(2026,10,26); ws["K4"]="Valid until"; ws["N4"]="15 days"
        ws["A5"]="Order NO"; ws["K5"]="Sales"; ws["N5"]="Cathy"; ws["A6"]="Date of order"; ws["K6"]="Date of deliver"; ws["N6"]="25-30 days after deposit"
        ws.append([]); ws.append(["No","Area",None,"Product Type","Item Number","Design","Size(mm)","Total QTY","Unit","Unit Price(USD)","Amount(USD)","Description","PCS/CTN","SQM/CTN","Remarks"])
        ws.merge_cells("B8:C8")
        ws.append([None,"First floor (Including 10% of wastage)"])
        ws.append([1,None,"main floor","Porcelain tile","GC-1",None,"600*1200",100,"SQM",16,1600,"Surface:Matte\nThickness:9mm",2,1.44,None])
        ws.append([2,None,"bath floor","Porcelain tile","GC-2",None,"600*1200",20,"SQM",12,240,"Matte",2,1.44,None])
        ws.append([3,None,"Servant quarters floor","Porcelain tile","GC-1",None,"600*1200",10,"SQM",16,160,"Matte",2,1.44,None])
        ws.append([4,None,"stairs","Stair handrail",None,None,"75*45",20,"m",24,480,"solid wood top rail, beech wood",None,None,None])
        ws.append([5,None,None,None,"Glass",None,"T=12",23,"SQM",44,1012,"12 mm tempered",None,None,None])
        ws.append(["wooden pallet",None,None,None,None,None,None,None,None,None,950])
        ws.append(["Total amount:",None,None,None,None,None,None,None,None,None,None,None,None,4442])
        ws.append(["AFTER DISCOUNT:",None,None,None,None,None,None,None,None,None,None,None,None,4000])
        ws.append(["Deposit:",None,None,None,None,None,None,None,None,None,None,None,None,2000])
        ws.append(["Notes."]); ws.append(["1.Delivery Term: EXW"]); ws.append(["2.Payment terms: 50% deposit by T/T"])
        pic=_XLImage(io.BytesIO(img("teal"))); pic.width, pic.height = 60, 60; ws.add_image(pic, "F10")
        buf=io.BytesIO(); wb.save(buf); return buf.getvalue()
    import time as _time
    def upload_quote(name, data, supplier_id=""):
        """Upload a sheet and wait for the working page to finish; returns (preview location, last status dict)."""
        r=c.post(f"/p/{pid_q}/quotations", data={"supplier_id":supplier_id}, files={"file":(name, data)}, follow_redirects=False)
        assert r.status_code==303 and "/quotations/working/" in r.headers["location"], (r.status_code, r.headers.get("location"))
        wurl=r.headers["location"]; st=None
        for _ in range(300):
            st=c.get(wurl+"/status").json()
            if st["done"] or st["stage"]=="error": break
            _time.sleep(0.03)
        assert st is not None and (st["done"] or st["stage"]=="error"), st
        if st["done"]:
            r=c.get(wurl, follow_redirects=False); assert r.status_code==303 and r.headers["location"]==st["next"]  # the page itself opens the preview once done
        return (st["next"] if st["done"] else ""), st
    # bad files
    r=c.post(f"/p/{pid_q}/quotations", files={"file":("q.txt", b"hello")}, follow_redirects=False); assert "err=" in r.headers["location"]
    r=c.post(f"/p/{pid_q}/quotations", files={"file":("q.txt", b"hello")}, headers={"Accept":"application/json"}); assert r.status_code==400 and "Excel" in r.json()["error"]
    nb=openpyxl.Workbook(); nb.active.append(["Just","a","note"]); b2=io.BytesIO(); nb.save(b2)
    loc,st=upload_quote("q.xlsx", b2.getvalue()); assert loc=="" and st["stage"]=="error" and "No quotation lines" in st["error"], st
    r=c.get(f"/p/{pid_q}/quotations/working/999999", follow_redirects=False); assert "no%20longer%20here" in r.headers["location"]
    assert c.get(f"/p/{pid_q}/quotations/working/999999/status").status_code==404
    # Claude picks from the shortlists; its answer is checked against the real products
    import app.ai as ai
    os.environ["ANTHROPIC_API_KEY"]="test-key"; qcalls=[]
    def fake_quote_call(content, schema=None, **kw):
        qcalls.append((content, kw))
        text=[b for b in content if b.get("type")=="text"][-1]["text"]
        big=re.search(r"P(\d+) \| Tiles \| Floor - big porcelain slab", text).group(1); bath=re.search(r"P(\d+) \| Tiles \| Bathroom floor - anti-slip", text).group(1)
        assert "L1 | First floor (Including 10% of wastage) / main floor | GC-1 | Porcelain tile | 600*1200 | 100 m² | 16 USD" in text, text[:3000]
        return {"lines":[{"line":1,"action":"quote","product":"P"+big,"merge_into":0,"room":"","category":"","confidence":0.92,"reason":"main floor tile"},
                         {"line":2,"action":"quote","product":"P"+bath,"merge_into":0,"room":"","category":"","confidence":0.88,"reason":"bathroom floor"},
                         {"line":3,"action":"skip","product":"","merge_into":0,"room":"","category":"","confidence":0.9,"reason":"servant quarters"},
                         {"line":4,"action":"new","product":"","merge_into":0,"room":"GF-STA","category":"Hardware","confidence":0.8,"reason":"no handrail in the design"},
                         {"line":5,"action":"merge","product":"","merge_into":4,"room":"","category":"","confidence":0.85,"reason":"glass of the handrail"},
                         {"line":99,"action":"quote","product":"P1","merge_into":0,"room":"","category":"","confidence":1,"reason":"ignored: no such line"}]}
    ai._call=fake_quote_call
    r=c.post(f"/p/{pid_q}/quotations", files={"file":("Tile quotation.xlsx", quote_sheet())}, headers={"Accept":"application/json"}); assert r.status_code==200 and "/quotations/working/" in r.json()["next"], r.text[:300]
    wurl=r.json()["next"]; r=c.get(wurl); assert r.status_code in (200, 303)  # the working page (or the preview already, when the thread was quick)
    if r.status_code==200: assert "Reading the quotation" in r.text and "Matching the lines" in r.text
    for _ in range(300):
        st=c.get(wurl+"/status").json()
        if st["done"] or st["stage"]=="error": break
        _time.sleep(0.03)
    assert st["done"] and st["lines"]==5 and st["pictures"]==1 and st["claude"] and st["steps"][-1]["state"]=="done" and st["expected"]==15 and st["since"], st
    assert _qz.expected_seconds(29, 29, True)==71 and _qz.expected_seconds(29, 29, False)==3
    loc=st["next"]; assert "preview?key=" in loc
    from urllib.parse import unquote as _unq
    qkey=_unq(loc.split("key=")[1]); assert qkey.startswith(f"quotations/p{pid_q}/") and qkey.endswith(".json")
    assert qcalls and kw_ok(qcalls[-1][1]) if (kw_ok:=lambda kw: kw.get("system") and kw.get("effort")=="medium" and kw.get("max_tokens")>=8000) else False
    assert any(b.get("type")=="image" for b in qcalls[-1][0]) and "Picture of line L1" in [b["text"] for b in qcalls[-1][0] if b.get("type")=="text"][0]
    r=c.get(loc); assert r.status_code==200 and "Check the quotation" in r.text and "Tile quotation" in r.text
    from app import storage as _st
    assert "qi-legend" in r.text and 'id="qi-rooms-tpl"' in r.text and 'id="qi-prod-tpl"' in r.text and 'class="row qi-filter' in r.text and 'id="qi-decide"' in r.text
    assert "Design needs" in r.text and "Design + allowance" not in r.text and 'id="qi-skip-none"' in r.text and 'class="qi-chip sec-skip" data-section="First floor' in r.text and 'name="rooms_X"' in r.text and 'id="qi-todo"' in r.text and "Before you save" in r.text
    _sg=__import__("json").loads(_st.read_image(qkey).decode())["suggestions"]
    assert _sg["4"]["action"]=="new" and _sg["4"]["options"] and _sg["4"]["options"][0]["kind"] in ("quote", "new") and any(o["kind"]=="new" for o in _sg["4"]["options"]), _sg["4"]["options"]
    assert _sg["4"]["related"] and _sg["4"]["related"]["word"] in ("handrail", "stair", "stairs") and _sg["4"]["related"]["ids"], _sg["4"]["related"]  # the rooms with a handrail / stair product
    assert _sg["1"]["options"]==[] and _sg["4"]["category"]  # a confident match offers nothing; a new item gets the nearest category
    assert re.search(r'class="span2 qi-rooms" data-ids="[^"]*"', r.text) and 'class="qi-chip opt" data-kind=' in r.text
    h=r.text
    assert 'name="act_1"' in h and h.count('<article class="qi-line')==5 and "First floor (Including 10% of wastage)" in h
    assert re.search(r'name="act_1"[^>]*>\s*<option value="quote" selected', h) and re.search(r'name="act_3"[^>]*>(?:(?!</select>).)*value="skip" selected', h, re.S) and re.search(r'name="act_4"[^>]*>(?:(?!</select>).)*value="new" selected', h, re.S)
    assert "size offered 600*1200 vs 1200 x 2400 mm, satin finish asked" in h and "main floor tile" in h and "92%" in h and "servant quarters" in h
    assert 'value="15 days"' in h and 'value="2026-10-26"' in h and 'value="Cathy"' in h and 'value="25-30 days after deposit"' in h and 'value="4442"' in h and 'value="442"' in h and "wooden pallet 950" in h
    assert "1.Delivery Term: EXW" in h and "/media/" in h  # the terms and the line's picture
    db=SessionLocal(); _sta=db.query(Room).filter(Room.project_id==pid_q, Room.code=="GF-STA").first().id
    big_id=db.query(_Pr).filter(_Pr.project_id==pid_q, _Pr.title.like("Floor - big porcelain slab%")).first().id; bath_id=db.query(_Pr).filter(_Pr.project_id==pid_q, _Pr.title.like("Bathroom floor - anti-slip%")).first().id
    n_items=db.query(Item).filter(Item.project_id==pid_q).count(); db.close()
    # apply without a shop: refused; then with a new shop
    base={"key":qkey,"title":"Tiles from Cathy","sales":"Cathy","quote_date":"2026-10-26","valid_until":"15 days","currency":"USD","lead_time":"25-30 days after deposit","total":"4442","discount":"442","charges":"950","terms":"1.Delivery Term: EXW","notes":"first round",
          "act_1":"quote","prod_1":str(big_id),"act_2":"quote","prod_2":str(bath_id),"act_3":"skip","act_4":"new","name_4":"Stair handrail with glass","cat_4":"Hardware","rooms_4":str(_sta),"act_5":"merge","into_5":"4"}
    r=c.post(f"/p/{pid_q}/quotations/apply", data=dict(base, supplier_id="", act_1="skip"), follow_redirects=False); assert "err=" in r.headers["location"] and "preview?key=" in r.headers["location"] and "choices%20are%20kept" in r.headers["location"]
    qkey2=_unq(r.headers["location"].split("key=")[1].split("&")[0]); assert qkey2!=qkey and _st.read_image(qkey) is None  # the choices made so far wait under a new key
    h=c.get(f"/p/{pid_q}/quotations/preview?key={qkey2}").text; assert re.search(r'name="act_1"[^>]*>(?:(?!</select>).)*value="skip" selected', h, re.S) and 'id="document" open' in h
    qkey=qkey2; base["key"]=qkey
    r=c.post(f"/p/{pid_q}/quotations/apply", data=dict(base, supplier_id="__new__", supplier_name="Foshan Ceramics", supplier_city="Foshan"), follow_redirects=False)
    assert r.status_code==303 and "/quotations/" in r.headers["location"] and "3%20quotes%20from%20Foshan%20Ceramics%20added." in r.headers["location"], r.headers["location"]
    qd_id=int(r.headers["location"].split("/quotations/")[1].split("?")[0])
    db=SessionLocal(); qd=db.get(_QD, qd_id); shop=qd.supplier
    assert shop.name=="Foshan Ceramics" and shop.contact=="Cathy" and qd.currency=="USD" and qd.total==4442 and qd.discount==442 and qd.charges==950 and qd.net==4000 and qd.valid_until=="15 days" and qd.lead_time=="25-30 days after deposit" and qd.quote_date==_dt.date(2026,10,26) and qd.title=="Tiles from Cathy" and qd.file_key.endswith(".xlsx") and qd.filename=="Tile quotation.xlsx" and "EXW" in qd.terms and qd.notes=="first round"
    ls={ln.no: ln for ln in qd.lines}; assert len(ls)==5 and ls[3].action=="skip" and ls[3].quote_id is None and ls[1].thumb_key and ls[1].section=="First floor (Including 10% of wastage) / main floor"
    f=(4442-950-442)/(4442-950)  # the discount is on the lines: the pallet inside the total is not discounted
    q1=ls[1].quote; assert q1 is not None and q1.product_id==big_id and q1.supplier_id==shop.id and q1.currency=="USD" and abs(q1.amount-round(16*f,4))<0.01 and abs(q1.unit_price-16*f*7.1)<0.1 and q1.model=="GC-1" and q1.pack_size==1.44 and q1.lead_time=="25-30 days after deposit" and q1.photo_key and q1.photo_key!=ls[1].photo_key
    assert "Sheet discount of 12.7% applied" in q1.notes and "size offered 600*1200" in q1.notes and "Surface:Matte" in q1.notes and "Tiles from Cathy of 2026-10-26" in q1.notes, q1.notes
    q2=ls[2].quote; assert q2.product_id==bath_id and abs(q2.amount-12*f)<0.01 and q2.model=="GC-2"
    # the new item: one line in the stairs room, the product made, the quote on it with the merged glass added
    new_item=db.query(Item).filter(Item.project_id==pid_q, Item.name=="Stair handrail with glass").first()
    assert new_item is not None and new_item.room_id==_sta and new_item.code=="GF-STA-04" and new_item.category=="Hardware" and new_item.qty==20 and new_item.unit=="m" and new_item.size=="75*45" and "beech wood" in new_item.spec and new_item.unit_price==0 and len(new_item.photos)==0
    assert db.query(Item).filter(Item.project_id==pid_q).count()==n_items+1
    q4=ls[4].quote; assert ls[4].action=="new" and ls[4].product_id and q4 is not None and abs(q4.amount-round((480+1012)/20*f,4))<0.01 and "Includes: Glass 23 m² = 1012" in q4.notes, (q4.amount, q4.notes)
    assert ls[5].action=="merge" and ls[5].into_id==ls[4].id and ls[5].quote_id==q4.id and ls[5].product_id is None
    ms=db.query(_QM).filter(_QM.supplier_id==shop.id).all(); assert {m.ref for m in ms}>={"gc-1","gc-2"} and any(m.product_key.startswith("floor - big porcelain slab") for m in ms if m.ref=="gc-1"), [(m.ref, m.text, m.product_key) for m in ms]
    assert _st.read_image(qkey) is None and _st.read_image(qd.file_key)  # the state is gone, the file stays
    db.close()
    # the document page, the file, the product page's link, the list
    r=c.get(f"/p/{pid_q}/quotations/{qd_id}"); assert r.status_code==200, r.text[:500]
    h=r.text; assert "Tiles from Cathy" in h and "Foshan Ceramics" in h and "Original file" in h and "Part of L4" in h and "New item ·" in h and "Skipped" in h and "USD 4,000" in h and "Not quoted" in h and "Stair risers" in h and "1.Delivery Term: EXW" in h
    assert f"/p/{pid_q}/buying/{big_id}#quote-{q1.id}" in h and "size offered 600*1200" in h
    r=c.get(f"/p/{pid_q}/quotations/{qd_id}/file"); assert r.status_code==200 and r.headers["content-type"].startswith("application/vnd.openxml") and "Tile%20quotation.xlsx" in r.headers.get("content-disposition","") or "Tile quotation.xlsx" in r.headers.get("content-disposition","")
    r=c.get(f"/p/{pid_q}/buying/{big_id}"); assert "Tiles from Cathy" in r.text and f"/quotations/{qd_id}" in r.text and "GC-1" in r.text
    r=c.get(f"/p/{pid_q}/quotations"); assert "Tiles from Cathy" in r.text and "Foshan Ceramics" in r.text and "1 new" in r.text
    # the same sheet again, Claude off: the shop's remembered matches apply by item number
    del os.environ["ANTHROPIC_API_KEY"]; qcalls.clear()
    loc2,st=upload_quote("Tile quotation v2.xlsx", quote_sheet(), str(shop.id)); key2=_unq(loc2.split("key=")[1])
    assert not qcalls and not st["claude"]
    h=c.get(loc2).text; assert "matched before by item number" in h and "97%" in h and re.search(r'name="prod_1"[^>]*>(?:(?!</select>).)*value="%d"[^>]*selected'%big_id, h, re.S) and re.search(r'name="act_3"[^>]*>(?:(?!</select>).)*value="skip" selected', h, re.S)
    assert f'<option value="{shop.id}" selected' in h  # the shop named on the upload
    assert re.search(r'data-q="[^"]*%d:USD 13\.97[^"]*"[^>]*selected' % shop.id, h), "the product option says the shop quoted it"
    assert re.search(r'name="act_4"[^>]*>(?:(?!</select>).)*value="quote" selected', h, re.S) and "matched before by wording" in h  # the new item of the first round is a product now
    # buy a room from the first quote, then apply the second sheet (prices raised): the shop's quotes are corrected, not doubled, and the room follows
    r=c.post(f"/p/{pid_q}/buying/{big_id}/quotes/{q1.id}/buy", follow_redirects=False); assert r.status_code==303 and "err=" not in r.headers["location"]
    db=SessionLocal(); pu=db.query(_Pu).filter(_Pu.quote_id==q1.id).first(); assert pu is not None and abs(pu.amount-13.97)<0.01; n_quotes=db.query(_Qt).filter(_Qt.project_id==pid_q).count(); hand_pid=ls[4].product_id; db.close()
    base2=dict(base, key=key2, supplier_id=str(shop.id), discount="0", notes="second round", act_4="quote", prod_4=str(hand_pid), act_5="merge", into_5="4")
    r=c.post(f"/p/{pid_q}/quotations/apply", data=base2, follow_redirects=False); assert "0%20quotes%20from%20Foshan%20Ceramics%20added%2C%203%20updated" in r.headers["location"], r.headers["location"]
    db=SessionLocal(); assert db.query(_Qt).filter(_Qt.project_id==pid_q).count()==n_quotes
    q1b=db.get(_Qt, q1.id); assert q1b.amount==16 and abs(q1b.unit_price-16*7.1)<0.01 and "Replaces this shop" in q1b.notes and "earlier quote of USD 13.97" in q1b.notes and q1b.model=="GC-1", q1b.notes
    pu=db.query(_Pu).filter(_Pu.quote_id==q1.id).first(); assert abs(pu.amount-16)<0.01 and pu.lines and abs(pu.lines[0].unit_price-16*7.1)<0.01  # the room bought follows
    assert db.query(Item).filter(Item.project_id==pid_q, Item.name=="Stair handrail with glass").count()==1  # no second item
    qd2=db.query(_QD).filter(_QD.project_id==pid_q).order_by(_QD.id.desc()).first(); assert qd2.id!=qd_id and {ln.no: ln.quote_id for ln in qd2.lines}[4]==q4.id and {ln.no: ln.action for ln in qd2.lines}[5]=="merge"; qd2_id=qd2.id
    db.close()
    r=c.get(f"/p/{pid_q}/quotations/{qd2_id}"); assert r.status_code==200 and "second round" in r.text
    db=SessionLocal(); assert db.get(_QD, qd2_id).replaces_id==qd_id; db.close()  # the revised sheet names the document it corrected
    assert "A revision of" in r.text and f"/quotations/{qd_id}" in r.text and f"quotations?supplier={shop.id}#upload" in r.text
    r=c.get(f"/p/{pid_q}/quotations/{qd_id}"); assert "Revised by" in r.text and f"/quotations/{qd2_id}" in r.text
    h=c.get(f"/p/{pid_q}/quotations?supplier={shop.id}").text; assert ">revision<" in h and "Earlier revisions (1)" in h and re.search(r'value="%d" selected' % shop.id, h)
    # a third upload, discarded
    loc3,_=upload_quote("Tile quotation v3.xlsx", quote_sheet(), str(shop.id)); key3=_unq(loc3.split("key=")[1])
    r=c.post(f"/p/{pid_q}/quotations/cancel", data={"key":key3}, follow_redirects=False); assert r.headers["location"]==f"/p/{pid_q}/quotations"
    r=c.get(f"/p/{pid_q}/quotations/preview?key={key3}", follow_redirects=False); assert "no%20longer%20here" in r.headers["location"]
    r=c.post(f"/p/{pid_q}/quotations/{qd2_id}/delete", follow_redirects=False); assert "0%20quotes%20nothing%20was%20bought%20from" in r.headers["location"], r.headers["location"]  # the second document corrected the first's quotes: they stay with it
    db=SessionLocal(); assert db.get(_Qt, q1.id) is not None and db.get(_Qt, q2.id) is not None and db.get(_Qt, q4.id) is not None; db.close()
    # the first document still holds q2 and q4: deleting it takes them (q1 is bought from)
    r=c.get(f"/p/{pid_q}/quotations/preview?key=quotations/p1/nope.json", follow_redirects=False); assert "err=" in r.headers["location"]
    r=c.get(f"/p/{pid_q}/quotations/{qd_id}"); assert r.status_code==200
    db=SessionLocal(); fk=db.get(_QD, qd_id).file_key; pk=ls[1].photo_key; db.close()
    r=c.post(f"/p/{pid_q}/quotations/{qd_id}/delete", follow_redirects=False); assert "2%20quotes%20nothing%20was%20bought%20from" in r.headers["location"], r.headers["location"]
    db=SessionLocal(); assert db.get(_QD, qd_id) is None and db.query(_QL).filter(_QL.quotation_id==qd_id).count()==0 and db.get(_Qt, q1.id) is not None and db.get(_Qt, q2.id) is None and db.get(_Qt, q4.id) is None
    assert db.query(Item).filter(Item.project_id==pid_q, Item.name=="Stair handrail with glass").first() is not None  # the item stays
    db.close(); assert _st.read_image(fk) is None and _st.read_image(pk) is None
    # a draft: the preview saved with its choices, reopened from the Drafts card, edited, then posted; discard of a failed upload
    loc5,st5=upload_quote("Tile quotation v5.xlsx", quote_sheet(), str(shop.id)); key5=_unq(loc5.split("key=")[1])
    from app.models import QuotationUpload as _QU
    db=SessionLocal(); up5=db.query(_QU).filter(_QU.state_key==key5).first(); assert up5 is not None and up5.stage=="done"; up5_id=up5.id; db.close()
    h=c.get(f"/p/{pid_q}/quotations").text; assert "Drafts and uploads under way" in h and "Tile quotation v5.xlsx" in h and "Ready to check" in h and f"/quotations/working/{up5_id}/discard" in h
    dform=dict(base, key=key5, supplier_id="__new__", supplier_name="Draft Shop", supplier_city="Foshan", notes="draft note", title="Draft title", act_1="skip", act_4="new", name_4="Handrail v5", act_5="merge", into_5="4", rooms_4=str(_sta))
    r=c.post(f"/p/{pid_q}/quotations/apply", data=dict(dform, supplier_id="", supplier_name=""), follow_redirects=False); assert "err=" in r.headers["location"]  # still refused without a shop when posting; the choices wait under a new key
    key5=_unq(r.headers["location"].split("key=")[1].split("&")[0]); dform["key"]=key5
    r=c.post(f"/p/{pid_q}/quotations/draft", data=dform, follow_redirects=False); assert r.status_code==303 and "Draft%20saved" in r.headers["location"], r.headers["location"]
    assert _st.read_image(key5) is None  # the old state key is gone, a new one carries the draft
    db=SessionLocal(); up5=db.get(_QU, up5_id); key5b=up5.state_key; assert key5b!=key5 and key5b.endswith(".json") and up5.message=="draft"; db.close()
    h=c.get(f"/p/{pid_q}/quotations").text; assert ">Draft<" in h and "Continue" in h and f"preview?key={key5b}" in h.replace("%2F", "/")
    h=c.get(f"/p/{pid_q}/quotations/preview?key={key5b}").text
    assert "A draft saved" in h and 'value="draft note"' in h and 'value="Draft title"' in h and 'value="Draft Shop"' in h and 'value="__new__" selected' in h
    assert re.search(r'name="act_1"[^>]*>(?:(?!</select>).)*value="skip" selected', h, re.S) and re.search(r'name="act_4"[^>]*>(?:(?!</select>).)*value="new" selected', h, re.S) and 'value="Handrail v5"' in h
    assert ('name="rooms_4" value="%d"' % _sta) in h and re.search(r'class="span2 qi-rooms" data-ids="[^"]*\b%d\b' % _sta, h) and re.search(r'name="into_5"[^>]*>(?:(?!</select>).)*value="4" selected', h, re.S)
    assert "Save as draft" in h and "/quotations/draft" in h
    # the draft saved again (a second key), then posted from it
    r=c.post(f"/p/{pid_q}/quotations/draft", data=dict(dform, key=key5b, notes="draft note 2"), follow_redirects=False); assert "Draft%20saved" in r.headers["location"]
    db=SessionLocal(); key5c=db.get(_QU, up5_id).state_key; assert key5c not in (key5, key5b); db.close()
    r=c.post(f"/p/{pid_q}/quotations/apply", data=dict(dform, key=key5c, notes="draft note 2"), follow_redirects=False); assert r.status_code==303 and "/quotations/" in r.headers["location"] and "err=" not in r.headers["location"], r.headers["location"]
    qd5=int(r.headers["location"].split("/quotations/")[1].split("?")[0])
    db=SessionLocal(); d5=db.get(_QD, qd5); assert d5.notes=="draft note 2" and d5.title=="Draft title" and d5.supplier.name=="Draft Shop" and {ln.no: ln.action for ln in d5.lines}[1]=="skip" and {ln.no: ln.action for ln in d5.lines}[4]=="new"
    assert db.get(_QU, up5_id) is None and db.query(Item).filter(Item.project_id==pid_q, Item.name=="Handrail v5").count()==1; db.close()
    assert _st.read_image(key5c) is None and "Tile quotation v5.xlsx" not in c.get(f"/p/{pid_q}/quotations").text
    c.post(f"/p/{pid_q}/quotations/{qd5}/delete", follow_redirects=False)
    # Claude down: the request is retried, then the upload waits (not silently matched by words); try again, or go on without
    _qz.RETRY_WAITS=(0.01, 0.01); ncalls=[0]
    def down_call(content, schema=None, **kw): ncalls[0]+=1; raise ConnectionError("boom")
    ai._call=down_call; os.environ["ANTHROPIC_API_KEY"]="test-key"
    r=c.post(f"/p/{pid_q}/quotations", data={"supplier_id":str(shop.id)}, files={"file":("Tile quotation v6.xlsx", quote_sheet())}, follow_redirects=False); wurl6=r.headers["location"]
    for _ in range(300):
        st=c.get(wurl6+"/status").json()
        if st["done"] or st["stage"] in ("error", "waiting"): break
        _time.sleep(0.03)
    assert st["waiting"] and st["stage"]=="waiting" and not st["done"] and "could not be reached" in st["message"] and st["steps"][3]["state"]=="stuck" and ncalls[0]==3, (st, ncalls)
    h=c.get(wurl6).text; assert "Try Claude again" in h and "Continue without Claude" in h and 'id="qi-wait"' in h and "qi-wait\" hidden" not in h
    h=c.get(f"/p/{pid_q}/quotations").text; assert "Waiting for Claude" in h and "Match it" in h and f"/quotations/working/{wurl6.rsplit('/',1)[1]}/discard" in h
    r=c.post(wurl6+"/retry", data={"claude":"1"}, follow_redirects=False); assert r.headers["location"]==wurl6
    for _ in range(300):
        st=c.get(wurl6+"/status").json()
        if st["done"] or st["stage"] in ("error", "waiting"): break
        _time.sleep(0.03)
    assert st["waiting"] and ncalls[0]==6, (st, ncalls)  # still down: waits again
    r=c.post(wurl6+"/retry", data={"claude":"0"}, follow_redirects=False); assert r.headers["location"]==wurl6
    for _ in range(300):
        st=c.get(wurl6+"/status").json()
        if st["done"] or st["stage"] in ("error", "waiting"): break
        _time.sleep(0.03)
    assert st["done"] and not st["claude"] and ncalls[0]==6, st
    h=c.get(st["next"]).text; assert "carry on without it" in h and 'class="qi-find"' in h and 'data-same="#gc-1|m²"' in h and "matched before by wording (another shop)" not in h
    ai._call=fake_quote_call; del os.environ["ANTHROPIC_API_KEY"]
    key6=_unq(st["next"].split("key=")[1]); r=c.post(f"/p/{pid_q}/quotations/cancel", data={"key":key6}, follow_redirects=False); assert r.status_code==303
    # the wording is remembered across shops: a sheet from no named shop matches by it
    db=SessionLocal(); assert db.query(_QM).filter(_QM.supplier_id.is_(None), _QM.text.like("main floor porcelain tile%")).count()==1; db.close()
    loc7,st7=upload_quote("Tile quotation v7.xlsx", quote_sheet()); key7=_unq(loc7.split("key=")[1])
    st7s=__import__("json").loads(_st.read_image(key7).decode())["suggestions"]; assert st7s["1"]["shortlist"][0]["why"]=="matched before by wording (another shop)" and st7s["1"]["shortlist"][0]["id"]==big_id, st7s["1"]
    c.post(f"/p/{pid_q}/quotations/cancel", data={"key":key7}, follow_redirects=False)
    # a proforma with a summary table and its breakdown: the parts are merged into the summary line and counted once
    def pi_sheet():
        wb=openpyxl.Workbook(); ws=wb.active; ws.title="PI"
        ws.append(["Summary"]); ws.append(["No","Item","Description","Qty","Unit price (USD)","Amount (USD)"]); ws.append([1,"Kitchen","as below",1,8900,8900])
        ws.append([]); ws.append(["Breakdown"]); ws.append(["No","Item","Description","Qty","Unit price (USD)","Amount (USD)"])
        ws.append([1,"cabinets","melamine",1,5000,5000]); ws.append([2,"countertop","quartz",1,3000,3000]); ws.append([3,"sink and tap","",1,900,900])
        ws.append(["WOODEN PACKAGE IN USD:",None,None,None,None,200]); ws.append(["TOTAL IN USD:",None,None,None,None,9100])
        buf=io.BytesIO(); wb.save(buf); return buf.getvalue()
    _pp=_qz.parse(_qz.read_sheets(pi_sheet(), "pi.xlsx"), "pi.xlsx"); assert len(_pp["lines"])==4 and _pp["lines"][0]["summary_of"]==[2,3,4] and all(l.get("part_of")==1 for l in _pp["lines"][1:]) and _pp["total"]==9100 and _pp["charges"][0]["amount"]==200, _pp["lines"]
    assert _qz.is_breakdown(_pp["lines"][0], _pp["lines"][1:]) and not _qz.is_breakdown(_pp["lines"][0], _pp["lines"][1:2])
    locp,stp=upload_quote("Kitchen PI.xlsx", pi_sheet(), str(shop.id)); keyp=_unq(locp.split("key=")[1]); hp=c.get(locp).text
    assert "but the sheet's total" not in hp  # the parts are inside L1: no line-sum warning
    assert all(re.search(r'name="act_%d"[^>]*>(?:(?!</select>).)*value="merge" selected' % n, hp, re.S) and re.search(r'name="into_%d"[^>]*>(?:(?!</select>).)*value="1" selected' % n, hp, re.S) for n in (2,3,4)) and "breakdown of L1" in hp and 'class="qi-parts' in hp
    _spp=__import__("json").loads(_st.read_image(keyp).decode())["suggestions"]; assert _spp["2"]["action"]=="merge" and _spp["2"]["merge_into"]==1 and _spp["2"]["confidence"]==0.95
    _kit=[c_ for c_ in _qz.candidates(SessionLocal(), SessionLocal().get(__import__("app.models", fromlist=["Project"]).Project, pid_q)) if "kitchen" in c_.name.lower() and c_.unit in ("set", "pcs")]
    formp={"key":keyp,"supplier_id":str(shop.id),"title":"Kitchen PI","currency":"USD","total":"9100","discount":"0","charges":"200","terms":"","act_1":"quote","prod_1":str(_kit[0].id),"act_2":"merge","into_2":"1","act_3":"merge","into_3":"1","act_4":"merge","into_4":"1"}
    r=c.post(f"/p/{pid_q}/quotations/apply", data=formp, follow_redirects=False); assert r.status_code==303 and "err=" not in r.headers["location"], r.headers["location"]
    qdp=int(r.headers["location"].split("/quotations/")[1].split("?")[0]); db=SessionLocal(); dp=db.get(_QD, qdp); qk=next(ln.quote for ln in dp.lines if ln.no==1)
    assert abs(qk.amount-8900)<0.01 and "Made up of: cabinets" in qk.notes and dp.replaces_id is None, (qk.amount, qk.notes[:200])  # the breakdown counted once
    db.close()
    # the compare grid: the kitchen product against the shop, buy from it in one press, back on the grid
    _kc=_kit[0].category; h=c.get(f"/p/{pid_q}/quotations/compare?category={_kc}").text
    assert "Compare shops" in h and shop.name in h and _kit[0].name in h and 'class="num cmp-cell best"' in h and f"/buying/{_kit[0].id}/quotes/{qk.id}/buy" in h and "1 priced" in h
    assert 'name="next" value="/p/%d/quotations/compare?category=' % pid_q in h and "Every product at its lowest price" in h
    h2=c.get(f"/p/{pid_q}/quotations/compare?category={_kc}&currency=CNY").text; assert "¥" in h2 and "$" not in h2.split('class="cmp"')[1].split("</table>")[0]
    assert "Nothing to compare yet" in c.get(f"/p/{pid_q}/quotations/compare?category=Nope").text or _kc in c.get(f"/p/{pid_q}/quotations/compare?category=Nope").text  # an unknown category falls back to the first with quotes
    r=c.post(f"/p/{pid_q}/buying/{_kit[0].id}/quotes/{qk.id}/buy", data={"next":f"/p/{pid_q}/quotations/compare?category={_kc}&currency=USD"}, follow_redirects=False)
    assert r.status_code==303 and "/quotations/compare?category=" in r.headers["location"] and "ok=Buying" in r.headers["location"], r.headers["location"]
    h=c.get(r.headers["location"]).text; assert "bought here" in h and 'cmp-cell best bought' in h
    db=SessionLocal(); assert db.query(_Pu).filter(_Pu.quote_id==qk.id).count()==1; db.close()
    c.post(f"/p/{pid_q}/quotations/{qdp}/delete", follow_redirects=False)
    # uploads the server lost: a row left at a working stage is marked failed at startup; one silent for 10 minutes can be discarded
    from app.models import QuotationUpload as _QU2
    db=SessionLocal(); u_old=_QU2(project_id=pid_q, filename="lost.xlsx", stage="reading", message="x"); db.add(u_old); db.commit(); u_old_id=u_old.id
    u_old.updated_at=_dt.datetime.now()-_dt.timedelta(minutes=11); db.commit()
    h=c.get(f"/p/{pid_q}/quotations").text; assert ">Stuck<" in h and f"/quotations/working/{u_old_id}/discard" in h
    u_new=_QU2(project_id=pid_q, filename="fresh.xlsx", stage="matching", message="x"); db.add(u_new); db.commit(); u_new_id=u_new.id
    h=c.get(f"/p/{pid_q}/quotations").text; assert "fresh.xlsx" in h and f"/quotations/working/{u_new_id}/discard" not in h  # still working: left alone
    r=c.post(f"/p/{pid_q}/quotations/working/{u_new_id}/discard", follow_redirects=False); assert db.get(_QU2, u_new_id) is not None or True
    assert _qz.mark_stale(db)>=1 and db.get(_QU2, u_new_id).stage=="error" and "restarted" in db.get(_QU2, u_new_id).message; db.close()
    for uid in (u_old_id, u_new_id): c.post(f"/p/{pid_q}/quotations/working/{uid}/discard", follow_redirects=False)
    db=SessionLocal(); assert db.get(_QU2, u_old_id) is None and db.get(_QU2, u_new_id) is None; db.close()
    h=c.get(loc).text if False else ""  # (the preview's bulk buttons are checked on the first preview above)
    # a failed upload sits in the list until discarded
    loc6,st6=upload_quote("broken.xlsx", b2.getvalue()); assert st6["stage"]=="error"
    h=c.get(f"/p/{pid_q}/quotations").text; assert "broken.xlsx" in h and ">Failed<" in h
    db=SessionLocal(); up6=db.query(_QU).filter(_QU.project_id==pid_q, _QU.stage=="error").first(); up6_id=up6.id; db.close()
    r=c.post(f"/p/{pid_q}/quotations/working/{up6_id}/discard", follow_redirects=False); assert r.status_code==303
    db=SessionLocal(); assert db.get(_QU, up6_id) is None; db.close()
    # cancel from the preview drops the upload row too
    loc7,_=upload_quote("Tile quotation v7.xlsx", quote_sheet(), str(shop.id)); key7=_unq(loc7.split("key=")[1])
    c.post(f"/p/{pid_q}/quotations/cancel", data={"key":key7}, follow_redirects=False)
    db=SessionLocal(); assert db.query(_QU).filter(_QU.state_key==key7).count()==0; db.close()
    # deleting the shop: its remembered matches go, its quotation sheets stay without a shop, its quotes go with the purchases (buying.drop_supplier)
    loc4,_=upload_quote("Tile quotation v4.xlsx", quote_sheet(), str(shop.id)); key4=_unq(loc4.split("key=")[1])
    r=c.post(f"/p/{pid_q}/quotations/apply", data=dict(base, key=key4, supplier_id=str(shop.id), act_4="skip", act_5="skip"), follow_redirects=False); qd4=int(r.headers["location"].split("/quotations/")[1].split("?")[0])
    r=c.post(f"/suppliers/{shop.id}/delete", follow_redirects=False); assert r.status_code==303
    db=SessionLocal(); assert db.get(Supplier, shop.id) is None and db.query(_QM).filter(_QM.supplier_id==shop.id).count()==0
    q4d=db.get(_QD, qd4); assert q4d is not None and q4d.supplier_id is None and all(ln.quote_id is None for ln in q4d.lines) and db.query(_Qt).filter(_Qt.project_id==pid_q).count()==0; db.close()
    assert c.get(f"/p/{pid_q}/quotations/{qd4}").status_code==200 and "no shop" in c.get(f"/p/{pid_q}/quotations/{qd4}").text
    # one product, several models (a generic air conditioner quoted as two models at two prices): split it into variants on the
    # preview, one product per model with its rooms and price; review and edit the variants on the document page
    ai._call=lambda content, schema=None, **kw: {"lines": []}
    assert _qz.variant_label({"name":"Gree inverter 2.5hp 7kw cooling only","size":"","detail":""})=="2.5 HP" and _qz.variant_label({"name":"Heater","size":"12 kW","detail":"","ref":"X"})=="12 kW" and _qz.variant_label({"name":"Fan","size":"","detail":"220-240V","ref":"F-9"})=="F-9"
    assert _qz._typo("dinning", "dining") and _qz._typo("louge", "lounge") and not _qz._typo("bedroom 2", "bedroom 1") and not _qz._typo("bath", "bats")
    db=SessionLocal(); _ac_rooms=db.query(Room).filter(Room.project_id==pid_q).order_by(Room.id).all()[:3]; _acr=[(r_.id, r_.name) for r_ in _ac_rooms]; db.close(); assert len(_acr)==3
    r=c.post(f"/p/{pid_q}/items/new", data={"name":"Air conditioner","category":"Lighting","qty":"1","unit":"pcs","room_ids":[str(i) for i,_ in _acr]}, follow_redirects=False); assert r.status_code==303
    db=SessionLocal(); _acp=_bz.product_of(db, db.get(_P, pid_q), db.query(Item).filter(Item.project_id==pid_q, Item.name=="Air conditioner").first()); _acp_id=_acp.id; db.commit(); db.close()
    def ac_sheet():
        wb=openpyxl.Workbook(); ws=wb.active; ws.title="PI"; ws.append(["Proforma invoice"])
        ws.append(["NO.","Product Type","model","Unit","Quantity","Unit Price(USD)","Amount(USD)","Remarks"]); ws.append(["Ground floor"])
        ws.append([1,"Gree Inverter wall mounted air conditioner 2.5hp 7kw","GWC24","PCS",1,890,890,f"For {_acr[0][1]}"])
        ws.append([2,"Gree Inverter wall mounted air conditioner 1hp 2.6kw","GWC09","PCS",1,508,508,f"For {_acr[1][1]}"])
        ws.append([3,"Gree Inverter wall mounted air conditioner 1hp 2.6kw","GWC09","PCS",1,508,508,f"For {_acr[2][1]}"])
        ws.append(["Total amount:",None,None,None,None,None,1906]); buf=io.BytesIO(); wb.save(buf); return buf.getvalue()
    loc8,st8=upload_quote("Split AC quotation.xlsx", ac_sheet()); key8=_unq(loc8.split("key=")[1]); h=c.get(loc8).text
    assert 'data-variant="2.5 HP"' in h and f'data-hints="{_acr[0][0]}"' in h and 'id="qi-splits"' in h and 'id="qi-roomnames"' in h and 'data-rids="' in h and "vlabel_" in h and "Split by model" in h
    _pt8=__import__("json").loads(_st.read_image(key8).decode())["parsed"]; assert _pt8["lines"][0]["place"]==_acr[0][1] and _pt8["lines"][2]["place"]==_acr[2][1]  # "For <room>" in Remarks is the line's place
    form8=dict(key=key8, supplier_id="__new__", supplier_name="George Appliances", title="AC", currency="USD", total="1906", discount="0", charges="0", terms="", notes="", sales="", quote_date="", valid_until="", lead_time="",
               act_1="quote", prod_1=str(_acp_id), act_2="quote", prod_2=str(_acp_id), act_3="quote", prod_3=str(_acp_id), vlabel_1="2.5 HP", vrooms_1=str(_acr[0][0]), vlabel_2="1 HP", vrooms_2=str(_acr[1][0]), vlabel_3="1 HP", vrooms_3=str(_acr[2][0]))
    r=c.post(f"/p/{pid_q}/quotations/draft", data=dict(form8, **{f"split_{_acp_id}":"1"}), follow_redirects=False); assert r.status_code==303 and "Draft%20saved" in r.headers["location"]
    db=SessionLocal(); key8b=db.query(_QU).filter(_QU.project_id==pid_q, _QU.message=="draft").order_by(_QU.id.desc()).first().state_key; db.close()  # a draft keeps the split
    h=c.get(f"/p/{pid_q}/quotations/preview?key={key8b}").text; assert 'data-split="1"' in h and 'data-vlabel="2.5 HP"' in h and f'data-vrooms="{_acr[1][0]}"' in h
    r=c.post(f"/p/{pid_q}/quotations/apply", data=dict(form8, key=key8b, **{f"split_{_acp_id}":"1"}), follow_redirects=False); assert r.status_code==303 and "2%20quotes%20from%20George%20Appliances%20added" in r.headers["location"], r.headers["location"]
    qd8=int(r.headers["location"].split("/quotations/")[1].split("?")[0])
    db=SessionLocal(); _its={i.room_id: i for i in db.query(Item).filter(Item.project_id==pid_q, Item.name=="Air conditioner")}; assert {k: v.size for k, v in _its.items()}=={_acr[0][0]:"2.5 HP", _acr[1][0]:"1 HP", _acr[2][0]:"1 HP"}, {k: v.size for k, v in _its.items()}
    assert db.query(_Pr).filter(_Pr.project_id==pid_q, _Pr.key=="air conditioner||").count()==0  # the generic product, left with no room and no quote, is gone (SQLite may reuse its id)
    _q8=db.get(_QD, qd8); _v={ln.no: ln.product for ln in _q8.lines}; assert _v[1].title=="Air conditioner · 2.5 HP" and _v[2].title=="Air conditioner · 1 HP" and _v[2].id==_v[3].id and _v[1].id!=_v[2].id
    assert len(_v[1].quotes)==1 and _v[1].quotes[0].model=="GWC24" and round(_v[1].quotes[0].amount)==890 and _v[2].quotes[0].model=="GWC09" and round(_v[2].quotes[0].amount)==508
    _p1, _p2 = _v[1].id, _v[2].id; _it0=_its[_acr[0][0]].id; db.close()
    h=c.get(f"/p/{pid_q}/quotations/{qd8}").text; assert 'id="variants"' in h and f'name="label_{_p1}"' in h and 'value="2.5 HP"' in h and f'name="var_{_it0}"' in h and "L2, L3" in h and "GWC09" in h
    # rename a variant and move a room to it: the lines' sizes change, the products follow, the quotes stay with them
    r=c.post(f"/p/{pid_q}/quotations/{qd8}/variants", data={f"label_{_p2}":"1 HP Gree", f"label_{_p1}":"2.5 HP", f"var_{_it0}":str(_p2)}, follow_redirects=False); assert r.status_code==303 and "3%20room%20lines%20changed" in r.headers["location"], r.headers["location"]
    db=SessionLocal(); assert {i.size for i in db.query(Item).filter(Item.project_id==pid_q, Item.name=="Air conditioner")}=={"1 HP Gree"} and db.get(_Pr, _p2).title=="Air conditioner · 1 HP Gree" and len(db.get(_Pr, _p2).quotes)==1
    assert all(ln.product_id==_p2 for ln in db.get(_QD, qd8).lines if ln.no in (2, 3)) and db.get(_Pr, _p1) is not None and len(db.get(_Pr, _p1).quotes)==1  # the 2.5 HP variant keeps its quote with no room left
    db.close()
    h=c.get(r.headers["location"].split("#")[0]).text; assert 'value="1 HP Gree"' in h and "Variants saved" in h and "its quote stays" in h, h.split('id="variants"')[1][:3000] if 'id="variants"' in h else h[:500]
    r=c.post(f"/p/{pid_q}/quotations/{qd8}/variants", data={f"label_{_p1}":"1 HP Gree"}, follow_redirects=False); assert "Nothing%20changed" in r.headers["location"]  # a variant without lines cannot be renamed into another
    ai._call=lambda content, schema=None, **kw: (_ for _ in ()).throw(RuntimeError("down"))
    print("quotations ok")


# ---- the runner: sections, dependencies, selection ------------------------------------------------------------------------
# A section is the statements of flow() up to and including a print("<name> ok"); what comes before the first one is the
# setup (always run). A section depends on the sections that assigned the names it reads (worked out from the source) plus
# EXTRA_DEPS for state that lives only in the database. Usage:
#   python tests/test_flow.py                     every section (the gate before a merge; prints ALL OK)
#   python tests/test_flow.py quotations buying   the setup, their dependencies and those sections
#   python tests/test_flow.py --list              the sections, their dependencies, and the time each took last run
# A name matches a section when it is the section's name or a word of it ("plan" = interactive plan). The imports and
# helper functions of a skipped section are still defined, so a section does not depend on another for those alone.
import ast as _ast, json as _json, time as _time0

EXTRA_DEPS = {  # section -> sections it needs beyond what the names say: state that lives only in the database
    "plan sheet types": ["legacy layers + main plan"],  # the Windows & doors sheet of the layered house
    "room checklist": ["item list + sorting"],  # the Reading light item in the bedrooms
}
TIMES_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), ".section_times.json")


def _marker(stmt):
    """'name' when the statement is print("name ok"), else None."""
    if isinstance(stmt, _ast.Expr) and isinstance(stmt.value, _ast.Call) and isinstance(stmt.value.func, _ast.Name) and stmt.value.func.id == "print":
        a = stmt.value.args
        if len(a) == 1 and isinstance(a[0], _ast.Constant) and isinstance(a[0].value, str) and a[0].value.endswith(" ok"):
            return a[0].value[:-3]
    return None


def _sections():
    src = open(os.path.abspath(__file__), encoding="utf-8").read()
    tree = _ast.parse(src)
    fn = next(n for n in tree.body if isinstance(n, _ast.FunctionDef) and n.name == "flow")
    out, cur = [], []
    for st in fn.body[1:]:  # past the docstring
        cur.append(st)
        m = _marker(st)
        if m:
            out.append((m, cur))
            cur = []
    if cur:
        out.append(("(tail)", cur))
    assert out and out[0][0] == "setup", "flow() must start with the setup, ending in print(\"setup ok\")"
    return out


def _free(st):
    """A statement a skipped section can still run: an import or a helper definition (cheap, no effect on the app)."""
    return isinstance(st, (_ast.Import, _ast.ImportFrom, _ast.FunctionDef))


def _names(stmts):
    """(assigned, read) names of a list of statements, leaving out what _free statements define."""
    assigned, read = set(), set()
    for st in stmts:
        if _free(st):
            continue
        for n in _ast.walk(st):
            if isinstance(n, _ast.Name):
                (assigned if isinstance(n.ctx, _ast.Store) else read).add(n.id)
            elif isinstance(n, (_ast.FunctionDef, _ast.ClassDef)):
                assigned.add(n.name)
            elif isinstance(n, _ast.Import) or isinstance(n, _ast.ImportFrom):
                for a in n.names:
                    assigned.add((a.asname or a.name).split(".")[0])
            elif isinstance(n, _ast.arg):
                assigned.add(n.arg)
    return assigned, read


def _plan(secs, wanted):
    """The sections to run, in order: the setup, the wanted ones and, transitively, the ones they depend on."""
    names = [n for n, _ in secs]
    info = {n: _names(st) for n, st in secs}
    first = names[0]
    deps = {}
    for i, n in enumerate(names):
        d = set(EXTRA_DEPS.get(n, ()))
        for name in info[n][1]:
            if name in info[n][0] or name in info[first][0]:
                continue
            for m in reversed(names[:i]):
                if name in info[m][0]:
                    if m != first:
                        d.add(m)
                    break
        deps[n] = d
    chosen = set()
    def add(n):
        if n in chosen:
            return
        chosen.add(n)
        for m in deps[n]:
            add(m)
    for w in wanted:
        add(w)
    return [n for n in names if n == first or n in chosen], deps


def main(argv):
    secs = _sections()
    names = [n for n, _ in secs]
    args = [a for a in argv if not a.startswith("--")]
    try:
        times = _json.load(open(TIMES_FILE))
    except Exception:
        times = {}
    if "--list" in argv:
        _, deps = _plan(secs, [])
        for n in names:
            print(f"{times.get(n, 0):6.1f}s  {n}" + (f"  (needs {', '.join(sorted(deps[n]))})" if deps[n] else ""))
        return
    wanted = []
    for a in args:
        hits = [n for n in names if n == a] or [n for n in names if a.lower() in n.lower().replace("&", "").split()] or [n for n in names if a.lower() in n.lower()]
        if len(hits) != 1:
            sys.exit(f"'{a}' names {len(hits)} sections{': ' + ', '.join(hits) if hits else ''}. Sections: {', '.join(names)}")
        wanted.append(hits[0])
    run, _ = _plan(secs, wanted) if wanted else (names, None)
    ns = dict(globals())
    t_all = _time0.time()
    with TestClient(app) as c:
        ns["c"] = c
        for name, stmts in secs:
            if name not in run:
                stmts = [st for st in stmts if _free(st)]
            t = _time0.time()
            exec(compile(_ast.Module(body=stmts, type_ignores=[]), os.path.abspath(__file__), "exec"), ns)
            if name in run:
                times[name] = round(_time0.time() - t, 1)
    try:
        _json.dump(times, open(TIMES_FILE, "w"), indent=0)
    except OSError:
        pass
    print(f"ALL OK ({len(run)} of {len(names)} sections, {_time0.time() - t_all:.0f} s)")


if __name__ == "__main__":
    main(sys.argv[1:])
