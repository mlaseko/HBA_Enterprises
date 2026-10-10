"""Import rooms and items from an Excel file: this app's template or export, or the Kinondoni procurement list format.
A Shopping List row with the Code of an existing item updates that item; a row without one adds a new item, in one room or
in several (a quick pick such as "All bedrooms", or rooms separated by ";"), exactly like ticking rooms on the item form."""
import io
import re
import uuid
from fastapi import APIRouter, Request, Depends, UploadFile, File, Form
from fastapi.responses import Response
from sqlalchemy.orm import Session
from openpyxl import load_workbook, Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from ..db import get_db
from ..models import Project, Room, Item, Supplier
from ..common import render, redirect, require_login, get_project, ffloat, content_disposition
from ..services import next_code, guess_kind, set_price
from .. import config, drawings, storage
from .items import new_photo
from PIL import Image as PILImage
import json

router = APIRouter(dependencies=[Depends(require_login)])

CAT_MAP = {"tiles": "Tiles", "sanitary": "Plumbing & Sanitary", "plumbing": "Plumbing & Sanitary", "lighting": "Lighting",
           "windows": "Windows & Doors", "switches & sockets": "Electrical", "electrical": "Electrical", "door hardware": "Hardware",
           "hardware": "Hardware", "kitchen": "Cabinets", "cabinets": "Cabinets", "dining": "Furniture", "living room": "Furniture",
           "master bedroom": "Furniture", "furniture": "Furniture", "water treatment (ro)": "Water Treatment",
           "water treatment": "Water Treatment", "tools & spares": "Other", "paint": "Paint", "flooring": "Flooring",
           "countertops": "Countertops", "appliances": "Appliances", "curtains & soft": "Curtains & Soft"}
APPLIANCE_WORDS = ("hob", "oven", "hood", "dishwasher", "refrigerator", "fridge", "microwave", "washing", "dryer", "water heater")
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def norm(s):
    return str(s or "").strip().lower()


def text(v) -> str:
    """A cell as typed: '' for an empty one. Line breaks inside a spec or a note are kept."""
    return str(v).strip() if v not in (None, "") else ""


def find_header(ws):
    for row in ws.iter_rows(min_row=1, max_row=12):
        vals = [norm(c.value) for c in row]
        if "item" in vals and "qty" in vals:
            return row[0].row, {v: i for i, v in enumerate(vals) if v}
    return None, {}


def col(h, row, *names):
    for n in names:
        if n in h:
            v = row[h[n]]
            return v if v is not None else ""
    return ""


@router.get("/p/{project_id}/import")
def import_page(request: Request, p: Project = Depends(get_project)):
    return render(request, "import.html", p=p, floors=drawings.floor_order(p), result=None)


# ---- the workbook: what the import reads, prefilled with this project's rooms (and, for the export, its items) ----------------
ROOM_COLS = [("Room", 30), ("Floor", 14), ("Kind", 10), ("Floor area m²", 14), ("Wall tile m²", 14), ("Notes", 42)]
ITEM_COLS = [("Code", 12), ("Room", 32), ("Category", 20), ("Item", 36), ("Photo", 17), ("Must-have spec", 46), ("Brand", 16), ("Size", 18),
             ("Finish", 18), ("Qty", 8), ("Unit", 8), ("Price (CNY)", 12), ("Supplier", 24), ("Lead time", 12), ("Status", 12), ("Optional", 10), ("Notes", 30)]
TOTAL_COLS = [("Total CNY", 12), ("Total USD", 12)]  # the filled export only: calculated, grey, ignored by the import
PHOTO_PX = (112, 84)  # the cover picture in the export's Photo column
PHOTO_ROW_PT = 66  # a row height that shows it


def _letter(cols, name: str) -> str:
    return get_column_letter([n for n, _ in cols].index(name) + 1)
PICK_KIND = "pick"  # the Kind of the quick-pick rows at the top of the Rooms sheet: not rooms, skipped by the import
HOW_TO = [
    ("How to fill this in", True),
    ("One row = one room or area on the Rooms sheet, one item on the Shopping List. Yellow cells are yours to type; do not rename the header rows.", False),
    ("", False),
    ("Rooms sheet", True),
    ("Room: CODE - Name, e.g. FF-BR2 - Bedroom 2. Short upper-case codes; GF for the ground floor, FF for the first floor is a good habit.", False),
    ("Your existing rooms and areas are already listed. Leave them, or correct a floor or kind; they are matched by their code, never duplicated.", False),
    ("The first rows (All rooms, All areas, All bedrooms, All bathrooms, All on ... floor) are quick picks for the Shopping List's Room column, not rooms. Leave them as they are.", False),
    ("Floor: Ground, First, Second, Roof, Site... It must match the floor you give the floor plans. Kind: room or area (entrance, corridors, stairs, balconies, carport, whole house).", False),
    ("Floor area and wall tile m² are optional and only help with tile quantities.", False),
    ("", False),
    ("Shopping List sheet", True),
    ("Code: leave it empty for a new item. A row with the code of an existing item (Import page → Export all items) updates that item: the cells you fill in change, blank cells leave the value alone, a single - clears a cell. The code itself never changes.", False),
    ("Room: pick one from the dropdown, pick All bedrooms / All bathrooms / All rooms / All areas / All on a floor, or type several rooms with ; between them (GF-LIV - Living room; FF-BR1 - Bedroom 1). A new item is created once per room, each with its own code. Leave it empty, or pick ALL - Whole house, for an item that belongs to no room.", False),
    ("Category decides which page of the client schedule the item prints on: one per item. Item is the name the client reads; Must-have spec is what the supplier must deliver.", False),
    ("Photo: the export shows each item's picture here. For a new item, place a picture over its Photo cell (Insert > Pictures > Place over Cells) and it is added with the item. Or upload picture files with the sheet, named after the item code (FF-MBR-04.jpg) or the item name (Bedside wall light.jpg, Bedside wall light-2.jpg). Pictures only go to items that have no photo yet.", False),
    ("Qty and Unit are per room. Price (CNY) is the unit price. Supplier: a new name is created in your supplier book. Status: To buy, Quoted, Ordered, Paid, Shipped, Received. Optional: yes for an alternative the client may skip.", False),
    ("", False),
    ("Importing", True),
    ("Import page → choose this file → Import. Rooms are matched by code and created when new. Item rows with a known code update that item; the others are added as new items.", False),
    ("To change many items at once: Import page → Export all items, edit the file, import it back. An update keeps the photos an item has; the pictures in the export are there to see, not re-imported.", False),
    ("Rows whose Code, Room or Item starts with (example) are ignored, so you can keep notes to yourself in the sheet.", False),
]
YELLOW = PatternFill("solid", fgColor="FFF8E6")
GREY = PatternFill("solid", fgColor="EFEBE4")
HEAD = PatternFill("solid", fgColor="1F3A5F")


def _sheet(wb, title, cols, n_rows=400):
    ws = wb.create_sheet(title)
    for i, (name, width) in enumerate(cols, start=1):
        c = ws.cell(row=1, column=i, value=name)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD
        c.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = width
        for r in range(2, n_rows + 2):
            ws.cell(row=r, column=i).fill = YELLOW
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 22
    return ws


def _list_validation(ws, col_letter, formula, rows=400):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=False)
    ws.add_data_validation(dv)
    dv.add(f"{col_letter}2:{col_letter}{rows + 1}")


def quick_picks(p) -> list[tuple[str, str]]:
    """The Room column's group entries, (label, note), listed at the top of the Rooms sheet. The import expands them to
    several rooms (RoomIndex.pick), exactly like the chips on the item form."""
    rooms = [r for r in p.rooms if r.code != "ALL"]
    out = [("All rooms", "every room"), ("All areas", "every area")]
    if any(r.group == "bedroom" for r in rooms):
        out.append(("All bedrooms", "every bedroom"))
    if any(r.group == "bathroom" for r in rooms):
        out.append(("All bathrooms", "every bathroom"))
    for f in drawings.floor_order(p):
        if any(drawings.floor_key(r.floor) == drawings.floor_key(f) for r in rooms):
            out.append((f"All on {drawings.floor_title(f, p)}", f"every room and area on the {drawings.floor_title(f, p).lower()}"))
    return [(label, f"Quick pick for the Shopping List: a new item in {what}. Not a room; leave this row as it is.") for label, what in out]


def build_workbook(p: Project, db: Session, items: list | None = None) -> bytes:
    """The Excel the import reads back. Without `items`: the template (Rooms prefilled, an empty Shopping List with dropdowns,
    the Lists behind them, a How-to). With `items`: the export for editing, the same sheets with every item filled in, its Code
    first, plus read-only Total columns."""
    wb = Workbook()
    how = wb.active
    how.title = "How to"
    how.column_dimensions["A"].width = 120
    for i, (t, bold) in enumerate(HOW_TO, start=1):
        c = how.cell(row=i, column=1, value=t)
        c.font = Font(bold=bold, size=12 if bold else 11)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    rooms = _sheet(wb, "Rooms", ROOM_COLS)
    row = 2
    for label, note in quick_picks(p):
        for j, v in enumerate((label, None, PICK_KIND, None, None, note), start=1):
            rooms.cell(row=row, column=j, value=v).font = Font(italic=True, color="857C72")
        row += 1
    for r in p.rooms:
        for j, v in enumerate((r.label, r.floor, r.kind, r.floor_area or None, r.wall_area or None, r.notes or None), start=1):
            rooms.cell(row=row, column=j, value=v)
        row += 1
    n_items = max(400, len(items or []) + 200)
    cols = ITEM_COLS + (TOTAL_COLS if items is not None else [])
    sheet = _sheet(wb, "Shopping List", cols, n_rows=n_items)
    if items is not None:
        for j in range(len(ITEM_COLS) + 1, len(cols) + 1):  # the totals are calculated, not typed: grey
            for r_ in range(2, n_items + 2):
                sheet.cell(row=r_, column=j).fill = GREY
        for i, it in enumerate(items, start=2):
            vals = [it.code, it.room.label if it.room else None, it.category, it.name, None, it.spec, it.brand, it.size, it.finish, it.qty, it.unit,
                    it.unit_price or None, it.supplier.name if it.supplier else None, it.lead_time, it.status, "yes" if it.optional else "no",
                    it.notes, it.total, round(it.total / (p.rate or 1), 2)]
            for j, v in enumerate(vals, start=1):
                sheet.cell(row=i, column=j, value=v if v != "" else None)
        # each item's cover picture, small, over its Photo cell (kept in memory until the workbook is saved)
        storage.prefetch([it.cover.thumb for it in items if it.cover])
        photo_col, buffers = _letter(ITEM_COLS, "Photo"), []
        for i, it in enumerate(items, start=2):
            data = storage.read_image(it.cover.thumb) if it.cover else None
            if not data:
                continue
            try:
                pic = PILImage.open(io.BytesIO(data)).convert("RGB")
                pic.thumbnail(PHOTO_PX)
            except Exception:
                continue
            b = io.BytesIO()
            pic.save(b, "JPEG", quality=78)
            b.seek(0)
            img = XLImage(b)
            img.width, img.height = pic.size
            sheet.add_image(img, f"{photo_col}{i}")
            sheet.row_dimensions[i].height = PHOTO_ROW_PT
            buffers.append(b)
    lists = wb.create_sheet("Lists")
    floors = list(dict.fromkeys([f for f in drawings.floor_order(p) if f] + ["Ground", "First", "Second", "Roof", "Site", "All"]))
    sups = [s.name for s in db.query(Supplier).order_by(Supplier.name)]
    columns = [("Categories", config.CATEGORIES), ("Units", config.UNITS), ("Statuses", config.STATUSES), ("Kind", config.ROOM_KINDS),
               ("Floors", floors), ("Suppliers", sups), ("Optional", ["yes", "no"])]
    refs = {}
    for j, (title, values) in enumerate(columns, start=1):
        lists.cell(row=1, column=j, value=title).font = Font(bold=True)
        for i, v in enumerate(values, start=2):
            lists.cell(row=i, column=j, value=v)
        lists.column_dimensions[get_column_letter(j)].width = 26
        if values:
            refs[title] = f"Lists!${get_column_letter(j)}$2:${get_column_letter(j)}${len(values) + 1}"
    _list_validation(rooms, _letter(ROOM_COLS, "Floor"), refs["Floors"])
    _list_validation(rooms, _letter(ROOM_COLS, "Kind"), refs["Kind"])
    _list_validation(sheet, _letter(ITEM_COLS, "Room"), "Rooms!$A$2:$A$401", rows=n_items)
    _list_validation(sheet, _letter(ITEM_COLS, "Category"), refs["Categories"], rows=n_items)
    _list_validation(sheet, _letter(ITEM_COLS, "Unit"), refs["Units"], rows=n_items)
    _list_validation(sheet, _letter(ITEM_COLS, "Status"), refs["Statuses"], rows=n_items)
    _list_validation(sheet, _letter(ITEM_COLS, "Optional"), refs["Optional"], rows=n_items)
    if "Suppliers" in refs:
        _list_validation(sheet, _letter(ITEM_COLS, "Supplier"), refs["Suppliers"], rows=n_items)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@router.get("/p/{project_id}/import/template.xlsx")
def import_template(p: Project = Depends(get_project), db: Session = Depends(get_db)):
    return Response(content=build_workbook(p, db), media_type=XLSX, headers=content_disposition(f"{p.name} - import template.xlsx"))


# ---- reading the Room column: one room, the whole house, a quick pick or several rooms -----------------------------------------
_PICK_RE = re.compile(r"^(?:all|every)\s+(rooms?|areas?|bedrooms?|bathrooms?)$")
_PICK_PREFIX_RE = re.compile(r"^(?:all|every|whole)\s+(?:on\s+|of\s+)?")
_SEP_RE = re.compile(r"[;,|/\n]")  # what separates several rooms in one cell


def is_pick(label: str, kind: str = "") -> bool:
    """A quick-pick row of the Rooms sheet (not a room): by its Kind, or by its label when the Kind cell went missing."""
    k = " ".join((label or "").lower().split())
    return kind == PICK_KIND or bool(_PICK_RE.match(k)) or k.startswith("all on ")


class RoomIndex:
    """The project's rooms by label ("GF-KIT - Kitchen"), code and name, and what a Room cell means. Rooms the Rooms sheet
    adds are fed in with add() as they are created (the session does not expire on commit, so p.rooms would be stale)."""

    def __init__(self, p: Project):
        self.rooms: list = []
        self.by_label, self.by_code, self.by_name = {}, {}, {}
        for r in p.rooms:
            self.add(r)

    def add(self, r: Room) -> None:
        if r not in self.rooms:
            self.rooms.append(r)
        self.by_label[r.label.lower()] = r
        self.by_code[r.code.lower()] = r
        self.by_name.setdefault(r.name.strip().lower(), r)

    def one(self, s: str):
        """One room; None for the whole house; False when the text names no room. A room is named by its label
        ("GF-KIT - Kitchen"), its name, its code, or its code with another name after " - " (a renamed room); the code
        shortcut is not taken for text that lists several rooms or says "All bedrooms"."""
        k = " ".join(s.lower().split())
        if not k:
            return False
        if k in ("all", "whole house", "all - whole house", "none", "no room"):
            return None
        r = self.by_label.get(k) or self.by_name.get(k)
        if r is None and not _SEP_RE.search(k) and (" " not in k or re.match(r"^\S+ - ", k)):
            r = self.by_code.get(k.split(" ")[0])
        if r is None:
            return False
        return None if r.code == "ALL" else r

    def pick(self, s: str):
        """The rooms a quick pick names ("All bedrooms", "All on First floor", "Ground floor"), or None when `s` is not one."""
        k = " ".join(s.lower().split())
        rooms = [r for r in self.rooms if r.code != "ALL"]
        m = _PICK_RE.match(k)
        if m:
            w = m.group(1).rstrip("s")
            if w in ("room", "area"):
                return [r for r in rooms if r.kind == w]
            return [r for r in rooms if r.group == w]
        fk = drawings.floor_key(_PICK_PREFIX_RE.sub("", k))
        if fk and not drawings.is_pseudo(fk) and any(drawings.floor_key(r.floor) == fk for r in rooms):
            return [r for r in rooms if drawings.floor_key(r.floor) == fk]
        return None

    def from_cell(self, s: str):
        """[None] for a blank cell or the whole house; the rooms the cell names (one room, a quick pick, or several separated
        by ; , | / or a line break); [] when a pick matches no room; None when the text names nothing known."""
        s = (s or "").strip()
        if not s:
            return [None]
        r = self.one(s)
        if r is not False:
            return [r]
        pk = self.pick(s)
        if pk is not None:
            return pk
        out, seen, known = [], set(), False
        for part in _SEP_RE.split(s):
            if not part:
                continue
            r = self.one(part)
            if r is not False:
                lst, known = [r], True
            else:
                pk = self.pick(part)
                lst, known = pk or [], known or pk is not None
            for x in lst:
                key = x.id if x is not None else 0
                if key not in seen:
                    seen.add(key)
                    out.append(x)
        return out if known else None


# ---- reading a Shopping List row, adding or updating ---------------------------------------------------------------------------
FIELDS = ("name", "category", "spec", "brand", "size", "finish", "qty", "unit", "unit_price", "supplier_id", "lead_time", "status",
          "optional", "notes", "room_id")


def _snapshot(item: Item) -> tuple:
    return tuple(getattr(item, f) for f in FIELDS)


def read_row(h, row) -> dict:
    """What a row says, as typed: '' for a blank cell; a '-' is kept so an update can clear the field."""
    name = text(col(h, row, "item", "product", "product name"))
    optional = None
    if "(optional)" in name.lower():
        optional, name = True, re.sub(r"(?i)\s*\(optional\)", "", name).strip()
    opt = norm(col(h, row, "optional"))
    if opt in ("yes", "y", "1", "true"):
        optional = True
    elif opt in ("no", "n", "0", "false"):
        optional = False
    cat_raw, cat = norm(col(h, row, "category")), ""
    if cat_raw:
        cat = CAT_MAP.get(cat_raw, cat_raw.title())
        if cat not in config.CATEGORIES:
            cat = "Other"
        if cat == "Cabinets" and any(w in name.lower() for w in APPLIANCE_WORDS):
            cat = "Appliances"
    status = text(col(h, row, "status"))
    return dict(code=text(col(h, row, "code", "id")), room=text(col(h, row, "room", "location")), name=name, category=cat,
                spec=text(col(h, row, "must-have spec", "spec", "specification")), brand=text(col(h, row, "brand")),
                size=text(col(h, row, "size")) or text(col(h, row, "size / finish")),
                finish=text(col(h, row, "finish", "colour / finish", "color / finish")), qty=col(h, row, "qty"), unit=text(col(h, row, "unit")),
                price=col(h, row, "price (cny)", "unit price cny", "unit price", "price"), supplier=text(col(h, row, "supplier")),
                lead_time=text(col(h, row, "lead time")), status=status if status in config.STATUSES else "", optional=optional,
                notes=text(col(h, row, "notes")))


def apply_update(item: Item, f: dict, rooms, sup, p: Project, result: dict) -> None:
    """Change an existing item from a row: filled cells change, blank cells leave the value alone, a '-' clears a text cell.
    The code is never rewritten. A changed room moves the item and drops its dots (they sat in the old room's box)."""
    def put(attr, val):
        if val == "-":
            setattr(item, attr, "")
        elif val:
            setattr(item, attr, val)
    if f["name"]:
        item.name = f["name"]
    if f["category"]:
        item.category = f["category"]
    for attr in ("spec", "brand", "size", "finish", "lead_time", "notes"):
        put(attr, f[attr])
    if f["qty"] not in (None, ""):
        item.qty = ffloat(f["qty"], item.qty)
    if f["unit"] and f["unit"] != "-":
        item.unit = f["unit"]
    if f["price"] not in (None, ""):
        set_price(item, ffloat(f["price"], item.unit_price), "CNY", p.rate)
    if f["supplier"] == "-":
        item.supplier_id = None
    elif sup is not None:
        item.supplier_id = sup.id
    if f["status"]:
        item.status = f["status"]
    if f["optional"] is not None:
        item.optional = f["optional"]
    if f["room"] and rooms:
        if len(rooms) == 1:
            rid = rooms[0].id if rooms[0] is not None else None
            if rid != item.room_id:
                item.room_id = rid
                item.pins.clear()
        else:
            result["warnings"].append(f"{item.code}: several rooms listed, but an existing item stays in its room. Add it to more rooms from its page.")


def _supplier(db: Session, sups: dict, name: str, result: dict):
    if not name or name == "-" or name.lower().startswith("(example)"):
        return None
    s = sups.get(name.lower())
    if s is None:
        s = Supplier(name=name)
        db.add(s)
        db.flush()
        sups[name.lower()] = s
        result["suppliers"] += 1
    return s


FIELD_LABELS = {"name": "Item", "category": "Category", "spec": "Spec", "brand": "Brand", "size": "Size", "finish": "Finish", "qty": "Qty",
                "unit": "Unit", "unit_price": "Price (CNY)", "supplier_id": "Supplier", "lead_time": "Lead time", "status": "Status",
                "optional": "Optional", "notes": "Notes", "room_id": "Room"}


def _fmt(field: str, v, rooms: dict, sups: dict) -> str:
    """A field value as the preview shows it: names for ids, yes / no, trimmed numbers, '' for blanks."""
    if field == "room_id":
        return rooms.get(v, "Whole house") if v else "Whole house"
    if field == "supplier_id":
        return sups.get(v, "") if v else ""
    if field == "optional":
        return "yes" if v else "no"
    if isinstance(v, float):
        return f"{v:,.2f}".rstrip("0").rstrip(".") if v else ""
    return str(v or "")


_SUFFIX_RE = re.compile(r"(?:[\s_-]+\d{1,2}|\s*\(\d{1,2}\))$")


def _photo_keys(filename: str) -> list[str]:
    """What a picture file may be named after, most exact first: the stem as typed, then without a '-2' / ' (2)' suffix,
    each as a code key (lower) and a name key (underscores as spaces)."""
    stem = re.sub(r"\.[A-Za-z0-9]+$", "", filename.replace("\\", "/").rsplit("/", 1)[-1]).strip()
    out = []
    for cand in (stem, _SUFFIX_RE.sub("", stem)):
        for k in (" ".join(cand.lower().split()), " ".join(cand.lower().replace("_", " ").split())):
            if k and k not in out:
                out.append(k)
    return out


def _embedded_pictures(ws) -> dict[int, list[bytes]]:
    """Pictures placed over cells of the sheet, by row number (openpyxl keeps them with their anchor)."""
    out: dict[int, list[bytes]] = {}
    for im in getattr(ws, "_images", []) or []:
        try:
            row = im.anchor._from.row + 1
            data = im._data()
        except Exception:
            continue
        if data:
            out.setdefault(row, []).append(data)
    return out


def run_import(p: Project, db: Session, wb, replace: bool, dry: bool, photos: list | None = None) -> dict:
    """Read the workbook into the project: rooms first, then the Shopping List (new items, one per room, and updates by
    code). Returns the counts, a list of what was (or would be) added, changed and skipped, and warnings. With `dry`,
    nothing is kept: the whole run is rolled back, so the Import page can show a preview before anything is saved.
    `photos` = [(file name, bytes)] uploaded with the sheet; with the pictures placed over the sheet's Photo cells they go
    to the items they name (code or item name) when those have no photo yet, after the commit."""
    result = {"rooms": 0, "items": 0, "updated": 0, "unchanged": 0, "suppliers": 0, "skipped": 0, "deleted": 0, "warnings": [],
              "new": [], "changes": [], "new_rooms": [], "new_floors": [], "floors": 0, "room_changes": [], "new_suppliers": [], "photos": 0, "photos_kept": 0, "error": None, "dry": dry}
    photos = photos or []
    warnings: list[str] = []

    def warn(msg: str):
        if msg not in warnings:
            warnings.append(msg)
    photos_to_drop = []
    try:
        if replace:
            doomed = list(p.items)
            result["deleted"] = len(doomed)
            photos_to_drop = [ph for i in doomed for ph in i.photos]
            for i in doomed:
                db.delete(i)
            db.flush()
            db.expire(p)  # the session keeps objects after a flush: reload the item collections without the deleted rows

        # ---- rooms ----
        rooms_by_code = {r.code.lower(): r for r in p.rooms}
        idx = RoomIndex(p)  # fed with the rooms the sheet adds, so an item row can name them
        if "Rooms" in wb.sheetnames:
            ws = wb["Rooms"]
            hdr_row = None
            for row in ws.iter_rows(min_row=1, max_row=8):
                vals = [norm(c.value) for c in row]
                if "room" in vals:
                    hdr_row, h = row[0].row, {v: i for i, v in enumerate(vals) if v}
                    break
            if hdr_row:
                sort = max([r.sort for r in p.rooms if r.code != "ALL"] + [0])
                for row in ws.iter_rows(min_row=hdr_row + 1, values_only=True):
                    label = text(col(h, row, "room"))
                    kind = norm(col(h, row, "kind", "type"))
                    if not label or label.upper().startswith("TOTAL") or label.lower().startswith("(example)") or is_pick(label, kind):
                        continue  # blank, a totals line, a note to self, or a quick pick of the Room dropdown (not a room)
                    if " - " in label:
                        code, name = label.split(" - ", 1)
                    else:
                        code, name = label[:12].upper().replace(" ", "-"), label
                    code = code.strip().upper()
                    r = rooms_by_code.get(code.lower())
                    new_room = r is None
                    if not r:
                        sort += 1
                        r = Room(project_id=p.id, code=code, name=name.strip(), sort=sort)
                        db.add(r)
                        result["rooms"] += 1
                    before = None if new_room else (r.floor, r.kind)
                    floor_cell = text(col(h, row, "floor"))
                    if floor_cell:
                        if not drawings.is_pseudo(floor_cell) and drawings.floor_entry(p, floor_cell) is None:
                            result["new_floors"].append(" ".join(floor_cell.split())[:40])  # joins the floor list (Rooms page)
                        r.floor = drawings.register_floor(db, p, floor_cell)
                    else:
                        r.floor = r.floor or ""
                    if kind in ("room", "area"):
                        r.kind = kind
                    elif new_room:
                        r.kind = guess_kind(r.name, r.floor)  # entrance, corridors, stairs, balconies... = area
                    r.floor_area = ffloat(col(h, row, "floor area m²", "floor area", "floor area m2"), r.floor_area)
                    r.wall_area = ffloat(col(h, row, "wall tile m²", "wall tile m2", "wall area"), r.wall_area)
                    r.notes = str(col(h, row, "notes") or r.notes or "")
                    db.flush()
                    rooms_by_code[r.code.lower()] = r
                    idx.add(r)
                    if new_room:
                        result["new_rooms"].append(f"{r.label} ({r.floor or 'no floor'}, {r.kind})")
                    elif (r.floor, r.kind) != before:  # an existing room moved or reclassified: say so in the preview
                        what = ([f"floor {before[0] or 'none'} → {r.floor or 'none'}"] if r.floor != before[0] else []) + ([f"{before[1]} → {r.kind}"] if r.kind != before[1] else [])
                        result["room_changes"].append(f"{r.label}: " + ", ".join(what))
                result["floors"] = len(result["new_floors"])

        # ---- items ----
        sheet = next((n for n in wb.sheetnames if n.lower() in ("shopping list", "items", "master list")), None)
        if not sheet:
            for n in wb.sheetnames:
                hr, _ = find_header(wb[n])
                if hr:
                    sheet = n
                    break
        if not sheet:
            result["error"] = "No sheet with Item and Qty columns found."
            db.rollback()
            return result
        ws = wb[sheet]
        hdr_row, h = find_header(ws)
        sups = {s.name.lower(): s for s in db.query(Supplier).all()}
        sup_names = {s.id: s.name for s in sups.values()}
        room_labels = {r.id: r.label for r in idx.rooms}
        live_by_code = {i.code.lower(): i for i in p.live_items if i.code}
        had_photos = {i.id for i in p.live_items if i.photos}  # pictures only go to items that have none yet: never a duplicate
        row_items: dict[int, list] = {}  # sheet row -> the items it added or updated (for a picture placed over that row)
        all_items = list(p.live_items)
        for rno, row in enumerate(ws.iter_rows(min_row=hdr_row + 1, values_only=True), start=hdr_row + 1):
            if not any(v not in (None, "") for v in row):
                continue  # a blank row (the template carries hundreds of styled empty ones): not worth counting
            f = read_row(h, row)
            if any(f[k].lower().startswith("(example)") for k in ("code", "room", "name")):
                result["skipped"] += 1
                continue
            existing = live_by_code.get(f["code"].lower()) if f["code"] else None
            if existing is None and not f["name"]:
                result["skipped"] += 1
                continue
            raw_status = text(col(h, row, "status"))
            if raw_status and not f["status"]:
                warn(f'Status "{raw_status}" is not one of {", ".join(config.STATUSES)}: ' + ("left as it is." if existing else '"To buy" is used.'))
            rooms = idx.from_cell(f["room"])
            if rooms is None:
                warn(f'Room "{f["room"]}" not found: ' + ("the item stays where it is." if existing else "the item goes to Whole house. Add the room on the Rooms sheet, or pick an existing one."))
                rooms = [] if existing else [None]
            elif not rooms:
                warn(f'"{f["room"]}" matches no room' + ("." if existing else f': "{f["name"]}" is not added.'))
                if existing is None:
                    result["skipped"] += 1
                    continue
            sup = _supplier(db, sups, f["supplier"], result)
            if sup is not None and sup.id not in sup_names:
                sup_names[sup.id] = sup.name
                result["new_suppliers"].append(sup.name)
            if existing is not None:
                row_items[rno] = [existing]
                before = _snapshot(existing)
                apply_update(existing, f, rooms, sup, p, result)
                after = _snapshot(existing)
                if after != before:
                    result["updated"] += 1
                    fields = [(FIELD_LABELS[name], _fmt(name, a, room_labels, sup_names), _fmt(name, b, room_labels, sup_names))
                              for name, a, b in zip(FIELDS, before, after) if a != b]
                    result["changes"].append({"code": existing.code, "name": existing.name, "fields": fields})
                else:
                    result["unchanged"] += 1
                continue
            for room in rooms:  # one new item per room, like the room checklist of the item form
                item = Item(project_id=p.id, room_id=room.id if room else None, category=f["category"] or "Other", name=f["name"],
                            brand=_clean(f["brand"]), spec=_clean(f["spec"]), size=_clean(f["size"]), finish=_clean(f["finish"]),
                            qty=ffloat(f["qty"], 1), unit=_clean(f["unit"]) or "pcs", supplier_id=sup.id if sup else None,
                            lead_time=_clean(f["lead_time"]), status=f["status"] or "To buy", optional=bool(f["optional"]), notes=_clean(f["notes"]))
                set_price(item, ffloat(f["price"]), "CNY", p.rate)
                code = f["code"]
                item.code = code if code and len(rooms) == 1 and not any(i.code == code for i in p.items) else next_code(db, p, room)
                db.add(item)
                db.flush()
                p.items.append(item)
                all_items.append(item)
                row_items.setdefault(rno, []).append(item)
                result["items"] += 1
                result["new"].append({"code": item.code, "name": item.name, "room": room.label if room else "Whole house", "category": item.category,
                                      "qty": _fmt("qty", item.qty, {}, {}) + " " + item.unit, "price": _fmt("unit_price", item.unit_price, {}, {}),
                                      "status": item.status, "supplier": sup.name if sup else ""})
        # ---- pictures: over the sheet's rows, or files named after a code or an item name ----
        pending: list[tuple] = []  # (item, bytes)
        by_code = {i.code.lower(): [i] for i in all_items if i.code}
        by_name = {}
        for i in all_items:
            by_name.setdefault(" ".join(i.name.lower().split()), []).append(i)
        def give(targets, data):
            for it in targets:
                if it.id in had_photos:
                    result["photos_kept"] += 1
                else:
                    pending.append((it, data))
        for rno, datas in _embedded_pictures(ws).items():
            for data in datas:
                give(row_items.get(rno, []), data)
        for name, data in photos:
            targets = []
            for k in _photo_keys(name):
                targets = by_code.get(k) or by_name.get(k) or []
                if targets:
                    break
            if not targets:
                warn(f'Picture "{name}" matches no item code or name.')
                continue
            give(targets, data)
        result["photos"] = len(pending)
        if result["photos_kept"]:
            warn(f'{result["photos_kept"]} picture{"s" if result["photos_kept"] != 1 else ""} not added: those items already have a photo (change photos on the item page).')
        counts: dict[str, int] = {}
        for it, _ in pending:
            counts[it.code] = counts.get(it.code, 0) + 1
        for e in result["new"]:
            e["photos"] = counts.get(e["code"], 0)
        listed = {e["code"] for e in result["changes"]}
        for e in result["changes"]:
            if counts.get(e["code"]):
                e["fields"].append(("Photo", "", f'{counts[e["code"]]} added'))
        for it, _ in pending:  # an existing item whose only change is the picture
            if it.code in counts and it.code not in listed and it.id in {x.id for x in live_by_code.values()}:
                result["changes"].append({"code": it.code, "name": it.name, "fields": [("Photo", "", f"{counts[it.code]} added")]})
                result["updated"] += 1
                listed.add(it.code)
        result["warnings"] = warnings[:12] + ([f"… and {len(warnings) - 12} more."] if len(warnings) > 12 else [])
        if dry:
            db.rollback()
        else:
            db.commit()
            for ph in photos_to_drop:  # the files of the items "replace" deleted, once the deletion is final
                storage.delete_photo(ph)
            for it, data in pending:  # pictures go through storage, so after the rows are final
                db.add(new_photo(it.id, data, p))
            db.commit()
        return result
    except Exception:
        db.rollback()
        raise


def _import_key(p: Project, key: str) -> str:
    """The storage key of a file kept between the preview and the apply step; '' unless it is this project's."""
    return key if key.startswith(f"imports/p{p.id}/") and storage.safe_key(key) else ""


def _manifest_key(key: str) -> str:
    return key[:-5] + ".json"


def _load_photos(key: str) -> list[tuple[str, bytes]]:
    """The pictures kept with a sheet between the preview and the apply step: [(file name, bytes)]."""
    raw = storage.read_image(_manifest_key(key))
    if not raw:
        return []
    out = []
    for entry in json.loads(raw.decode("utf-8")):
        data = storage.read_image(entry["key"])
        if data:
            out.append((entry["name"], data))
    return out


def _drop_kept(key: str) -> None:
    raw = storage.read_image(_manifest_key(key))
    if raw:
        for entry in json.loads(raw.decode("utf-8")):
            storage.delete_image(entry["key"])
        storage.delete_image(_manifest_key(key))
    storage.delete_image(key)


@router.post("/p/{project_id}/import")
async def do_import(request: Request, p: Project = Depends(get_project), db: Session = Depends(get_db),
                    file: UploadFile | None = File(None), photos: list[UploadFile] = File(default=[]), replace: str = Form(""),
                    apply: str = Form(""), key: str = Form("")):
    """Step one (a file, and optionally pictures): keep them, run the import without saving, show what it would do. Step two
    (the kept file's key with apply=1): run it for real. A file posted with apply=1 skips the preview."""
    key = _import_key(p, key)
    pics: list[tuple[str, bytes]] = []
    if key:
        data = storage.read_image(key)
        if not data:
            return render(request, "import.html", p=p, floors=drawings.floor_order(p), result={"error": "The uploaded file is no longer here. Choose it again."})
        pics = _load_photos(key)
    else:
        data = await file.read() if file is not None else b""
        if not data:
            return render(request, "import.html", p=p, floors=drawings.floor_order(p), result={"error": "Choose an .xlsx file first."})
        for up in photos:
            pdata = await up.read()
            if pdata and up.filename:
                pics.append((up.filename, pdata))
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True)
    except Exception as e:
        return render(request, "import.html", p=p, floors=drawings.floor_order(p), result={"error": f"Could not read the file: {e}"})
    if apply == "1":
        result = run_import(p, db, wb, replace == "1", dry=False, photos=pics)
        if key:
            _drop_kept(key)
        return render(request, "import.html", p=p, floors=drawings.floor_order(p), result=result)
    if not key:
        key = storage.save_blob(data, f"imports/p{p.id}/{uuid.uuid4().hex}.xlsx", XLSX)
        if pics:
            kept = []
            for n, (name, pdata) in enumerate(pics):
                ext = (name.rsplit(".", 1)[-1].lower() if "." in name else "jpg")[:5]
                pkey = storage.save_blob(pdata, f"{key[:-5]}-{n:03d}.{re.sub(r'[^a-z0-9]', '', ext) or 'jpg'}", storage.content_type(name))
                kept.append({"key": pkey, "name": name})
            storage.save_blob(json.dumps(kept).encode("utf-8"), _manifest_key(key), "application/json")
    preview = run_import(p, db, wb, replace == "1", dry=True, photos=pics)
    if preview["error"]:
        _drop_kept(key)
        return render(request, "import.html", p=p, floors=drawings.floor_order(p), result=preview)
    return render(request, "import.html", p=p, floors=drawings.floor_order(p), result=None, preview=preview, key=key, replace=replace == "1")


@router.post("/p/{project_id}/import/cancel")
def cancel_import(p: Project = Depends(get_project), key: str = Form("")):
    key = _import_key(p, key)
    if key:
        _drop_kept(key)
    return redirect(f"/p/{p.id}/import")


def _clean(v: str) -> str:
    return "" if v == "-" else v
