import io
from fastapi import APIRouter, Request, Depends, HTTPException
from fastapi.responses import Response
from sqlalchemy.orm import Session
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from ..db import get_db
from ..models import Project, Supplier, Room, Payment
from ..common import require_login, get_project, get_settings
from ..pdf.schedule import build_schedule
from ..pdf.packing import packing_list, labels, room_checklist, purchase_order
from .items import sort_items

router = APIRouter()


def pdf(data: bytes, name: str, inline=True):
    disp = "inline" if inline else "attachment"
    return Response(content=data, media_type="application/pdf", headers={"Content-Disposition": f'{disp}; filename="{name}"'})


@router.get("/p/{project_id}/export/schedule.pdf", dependencies=[Depends(require_login)])
def schedule_pdf(p: Project = Depends(get_project), db: Session = Depends(get_db), currency: str = "USD", prices: int = 1, photos: int = 1,
                 layout: str = "category"):
    by_floor = layout == "floor"
    data = build_schedule(p, get_settings(db), currency=("CNY" if currency == "CNY" else "USD"), show_prices=bool(prices),
                          include_photos=bool(photos), layout="floor" if by_floor else "category")
    return pdf(data, f"FFE-Schedule-{p.name}{'-by-floor' if by_floor else ''}.pdf")


@router.get("/c/{token}/schedule.pdf")
def client_schedule_pdf(token: str, db: Session = Depends(get_db), currency: str = "USD", layout: str = "category"):
    p = db.query(Project).filter(Project.client_token == token).first()
    if not p:
        raise HTTPException(404)
    by_floor = layout == "floor"
    data = build_schedule(p, get_settings(db), currency=("CNY" if currency == "CNY" else "USD"), layout="floor" if by_floor else "category")
    return pdf(data, f"FFE-Schedule-{p.name}{'-by-floor' if by_floor else ''}.pdf")


@router.get("/p/{project_id}/export/packing.pdf", dependencies=[Depends(require_login)])
def packing_pdf(p: Project = Depends(get_project), db: Session = Depends(get_db), supplier: int = 0):
    cartons = [c for c in p.cartons if (not supplier or c.supplier_id == supplier)]
    return pdf(packing_list(p, get_settings(db), cartons), f"Packing-{p.name}.pdf")


@router.get("/p/{project_id}/export/labels.pdf", dependencies=[Depends(require_login)])
def labels_pdf(p: Project = Depends(get_project), db: Session = Depends(get_db), supplier: int = 0, room: int = 0):
    cartons = [c for c in p.cartons if (not supplier or c.supplier_id == supplier) and (not room or c.room_id == room)]
    return pdf(labels(p, get_settings(db), cartons), f"Labels-{p.name}.pdf")


@router.get("/s/{token}/labels.pdf")
def supplier_labels_pdf(token: str, db: Session = Depends(get_db)):
    from ..models import SupplierLink
    l = db.query(SupplierLink).filter(SupplierLink.token == token, SupplierLink.active == True).first()  # noqa: E712
    if not l:
        raise HTTPException(404)
    cartons = [c for c in l.project.cartons if c.supplier_id == l.supplier_id]
    return pdf(labels(l.project, get_settings(db), cartons), "Labels.pdf")


@router.get("/s/{token}/packing.pdf")
def supplier_packing_pdf(token: str, db: Session = Depends(get_db)):
    from ..models import SupplierLink
    l = db.query(SupplierLink).filter(SupplierLink.token == token, SupplierLink.active == True).first()  # noqa: E712
    if not l:
        raise HTTPException(404)
    cartons = [c for c in l.project.cartons if c.supplier_id == l.supplier_id]
    return pdf(packing_list(l.project, get_settings(db), cartons, title=f"Packing List - {l.supplier.name}"), "Packing.pdf")


@router.get("/p/{project_id}/export/room/{room_id}.pdf", dependencies=[Depends(require_login)])
def room_pdf(room_id: int, p: Project = Depends(get_project), db: Session = Depends(get_db)):
    room = db.get(Room, room_id) if room_id else None
    if room is not None and room.project_id != p.id:
        raise HTTPException(404)
    return pdf(room_checklist(p, get_settings(db), room), f"Room-{room.code if room else 'ALL'}.pdf")


@router.get("/p/{project_id}/export/po/{supplier_id}.pdf", dependencies=[Depends(require_login)])
def po_pdf(supplier_id: int, p: Project = Depends(get_project), db: Session = Depends(get_db)):
    s = db.get(Supplier, supplier_id)
    if not s:
        raise HTTPException(404)
    items = sort_items([i for i in p.live_items if i.supplier_id == s.id], p)
    pays = [x for x in p.payments if x.supplier_id == s.id]
    return pdf(purchase_order(p, get_settings(db), s, items, pays), f"PO-{s.name}.pdf")


@router.get("/p/{project_id}/export/items.xlsx", dependencies=[Depends(require_login)])
def items_xlsx(p: Project = Depends(get_project), db: Session = Depends(get_db)):
    wb = Workbook()
    ws = wb.active
    ws.title = "Shopping List"
    heads = ["Code", "Room", "Category", "Item", "Brand", "Spec", "Size", "Finish", "Qty", "Unit", "Unit price CNY", "Total CNY",
             "Total USD", "Supplier", "Lead time", "Status", "Optional", "Notes"]
    ws.append(heads)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F3A5F")
    for i in sort_items(p.live_items, p):
        ws.append([i.code, i.room.label if i.room else "", i.category, i.name, i.brand, i.spec, i.size, i.finish, i.qty, i.unit,
                   i.unit_price, i.total, round(i.total / (p.rate or 1), 2), i.supplier.name if i.supplier else "", i.lead_time,
                   i.status, "Yes" if i.optional else "", i.notes])
    for col, w in zip("ABCDEFGHIJKLMNOPQR", [12, 26, 16, 36, 14, 44, 20, 18, 7, 6, 12, 12, 12, 20, 10, 10, 8, 30]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    ws2 = wb.create_sheet("Rooms")
    ws2.append(["Room", "Floor", "Floor area m²", "Wall tile m²", "Notes"])
    for r in p.rooms:
        ws2.append([r.label, r.floor, r.floor_area, r.wall_area, r.notes])
    buf = io.BytesIO()
    wb.save(buf)
    return Response(content=buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{p.name}-items.xlsx"'})
